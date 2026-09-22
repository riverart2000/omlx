#!/usr/bin/env python3
"""oMLX Kindle Book Studio — Grok-powered story, illustration and KDP exporter."""
from __future__ import annotations

import base64
import hashlib
import html
import io
import json
import math
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import traceback
import urllib.error
import urllib.parse
import urllib.request
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

COMMON_DIR = Path(os.environ.get("OMLX_COMMON_DIR", "/Users/joebains/omlx-common"))
if str(COMMON_DIR) not in sys.path:
    sys.path.insert(0, str(COMMON_DIR))
from studio_logging import EventLogger, read_events
from model_registry import get_model, public_models
from model_memory_coordinator import (
    acquire_lease, release_lease, unload_omlx_models,
)

HOST = os.environ.get("KINDLE_HOST", "127.0.0.1")
PORT = int(os.environ.get("KINDLE_PORT", "8800"))
BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
ENV_FILE = Path(os.environ.get("OMLX_ENV_FILE", "/Users/joebains/.omlx/.env"))
LIBRARY_DIR = Path(os.environ.get(
    "KINDLE_LIBRARY_DIR", "/Users/joebains/Documents/Kindle Books/Projects"))
EXPORT_DIR = Path(os.environ.get(
    "KINDLE_EXPORT_DIR", "/Users/joebains/Documents/Kindle Books/Exports"))
SERIES_CATALOG_FILE = Path(os.environ.get(
    "KINDLE_SERIES_CATALOG_FILE",
    str(LIBRARY_DIR.parent / "28-day-series-catalog.json")))
TEXT_MODEL = os.environ.get("GROK_BOOK_MODEL", "grok-4.5")
IMAGE_MODEL = os.environ.get("GROK_IMAGE_MODEL", "grok-imagine-image")
HUMAN_EDITORIAL_STANDARD = """
HUMAN EDITORIAL STANDARD — apply this to every reader-facing word:
- Write with the judgement of a senior developmental editor and book doctor who
  has spent years shaping successful human-authored manuscripts in this genre.
  Match the actual reader, subject and form instead of using a generic house voice.
- Prefer concrete nouns, active verbs, specific observations and earned emotion.
  Vary sentence length and cadence. Use contractions when they belong naturally.
  Allow a little personality, wit, surprise or restraint where the material earns it.
- Remove the tells of machine-written prose: symmetrical paragraph templates,
  constant three-item lists, repeated mini-summaries, excessive signposting,
  generic reassurance, inflated adjectives and a concluding slogan on every page.
- Never use vague promotional filler such as “embark on a journey”, “unlock your
  potential”, “discover the power of”, “in today's fast-paced world”, “delve
  into”, “a transformative experience”, “more than just”, “not just … but …”,
  “whether you're … or …”, “designed to empower”, or “a tapestry of”. Ordinary
  words such as journey or transform are allowed only when they are literal and
  genuinely the clearest words—not as decoration.
- Do not explain that the writing is helpful, engaging, comprehensive, meaningful
  or inspiring. Make it those things through the substance. Do not address the
  reader by name or invent testimonials, credentials, experiences or certainty.
- Read the result once as a human copy editor before returning it. Cut anything a
  skilled editor would call canned, padded, repetitive, over-polished or unnatural.
"""
HUMAN_LISTING_STANDARD = """
HUMAN BOOK-DESCRIPTION STANDARD:
- Write the description as a veteran jacket-copy editor and bookseller would:
  specific, confident, readable and honest. It must sound written for this exact
  book, not filled from a marketing template.
- Lead with the reader's real situation, the story's concrete hook, or a crisp
  question the book genuinely answers. Do not merely repeat the title followed by
  “is a practical guide/workbook/story”.
- Describe what actually happens or what the reader actually does. Use tangible
  details selected from the manuscript. For practical nonfiction, a short bullet
  list is optional only when it makes the contents easier to scan.
- Avoid breathless claims, generic benefit stacks, fake intimacy, sales clichés,
  repetitive “you will” sentences and formulaic final lines beginning “Perfect
  for”, “Ideal for”, “Whether”, “Step into”, “Get ready to”, or “By the end”.
- Do not mention prompts, AI, generation, illustration models, trim, metadata or
  production technique in the public description. AI disclosure belongs only in
  its separate publishing field.
- Aim for the shortest copy that sells the book accurately—normally 150–300 words,
  with natural paragraphs and no repeated conclusion.
"""
XAI_BASE = os.environ.get("XAI_BASE_URL", "https://api.x.ai/v1").rstrip("/")
AMAZON_CHILDRENS_BESTSELLERS_URL = (
    "https://www.amazon.com/Best-Sellers-Children's-eBooks/zgbs/"
    "digital-text/155009011"
)
LOCAL_IMAGE_BASE = os.environ.get(
    "OMLX_IMAGE_BASE_URL", "http://127.0.0.1:8400").rstrip("/")
TTS_BASE = os.environ.get("OMLX_TTS_BASE_URL", "http://127.0.0.1:8200").rstrip("/")
ORCHESTRATOR_BASE = os.environ.get(
    "OMLX_ORCHESTRATOR_BASE_URL", "http://127.0.0.1:8700").rstrip("/")
SFX_CLI = Path(os.environ.get(
    "OMLX_SFX_CLI",
    "/Users/joebains/stable-audio-3/optimized/mlx/sa3",
))
SFX_PYTHON = Path(os.environ.get(
    "OMLX_SFX_PYTHON", str(SFX_CLI.parent / ".venv" / "bin" / "python"),
))
SFX_SCRIPT = Path(os.environ.get(
    "OMLX_SFX_SCRIPT", str(SFX_CLI.parent / "scripts" / "sa3_mlx.py"),
))
SFX_MODEL_LABEL = "Stable Audio 3 Small SFX · local MLX"
SFX_MASTERING_VERSION = 2
SFX_MIX_VERSION = 4
NARRATION_SCRIPT_VERSION = 3
NARRATION_AUDIO_VERSION = 6
NARRATION_TTS_ENGINE = "qwen3_clone"
NARRATION_TTS_BATCH_SIZE = max(1, min(4, int(os.environ.get(
    "KINDLE_TTS_BATCH_SIZE", "3"))))
PUBLISHING_BUILD_VERSION = 6
TRASH_DIR = Path(os.environ.get(
    "KINDLE_TRASH_DIR", "/Users/joebains/.Trash/Kindle Book Studio"))
IMAGE_TEXT_DETECTOR = Path(os.environ.get(
    "OMLX_IMAGE_TEXT_DETECTOR",
    "/Users/joebains/omlx-image/bin/image-text-detector",
))
IMAGE_TEXT_REPAIR_PYTHON = Path(os.environ.get(
    "OMLX_IMAGE_TEXT_REPAIR_PYTHON",
    "/Users/joebains/omlx-image/.venv/bin/python",
))
IMAGE_TEXT_REPAIR_SCRIPT = Path(os.environ.get(
    "OMLX_IMAGE_TEXT_REPAIR_SCRIPT",
    "/Users/joebains/omlx-image/repair_image_text.py",
))
VIDEO_FOCUS_DETECTOR = Path(os.environ.get(
    "KINDLE_VIDEO_FOCUS_DETECTOR",
    str(BASE_DIR / "bin" / "image-focus-detector"),
))
LIBRARY_DIR.mkdir(parents=True, exist_ok=True)
EXPORT_DIR.mkdir(parents=True, exist_ok=True)
COST_SETTINGS_FILE = LIBRARY_DIR.parent / "cloud-cost-settings.json"
COST_LEDGER_FILE = LIBRARY_DIR.parent / "cloud-cost-ledger.jsonl"
COST_SESSION_ID = "session_" + uuid.uuid4().hex[:12]

EVENT_LOG = EventLogger("kindle-studio", "kindle-events.jsonl")
IMAGE_EVENT_LOG = EVENT_LOG.log_dir / "image-events.jsonl"


def studio_event(event: str, level="info", **fields):
    """Write a durable event with the active book/job correlation IDs."""
    fields.setdefault("job_id", getattr(_job_context, "job_id", ""))
    fields.setdefault("project_id", getattr(_job_context, "project_id", ""))
    return EVENT_LOG.event(event, level=level, **fields)

IMAGE_ENGINES = public_models("image", consumer="kindle")
_grok_registry_model = get_model("grok")
if _grok_registry_model:
    IMAGE_MODEL = _grok_registry_model["model"]
IMAGE_ENGINE_IDS = {x["id"] for x in IMAGE_ENGINES}
MAX_CHARACTER_REFERENCES = 4
MAX_GROK_REFERENCES = 3
KLEIN_IMAGE_ENGINES = {"flux_2_klein_4b", "flux_2_klein_4b_local"}
_project_create_lock = threading.Lock()
_library_maintenance_lock = threading.Lock()
_series_catalog_lock = threading.Lock()
_series_batch_lock = threading.Lock()
_export_name_locks_guard = threading.Lock()
_export_name_locks: dict[str, threading.Lock] = {}
_cost_lock = threading.RLock()
_cost_reservations: dict[str, dict] = {}

BOOK_TYPES = [
    ("picture_book", "Children’s picture book"),
    ("early_reader", "Early reader"),
    ("chapter_book", "Illustrated chapter book"),
    ("ancient_wisdom", "Ancient wisdom · faithful retelling"),
    ("comic", "Comic / graphic novel"),
    ("activity", "Activity book"),
    ("colouring", "Colouring book"),
    ("educational", "Educational workbook"),
    ("fiction", "Fiction novel"),
    ("nonfiction", "Non-fiction guide"),
    ("cookbook", "Cookbook"),
    ("poetry", "Poetry collection"),
    ("journal", "Journal / planner"),
]
PICTURE_BOOK_TYPES = [
    ("classic_story", "Classic illustrated story"),
    ("bedtime", "Bedtime story"),
    ("personalised_adventure", "Personalised child adventure"),
    ("educational_concept", "Educational concept book"),
    ("social_emotional", "Social and emotional learning"),
    ("values_lesson", "Values or life-lesson story"),
    ("rhyming", "Rhyming picture book"),
    ("alphabet", "Alphabet / ABC book"),
    ("counting", "Counting / numbers book"),
    ("seek_find", "Seek-and-find picture book"),
    ("wordless", "Wordless visual story"),
    ("holiday", "Holiday or seasonal story"),
    ("nature_animals", "Nature or animal story"),
    ("funny", "Funny picture book"),
    ("family_keepsake", "Family keepsake story"),
]
POPULAR_KIDS_BLUEPRINTS = [
    {
        "id": "funny_interactive", "label": "Funny interactive read-aloud",
        "formula": "A mischievous original character directly involves the child through predictions, choices, actions or repeated responses, with escalating comedy and a warm payoff.",
    },
    {
        "id": "everyday_bravery", "label": "Everyday worry to confidence",
        "formula": "A relatable small worry becomes a safe, funny adventure that gives the child a memorable coping idea and an emotionally satisfying victory.",
    },
    {
        "id": "animal_discovery", "label": "Animal curiosity and surprising facts",
        "formula": "An original animal-led story delivers surprising, accurate facts through play, comparison and imagination rather than a textbook voice.",
    },
    {
        "id": "rhythmic_repeat", "label": "Rhythmic repetition for preschoolers",
        "formula": "A strong original refrain, cumulative pattern and page-turn reveal invite participation and make the book enjoyable to read aloud repeatedly.",
    },
    {
        "id": "silly_cause_effect", "label": "Silly cause-and-effect chain",
        "formula": "One tiny choice triggers an increasingly surprising chain of visual consequences before resolving in a clever circular ending.",
    },
    {
        "id": "learning_adventure", "label": "Early learning inside an adventure",
        "formula": "Counting, alphabet, colours, shapes, first words or simple maths are necessary to solve an entertaining story problem.",
    },
    {
        "id": "how_it_works", "label": "How-it-works science adventure",
        "formula": "A curious original character investigates one concrete science or nature question through an accurate, visual, age-appropriate adventure.",
    },
    {
        "id": "friendship_kindness", "label": "Friendship and emotional skills",
        "formula": "A specific playground, sibling or friendship problem is resolved through believable choices, humour and a practical social-emotional insight.",
    },
    {
        "id": "gentle_gross_out", "label": "Gentle gross-out humour",
        "formula": "Child-safe mess, smells, noises or animal silliness drive a playful story with a clear plot, interactive moments and a reassuring ending.",
    },
    {
        "id": "repeatable_hero", "label": "Repeatable character-led series",
        "formula": "A distinctive original hero with a recognisable flaw, visual identity and recurring world solves one self-contained child-sized problem per book.",
    },
]
GENRES = [
    "Adventure", "Fantasy", "Mystery", "Comedy", "Bedtime", "Friendship",
    "Animals", "Science fiction", "Educational", "Historical", "Romance",
    "Thriller", "Self-help", "Business", "Health and wellness", "Custom",
]
IMAGE_STYLES = [
    ("cinematic_3d", "Cinematic 3D family animation"),
    ("claymation", "Claymation"),
    ("watercolour", "Watercolour storybook"),
    ("paper_cut", "Layered paper cut-out"),
    ("comic", "Bold comic-book art"),
    ("manga", "Manga / anime"),
    ("ink", "Pen-and-ink illustration"),
    ("pencil", "Coloured pencil"),
    ("retro", "Retro children’s cartoon"),
    ("photoreal", "Photorealistic"),
    ("colouring", "Clean colouring-book line art"),
    ("custom", "Custom visual style"),
]
TRIMS = [
    ("8x10", "8 × 10 in portrait", 8.0, 10.0),
    ("8.5x8.5", "8.5 × 8.5 in square", 8.5, 8.5),
    ("6x9", "6 × 9 in novel", 6.0, 9.0),
    ("7x10", "7 × 10 in", 7.0, 10.0),
    ("landscape", "10 × 8 in landscape", 10.0, 8.0),
]
PUBLISHING_PLATFORMS = {
    "kdp": {
        "label": "Amazon KDP",
        "url": "https://kdp.amazon.com/bookshelf",
        "purpose": "Amazon Kindle ebook and direct Amazon paperback",
    },
    "ingram": {
        "label": "IngramSpark",
        "url": "https://myaccount.ingramspark.com/",
        "template_url": (
            "https://myaccount.ingramspark.com/Portal/Tools/"
            "CoverTemplateGenerator"
        ),
        "purpose": "Bookshop, library and wide print distribution",
    },
    "d2d": {
        "label": "Draft2Digital",
        "url": "https://books2read.com/account/login",
        "purpose": "Wide ebook distribution outside Amazon",
    },
    "blurb": {
        "label": "Blurb",
        "url": "https://www.blurb.com/my/dashboard",
        "spec_url": "https://www.blurb.com/pdf-to-book",
        "purpose": "Optional premium direct-sale photo-book edition",
    },
    "lulu": {
        "label": "Lulu Direct",
        "url": "https://www.lulu.com/",
        "template_url": "https://developers.lulu.com/price-calculator",
        "spec_url": (
            "https://help.luludirect.lulu.com/en/support/solutions/articles/"
            "64000294595-pdf-creation-settings"
        ),
        "purpose": (
            "Direct-store print-on-demand fulfilment through Shopify, "
            "WooCommerce, Wix or the Lulu Print API"
        ),
    },
}
PUBLISHING_STATUSES = {
    "not_started", "account_ready", "files_ready", "proof_ordered",
    "proof_approved", "submitted", "under_review", "published",
}
QWEN3_NARRATORS = {
    "Eleanor": {"label": "Eleanor · warm British female", "accent": "british", "gender": "female", "quality": "Stable Clone BF16", "engine": "qwen3_clone"},
    "Poppy": {"label": "Poppy · lively British female", "accent": "british", "gender": "female", "quality": "Stable Clone BF16", "engine": "qwen3_clone"},
    "Maya": {"label": "Maya · warm American female", "accent": "american", "gender": "female", "quality": "Stable Clone BF16", "engine": "qwen3_clone"},
    "Harper": {"label": "Harper · lively American female", "accent": "american", "gender": "female", "quality": "Stable Clone BF16", "engine": "qwen3_clone"},
    "Arthur": {"label": "Arthur · warm British male", "accent": "british", "gender": "male", "quality": "Stable Clone BF16", "engine": "qwen3_clone"},
    "Oliver": {"label": "Oliver · lively British male", "accent": "british", "gender": "male", "quality": "Stable Clone BF16", "engine": "qwen3_clone"},
    "Noah": {"label": "Noah · warm American male", "accent": "american", "gender": "male", "quality": "Stable Clone BF16", "engine": "qwen3_clone"},
    "Jack": {"label": "Jack · lively American male", "accent": "american", "gender": "male", "quality": "Stable Clone BF16", "engine": "qwen3_clone"},
}
LEGACY_KOKORO_TO_QWEN3 = {
    "af_heart": "Maya", "af_bella": "Maya", "af_nicole": "Maya",
    "bf_emma": "Eleanor", "bf_isabella": "Eleanor",
    "am_fenrir": "Noah", "am_michael": "Noah", "am_puck": "Noah",
    "bm_fable": "Arthur", "bm_george": "Arthur",
}
BLURB_EDITIONS = {
    "7x7": {
        "label": "Small Square 7 × 7 photo book",
        "layout_trim": (7.0, 7.0),
    },
    "12x12": {
        "label": "Large Square 12 × 12 photo book",
        "layout_trim": (12.0, 12.0),
    },
}
READING_LEVELS = [
    "Ages 2–4", "Ages 4–6", "Ages 6–8", "Ages 8–12",
    "Young adult", "Adult general", "Professional / specialist",
]
PAGE_COUNTS = [4, 6, 12, 16, 20, 24, 28, 32, 40, 48, 64, 96, 128]
SERIES_NAME = "The 28-Day Inner Transformation Series"
SERIES_BOOKS = [
    {
        "id": "introspection", "title": "28 Days of Introspection",
        "aliases": ["Daily Introspection"],
        "focus": "deeper self-awareness through reflection, gratitude, intentions and honest daily review",
        "weekly_arc": ["observe yourself", "understand your patterns", "make conscious changes", "integrate and continue"],
        "palette": "warm ivory, sunrise gold, charcoal and soft sage",
    },
    {
        "id": "gratitude", "title": "28 Days of Gratitude",
        "focus": "building a sustainable gratitude practice that notices ordinary, relational and personal gifts",
        "weekly_arc": ["notice everyday gifts", "appreciate people and support", "find learning in difficulty", "live and express gratitude"],
        "palette": "honey gold, blush, warm cream and gentle green",
    },
    {
        "id": "self-confidence", "title": "28 Days of Self-Confidence",
        "focus": "recognising strengths, quieting self-doubt and taking increasingly courageous action",
        "weekly_arc": ["recognise existing strengths", "challenge limiting stories", "practise visible courage", "embody steady confidence"],
        "palette": "deep blue, amber, warm white and confident coral",
    },
    {
        "id": "calm", "title": "28 Days of Calm",
        "focus": "creating practical moments of calm, steadiness and restoration in ordinary daily life",
        "weekly_arc": ["settle the body", "quiet mental noise", "respond calmly to pressure", "build a sustainable calm rhythm"],
        "palette": "mist blue, sea glass, pale sand and soft lavender",
    },
    {
        "id": "self-love", "title": "28 Days of Self-Love",
        "focus": "developing kinder self-talk, healthy boundaries, self-respect and compassionate daily care",
        "weekly_arc": ["meet yourself kindly", "accept the whole self", "protect your needs", "live from self-respect"],
        "palette": "rose, plum, warm cream and muted gold",
    },
    {
        "id": "emotional-resilience", "title": "28 Days of Emotional Resilience",
        "focus": "understanding emotions, recovering from setbacks and responding with flexibility and self-compassion",
        "weekly_arc": ["name and allow emotions", "discover coping strengths", "reframe setbacks", "create a personal resilience plan"],
        "palette": "storm blue, fresh green, copper and clear sky",
    },
    {
        "id": "purpose-clarity", "title": "28 Days of Purpose and Clarity",
        "focus": "clarifying values, priorities, meaningful direction and the next practical steps",
        "weekly_arc": ["clear the noise", "identify values and strengths", "shape a meaningful vision", "commit to aligned action"],
        "palette": "indigo, parchment, sunlit yellow and forest green",
    },
    {
        "id": "better-habits", "title": "28 Days of Better Habits",
        "focus": "designing realistic routines through small actions, useful cues, reflection and compassionate consistency",
        "weekly_arc": ["understand current patterns", "design tiny changes", "strengthen consistency", "make the habits sustainable"],
        "palette": "fresh teal, tangerine, clean white and graphite",
    },
    {
        "id": "focus-productivity", "title": "28 Days of Focus and Productivity",
        "focus": "choosing meaningful priorities, reducing distraction and completing important work without burnout",
        "weekly_arc": ["discover attention patterns", "simplify priorities", "practise deep focus", "build a balanced productivity system"],
        "palette": "navy, electric blue, citrus and cool white",
    },
    {
        "id": "mindful-living", "title": "28 Days of Mindful Living",
        "focus": "bringing non-judgemental awareness into the senses, routines, relationships and choices",
        "weekly_arc": ["arrive in the senses", "bring presence to routines", "relate with awareness", "carry mindfulness forward"],
        "palette": "moss, stone, water blue and natural linen",
    },
    {
        "id": "creativity", "title": "28 Days of Creativity",
        "focus": "reawakening curiosity, overcoming creative inhibition and establishing a playful creative practice",
        "weekly_arc": ["notice and collect inspiration", "play without judgement", "develop original ideas", "complete and share something"],
        "palette": "magenta, turquoise, sunshine yellow and ink",
    },
    {
        "id": "positive-thinking", "title": "28 Days of Positive Thinking",
        "focus": "building realistic optimism by noticing helpful possibilities without denying difficult feelings",
        "weekly_arc": ["notice thought patterns", "find balanced alternatives", "practise possibility and appreciation", "make optimism actionable"],
        "palette": "sunflower, sky blue, fresh white and warm orange",
    },
    {
        "id": "personal-growth", "title": "28 Days of Personal Growth",
        "focus": "reviewing identity, values, courage, relationships and goals to create an integrated growth plan",
        "weekly_arc": ["take an honest inventory", "stretch beyond old limits", "strengthen relationships and choices", "design the next chapter"],
        "palette": "emerald, midnight blue, warm gold and ivory",
    },
]


def normalize_series_catalog_entry(value: dict, custom: bool = False) -> dict:
    title = re.sub(r"\s+", " ", str(value.get("title") or "").strip())[:140]
    if title and not re.match(r"^28\s+days\s+of\b", title, re.I):
        title = "28 Days of " + title
    weekly = value.get("weekly_arc") or []
    if isinstance(weekly, str):
        weekly = [line.strip() for line in weekly.splitlines() if line.strip()]
    defaults = [
        "notice current patterns",
        "build understanding and useful tools",
        "practise small changes in daily life",
        "integrate the learning and continue",
    ]
    weekly = [str(item).strip()[:180] for item in weekly if str(item).strip()]
    weekly = (weekly + defaults[len(weekly):])[:4]
    topic = re.sub(r"^28\s+days\s+of\s+", "", title, flags=re.I).strip()
    identifier = str(value.get("id") or "").strip()
    if not identifier:
        identifier = ("custom-" if custom else "") + slug(topic or title, "workbook")
    return {
        "id": slug(identifier, "workbook"),
        "title": title,
        "aliases": [
            str(alias).strip()[:140] for alias in value.get("aliases") or []
            if str(alias).strip()
        ][:8],
        "focus": str(value.get("focus") or
                     f"building practical, sustainable growth in {topic.lower()}").strip()[:800],
        "approach": str(value.get("approach") or "").strip()[:1000],
        "weekly_arc": weekly,
        "palette": str(value.get("palette") or
                       "warm ivory, sunrise gold, soft sage and deep blue").strip()[:300],
        "custom": bool(value.get("custom", custom)),
        "created_at": str(value.get("created_at") or (now() if custom else "")),
    }


def custom_series_books() -> list[dict]:
    try:
        payload = json.loads(SERIES_CATALOG_FILE.read_text(encoding="utf-8"))
        rows = payload.get("books") if isinstance(payload, dict) else payload
        if not isinstance(rows, list):
            return []
        return [normalize_series_catalog_entry(row, custom=True)
                for row in rows if isinstance(row, dict) and row.get("title")]
    except (OSError, ValueError, TypeError):
        return []


def series_books_catalog() -> list[dict]:
    rows = [normalize_series_catalog_entry(row) for row in SERIES_BOOKS]
    seen = {row["title"].casefold() for row in rows}
    for row in custom_series_books():
        if row["title"].casefold() not in seen:
            rows.append(row)
            seen.add(row["title"].casefold())
    return rows


def add_series_book(data: dict) -> dict:
    entry = normalize_series_catalog_entry(data, custom=True)
    if not entry["title"]:
        raise ValueError("Enter a title or topic for the new 28-day workbook")
    with _series_catalog_lock:
        existing = series_books_catalog()
        if any(row["title"].casefold() == entry["title"].casefold()
               for row in existing):
            raise ValueError("That title is already in the 28-day series list")
        used_ids = {row["id"] for row in existing}
        base_id = entry["id"]
        suffix = 2
        while entry["id"] in used_ids:
            entry["id"] = f"{base_id}-{suffix}"
            suffix += 1
        custom = custom_series_books()
        custom.append(entry)
        SERIES_CATALOG_FILE.parent.mkdir(parents=True, exist_ok=True)
        tmp = SERIES_CATALOG_FILE.with_suffix(".json.tmp")
        tmp.write_text(json.dumps({"books": custom}, ensure_ascii=False, indent=2),
                       encoding="utf-8")
        tmp.replace(SERIES_CATALOG_FILE)
    studio_event("series.catalog_title_added", title=entry["title"],
                 series_entry_id=entry["id"])
    return entry

_jobs: dict[str, dict] = {}
_jobs_lock = threading.Lock()
_job_context = threading.local()


class ImageJobCanceled(RuntimeError):
    pass


def current_job_canceled() -> bool:
    job_id = getattr(_job_context, "job_id", "")
    if not job_id:
        return False
    with _jobs_lock:
        return bool((_jobs.get(job_id) or {}).get("cancel_requested"))


def cancel_image_jobs() -> dict:
    """Stop the local generator and mark all active Kindle image jobs canceled."""
    local_result = {"active_canceled": False, "queued_canceled": 0}
    try:
        local_result = local_image_json("/cancel", {}, timeout=15)
    except Exception:
        pass
    canceled = 0
    with _jobs_lock:
        for job in _jobs.values():
            if job.get("kind") == "image" and job.get("status") == "running":
                job.update(status="canceled", stage="stopped", progress=0,
                           error=None, cancel_requested=True,
                           finished_at=now(), updated_at=now())
                canceled += 1
    return {"ok": True, "kindle_jobs_canceled": canceled, **local_result}


def update_current_job(stage: str, progress: int | None = None) -> None:
    """Publish sub-job progress to the browser-visible studio activity."""
    job_id = getattr(_job_context, "job_id", "")
    if not job_id:
        return
    prefix = str(getattr(_job_context, "stage_prefix", "") or "").strip()
    if prefix:
        stage = prefix + " · " + str(stage or "working")
    progress_range = getattr(_job_context, "progress_range", None)
    if progress is not None and isinstance(progress_range, tuple):
        start, finish = progress_range
        progress = int(start + (finish - start) * max(
            0, min(100, int(progress or 0))) / 100)
    log_stage = False
    logged_progress = None
    with _jobs_lock:
        job = _jobs.get(job_id)
        if job and job.get("status") == "running":
            previous_stage = job.get("stage")
            previous_bucket = int(job.get("_logged_progress_bucket", -1))
            job["stage"] = str(stage or "working")
            if progress is not None:
                next_progress = max(0, min(99, int(progress or 0)))
                if getattr(_job_context, "progress_monotonic", False):
                    next_progress = max(int(job.get("progress") or 0), next_progress)
                job["progress"] = next_progress
            bucket = int(job.get("progress") or 0) // 10
            log_stage = previous_stage != job["stage"] or bucket != previous_bucket
            if log_stage:
                job["_logged_progress_bucket"] = bucket
                logged_progress = int(job.get("progress") or 0)
            job["updated_at"] = now()
    if log_stage:
        studio_event("job.progress", stage=str(stage or "working"),
                     progress=logged_progress)


def now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


class CloudBudgetExceeded(RuntimeError):
    pass


def cloud_cost_settings() -> dict:
    defaults = {
        "enabled": True,
        "daily_limit_usd": 5.0,
        "session_limit_usd": 2.0,
        "single_call_limit_usd": 0.5,
    }
    try:
        saved = json.loads(COST_SETTINGS_FILE.read_text())
        if isinstance(saved, dict):
            defaults.update(saved)
    except (OSError, ValueError, TypeError):
        pass
    defaults["enabled"] = bool(defaults.get("enabled", True))
    for key, fallback in (
            ("daily_limit_usd", 5.0), ("session_limit_usd", 2.0),
            ("single_call_limit_usd", 0.5)):
        try:
            defaults[key] = max(0.01, min(10000.0, float(defaults.get(key))))
        except (TypeError, ValueError):
            defaults[key] = fallback
    return defaults


def save_cloud_cost_settings(data: dict) -> dict:
    current = cloud_cost_settings()
    current["enabled"] = bool(data.get("enabled", current["enabled"]))
    for key in ("daily_limit_usd", "session_limit_usd", "single_call_limit_usd"):
        if key in data:
            current[key] = max(0.01, min(10000.0, float(data[key])))
    tmp = COST_SETTINGS_FILE.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(current, ensure_ascii=False, indent=2))
    tmp.replace(COST_SETTINGS_FILE)
    studio_event("cost.settings_updated", settings=current)
    return current


def _cloud_cost_rows() -> list[dict]:
    rows = []
    try:
        with COST_LEDGER_FILE.open() as handle:
            for line in handle:
                try:
                    row = json.loads(line)
                    if isinstance(row, dict):
                        rows.append(row)
                except (ValueError, TypeError):
                    continue
    except OSError:
        pass
    return rows


def _today_local() -> str:
    return datetime.now().astimezone().date().isoformat()


def cloud_cost_snapshot(project_id: str = "") -> dict:
    with _cost_lock:
        rows = _cloud_cost_rows()
        pending = list(_cost_reservations.values())
    today = _today_local()

    def amount(items):
        return round(sum(float(row.get("amount_usd") or 0) for row in items), 6)

    today_rows = [row for row in rows if row.get("local_date") == today]
    session_rows = [row for row in rows if row.get("session_id") == COST_SESSION_ID]
    book_rows = [row for row in rows if project_id and row.get("project_id") == project_id]
    book_categories = {
        key: amount([row for row in book_rows if row.get("category") == key])
        for key in ("writing", "grok_image", "replicate_image", "local_image")
    }
    pending_today = amount([row for row in pending if row.get("local_date") == today])
    pending_session = amount(
        [row for row in pending if row.get("session_id") == COST_SESSION_ID])
    pending_book = amount(
        [row for row in pending if project_id and row.get("project_id") == project_id])
    return {
        "currency": "USD",
        "settings": cloud_cost_settings(),
        "today_usd": amount(today_rows),
        "session_usd": amount(session_rows),
        "book_usd": amount(book_rows),
        "book_project_id": project_id,
        "book_breakdown": book_categories,
        "book_recent": list(reversed(book_rows[-20:])),
        "pending_today_usd": pending_today,
        "pending_session_usd": pending_session,
        "pending_book_usd": pending_book,
        "recent": list(reversed(rows[-20:])),
        "session_id": COST_SESSION_ID,
        "note": (
            "xAI charges use the provider's returned billed cost when available; "
            "Replicate amounts are conservative estimates. Local MLX work costs $0."
        ),
    }


def reserve_cloud_cost(provider: str, model: str, category: str,
                       estimate_usd: float, detail: str = "") -> str:
    estimate = max(0.0, round(float(estimate_usd or 0), 6))
    settings = cloud_cost_settings()
    with _cost_lock:
        snapshot = cloud_cost_snapshot(
            str(getattr(_job_context, "project_id", "") or ""))
        projected_daily = (
            snapshot["today_usd"] + snapshot["pending_today_usd"] + estimate)
        projected_session = (
            snapshot["session_usd"] + snapshot["pending_session_usd"] + estimate)
        if settings["enabled"]:
            reasons = []
            if estimate > settings["single_call_limit_usd"] + 1e-9:
                reasons.append(
                    f"the estimated ${estimate:.3f} call exceeds the "
                    f"${settings['single_call_limit_usd']:.2f} per-call limit")
            if projected_daily > settings["daily_limit_usd"] + 1e-9:
                reasons.append(
                    f"today would reach ${projected_daily:.3f}, above the "
                    f"${settings['daily_limit_usd']:.2f} daily limit")
            if projected_session > settings["session_limit_usd"] + 1e-9:
                reasons.append(
                    f"this session would reach ${projected_session:.3f}, above the "
                    f"${settings['session_limit_usd']:.2f} session limit")
            if reasons:
                message = (
                    "Cloud request blocked before it was sent: " + "; ".join(reasons)
                    + ". Open Cloud cost safety in Plan to change the limits. "
                    "No provider charge was created."
                )
                studio_event(
                    "cost.request_blocked", level="warning", provider=provider,
                    model=model, category=category, estimate_usd=estimate,
                    projected_daily_usd=projected_daily,
                    projected_session_usd=projected_session, reason=message,
                )
                raise CloudBudgetExceeded(message)
        reservation_id = "cost_" + uuid.uuid4().hex[:12]
        _cost_reservations[reservation_id] = {
            "reservation_id": reservation_id,
            "timestamp": now(), "local_date": _today_local(),
            "session_id": COST_SESSION_ID,
            "project_id": str(getattr(_job_context, "project_id", "") or ""),
            "job_id": str(getattr(_job_context, "job_id", "") or ""),
            "provider": provider, "model": model, "category": category,
            "amount_usd": estimate, "detail": detail,
        }
    studio_event(
        "cost.request_reserved", provider=provider, model=model,
        category=category, estimate_usd=estimate,
        reservation_id=reservation_id, detail=detail,
    )
    return reservation_id


def settle_cloud_cost(reservation_id: str, actual_usd: float | None = None,
                      estimated=False, status="charged") -> dict | None:
    with _cost_lock:
        row = _cost_reservations.pop(reservation_id, None)
        if not row:
            return None
        if actual_usd is not None:
            row["amount_usd"] = max(0.0, round(float(actual_usd), 6))
        row.update(
            settled_at=now(), estimated=bool(estimated), status=str(status))
        COST_LEDGER_FILE.parent.mkdir(parents=True, exist_ok=True)
        with COST_LEDGER_FILE.open("a") as handle:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
    studio_event(
        "cost.request_settled", provider=row.get("provider"),
        model=row.get("model"), category=row.get("category"),
        amount_usd=row.get("amount_usd"), estimated=row.get("estimated"),
        reservation_id=reservation_id, status=status,
    )
    return row


def release_cloud_cost(reservation_id: str, reason="not charged") -> None:
    with _cost_lock:
        row = _cost_reservations.pop(reservation_id, None)
    if row:
        studio_event(
            "cost.request_released", provider=row.get("provider"),
            model=row.get("model"), category=row.get("category"),
            reservation_id=reservation_id, reason=reason,
        )


def record_local_image_use(model: str, detail: str = "") -> None:
    reservation = reserve_cloud_cost(
        "local", model, "local_image", 0.0, detail=detail)
    settle_cloud_cost(reservation, 0.0, estimated=False, status="local")


def slug(value: str, fallback="book") -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "-", (value or "").strip()).strip("-").lower()
    return (s[:70].rstrip("-") or fallback)


def load_env() -> dict:
    vals = {}
    try:
        for line in ENV_FILE.read_text().splitlines():
            if "=" not in line or line.lstrip().startswith("#"):
                continue
            k, v = line.split("=", 1)
            vals[k.strip()] = v.strip().strip('"').strip("'")
    except OSError:
        pass
    return vals


def api_key() -> str:
    vals = load_env()
    return (os.environ.get("XAI_API_KEY") or os.environ.get("GROK_API_KEY")
            or vals.get("XAI_API_KEY") or vals.get("GROK_API_KEY") or "").strip()


def api_key_tag() -> str:
    return hashlib.sha256(api_key().encode()).hexdigest()[:12]


def project_dir(project_id: str) -> Path:
    if not re.fullmatch(r"[a-zA-Z0-9_-]{4,80}", project_id or ""):
        raise ValueError("invalid project id")
    return LIBRARY_DIR / project_id


def project_file(project_id: str) -> Path:
    return project_dir(project_id) / "project.json"


def normalize_series_bible(value: dict | None, project: dict | None = None) -> dict:
    value = dict(value or {})
    project = project or {}
    settings = project.get("settings") or {}
    characters = value.get("characters")
    if not isinstance(characters, list):
        characters = project.get("character_bible") or []
    character_rules = []
    for character in characters[:MAX_CHARACTER_REFERENCES]:
        if not isinstance(character, dict):
            continue
        character_rules.append({
            key: str(character.get(key) or "")
            for key in ("name", "role", "appearance", "personality",
                        "continuity_rules", "reference_image_prompt")
        })
    locks = dict(value.get("locks") or {})
    normalized_locks = {
        key: bool(locks.get(key, True))
        for key in ("characters", "world", "visual_style", "palette",
                    "audience_tone", "typography", "trim_print")
    }
    return {
        "enabled": bool(value.get("enabled", bool(
            settings.get("series_name") or value.get("series_name")))),
        "series_name": str(value.get("series_name") or
                           settings.get("series_name") or ""),
        "series_promise": str(value.get("series_promise") or
                              settings.get("series_hook") or ""),
        "audience": str(value.get("audience") or settings.get("audience") or ""),
        "genre": str(value.get("genre") or settings.get("genre") or ""),
        "tone": str(value.get("tone") or settings.get("tone") or ""),
        "image_style": str(value.get("image_style") or
                           settings.get("image_style") or "cinematic_3d"),
        "custom_style": str(value.get("custom_style") or
                            settings.get("custom_style") or ""),
        "palette": str(value.get("palette") or ""),
        "typography": str(value.get("typography") or
                           settings.get("font_style") or ""),
        "trim": str(value.get("trim") or settings.get("trim") or "8x10"),
        "layout": str(value.get("layout") or settings.get("layout") or "fixed"),
        "print_interior": str(value.get("print_interior") or
                              settings.get("print_interior") or "premium_colour"),
        "world_rules": str(value.get("world_rules") or
                           project.get("world_bible") or ""),
        "story_rules": str(value.get("story_rules") or ""),
        "characters": character_rules,
        "locks": normalized_locks,
        "updated_at": str(value.get("updated_at") or now()),
    }


def apply_series_bible(project: dict, bible: dict | None = None) -> dict:
    bible = normalize_series_bible(bible or project.get("series_bible"), project)
    if not bible.get("enabled"):
        project["series_bible"] = bible
        return project
    settings = project.setdefault("settings", {})
    locks = bible.get("locks") or {}
    if bible.get("series_name"):
        settings["series_name"] = bible["series_name"]
    if bible.get("series_promise"):
        settings["series_hook"] = bible["series_promise"]
    if locks.get("audience_tone"):
        for key in ("audience", "genre", "tone"):
            if bible.get(key):
                settings[key] = bible[key]
    if locks.get("visual_style"):
        for key in ("image_style", "custom_style"):
            if bible.get(key):
                settings[key] = bible[key]
    if locks.get("typography") and bible.get("typography"):
        settings["font_style"] = bible["typography"]
    if locks.get("trim_print"):
        for source, target in (
                ("trim", "trim"), ("layout", "layout"),
                ("print_interior", "print_interior")):
            if bible.get(source):
                settings[target] = bible[source]
    if locks.get("world") and bible.get("world_rules"):
        project["world_bible"] = bible["world_rules"]
    if locks.get("characters") and bible.get("characters"):
        existing = project.get("character_bible") or []
        preserved = {
            str(c.get("name") or "").casefold(): c for c in existing}
        locked_characters = json.loads(json.dumps(
            bible["characters"]))[:MAX_CHARACTER_REFERENCES]
        for index, character in enumerate(locked_characters):
            old = preserved.get(str(character.get("name") or "").casefold(),
                                existing[index] if index < len(existing) else {})
            for key in ("reference_image", "reference_file_id", "reference_slot"):
                if old.get(key):
                    character[key] = old[key]
        project["character_bible"] = locked_characters
    bible["updated_at"] = now()
    project["series_bible"] = bible
    return project


def save_project(project: dict) -> dict:
    pid = project["id"]
    project["publication"] = normalize_publication(project.get("publication"))
    project["review_approval"] = normalize_review_approval(
        project.get("review_approval"))
    project["publishing"] = normalize_publishing(
        project.get("publishing"), project)
    project["video_audio"] = normalize_video_audio(
        project.get("video_audio"), project)
    if project.get("series_bible"):
        project["series_bible"] = normalize_series_bible(
            project.get("series_bible"), project)
    if project.get("settings", {}).get("image_engine") in (
            "flux_2_klein_4b", "flux_2_klein_4b_local"):
        cover = project.get("cover") or {}
        if cover.get("image_prompt"):
            cover["image_prompt"] = local_wordless_description(
                cover["image_prompt"])
        for page in project.get("pages") or []:
            if page.get("image_prompt"):
                page["image_prompt"] = local_wordless_description(
                    page["image_prompt"])
    folder = project_dir(pid)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "images").mkdir(exist_ok=True)
    (folder / "revisions").mkdir(exist_ok=True)
    project["updated_at"] = now()
    tmp = folder / "project.json.tmp"
    tmp.write_text(json.dumps(project, ensure_ascii=False, indent=2))
    tmp.replace(folder / "project.json")
    return project


def asset_epoch(project: dict) -> int:
    try:
        return max(0, int(project.get("asset_epoch") or 0))
    except (TypeError, ValueError):
        return 0


def bump_asset_epoch(project: dict) -> int:
    project["asset_epoch"] = asset_epoch(project) + 1
    return project["asset_epoch"]


def preserve_server_assets(current: dict, incoming: dict) -> dict:
    """Do not let a stale browser tab erase private references or generated art."""
    # These are server-authored library records. A tab opened before a bulk
    # build or publication update must never erase them on its next save.
    incoming["build"] = current.get("build") or {}
    incoming["publishing_outputs"] = current.get("publishing_outputs") or {}
    incoming["series_automation"] = current.get("series_automation") or {}
    # Review approval is server-authored and tied to a content signature. An
    # older browser tab may make it stale, but must never erase its audit trail.
    incoming["review_approval"] = current.get("review_approval") or {}
    if "publication" not in incoming:
        incoming["publication"] = current.get("publication") or {}
    server_epoch = asset_epoch(current)
    client_epoch = asset_epoch(incoming)
    assets_changed_on_server = client_epoch != server_epoch
    incoming["asset_epoch"] = server_epoch
    incoming["reference_images"] = current.get("reference_images", [])
    current_cover = current.get("cover") or {}
    incoming_cover = incoming.setdefault("cover", {})
    if assets_changed_on_server:
        incoming_cover["image"] = current_cover.get("image", "")
    elif current_cover.get("image") and not incoming_cover.get("image"):
        incoming_cover["image"] = current_cover["image"]
    for key in ("image_warning", "image_versions", "image_repair"):
        if assets_changed_on_server:
            if key in current_cover:
                incoming_cover[key] = current_cover[key]
            else:
                incoming_cover.pop(key, None)
        elif current_cover.get(key) and key not in incoming_cover:
            incoming_cover[key] = current_cover[key]
    current_pages = {int(p.get("number", 0)): p for p in current.get("pages", [])}
    for page in incoming.get("pages", []):
        old = current_pages.get(int(page.get("number", 0)))
        if assets_changed_on_server:
            page["image"] = old.get("image", "") if old else ""
        elif old and old.get("image") and not page.get("image"):
            page["image"] = old["image"]
        for key in ("image_warning", "image_versions", "image_repair"):
            if assets_changed_on_server:
                if old and key in old:
                    page[key] = old[key]
                else:
                    page.pop(key, None)
            elif old:
                if old.get(key) and key not in page:
                    page[key] = old[key]
        # These server-authored scene contracts are not yet editable fields in
        # the browser. Keep them when an already-open tab saves an older copy.
        if old:
            for key in (
                    "character_directions", "supporting_characters",
                    "cast_complete", "continuity_objects"):
                if key in old and key not in page:
                    page[key] = old[key]
    current_character_list = current.get("character_bible", [])
    current_characters = {}
    for old_character in current_character_list:
        name_key = str(old_character.get("name", "")).strip().casefold()
        if name_key and name_key not in current_characters:
            current_characters[name_key] = old_character
    for index, character in enumerate(incoming.get("character_bible", [])):
        # The card position is the durable identity for an attached reference.
        # Names are editable and may temporarily be duplicated while designing
        # characters, so they must never decide which private image is kept.
        old = current_character_list[index] if index < len(current_character_list) else {}
        if not old:
            old = current_characters.get(
                str(character.get("name", "")).strip().casefold(), {})
        for key in ("reference_image", "reference_file_id"):
            if assets_changed_on_server:
                if old.get(key):
                    character[key] = old[key]
                else:
                    character.pop(key, None)
            elif old.get(key) and not character.get(key):
                character[key] = old[key]
    return incoming


def protected_series_identity_error(current: dict, incoming: dict) -> str:
    """Reject a stale browser save that is carrying another series book."""
    template = current.get("series_template") or {}
    if not template.get("lock_title"):
        return ""
    canonical_id = str(template.get("id") or "").strip()
    canonical_title = str(template.get("title") or current.get("title") or "").strip()
    incoming_template = incoming.get("series_template") or {}
    incoming_id = str(incoming_template.get("id") or "").strip()
    incoming_title = str(incoming.get("title") or "").strip()
    if incoming_id != canonical_id or incoming_title != canonical_title:
        return (
            f'This browser tab tried to save “{incoming_title or "another book"}” '
            f'over the protected series book “{canonical_title}”. Nothing was saved. '
            "Reload the intended book and try again.")
    return ""


def load_project(project_id: str) -> dict:
    path = project_file(project_id)
    if not path.exists():
        raise FileNotFoundError(project_id)
    project = json.loads(path.read_text())
    project["publication"] = normalize_publication(project.get("publication"))
    project["review_approval"] = normalize_review_approval(
        project.get("review_approval"))
    project["publishing"] = normalize_publishing(
        project.get("publishing"), project)
    project["video_audio"] = normalize_video_audio(
        project.get("video_audio"), project)
    settings = project.setdefault("settings", {})
    if settings.get("image_engine") == "flux_2_pro":
        settings["image_engine"] = "flux_2_klein_4b"
    if settings.get("image_engine") in {"hidream", "flux", "kontext"}:
        settings["image_engine"] = "flux_2_klein_4b_local"
    references = project.get("reference_images") or []
    characters = project.setdefault("character_bible", [])
    used_slots = set()
    normalized_references = []
    for fallback_slot, reference in enumerate(
            references[:MAX_CHARACTER_REFERENCES]):
        try:
            slot = int(reference.get("slot", fallback_slot))
        except (TypeError, ValueError):
            slot = fallback_slot
        if slot not in range(MAX_CHARACTER_REFERENCES) or slot in used_slots:
            slot = next((candidate for candidate in range(MAX_CHARACTER_REFERENCES)
                         if candidate not in used_slots), -1)
        if slot < 0:
            continue
        used_slots.add(slot)
        reference["slot"] = slot
        name = str(reference.get("character_name") or f"Character {slot + 1}").strip()
        while (len(characters) <= slot
               and len(characters) < MAX_CHARACTER_REFERENCES):
            character = {
                "name": name, "role": "", "appearance": "", "personality": "",
                "continuity_rules": "",
            }
            characters.append(character)
        character = characters[slot] if slot < len(characters) else None
        if character is not None:
            reference["character_name"] = str(character.get("name") or name).strip()
            character["reference_image"] = reference.get("path", "")
            character["reference_file_id"] = reference.get("file_id", "")
            character["reference_slot"] = slot
        normalized_references.append(reference)
    project["reference_images"] = sorted(
        normalized_references, key=lambda reference: reference["slot"])
    return project


def revision_snapshot(project: dict, label: str) -> None:
    folder = project_dir(project["id"]) / "revisions"
    folder.mkdir(exist_ok=True)
    name = datetime.now().strftime("%Y%m%d-%H%M%S") + "-" + slug(label) + ".json"
    (folder / name).write_text(json.dumps(project, ensure_ascii=False, indent=2))


def detach_active_artwork(project: dict) -> int:
    """Detach active illustrations without deleting their recoverable files."""
    detached = 0
    records = [project.setdefault("cover", {})] + list(project.get("pages") or [])
    for record in records:
        if record.get("image"):
            detached += 1
        record["image"] = ""
        record["approved"] = False
        for key in ("image_warning", "image_versions", "image_repair"):
            record.pop(key, None)
    if detached:
        bump_asset_epoch(project)
    return detached


def clear_project_artwork(project_id: str, reason: str = "",
                          disconnect_inherited_series: bool = False) -> dict:
    project = load_project(project_id)
    count = int(bool((project.get("cover") or {}).get("image")))
    count += sum(bool(page.get("image")) for page in project.get("pages") or [])
    if not count and not disconnect_inherited_series:
        return project
    revision_snapshot(project, "before-clearing-all-artwork")
    detached = detach_active_artwork(project)
    if disconnect_inherited_series:
        project["series_bible"] = {}
        project["series_template"] = {}
        for key in ("series_parent_id", "series_root_id", "series_topic",
                    "series_adventure"):
            project.pop(key, None)
        settings = project.setdefault("settings", {})
        settings["book_number"] = "1" if settings.get("series_name") else ""
        project["creative_identity_id"] = "concept_" + uuid.uuid4().hex[:12]
        if not detached:
            bump_asset_epoch(project)
    project["stage"] = "story" if project.get("pages") else "setup"
    project.setdefault("history", []).append({
        "at": now(),
        "action": (
            f"Detached {detached} active illustration{'s' if detached != 1 else ''}"
            + (f" — {str(reason).strip()}" if str(reason).strip() else "")
            + ("; disconnected inherited series identity"
               if disconnect_inherited_series else "")
        ),
    })
    saved = save_project(project)
    studio_event(
        "illustration.all_detached", project_id=project_id,
        image_count=detached, reason=str(reason or "manual reset"),
        asset_epoch=asset_epoch(saved),
        disconnected_inherited_series=disconnect_inherited_series,
    )
    return saved


def normalize_publication(value: dict | None) -> dict:
    value = dict(value or {})
    return {
        "published": bool(value.get("published")),
        "url": str(value.get("url") or "").strip()[:2000],
        "published_at": str(value.get("published_at") or "")[:80],
        "updated_at": str(value.get("updated_at") or "")[:80],
    }


def normalize_review_approval(value: dict | None) -> dict:
    value = dict(value or {})
    return {
        "approved": bool(value.get("approved")),
        "signature": str(value.get("signature") or "")[:80],
        "approved_at": str(value.get("approved_at") or "")[:80],
        "updated_at": str(value.get("updated_at") or "")[:80],
    }


def reader_visible_metadata(project: dict) -> dict:
    """Return publishing copy without internal authorship/provenance markers."""
    metadata = dict(project.get("metadata") or {})
    for key in (
            "description_source", "description_edited_at",
            "description_generated_at"):
        metadata.pop(key, None)
    return metadata


def book_review_signature(project: dict) -> str:
    """Fingerprint the reader-visible book, independent of build-tool versions."""
    assets = []
    rels = [str((project.get("cover") or {}).get("image") or "")]
    rels.extend(str(page.get("image") or "")
                for page in project.get("pages") or [])
    for rel in rels:
        path = image_abs(project, rel)
        try:
            stat = path.stat() if path else None
        except OSError:
            stat = None
        assets.append({
            "path": rel,
            "size": stat.st_size if stat else 0,
            "modified": stat.st_mtime_ns if stat else 0,
        })
    settings = dict(project.get("settings") or {})
    # Merely choosing another generator for a future regeneration is not a
    # reader-visible book change. The resulting image will invalidate approval.
    settings.pop("image_engine", None)
    transient_record_keys = {
        "approved", "text_approved", "image_warning", "image_versions",
        "image_repair",
    }
    cover = {
        key: value for key, value in (project.get("cover") or {}).items()
        if key not in transient_record_keys
    }
    pages = [
        {key: value for key, value in page.items()
         if key not in transient_record_keys}
        for page in (project.get("pages") or [])
    ]
    payload = {
        "title": project.get("title"),
        "subtitle": project.get("subtitle"),
        "story_summary": project.get("story_summary"),
        "settings": settings,
        "cover": cover,
        "pages": pages,
        "metadata": reader_visible_metadata(project),
        "assets": assets,
    }
    return hashlib.sha256(json.dumps(
        payload, sort_keys=True, ensure_ascii=False,
        separators=(",", ":")).encode()).hexdigest()[:24]


def project_review_status(project: dict) -> dict:
    record = normalize_review_approval(project.get("review_approval"))
    expected = book_review_signature(project)
    current = bool(record["approved"] and record["signature"] == expected)
    stale = bool(record["approved"] and record["signature"] != expected)
    return {
        "approved": current,
        "stale": stale,
        "approved_at": record["approved_at"],
        "expected_signature": expected,
    }


def publishing_build_signature(project: dict) -> str:
    """Fingerprint everything that can alter a publishing package."""
    assets = []
    rels = [str((project.get("cover") or {}).get("image") or "")]
    rels.extend(str(page.get("image") or "")
                for page in project.get("pages") or [])
    for rel in rels:
        path = image_abs(project, rel)
        try:
            stat = path.stat() if path else None
        except OSError:
            stat = None
        assets.append({
            "path": rel,
            "size": stat.st_size if stat else 0,
            "modified": stat.st_mtime_ns if stat else 0,
        })
    payload = {
        "build_version": PUBLISHING_BUILD_VERSION,
        "title": project.get("title"), "subtitle": project.get("subtitle"),
        "settings": project.get("settings") or {},
        "cover": project.get("cover") or {},
        "pages": project.get("pages") or [],
        "metadata": reader_visible_metadata(project),
        "publishing": normalize_publishing(project.get("publishing"), project),
        "assets": assets,
    }
    return hashlib.sha256(json.dumps(
        payload, sort_keys=True, ensure_ascii=False,
        separators=(",", ":")).encode()).hexdigest()[:24]


def publishing_build_status(project: dict) -> dict:
    pages = project.get("pages") or []
    missing = []
    if not pages:
        missing.append("story pages")
    if not image_abs(project, str((project.get("cover") or {}).get("image") or "")):
        missing.append("cover picture")
    if (project.get("settings") or {}).get("layout", "fixed") == "fixed":
        missing_pages = [str(page.get("number") or index + 1)
                         for index, page in enumerate(pages)
                         if not image_abs(project, str(page.get("image") or ""))]
        if missing_pages:
            missing.append(
                "page pictures " + ", ".join(missing_pages[:6])
                + ("…" if len(missing_pages) > 6 else ""))
    build = dict(project.get("build") or {})
    expected = publishing_build_signature(project)
    bundle = Path(str(build.get("bundle") or ""))
    if missing:
        state = "not_ready"
    elif (int(build.get("version") or 0) != PUBLISHING_BUILD_VERSION
          or build.get("signature") != expected
          or not bundle.is_file()):
        state = "outdated"
    else:
        state = "current"
    return {
        "state": state, "missing": missing,
        "expected_signature": expected,
        "build_version": int(build.get("version") or 0),
        "latest_version": PUBLISHING_BUILD_VERSION,
        "built_at": str(build.get("built_at") or ""),
        "bundle": str(build.get("bundle") or ""),
    }


def list_projects() -> list[dict]:
    out = []
    for path in LIBRARY_DIR.glob("*/project.json"):
        try:
            p = json.loads(path.read_text())
            publication = normalize_publication(p.get("publication"))
            review = project_review_status(p)
            build_status = publishing_build_status(p)
            out.append({
                "id": p["id"], "title": p.get("title") or "Untitled",
                "type": p.get("settings", {}).get("book_type"),
                "stage": p.get("stage", "setup"),
                "updated_at": p.get("updated_at", ""),
                "page_count": len(p.get("pages") or []),
                "published": publication["published"],
                "published_url": publication["url"],
                "published_at": publication["published_at"],
                "review_approved": review["approved"],
                "review_stale": review["stale"],
                "review_approved_at": review["approved_at"],
                "build_state": build_status["state"],
                "built_at": build_status["built_at"],
            })
        except Exception:
            continue
    def stable_title_key(row: dict):
        parts = re.split(r"(\d+)", str(row.get("title") or "").casefold())
        natural = tuple(
            (0, int(part)) if part.isdigit() else (1, part)
            for part in parts if part)
        return natural, str(row.get("id") or "")
    return sorted(out, key=stable_title_key)


def set_project_publication(project_id: str, data: dict) -> dict:
    project = load_project(project_id)
    published = bool(data.get("published"))
    url = str(data.get("url") or "").strip()
    if published and url:
        parsed = urllib.parse.urlparse(url)
        if parsed.scheme not in {"http", "https"} or not parsed.netloc:
            raise ValueError(
                "If you enter a published-book link, use the full address beginning with https://")
    previous = normalize_publication(project.get("publication"))
    project["publication"] = {
        "published": published,
        "url": url[:2000],
        "published_at": (
            previous["published_at"] or now() if published else ""),
        "updated_at": now(),
    }
    if (previous["published"] != published or previous["url"] != url):
        project.setdefault("history", []).append({
            "at": now(),
            "action": ("Marked as published" if published
                       else "Marked as not yet published")
                      + (" — " + url if published else ""),
        })
    return save_project(project)


def set_project_review_approval(project_id: str, data: dict) -> dict:
    project = load_project(project_id)
    approved = bool(data.get("approved"))
    previous = project_review_status(project)
    if approved:
        readiness = publishing_build_status(project)
        if readiness["missing"]:
            raise ValueError(
                "Finish the book before approving it: "
                + "; ".join(readiness["missing"]))
    timestamp = now()
    project["review_approval"] = {
        "approved": approved,
        "signature": book_review_signature(project) if approved else "",
        "approved_at": (
            timestamp if approved else ""),
        "updated_at": timestamp,
    }
    if previous["approved"] != approved or previous["stale"]:
        project.setdefault("history", []).append({
            "at": timestamp,
            "action": ("Book review approved" if approved
                       else "Book review approval removed"),
        })
    saved = save_project(project)
    studio_event(
        "project.review_approval_changed", project_id=project_id,
        approved=approved,
        signature=saved.get("review_approval", {}).get("signature", ""),
    )
    return saved


def _trash_destination(category: str, label: str) -> Path:
    folder = TRASH_DIR / category
    folder.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    destination = folder / f"{stamp}-{slug(label)}"
    if destination.exists():
        destination = folder / f"{stamp}-{slug(label)}-{uuid.uuid4().hex[:6]}"
    return destination


def delete_project_to_trash(project_id: str, confirm_title: str) -> dict:
    project = load_project(project_id)
    title = str(project.get("title") or "Untitled Book")
    if str(confirm_title or "").strip() != title:
        raise ValueError("The book title confirmation did not match")
    with _jobs_lock:
        active = [job for job in _jobs.values()
                  if job.get("status") == "running"
                  and (job.get("project_id") == project_id
                       or project_id in (job.get("project_ids") or []))]
    if active:
        raise RuntimeError("This book is currently being processed. Try again when it finishes.")
    source = project_dir(project_id)
    destination = _trash_destination("Deleted Books", title + "-" + project_id)
    shutil.move(str(source), str(destination))
    studio_event(
        "project.moved_to_trash", project_id=project_id, title=title,
        destination=str(destination))
    return {"ok": True, "id": project_id, "title": title,
            "recoverable": True, "trash_path": str(destination)}


_EXPORT_VERSION_RE = re.compile(
    r"^(?P<title>.+)-(?P<stamp>\d{8}-\d{6})(?P<suffix>.*)$")
_EXPORT_COMPONENTS = tuple(PUBLISHING_PLATFORMS)


def old_build_cleanup_plan() -> dict:
    groups: dict[tuple[str, str], dict[str, list[Path]]] = {}
    for path in EXPORT_DIR.iterdir():
        if not path.is_file():
            continue
        match = _EXPORT_VERSION_RE.match(path.name)
        if not match:
            continue
        suffix = match.group("suffix")
        component = "master"
        for platform in _EXPORT_COMPONENTS:
            if suffix.startswith("-" + platform):
                component = platform
                break
        key = (match.group("title"), component)
        groups.setdefault(key, {}).setdefault(match.group("stamp"), []).append(path)
    candidates = []
    kept_sets = 0
    for versions in groups.values():
        latest = max(versions)
        kept_sets += 1
        for stamp, paths in versions.items():
            if stamp != latest:
                candidates.extend(paths)
    total_bytes = 0
    for path in candidates:
        try:
            total_bytes += path.stat().st_size
        except OSError:
            pass
    return {
        "files": sorted(candidates, key=lambda path: path.name),
        "file_count": len(candidates), "bytes": total_bytes,
        "kept_build_sets": kept_sets,
    }


def cleanup_old_builds() -> dict:
    if not _library_maintenance_lock.acquire(blocking=False):
        raise RuntimeError("Another library maintenance task is already running")
    try:
        plan = old_build_cleanup_plan()
        if not plan["files"]:
            return {"ok": True, "moved_files": 0, "bytes": 0,
                    "recoverable": True, "trash_path": ""}
        destination = _trash_destination(
            "Old Publishing Builds", "publishing-build-cleanup")
        destination.mkdir(parents=True, exist_ok=True)
        moved = []
        for index, path in enumerate(plan["files"]):
            update_current_job(
                f"Moving old build {index + 1} of {plan['file_count']} to Trash",
                8 + int(index / max(1, plan["file_count"]) * 84))
            target = destination / path.name
            if target.exists():
                target = destination / (path.stem + "-" + uuid.uuid4().hex[:6]
                                        + path.suffix)
            shutil.move(str(path), str(target))
            moved.append(path.name)
        studio_event(
            "library.old_builds_moved_to_trash", file_count=len(moved),
            bytes=plan["bytes"], destination=str(destination))
        return {"ok": True, "moved_files": len(moved),
                "bytes": plan["bytes"], "recoverable": True,
                "trash_path": str(destination), "files": moved[:30]}
    finally:
        _library_maintenance_lock.release()


def library_status() -> dict:
    projects = list_projects()
    cleanup = old_build_cleanup_plan()
    return {
        "book_count": len(projects),
        "current": sum(row["build_state"] == "current" for row in projects),
        "outdated": sum(row["build_state"] == "outdated" for row in projects),
        "not_ready": sum(row["build_state"] == "not_ready" for row in projects),
        "published": sum(bool(row["published"]) for row in projects),
        "review_approved": sum(bool(row["review_approved"]) for row in projects),
        "review_stale": sum(bool(row["review_stale"]) for row in projects),
        "review_needed": sum(not row["review_approved"] for row in projects),
        "cleanup_files": cleanup["file_count"],
        "cleanup_bytes": cleanup["bytes"],
        "latest_build_version": PUBLISHING_BUILD_VERSION,
    }


def rebuild_outdated_books(parallel_books: int = 2) -> dict:
    """Rebuild stale books concurrently without mixing their project files."""
    if not _library_maintenance_lock.acquire(blocking=False):
        raise RuntimeError("Another library maintenance task is already running")
    parent_job_id = str(getattr(_job_context, "job_id", "") or "")
    try:
        rows = list_projects()
        total = len(rows)
        with _jobs_lock:
            job = _jobs.get(parent_job_id) or {}
            job["details"] = {
                "type": "book_rebuild",
                "phase": "checking",
                "checked": 0,
                "book_count": total,
                "active_titles": [],
                "active": 0,
                "completed": 0,
                "waiting": 0,
                "total": 0,
                "parallel_books": max(2, min(4, int(parallel_books or 2))),
            }
        rebuilt = []
        current = []
        not_ready = []
        failed = []
        targets = []
        for index, row in enumerate(rows):
            title = str(row.get("title") or "Untitled Book")
            update_current_job(
                f"Checking book {index + 1} of {total}: {title}",
                3 + int(index / max(1, total) * 7))
            with _jobs_lock:
                job = _jobs.get(parent_job_id) or {}
                details = job.setdefault("details", {})
                details.update({
                    "type": "book_rebuild", "phase": "checking",
                    "checked": index + 1, "book_count": total,
                    "active_titles": [title],
                })
            try:
                project = load_project(str(row["id"]))
                status = publishing_build_status(project)
                if status["state"] == "current":
                    current.append(title)
                    continue
                if status["state"] == "not_ready":
                    not_ready.append({
                        "id": row["id"], "title": title,
                        "missing": status["missing"],
                    })
                    continue
                targets.append({"id": str(row["id"]), "title": title})
            except Exception as exc:
                failed.append({
                    "id": row.get("id"), "title": title,
                    "error": f"{type(exc).__name__}: {exc}",
                })
                EVENT_LOG.exception(
                    "library.book_rebuild_failed", exc,
                    project_id=str(row.get("id") or ""), title=title)

        concurrency = min(
            max(2, min(4, int(parallel_books or 2))),
            max(1, len(targets)),
        )
        state_lock = threading.Lock()
        active: dict[str, str] = {}
        finished = 0

        def publish_state() -> None:
            with state_lock:
                active_titles = list(active.values())
                done = finished
                rebuilt_count = len(rebuilt)
                failed_count = len(failed)
            remaining = max(0, len(targets) - done - len(active_titles))
            stage = (f"Completed {done} of {len(targets)} book rebuilds"
                     if targets and done >= len(targets) else
                "Rebuilding " + ", ".join(active_titles[:2])
                + (f" and {len(active_titles) - 2} more" if len(active_titles) > 2 else "")
                if active_titles else "Preparing parallel book rebuilds"
            )
            with _jobs_lock:
                job = _jobs.get(parent_job_id) or {}
                job["stage"] = stage
                job["progress"] = 10 + int(done / max(1, len(targets)) * 87)
                job["queue_label"] = (
                    f"{len(active_titles)} active · {done} complete · "
                    f"{remaining} waiting · up to {concurrency} in parallel"
                )
                job["details"] = {
                    "type": "book_rebuild",
                    "phase": "rebuilding",
                    "active_titles": active_titles[:4],
                    "active": len(active_titles),
                    "completed": done,
                    "rebuilt": rebuilt_count,
                    "failed": failed_count,
                    "waiting": remaining,
                    "total": len(targets),
                    "parallel_books": concurrency,
                }
                job["updated_at"] = now()

        def rebuild_one(target: dict) -> None:
            nonlocal finished
            pid, title = target["id"], target["title"]
            _job_context.job_id = ""
            _job_context.project_id = pid
            with state_lock:
                active[pid] = title
            publish_state()
            try:
                export_project(pid)
                with state_lock:
                    rebuilt.append(title)
            except Exception as exc:
                with state_lock:
                    failed.append({
                        "id": pid, "title": title,
                        "error": f"{type(exc).__name__}: {exc}",
                    })
                EVENT_LOG.exception(
                    "library.book_rebuild_failed", exc,
                    project_id=pid, title=title)
            finally:
                with state_lock:
                    active.pop(pid, None)
                    finished += 1
                publish_state()
                _job_context.job_id = ""
                _job_context.project_id = ""

        if targets:
            publish_state()
            with ThreadPoolExecutor(
                    max_workers=concurrency,
                    thread_name_prefix="kindle-rebuild") as executor:
                futures = [executor.submit(rebuild_one, target)
                           for target in targets]
                for future in as_completed(futures):
                    future.result()
        rebuilt.sort(key=str.casefold)
        current.sort(key=str.casefold)
        not_ready.sort(key=lambda row: str(row["title"]).casefold())
        failed.sort(key=lambda row: str(row["title"]).casefold())
        with _jobs_lock:
            job = _jobs.get(parent_job_id) or {}
            details = job.setdefault("details", {})
            details.update({
                "type": "book_rebuild",
                "phase": "complete",
                "active_titles": [],
                "active": 0,
                "completed": len(targets),
                "waiting": 0,
                "total": len(targets),
                "checked": total,
                "already_current": len(current),
                "not_ready": len(not_ready),
                "rebuilt": len(rebuilt),
                "failed": len(failed),
            })
        studio_event(
            "library.outdated_books_rebuilt", checked=total,
            rebuilt=len(rebuilt), current=len(current),
            not_ready=len(not_ready), failed=len(failed),
            parallel_books=concurrency)
        return {
            "ok": not failed, "checked": total, "rebuilt": rebuilt,
            "current": current, "not_ready": not_ready, "failed": failed,
            "parallel_books": concurrency,
        }
    finally:
        with _jobs_lock:
            job = _jobs.get(parent_job_id) or {}
            job["queue_label"] = "No queued book rebuilds"
        _library_maintenance_lock.release()


def _series_entry_names(entry: dict) -> set[str]:
    return {
        str(name).strip().casefold()
        for name in [entry.get("title"), *(entry.get("aliases") or [])]
        if str(name or "").strip()
    }


def _latest_kdp_bundle(project: dict) -> str:
    tracked = str(((project.get("publishing_outputs") or {}).get("kdp") or {}).get(
        "bundle") or "")
    if tracked and Path(tracked).is_file():
        return tracked
    candidates = list(EXPORT_DIR.glob(f"{slug(project.get('title', 'book'))}-*-kdp.zip"))
    if not candidates:
        return ""
    try:
        return str(max(candidates, key=lambda path: path.stat().st_mtime_ns))
    except OSError:
        return ""


def _series_project_state(project: dict | None) -> dict:
    if not project:
        return {
            "created": False, "workflow_state": "not_created",
            "kdp_ready": False, "missing_images": 29,
            "project_id": "", "project_title": "", "automation": False,
        }
    pages = project.get("pages") or []
    story_ready = len(pages) == 28 and all(
        str(page.get("text") or "").strip() for page in pages)
    missing_images = int(not bool(image_abs(
        project, str((project.get("cover") or {}).get("image") or ""))))
    missing_images += sum(not bool(image_abs(project, str(page.get("image") or "")))
                          for page in pages)
    metadata = project.get("metadata") or {}
    metadata_ready = (
        bool(str(metadata.get("description") or "").strip())
        and len(metadata.get("keywords") or []) == 7
        and len(metadata.get("categories") or []) == 3
    )
    kdp_bundle = _latest_kdp_bundle(project)
    kdp_ready = bool(story_ready and not missing_images and metadata_ready
                     and kdp_bundle)
    automation = dict(project.get("series_automation") or {})
    if kdp_ready:
        workflow_state = "kdp_ready"
    elif not story_ready:
        workflow_state = "needs_story"
    elif missing_images:
        workflow_state = "needs_images"
    elif not metadata_ready:
        workflow_state = "needs_metadata"
    else:
        workflow_state = "needs_kdp_package"
    return {
        "created": True, "workflow_state": workflow_state,
        "kdp_ready": kdp_ready, "missing_images": missing_images,
        "project_id": str(project.get("id") or ""),
        "project_title": str(project.get("title") or ""),
        "automation": bool(automation.get("started_at")),
        "automation_status": str(automation.get("status") or ""),
        "automation_error": str(automation.get("error") or "")[:500],
        "kdp_bundle": kdp_bundle,
    }


def _series_project_for_entry(entry: dict) -> dict | None:
    names = _series_entry_names(entry)
    candidates = []
    for path in LIBRARY_DIR.glob("*/project.json"):
        try:
            project = load_project(path.parent.name)
        except Exception:
            continue
        if str(project.get("title") or "").strip().casefold() not in names:
            continue
        state = _series_project_state(project)
        score = (
            int(state["kdp_ready"]),
            int(state["automation"]),
            int(bool(project.get("pages"))),
            str(project.get("updated_at") or ""),
        )
        candidates.append((score, project))
    return max(candidates, key=lambda item: item[0])[1] if candidates else None


def series_catalog_status() -> dict:
    books = []
    for index, entry in enumerate(series_books_catalog(), 1):
        state = _series_project_state(_series_project_for_entry(entry))
        batch_pending = not state["created"] or (
            state["automation"] and not state["kdp_ready"])
        books.append({
            **entry, **state, "book_number": index,
            "batch_pending": batch_pending,
        })
    return {
        "series_name": SERIES_NAME, "books": books,
        "total": len(books),
        "created": sum(book["created"] for book in books),
        "kdp_ready": sum(book["kdp_ready"] for book in books),
        "not_created": sum(not book["created"] for book in books),
        "batch_pending": sum(book["batch_pending"] for book in books),
        "estimated_local_images": sum(
            book["missing_images"] for book in books if book["batch_pending"]),
        "image_engine": "flux_2_klein_4b_local",
        "image_engine_label": "FLUX 2 Klein 4B · local MLX",
    }


def _series_project_payload(entry: dict, source_settings: dict,
                            book_number: int, series_name: str) -> dict:
    settings = dict(source_settings or {})
    focus = entry.get("focus") or entry["title"]
    weekly = entry.get("weekly_arc") or []
    approach = str(entry.get("approach") or "").strip()
    prompt = (
        f'Create an entirely original guided adult workbook titled "{entry["title"]}", '
        f'book {book_number} in "{series_name}". It is a 28-day practical journey '
        f'focused on {focus}. The four weekly phases are: {"; ".join(weekly)}. '
        "Create exactly one day per page and exactly 28 content pages. Every day "
        "must contain 4 to 8 concise exercises using a useful mixture of reflection, "
        "gratitude, intention, practical action, journaling and evening review. Put "
        "exactly two completely blank lines between exercises to provide writing "
        "space. Make every heading, exercise, example, quotation and illustration "
        "concept distinct. Keep an encouraging, grounded adult tone without medical "
        f'claims. {"Distinctive approach: " + approach + ". " if approach else ""}'
        f'Use this title\'s {entry.get("palette") or "warm balanced"} palette.'
    )
    return {
        "title": entry["title"], "prompt": prompt,
        "book_type": "educational", "genre": "Self-help",
        "audience": "Adult general", "language": settings.get("language") or "English",
        "primary_marketplace": settings.get("primary_marketplace") or "Amazon.co.uk",
        "tone": "Warm, reflective, practical and encouraging",
        "page_count": 28, "image_engine": "flux_2_klein_4b_local",
        "image_style": settings.get("image_style") or "cinematic_3d",
        "custom_style": settings.get("custom_style") or "",
        "trim": settings.get("trim") or "8x10", "layout": "fixed",
        "print_interior": settings.get("print_interior") or "premium_colour",
        "author": settings.get("author") or "",
        "font_style": settings.get("font_style") or "Friendly storybook",
        "dedication": "", "series_name": series_name,
        "book_number": str(book_number), "personalisation": "",
        "character_bible": [], "series_template": {**entry, "lock_title": True},
        "series_bible": {
            "enabled": True, "series_name": series_name,
            "series_promise": "A consistent 28-day guided workbook series with practical reflection and generous writing space",
            "audience": "Adult general", "genre": "Self-help",
            "tone": "Warm, reflective, practical and encouraging",
            "image_style": settings.get("image_style") or "cinematic_3d",
            "custom_style": settings.get("custom_style") or "",
            "palette": entry.get("palette") or "",
            "typography": settings.get("font_style") or "Friendly storybook",
            "trim": settings.get("trim") or "8x10", "layout": "fixed",
            "print_interior": settings.get("print_interior") or "premium_colour",
            "world_rules": "Maintain the recognisable 28 Days series visual identity and calm workbook layout.",
            "story_rules": "Exactly 28 distinct daily pages with practical exercises, blank writing space, topic-specific wisdom and a four-week progression.",
            "characters": [],
            "locks": {"characters": False, "world": True,
                      "visual_style": True, "palette": True,
                      "audience_tone": True, "typography": True,
                      "trim_print": True},
        },
        "prevent_duplicate": True,
    }


def _configure_series_kdp(project: dict, entry: dict) -> dict:
    project["settings"]["image_engine"] = "flux_2_klein_4b_local"
    project["settings"]["page_count"] = 28
    project["series_template"] = {**entry, "lock_title": True}
    publishing = normalize_publishing(project.get("publishing"), project)
    publishing["route"].update({
        "kdp_ebook": True, "kdp_paperback": True,
        "kdp_expanded": False, "kdp_select": False,
        "ingram_print": False, "d2d_ebook": False,
        "d2d_amazon": False, "blurb_direct": False,
        "blurb_global": False, "lulu_direct": False,
    })
    project["publishing"] = publishing
    return project


def complete_missing_series_books(source_project_id: str = "",
                                  series_name: str = "",
                                  parallel_books: int = 2) -> dict:
    """Create missing titles in parallel with one fair FLUX request per book."""
    if not _series_batch_lock.acquire(blocking=False):
        raise RuntimeError("The 28-day series is already being generated")
    if not _library_maintenance_lock.acquire(blocking=False):
        _series_batch_lock.release()
        raise RuntimeError("Another library maintenance task is already running")
    original_context_project = getattr(_job_context, "project_id", "")
    parent_job_id = str(getattr(_job_context, "job_id", "") or "")
    parent_function = str(getattr(_job_context, "function_name", "") or "")
    try:
        status = series_catalog_status()
        targets = [book for book in status["books"] if book["batch_pending"]]
        if not targets:
            return {"ok": True, "completed": [], "message": "No missing series titles"}
        concurrency = max(2, min(4, int(parallel_books or 2), len(targets)))
        source = None
        if source_project_id:
            try:
                source = load_project(source_project_id)
            except Exception:
                source = None
        if not source:
            source = _series_project_for_entry(series_books_catalog()[0])
        source_settings = dict((source or {}).get("settings") or {})
        selected_series_name = str(series_name or
                                   source_settings.get("series_name") or
                                   SERIES_NAME).strip()[:180]
        total_steps = max(1, sum(max(1, int(book["missing_images"]) + 4)
                                 for book in targets))
        completed = []
        failures = []
        progress_lock = threading.Lock()
        result_lock = threading.Lock()
        stop_event = threading.Event()
        completed_steps = 0
        active_books: dict[str, str] = {}
        image_books: set[str] = set()

        def update_queue_state(title: str = "", stage: str = "") -> None:
            with _jobs_lock:
                job = _jobs.get(parent_job_id) or {}
                if title:
                    active_books[title] = stage
                job["queue_label"] = (
                    f"{len(active_books)} book{'s' if len(active_books) != 1 else ''} active"
                    + (f" · {len(image_books)} FLUX image request"
                       f"{'s' if len(image_books) != 1 else ''} queued/running"
                       if image_books else "")
                    + f" · up to {concurrency} in parallel"
                )
                job["updated_at"] = now()

        def process_book(book: dict) -> dict:
            nonlocal completed_steps
            _job_context.job_id = parent_job_id
            _job_context.function_name = parent_function
            _job_context.progress_monotonic = True
            entry = next(row for row in series_books_catalog()
                         if row["id"] == book["id"])
            title = entry["title"]
            if stop_event.is_set() or current_job_canceled():
                raise ImageJobCanceled("Series generation stopped")
            project = _series_project_for_entry(entry)
            if not project:
                project = new_project(_series_project_payload(
                    entry, source_settings, int(book["book_number"]),
                    selected_series_name))
            pid = project["id"]
            _job_context.project_id = pid
            with _jobs_lock:
                active = _jobs.get(parent_job_id) or {}
                active.setdefault("project_ids", [])
                if pid not in active["project_ids"]:
                    active["project_ids"].append(pid)

            def run_step(description: str, function, *args, image_step=False):
                nonlocal completed_steps
                if stop_event.is_set() or current_job_canceled():
                    raise ImageJobCanceled("Series generation stopped")
                with progress_lock:
                    start = 3 + int(completed_steps / total_steps * 94)
                    finish = 3 + int((completed_steps + 1) / total_steps * 94)
                    active_books[title] = description
                    if image_step:
                        image_books.add(title)
                    update_queue_state()
                _job_context.stage_prefix = (
                    f"{len(active_books)} books parallel · {title} · {description}")
                _job_context.progress_range = (start, max(start + 1, finish))
                try:
                    result = function(*args)
                    if stop_event.is_set() or current_job_canceled():
                        raise ImageJobCanceled("Series generation stopped")
                    with progress_lock:
                        completed_steps += 1
                    return result
                finally:
                    with progress_lock:
                        if image_step:
                            image_books.discard(title)
                        update_queue_state()
                    _job_context.stage_prefix = ""
                    _job_context.progress_range = None

            with progress_lock:
                active_books[title] = "preparing"
                update_queue_state()
            try:
                project = _configure_series_kdp(load_project(pid), entry)
                automation = dict(project.get("series_automation") or {})
                automation.update({
                    "catalog_id": entry["id"], "status": "working",
                    "started_at": automation.get("started_at") or now(),
                    "updated_at": now(), "error": "",
                    "image_engine": "flux_2_klein_4b_local",
                })
                project["series_automation"] = automation
                save_project(project)

                if len(project.get("pages") or []) != 28:
                    project = run_step("writing 28 days", generate_story, pid)
                missing_quotes = [page for page in project.get("pages") or []
                                  if not str(page.get("quote") or "").strip()]
                if missing_quotes:
                    project = run_step("adding page wisdom", populate_quotes, pid, 0)
                project = load_project(pid)
                image_targets = []
                if not image_abs(project, str((project.get("cover") or {}).get("image") or "")):
                    image_targets.append("cover")
                image_targets.extend(str(page.get("number") or index + 1)
                                     for index, page in enumerate(project.get("pages") or [])
                                     if not image_abs(project, str(page.get("image") or "")))
                for image_index, image_target in enumerate(image_targets, 1):
                    run_step(
                        f"local FLUX picture {image_index} of {len(image_targets)}",
                        generate_image, pid, image_target,
                        "flux_2_klein_4b_local", image_step=True)
                project = load_project(pid)
                metadata = project.get("metadata") or {}
                if (not str(metadata.get("description") or "").strip()
                        or len(metadata.get("keywords") or []) != 7
                        or len(metadata.get("categories") or []) != 3):
                    project = run_step("creating KDP listing", generate_metadata, pid)
                result = run_step("building KDP package", export_platform, pid, "kdp")
                project = load_project(pid)
                project.setdefault("series_automation", {}).update({
                    "status": "complete", "completed_at": now(),
                    "updated_at": now(), "error": "",
                    "kdp_bundle": result.get("bundle", ""),
                })
                save_project(project)
                output = {"id": pid, "title": title,
                          "kdp_bundle": result.get("bundle", "")}
                with result_lock:
                    completed.append(output)
                return output
            except Exception as exc:
                stop_event.set()
                try:
                    failed_project = load_project(pid)
                    failed_project.setdefault("series_automation", {}).update({
                        "status": "stopped" if isinstance(exc, ImageJobCanceled) else "error",
                        "updated_at": now(), "error": str(exc)[:1000],
                    })
                    save_project(failed_project)
                except Exception:
                    pass
                if isinstance(exc, ImageJobCanceled):
                    raise
                raise RuntimeError(
                    f'{title} stopped at its current resumable stage: {exc}') from exc
            finally:
                with progress_lock:
                    active_books.pop(title, None)
                    image_books.discard(title)
                    update_queue_state()
                _job_context.stage_prefix = ""
                _job_context.progress_range = None
                _job_context.progress_monotonic = False

        update_queue_state()
        with ThreadPoolExecutor(
                max_workers=concurrency,
                thread_name_prefix="kindle-series") as executor:
            future_map = {executor.submit(process_book, book): book
                          for book in targets}
            for future in as_completed(future_map):
                if future.cancelled():
                    continue
                try:
                    future.result()
                except Exception as exc:
                    failures.append(exc)
                    stop_event.set()
                    for pending in future_map:
                        if not pending.done():
                            pending.cancel()
        if failures:
            canceled = next((exc for exc in failures
                             if isinstance(exc, ImageJobCanceled)), None)
            substantive = next((exc for exc in failures
                                if not isinstance(exc, ImageJobCanceled)), None)
            if substantive:
                raise substantive
            raise canceled or ImageJobCanceled("Series generation stopped")
        studio_event("series.batch_completed", completed=len(completed),
                     parallel_books=concurrency,
                     image_engine="flux_2_klein_4b_local")
        return {"ok": True, "completed": completed,
                "parallel_books": concurrency,
                "image_engine": "flux_2_klein_4b_local"}
    finally:
        _job_context.project_id = original_context_project
        _job_context.stage_prefix = ""
        _job_context.progress_range = None
        _job_context.progress_monotonic = False
        with _jobs_lock:
            job = _jobs.get(parent_job_id) or {}
            job["queue_label"] = "No queued series work"
        _library_maintenance_lock.release()
        _series_batch_lock.release()


def _recent_duplicate_project(data: dict, window_seconds: int = 45) -> dict | None:
    """Return a just-created matching project when a UI action was submitted twice."""
    if not data.get("prevent_duplicate"):
        return None
    title = str(data.get("title") or "Untitled Book").strip().casefold()
    series_name = str(data.get("series_name") or "").strip().casefold()
    book_number = str(data.get("book_number") or "").strip()
    template_id = str((data.get("series_template") or {}).get("id") or "")
    cutoff = time.time() - window_seconds
    for path in LIBRARY_DIR.glob("*/project.json"):
        try:
            if path.stat().st_mtime < cutoff:
                continue
            candidate = json.loads(path.read_text())
            settings = candidate.get("settings") or {}
            candidate_template = candidate.get("series_template") or {}
            if (
                str(candidate.get("title") or "").strip().casefold() == title
                and str(settings.get("series_name") or "").strip().casefold() == series_name
                and str(settings.get("book_number") or "").strip() == book_number
                and str(candidate_template.get("id") or "") == template_id
            ):
                return load_project(candidate["id"])
        except Exception:
            continue
    return None


def new_project(data: dict) -> dict:
    # The lock makes the duplicate check and project creation one atomic action.
    with _project_create_lock:
        duplicate = _recent_duplicate_project(data)
        if duplicate:
            return duplicate
        return _new_project(data)


def _new_project(data: dict) -> dict:
    pid = "book_" + uuid.uuid4().hex[:12]
    settings = {
        "book_type": data.get("book_type", "picture_book"),
        "picture_book_type": data.get("picture_book_type", "classic_story"),
        "popular_blueprint": data.get("popular_blueprint", "funny_interactive"),
        "market_topic": data.get("market_topic", ""),
        "series_hook": data.get("series_hook", ""),
        "genre": data.get("genre", "Adventure"),
        "audience": data.get("audience", "Ages 4–6"),
        "language": data.get("language", "English"),
        "tone": data.get("tone", "Warm, engaging and imaginative"),
        "page_count": int(data.get("page_count", 24)),
        "image_style": data.get("image_style", "cinematic_3d"),
        "image_engine": (
            data.get("image_engine", "grok")
            if data.get("image_engine", "grok") in IMAGE_ENGINE_IDS else "grok"
        ),
        "custom_style": data.get("custom_style", ""),
        "trim": data.get("trim", "8x10"),
        "layout": data.get("layout", "fixed"),
        "author": data.get("author", ""),
        "primary_marketplace": data.get("primary_marketplace", "Amazon.co.uk"),
        "print_interior": data.get("print_interior", "premium_colour"),
        "font_style": data.get("font_style", "Friendly storybook"),
        "dedication": data.get("dedication", ""),
        "series_name": data.get("series_name", ""),
        "book_number": data.get("book_number", ""),
        "personalisation": data.get("personalisation", ""),
    }
    p = {
        "id": pid, "title": data.get("title", "Untitled Book"),
        "subtitle": "", "title_prompt": data.get("title_prompt", ""),
        "prompt": data.get("prompt", ""),
        "stage": "setup", "created_at": now(), "updated_at": now(),
        "asset_epoch": 0,
        "settings": settings, "story_summary": "",
        "series_template": data.get("series_template") or {},
        "series_bible": data.get("series_bible") or {},
        "character_bible": data.get("character_bible") or [],
        "world_bible": "", "cover": {"title": "", "subtitle": "", "image_prompt": "",
                                     "image": "", "approved": False},
        "pages": [], "metadata": {}, "publishing": {}, "video_audio": {},
        "publication": normalize_publication(None), "build": {},
        "review_approval": normalize_review_approval(None),
        "notes": "", "history": [],
    }
    p["publishing"] = normalize_publishing(p.get("publishing"), p)
    p["video_audio"] = normalize_video_audio(p.get("video_audio"), p)
    if p.get("series_bible"):
        apply_series_bible(p)
    saved = save_project(p)
    studio_event(
        "project.created", project_id=pid, title=saved.get("title", ""),
        page_count=settings.get("page_count"), book_type=settings.get("book_type"),
        image_engine=settings.get("image_engine"),
    )
    return saved


def _xai_cost_estimate(url: str, payload: dict) -> tuple[str, float, str]:
    endpoint = urllib.parse.urlparse(url).path
    model = str(payload.get("model") or TEXT_MODEL)
    if "/images/" in endpoint:
        inputs = payload.get("images") or []
        input_count = len(inputs) if isinstance(inputs, list) else int(bool(inputs))
        estimate = 0.02 + 0.002 * input_count
        return "grok_image", estimate, (
            f"{endpoint} · one output image · {input_count} input image(s)")
    payload_chars = len(json.dumps(payload, ensure_ascii=False))
    input_tokens = max(1, math.ceil(payload_chars / 4))
    expected_output_tokens = max(2500, min(18000, int(input_tokens * 1.5)))
    estimate = input_tokens * 2.0 / 1_000_000
    estimate += expected_output_tokens * 6.0 / 1_000_000
    if payload.get("tools"):
        estimate += 0.08
    return "writing", max(0.03, round(estimate, 6)), (
        f"{endpoint} · approximately {input_tokens:,} input and "
        f"{expected_output_tokens:,} output tokens reserved")


def _xai_actual_cost(data: dict) -> float | None:
    usage = data.get("usage") if isinstance(data, dict) else None
    if not isinstance(usage, dict):
        return None
    ticks = usage.get("cost_in_usd_ticks")
    try:
        return max(0.0, float(ticks) / 10_000_000_000)
    except (TypeError, ValueError):
        return None


def xai_json(url: str, payload: dict, timeout=300) -> dict:
    key = api_key()
    if not key:
        raise RuntimeError(f"GROK_API_KEY/XAI_API_KEY is missing from {ENV_FILE}")
    req = urllib.request.Request(
        url, data=json.dumps(payload).encode(), method="POST",
        headers={"Authorization": "Bearer " + key, "Content-Type": "application/json",
                 "User-Agent": "oMLX-Kindle-Studio/1.0"})
    category, estimate, cost_detail = _xai_cost_estimate(url, payload)
    for attempt in range(3):
        reservation = reserve_cloud_cost(
            "xAI", str(payload.get("model") or TEXT_MODEL), category,
            estimate, detail=cost_detail + f" · attempt {attempt + 1}")
        started = time.monotonic()
        studio_event(
            "provider.request", provider="xai", endpoint=urllib.parse.urlparse(url).path,
            model=payload.get("model", ""), attempt=attempt + 1, timeout_seconds=timeout,
        )
        try:
            with urllib.request.urlopen(req, timeout=timeout) as response:
                raw = response.read() or b"{}"
                data = json.loads(raw)
                actual_cost = _xai_actual_cost(data)
                settle_cloud_cost(
                    reservation, actual_cost if actual_cost is not None else estimate,
                    estimated=actual_cost is None)
                studio_event(
                    "provider.response", provider="xai",
                    endpoint=urllib.parse.urlparse(url).path,
                    status=getattr(response, "status", 200), attempt=attempt + 1,
                    duration_ms=round((time.monotonic() - started) * 1000),
                    response_bytes=len(raw),
                )
                return data
        except urllib.error.HTTPError as e:
            release_cloud_cost(reservation, f"HTTP {e.code}")
            detail = e.read().decode("utf-8", "replace")[:1000]
            studio_event(
                "provider.http_error", level="warning", provider="xai",
                endpoint=urllib.parse.urlparse(url).path, status=e.code,
                attempt=attempt + 1,
                duration_ms=round((time.monotonic() - started) * 1000),
                detail=detail[:500],
            )
            if e.code not in (408, 425, 429, 500, 502, 503, 504) or attempt == 2:
                raise RuntimeError(f"xAI API {e.code}: {detail}") from e
            update_current_job(
                f"xAI cloud is busy; waiting to retry {attempt + 2} of 3")
        except (urllib.error.URLError, TimeoutError, ConnectionError, OSError) as e:
            # The request may have reached the provider even when its response
            # was lost. Count the conservative estimate instead of hiding risk.
            settle_cloud_cost(
                reservation, estimate, estimated=True, status="uncertain")
            studio_event(
                "provider.connection_error", level="warning", provider="xai",
                endpoint=urllib.parse.urlparse(url).path, attempt=attempt + 1,
                duration_ms=round((time.monotonic() - started) * 1000),
                error=f"{type(e).__name__}: {e}",
            )
            if attempt == 2:
                raise RuntimeError(
                    "The xAI cloud connection failed after three attempts: " + str(e)) from e
            update_current_job(
                f"xAI connection interrupted; waiting to retry {attempt + 2} of 3")
        except Exception:
            settle_cloud_cost(
                reservation, estimate, estimated=True, status="uncertain")
            raise
        time.sleep(2 ** attempt)
    raise RuntimeError("The xAI request could not be completed")


def xai_upload_file(path: Path) -> str:
    boundary = "----omlx" + uuid.uuid4().hex
    content_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    parts = [
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"purpose\"\r\n\r\nassistants\r\n".encode(),
        (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; "
         f"filename=\"{path.name}\"\r\nContent-Type: {content_type}\r\n\r\n").encode(),
        path.read_bytes(), f"\r\n--{boundary}--\r\n".encode(),
    ]
    req = urllib.request.Request(
        XAI_BASE + "/files", data=b"".join(parts), method="POST",
        headers={"Authorization": "Bearer " + api_key(),
                 "Content-Type": "multipart/form-data; boundary=" + boundary})
    try:
        with urllib.request.urlopen(req, timeout=180) as response:
            result = json.loads(response.read())
        if not result.get("id"):
            raise RuntimeError("xAI did not return a file ID")
        return result["id"]
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"xAI file upload {e.code}: "
                           + e.read().decode("utf-8", "replace")[:1000]) from e


def extract_response_text(data: dict) -> str:
    if not isinstance(data, dict):
        return ""
    candidates = []
    if data.get("output_text"):
        candidates.append(str(data["output_text"]))
    for item in data.get("output") or []:
        if not isinstance(item, dict):
            continue
        if item.get("type") == "message":
            for content in item.get("content") or []:
                if not isinstance(content, dict):
                    continue
                if content.get("type") in ("output_text", "text") and content.get("text"):
                    candidates.append(str(content["text"]))
    choices = data.get("choices") or []
    if choices and isinstance(choices[0], dict):
        message = choices[0].get("message") or {}
        if isinstance(message, dict):
            if message.get("content"):
                candidates.append(str(message["content"]))
    # Tool-using responses may contain an early assistant fragment before the
    # final post-tool answer. Prefer the last candidate that resembles JSON.
    for candidate in reversed(candidates):
        cleaned = candidate.strip()
        if "{" in cleaned and "}" in cleaned:
            return candidate
    return candidates[-1] if candidates else ""


def parse_grok_json_object(raw: str) -> dict:
    cleaned = str(raw or "").strip()
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", cleaned, flags=re.I)
    cleaned = re.sub(r"(?:<\|eos\|>|<\|endoftext\|>)\s*$", "", cleaned).strip()
    candidates = [cleaned]
    start, end = cleaned.find("{"), cleaned.rfind("}")
    if start >= 0 and end > start:
        candidates.append(cleaned[start:end + 1])
    for candidate in candidates:
        try:
            result = json.loads(candidate)
        except (json.JSONDecodeError, TypeError):
            continue
        if isinstance(result, dict):
            return result
    raise ValueError("response did not contain one valid JSON object")


def grok_structured(instruction: str, schema_hint: dict, timeout=600,
                    web_search=False) -> dict:
    base_prompt = (
        instruction + "\n\nReturn ONLY one valid JSON object. Follow this shape exactly:\n"
        + json.dumps(schema_hint, ensure_ascii=False, indent=2)
        + "\nDo not use Markdown fences."
    )
    last_raw = ""
    schema_fields = sorted(schema_hint.keys())
    for attempt in range(3):
        retry_note = ""
        if attempt:
            retry_note = (
                "\n\nCRITICAL RETRY: The preceding response was empty or malformed. "
                "Generate the complete requested JSON object again from the original "
                "instructions. Do not return a boolean, number, control token, "
                "commentary, or partial object."
            )
        payload = {
            "model": TEXT_MODEL, "input": base_prompt + retry_note, "store": False,
            "include": ["no_inline_citations"],
            "text": {"format": {"type": "json_object"}},
        }
        if web_search:
            payload["tools"] = [{"type": "web_search"}]
            payload["tool_choice"] = "auto"
        update_current_job(
            f"Waiting for {TEXT_MODEL} "
            + ("web research" if web_search else "cloud response")
            + (f" · retry {attempt} of 2" if attempt else ""))
        data = xai_json(XAI_BASE + "/responses", payload, timeout)
        update_current_job("Checking Grok's response")
        last_raw = extract_response_text(data).strip()
        try:
            parsed = parse_grok_json_object(last_raw)
            studio_event(
                "grok.structured_parsed", attempt=attempt + 1,
                web_search=bool(web_search), response_chars=len(last_raw),
                schema_fields=schema_fields, returned_fields=sorted(parsed.keys()),
            )
            return parsed
        except ValueError as exc:
            studio_event(
                "grok.malformed_response", level="warning", attempt=attempt + 1,
                web_search=bool(web_search), response_chars=len(last_raw),
                response_sha256=hashlib.sha256(last_raw.encode()).hexdigest()[:16],
                error=str(exc),
            )
            if attempt < 2:
                update_current_job(
                    "Grok response was malformed; retrying automatically")
                time.sleep(1)
    raise RuntimeError(
        "Grok returned an incomplete response after three automatic attempts. "
        "Please try again in a moment."
        + ((" Response began: " + last_raw[:160]) if last_raw else ""))


def style_label(settings: dict) -> str:
    sid = settings.get("image_style", "cinematic_3d")
    label = dict(IMAGE_STYLES).get(sid, sid)
    if sid == "custom" and settings.get("custom_style"):
        label = settings["custom_style"]
    return label


def selected_story_page_count(project: dict) -> int:
    return int((project.get("settings") or {}).get("page_count") or 24)


def expected_book_page_count(project: dict) -> int:
    story_pages = selected_story_page_count(project)
    return story_pages + (
        2 if (project.get("settings") or {}).get("book_type") == "ancient_wisdom"
        else 0
    )


def ancient_context_page_rules(project: dict) -> str:
    if (project.get("settings") or {}).get("book_type") != "ancient_wisdom":
        return ""
    story_pages = selected_story_page_count(project)
    total_pages = story_pages + 2
    return f"""
ANCIENT WISDOM INTRODUCTION AND AFTERWORD:
- The pages array must contain exactly {total_pages} pages: two additional context
  pages surrounding exactly {story_pages} narrative pages selected by the creator.
- Page 1 must have section="introduction". Write a substantial, approximately
  one-page, spoiler-light introduction. Explain the source work, author or ancient
  tradition, historical and cultural setting, why the story matters, and useful
  themes to notice. Do not reveal the climax or ending.
- Pages 2 through {story_pages + 1} must have section="story" and contain the
  faithful narrative itself, in the source's correct sequence.
- Page {total_pages} must have section="afterword". Write a substantial,
  approximately one-page closing summary that may discuss the complete plot and
  ending. Explain the wisdom it imparts, responsible connections to modern life,
  and how the story has influenced literature, theatre, art, language, philosophy,
  psychology, politics or culture where evidence supports that connection. Avoid
  vague claims of influence and do not force the work into one simplistic moral.
- The afterword must name the principal source version followed and briefly note
  any important disputed variant or later tradition that was deliberately excluded.
- Give both context pages a detailed full-page illustration prompt. The
  introduction illustration should establish the era, place and central symbols
  without spoiling the ending. The afterword illustration should be a reflective,
  historically appropriate visual synthesis of the work's consequences, wisdom
  and legacy. Neither is a plain text-only page.
"""


def ancient_wisdom_rules(project: dict) -> str:
    if (project.get("settings") or {}).get("book_type") != "ancient_wisdom":
        return ""
    return """
ANCIENT WISDOM FIDELITY STANDARD - this overrides ordinary commercial-story
formulas wherever they conflict:
- Treat this as a faithful, freshly worded retelling of an ancient source, not an
  original plot merely inspired by it. Identify the principal surviving source,
  author/tradition and version implied by the creator's brief before writing.
- Research the primary ancient text and reputable classical scholarship. Preserve
  the best-attested names, identities, kinship, motives, places, sequence of
  events, divine interventions, cause and consequence, central themes and ending.
- Do not combine incompatible variants or import details from later retellings
  without saying so. When sources genuinely differ, follow the version closest to
  the named work; record the material uncertainty or chosen variant in
  quality_report.notes instead of inventing certainty.
- Never add a happy ending, redemption, romance, villain, moral lesson, rescue,
  modern psychology or modern values that change what the source says. Do not
  remove difficult consequences merely to satisfy a standard story formula.
- Accuracy includes the literal identity and physical state of plot-critical
  remains, sacrificial victims, weapons and ritual objects. A non-graphic visual
  treatment may conceal wounds or minimise blood, but must never replace a human
  head or other essential object with a cup, abstract trophy, symbol or living
  uninjured character. Never resurrect a character after the source has killed them.
- Compression and age-appropriate wording are allowed only to fit the selected
  page count and audience. They must not alter who did what, why it happened, the
  order of decisive events, or the original outcome. Preserve essential causal
  links even when shortening.
- Write all narration and dialogue in fresh language. Do not fabricate quotations
  or copy wording from a modern copyrighted translation. A line may be identified
  as a paraphrase, but must not be presented as a verbatim ancient quotation.
- Keep material culture and illustrations free of anachronisms. Use historically
  and geographically appropriate clothing, armour, buildings, vessels, tools,
  ritual objects, social roles and customs, while distinguishing a mythic setting
  from a later performance or artistic tradition.
- Image prompts must depict the exact event on that page and preserve the correct
  identity, role, status, action and relationship of every ancient character.
- Before returning the book, perform a source-fidelity, chronology, genealogy,
  geography, material-culture and anachronism check. Put the principal source and
  any remaining disputed detail in quality_report.notes.
"""


def story_instruction(p: dict) -> str:
    s = p["settings"]
    kind = dict(BOOK_TYPES).get(s["book_type"], s["book_type"])
    picture_kind = dict(PICTURE_BOOK_TYPES).get(
        s.get("picture_book_type", "classic_story"),
        s.get("picture_book_type", "classic_story"))
    picture_direction = (
        f"\nPicture-book format: {picture_kind}"
        if s.get("book_type") == "picture_book" else "")
    story_count = selected_story_page_count(p)
    count = expected_book_page_count(p)
    audience = str(s.get("audience") or "Ages 4–6")
    if "2–4" in audience or "2-4" in audience:
        words_per_page = "12 to 35 words"
    elif "4–6" in audience or "4-6" in audience:
        words_per_page = "30 to 70 words"
    elif "6–8" in audience or "6-8" in audience:
        words_per_page = "50 to 110 words"
    else:
        words_per_page = "enough complete prose for the selected reading level"
    series = p.get("series_template") or {}
    series_rules = ""
    if series:
        weekly_arc = series.get("weekly_arc") or []
        approach_rule = (f"- Apply this book's distinctive approach: "
                         f"{series.get('approach')}.\n"
                         if series.get("approach") else "")
        series_rules = f"""
SERIES PRODUCTION BIBLE - follow this exactly:
- This is book {s.get('book_number') or '1'} in "{s.get('series_name') or SERIES_NAME}".
- Keep the main title exactly "{p.get('title')}"; do not rename it.
- This book's unique focus is: {series.get('focus')}.
- Its four weekly phases are: {'; '.join(weekly_arc)}.
{approach_rule}- Create exactly 28 daily workbook pages, one complete day per page.
- Give every day 4 to 8 concise, useful exercises combining reflection,
  gratitude, intention, a practical action, journaling and an end-of-day review
  when relevant to the focus.
- In each page's text, place exactly two empty lines between exercises so the
  printed workbook provides intentional writing space. Preserve those blank lines.
- Progress gently across the four phases without repeating prompts, sentences or
  exercises from another day or another title in the series.
- Keep the shared series structure and polished adult tone, while making all
  content, headings, examples, metadata and illustrations original to this topic.
- This is a reflective wellbeing workbook, not diagnosis, treatment or a promise
  of medical results.
- Use this book's identifying visual palette: {series.get('palette')}.
- Give every daily page one concise, topic-specific quote or piece of wisdom for
  the top of the page. Prefer an accurately worded, clearly attributed public-domain
  quotation when confidently known. Otherwise write original wisdom and leave its
  author blank. Never invent or guess an attribution. Keep it under 28 words and
  make it distinct from the quotations on every other day.
"""
    series_bible = normalize_series_bible(p.get("series_bible"), p)
    reusable_bible_rules = ""
    if series_bible.get("enabled"):
        reusable_bible_rules = (
            "\nREUSABLE SERIES BIBLE - honour every enabled lock:\n"
            + json.dumps(series_bible, ensure_ascii=False, indent=2)
            + "\nCreate a genuinely new book while retaining this locked series "
              "identity.\n"
        )
    ancient_rules = ancient_wisdom_rules(p)
    ancient_context_rules = ancient_context_page_rules(p)
    creation_task = (
        f"a faithful, freshly worded {kind} retelling"
        if ancient_rules else f"a complete original {kind}"
    )
    narrative_rules = (
        f"""For this ancient retelling, make the pages one clear causal narrative
while retaining the source's actual structure and outcome. Adapt the language to
the selected audience without simplifying away decisive motives, relationships,
reversals, recognition, suffering or consequences. The original source—not a
generic modern story arc—is authoritative. Page text should normally contain
{words_per_page}, using complete natural prose suited to the selected reading
level."""
        if ancient_rules else
        f"""For a children's story, write one connected narrative—not a list of isolated
moments. Establish the protagonist's desire and problem, make every page follow
causally from the previous page, escalate meaningful obstacles, deliver a real
choice or climax, and resolve both plot and emotion. Give the ending a satisfying
payoff to details planted earlier. Check physical logic, pronouns, names, tense,
motivations and transitions. Do not introduce unexplained people, objects or
solutions. Page text should normally contain {words_per_page}; use complete,
natural, read-aloud sentences and age-appropriate vocabulary rather than thin
        captions."""
    )
    character_prompt_rule = (
        "For FLUX 2 Klein, name each visible recurring character in image_prompt "
        "but do not repeat its full appearance there. Put its current physical "
        "form, wardrobe, action and position once in character_directions; the "
        "numbered master image anchors identity."
        if s.get("image_engine") in KLEIN_IMAGE_ENGINES else
        "Repeat each present character's identifying appearance from the "
        "character bible inside image_prompt."
    )
    return f"""You are an expert commercially published Kindle author, developmental
editor, book designer, and art director. Create {creation_task}.

Book description and story brief: {p.get('prompt')}
Working title: {p.get('title')}
Title prompt/direction: {p.get('title_prompt') or 'Use the working title, or invent the strongest commercially suitable title.'}
If the working title is blank or "Untitled Book", invent a compelling title from
the title prompt and book description. Otherwise retain the creator's working title.{picture_direction}
Genre: {s.get('genre')}
Audience/reading level: {s.get('audience')}
Language: {s.get('language')}
Tone: {s.get('tone')}
Selected narrative pages: {story_count}
Exact total interior-plan pages: {count}
Visual direction: {style_label(s)}
Layout: {s.get('layout')}
Author: {s.get('author')}
Primary Amazon marketplace: {s.get('primary_marketplace', 'Amazon.co.uk')}
Dedication or personal foreword: {s.get('dedication')}
Series: {s.get('series_name')} {s.get('book_number')}
Personalisation and character notes: {s.get('personalisation')}
Characters supplied by the creator: {json.dumps(p.get('character_bible'), ensure_ascii=False)}
{series_rules}
{reusable_bible_rules}
{ancient_rules}
{ancient_context_rules}
{HUMAN_EDITORIAL_STANDARD}
{HUMAN_LISTING_STANDARD}

{narrative_rules}
Write the reader-facing prose in the natural voice of a deeply experienced human
expert in this exact subject. For meditation, contemplative practice or yoga,
write with the quiet authority, practical nuance and restraint of a master who
has practised and taught for decades. For another subject, adopt the equivalent
credible domain expert. Never claim personal qualifications or invent a biography;
let expertise show through precise distinctions, grounded examples, useful
judgement and warm, unforced language. Avoid formulaic summaries, repetitive
transitions, generic motivational filler and conspicuously AI-like phrasing.
Begin every page directly with useful story, fact, instruction or insight. Never
use a recurring guide, teacher, narrator or named character merely as a wrapper
for exposition. Omit presenter stage directions and delivery filler such as a
person continuing in an even voice, gazing kindly, speaking with quiet precision,
smiling, nodding, pausing, settling into position or introducing the next point.
A character's physical action belongs in the prose only when it changes the plot,
demonstrates an instruction the reader must understand, reveals necessary emotion,
or interacts with something important. A visual guide may remain in the image
prompts without being named at the start of every page's reader-facing text.
Reader-facing prose must stay inside the subject itself. Never mention image
prompts, illustration style, claymation, rendering, generated pictures, page
layout, video production, AI, metadata, publishing workflow or how the book was
made unless the creator's requested subject is genuinely one of those things.
For comics, provide concise narration/dialogue and panel-aware image direction. For
non-fiction, build a useful factual progression and flag anything requiring
fact checking. Maintain exact character visual continuity. Image prompts must
be a literal visual contract for the words on that same page. Each one must state
the exact visible action, which named characters are present, their poses and
emotions, all plot-important props, the setting and the consequence of the action.
{character_prompt_rule}
Never use a generic location-only prompt. Never add a child, adult, animal or
crowd that the page text does not contain. Then specify composition, camera,
lighting, palette, and safe negative space for overlaid book copy.
{image_text_policy(
    s.get('image_engine', 'grok'),
    bool(p.get('series_template')) or s.get('book_type') == 'educational')}
{grok_local_image_prompt_rules(s)}

Also create cover direction and KDP metadata. The pages array MUST contain
exactly {count} objects numbered 1 through {count}. Use section="story" on
ordinary narrative/content pages.

Create a complete Amazon KDP listing worksheet. The description must accurately
sell this specific book without reviews, unverifiable claims, keyword stuffing,
URLs, prices, or promotional language, and must not exceed 4,000 characters.
Return exactly seven useful multi-word customer search phrases and exactly three
highly relevant Amazon category-path suggestions. Include the reading age and
responsible recommendations for marketplace, rights, territories, DRM, KDP
Select, royalty and price. Because Grok creates the book text and illustrations,
the AI disclosure must say that both text and images are AI-generated.

Required titling:
- Supply a compelling book subtitle that complements the main title.
- Supply a cover subtitle; it may match the book subtitle when appropriate.
- Every page must have a short, meaningful, unique heading suited to its content.
- Never leave title, subtitle, cover title, cover subtitle, or page heading blank.
- Do not use placeholders such as "Page 1", "Untitled", or "Chapter".

NARRATOR CASTING FOR THE VIDEO/AUDIO EDITION:
Choose whether this specific book is best read in a British or American accent
and by a female or male narrator. Base it on the story, setting, tone, target age,
rhythm and emotional range—not on a generic default. Select exactly one native-
English Qwen3-TTS stable cloned narrator whose accent and gender match:
- Eleanor (warm) or Poppy (lively): British female
- Maya (warm) or Harper (lively): American female
- Arthur (warm) or Oliver (lively): British male
- Noah (warm) or Jack (lively): American male
Set an expressive delivery direction and a natural speed between 0.86 and 1.05;
younger or more emotional read-aloud stories should normally be a little slower.
The voice, accent and gender must agree exactly. Explain the creative choice."""


def generate_market_concept(project_id: str) -> dict:
    p = load_project(project_id)
    s = p.get("settings") or {}
    blueprint_id = s.get("popular_blueprint", "funny_interactive")
    blueprint = next(
        (x for x in POPULAR_KIDS_BLUEPRINTS if x["id"] == blueprint_id),
        POPULAR_KIDS_BLUEPRINTS[0])
    topic = str(s.get("market_topic") or p.get("prompt") or "").strip()
    if not topic:
        raise ValueError("Add a topic, lesson or child-sized problem first")
    update_current_job("Preparing live children's-market research", 10)
    result = grok_structured(
        f"""Act as a commercially experienced children's picture-book market analyst,
concept editor and series architect. Use web search now. Inspect this current
Amazon Children's eBooks bestseller page and useful linked/category context:
{AMAZON_CHILDRENS_BESTSELLERS_URL}

Identify broad, current buyer and reader patterns: age bands, emotional needs,
formats, hooks, read-aloud qualities, visual storytelling, page-turn devices and
series opportunities. The rankings change, access may be personalised, and sales
cannot be predicted, so state those limitations. Then create one highly appealing,
completely original and practical book concept that an independent creator can
develop into a polished illustrated book.

Selected proven structural blueprint: {blueprint['label']}
Blueprint formula: {blueprint['formula']}
Creator's topic, lesson or child-sized problem: {topic}
Creator's recurring character or series hook: {s.get('series_hook') or 'Invent an original one'}
Target readers: {s.get('audience', 'Ages 4–6')}
Genre: {s.get('genre', 'Adventure')}
Picture-book type: {dict(PICTURE_BOOK_TYPES).get(s.get('picture_book_type'), 'Classic illustrated story')}
Requested pages: {s.get('page_count', 24)}

The concept needs an immediately understandable hook for the adult buyer, delight
and participation for the child, strong page turns, visual variety, emotional
payoff and genuine series potential. Recommend the best precise reading age,
picture-book type, genre, tone and page count based on the topic and research.
Keep the scope achievable. Do not promise sales or bestseller status.

Create between one and three recurring characters. Every character must be
visually distinct and fully specified for consistent image generation: species or
human type, apparent age, body proportions, face, hair/fur/skin, eyes, clothing,
exact colours, signature object, personality, mannerisms and never-changing rules.
Write a standalone reference-image prompt for each: one character only, full body,
front three-quarter view, neutral uncluttered background, no lettering.

ORIGINALITY IS MANDATORY: do not reuse, paraphrase, imitate or evoke any existing
book title, character, franchise, trademark, signature refrain, plot, cover or
illustration style. Do not mention a competing book. Invent fresh intellectual
property with a distinctive title and character identity.

{HUMAN_EDITORIAL_STANDARD}""",
        {
            "market_research": {
                "summary": "Current market-pattern summary grounded in web research",
                "patterns": ["Broad, non-copying pattern with buyer/reader relevance"],
                "source_urls": [AMAZON_CHILDRENS_BESTSELLERS_URL],
                "limitations": "Rankings change and no sales outcome can be guaranteed",
            },
            "recommended": {
                "age_range": "One available reading-level label, e.g. Ages 4–6",
                "picture_book_type": "One available picture-book type id, e.g. classic_story",
                "genre": "One suitable genre",
                "tone": "Specific read-aloud tone",
                "page_count": 24,
                "format_reason": "Why this combination suits the concept and audience",
            },
            "narration": {
                "accent": "british or american",
                "gender": "female or male",
                "voice": "One exact stable narrator: Eleanor, Poppy, Maya, Harper, Arthur, Oliver, Noah, or Jack",
                "delivery": "Expressive read-aloud direction",
                "speed": 0.92,
                "reason": "Why the casting suits the story and recommended age",
            },
            "title": "Short, distinctive, memorable original title",
            "title_prompt": "A reusable direction for generating alternative original titles",
            "book_description": "Detailed story brief with beginning, escalation, payoff, child participation and ending",
            "series_name": "Optional original series name",
            "series_hook": "How the hero and world can sustain multiple distinct books",
            "commercial_hook": "One sentence explaining the immediate buyer-facing appeal",
            "child_delight": "One sentence explaining what children will want to experience again",
            "originality_check": "Short confirmation of how the concept avoids existing IP",
            "characters": [{
                "name": "Distinctive original name",
                "role": "Narrative role and species/human type",
                "appearance": "Exact complete visual design",
                "personality": "Traits, desires and mannerisms",
                "continuity_rules": "Never-changing identity, clothing, colours and proportions",
                "reference_image_prompt": "Standalone one-character reference-sheet prompt",
            }],
        }, timeout=900, web_search=True)
    previous_title = str(p.get("title") or "Untitled Book")
    previous_images = int(bool((p.get("cover") or {}).get("image")))
    previous_images += sum(
        bool(page.get("image")) for page in p.get("pages") or [])
    previous_references = len(p.get("reference_images") or [])
    had_previous_book = bool(
        p.get("pages") or previous_images or previous_references
        or p.get("market_concept") or p.get("series_parent_id"))
    if had_previous_book:
        revision_snapshot(p, "before-original-concept-reset")
    detached_images = detach_active_artwork(p)
    if previous_references and not detached_images:
        bump_asset_epoch(p)

    # "Create original concept" starts a genuinely independent creative identity.
    # Never carry an earlier story, series ancestry, reference image or illustration
    # into the new concept merely because it reused the same project workspace.
    p["subtitle"] = ""
    p["story_summary"] = ""
    p["world_bible"] = ""
    p["cover"] = {
        "title": "", "subtitle": "", "image_prompt": "",
        "image": "", "approved": False,
    }
    p["pages"] = []
    p["metadata"] = {}
    p["quality_report"] = {}
    p["reference_images"] = []
    p["series_template"] = {}
    p["series_bible"] = {}
    for key in ("series_parent_id", "series_root_id", "series_topic",
                "series_adventure"):
        p.pop(key, None)
    p["creative_identity_id"] = "concept_" + uuid.uuid4().hex[:12]
    p["stage"] = "setup"

    p["title"] = str(result.get("title") or p.get("title") or "Untitled Book")
    p["title_prompt"] = str(result.get("title_prompt") or "")
    p["prompt"] = str(result.get("book_description") or "")
    s["series_name"] = str(result.get("series_name") or "")
    s["series_hook"] = str(result.get("series_hook") or "")
    s["book_number"] = "1" if s["series_name"] else ""
    s["dedication"] = ""
    s["personalisation"] = ""
    recommended = result.get("recommended") or {}
    if not isinstance(recommended, dict):
        recommended = {}
    available_ages = set(READING_LEVELS)
    age = str(recommended.get("age_range") or "")
    if age in available_ages:
        s["audience"] = age
    available_types = {key for key, _ in PICTURE_BOOK_TYPES}
    picture_type = str(recommended.get("picture_book_type") or "")
    if picture_type in available_types:
        s["picture_book_type"] = picture_type
    if recommended.get("genre"):
        s["genre"] = str(recommended["genre"])
    if recommended.get("tone"):
        s["tone"] = str(recommended["tone"])
    pages = int(recommended.get("page_count") or s.get("page_count") or 24)
    if pages in PAGE_COUNTS:
        s["page_count"] = pages
    concept_characters = []
    raw_characters = result.get("characters") or []
    if not isinstance(raw_characters, list):
        raw_characters = []
    for index, character in enumerate(raw_characters[:3]):
        if not isinstance(character, dict):
            continue
        if not str(character.get("name") or "").strip():
            continue
        clean = {
            "name": str(character.get("name") or ""),
            "role": str(character.get("role") or ""),
            "appearance": str(character.get("appearance") or ""),
            "personality": str(character.get("personality") or ""),
            "continuity_rules": str(character.get("continuity_rules") or ""),
            "reference_image_prompt": str(character.get("reference_image_prompt") or ""),
        }
        clean.setdefault("reference_slot", index)
        concept_characters.append(clean)
    p["character_bible"] = concept_characters
    p["settings"] = s
    if isinstance(result.get("narration"), dict):
        video_audio = dict(p.get("video_audio") or {})
        narrator = dict(result["narration"])
        narrator["selected_by"] = "grok"
        video_audio["narrator"] = narrator
        p["video_audio"] = normalize_video_audio(video_audio, p)
    p["market_concept"] = result
    p.setdefault("history", []).append({
        "at": now(),
        "action": (
            "Created an original market-led kids book concept"
            + (
                f"; detached {detached_images} earlier illustration"
                f"{'s' if detached_images != 1 else ''}, removed "
                f"{previous_references} inherited character reference"
                f"{'s' if previous_references != 1 else ''}, and disconnected "
                f"the earlier {previous_title} series identity"
                if had_previous_book else ""
            )
        ),
    })
    saved = save_project(p)
    studio_event(
        "market_concept.identity_reset", project_id=project_id,
        previous_title=previous_title, new_title=saved.get("title", ""),
        detached_images=detached_images,
        detached_references=previous_references,
        asset_epoch=asset_epoch(saved),
    )
    return saved


def _next_series_number(parent: dict, series_root_id: str) -> int:
    numbers = []
    try:
        numbers.append(int((parent.get("settings") or {}).get("book_number") or 1))
    except (TypeError, ValueError):
        numbers.append(1)
    for path in LIBRARY_DIR.glob("*/project.json"):
        try:
            candidate = json.loads(path.read_text())
        except Exception:
            continue
        candidate_root = str(candidate.get("series_root_id") or "")
        if candidate.get("id") == series_root_id or candidate_root == series_root_id:
            try:
                numbers.append(int(
                    (candidate.get("settings") or {}).get("book_number") or 1))
            except (TypeError, ValueError):
                pass
    return max(numbers or [1]) + 1


def _copy_series_references(parent: dict, child: dict) -> None:
    """Copy private master references into a new independent project folder."""
    source_folder = project_dir(parent["id"])
    target_folder = project_dir(child["id"]) / "references"
    target_folder.mkdir(exist_ok=True)
    copied = []
    characters = child.get("character_bible") or []
    for fallback_slot, reference in enumerate(
            (parent.get("reference_images") or [])[:MAX_CHARACTER_REFERENCES]):
        try:
            slot = int(reference.get("slot", fallback_slot))
        except (TypeError, ValueError):
            slot = fallback_slot
        if slot not in range(MAX_CHARACTER_REFERENCES):
            continue
        source = source_folder / str(reference.get("path") or "")
        if not source.exists() or not source.is_file():
            continue
        filename = (
            f"reference-{slot + 1}-{uuid.uuid4().hex[:10]}-"
            f"{slug(source.stem, 'character')}{source.suffix.lower()}"
        )
        destination = target_folder / filename
        shutil.copy2(source, destination)
        cloned = dict(reference)
        cloned.update(path="references/" + filename, slot=slot)
        copied.append(cloned)
        if slot < len(characters):
            characters[slot]["reference_image"] = cloned["path"]
            characters[slot]["reference_file_id"] = cloned.get("file_id", "")
            characters[slot]["reference_slot"] = slot
    child["reference_images"] = sorted(copied, key=lambda row: row["slot"])
    if copied:
        bump_asset_epoch(child)


def create_series_continuation(parent_id: str, topic: str,
                               adventure: str = "") -> dict:
    """Create a fresh book project while preserving reusable series continuity."""
    parent = load_project(parent_id)
    topic = str(topic or "").strip()
    adventure = str(adventure or "").strip()
    if not topic:
        raise ValueError("Add the new book's topic, lesson or theme")
    settings = dict(parent.get("settings") or {})
    original_concept = parent.get("market_concept") or {}
    bible = normalize_series_bible(parent.get("series_bible"), parent)
    if bible.get("enabled"):
        locked_source = (
            bible.get("characters")
            if (bible.get("locks") or {}).get("characters")
            else parent.get("character_bible"))
    else:
        concept_characters = original_concept.get("characters")
        locked_source = (
            concept_characters
            if isinstance(concept_characters, list) and concept_characters
            else parent.get("character_bible"))
    characters = json.loads(json.dumps(
        locked_source or []))[:MAX_CHARACTER_REFERENCES]
    for character in characters:
        # Local paths are replaced only after the new project exists.
        character.pop("reference_image", None)
        character.pop("reference_file_id", None)
    root_id = str(parent.get("series_root_id") or parent.get("id"))
    next_number = _next_series_number(parent, root_id)
    concept_summary = {
        "title": parent.get("title"),
        "subtitle": parent.get("subtitle"),
        "story_summary": parent.get("story_summary"),
        "series_name": settings.get("series_name") or original_concept.get("series_name"),
        "series_hook": settings.get("series_hook") or original_concept.get("series_hook"),
        "commercial_hook": original_concept.get("commercial_hook"),
        "child_delight": original_concept.get("child_delight"),
        "world_bible": parent.get("world_bible"),
        "series_bible": bible if bible.get("enabled") else {},
        "characters": [{
            key: character.get(key, "") for key in (
                "name", "role", "appearance", "personality",
                "continuity_rules", "reference_image_prompt")
        } for character in characters],
    }
    update_current_job("Grok is designing the next distinct series adventure", 15)
    result = grok_structured(
        f"""Act as a children's picture-book series editor. Create the next original,
standalone book in the same series as the source book below.

SOURCE SERIES DNA:
{json.dumps(concept_summary, ensure_ascii=False, indent=2)}

LOCKED SERIES BIBLE:
{json.dumps(bible if bible.get('enabled') else {}, ensure_ascii=False, indent=2)}

NEW BOOK NUMBER: {next_number}
NEW TOPIC, LESSON OR THEME: {topic}
NEW ADVENTURE REQUEST: {adventure or 'Invent a visually exciting, age-appropriate new adventure.'}
TARGET READERS: {settings.get('audience', 'Ages 4–6')}
TARGET LENGTH: {settings.get('page_count', 24)} pages

Keep the recurring main characters, their exact identity, relationships,
personalities, clothing and permanent accessories unchanged. Keep the recognisable
world and series promise. Create an entirely new central problem, locations,
escalation, visual set-pieces, climax, emotional learning and ending. Do not retell,
rephrase or lightly reskin the source plot. A reader must understand this book
without owning the earlier one. Minor story-only supporting characters are allowed,
but do not replace or redesign the recurring cast.

The title must be fresh, memorable and clearly different from the source title.
The detailed book_description will become the creator brief for generating the
complete page-by-page story next. Do not generate pages yet.

{HUMAN_EDITORIAL_STANDARD}""",
        {
            "title": "Fresh series-consistent title for this book",
            "subtitle": "Concise book subtitle",
            "title_prompt": "Reusable direction for alternative titles",
            "series_name": "The unchanged existing series name, or a suitable name if missing",
            "series_hook": "The unchanged recurring promise of the series",
            "book_description": "Detailed new story brief with beginning, escalating adventure, climax, lesson and ending",
            "commercial_hook": "One sentence buyer-facing hook for this specific book",
            "child_delight": "What makes this adventure fun to read repeatedly",
            "adventure_outline": [
                "Beginning setup", "Escalating attempt or discovery",
                "Major setback", "Climax and choice", "Earned resolution"
            ],
            "world_continuity": "How the familiar world continues while adding fresh locations",
            "cover_direction": "Distinctive cover scene for this adventure",
        }, timeout=900)
    update_current_job("Creating the new independent book project", 75)
    series_name = str(result.get("series_name") or settings.get("series_name")
                      or f"{parent.get('title', 'Book')} Series")
    creation = {
        **settings,
        "title": str(result.get("title") or f"{parent.get('title')} — {topic}"),
        "title_prompt": str(result.get("title_prompt") or ""),
        "prompt": str(result.get("book_description") or ""),
        "market_topic": topic,
        "series_hook": str(result.get("series_hook") or settings.get("series_hook") or ""),
        "series_name": series_name,
        "book_number": str(next_number),
        "character_bible": characters,
        "series_bible": bible if bible.get("enabled") else {},
        "prevent_duplicate": True,
    }
    child = new_project(creation)
    child["subtitle"] = str(result.get("subtitle") or "")
    child["settings"]["market_topic"] = topic
    child["settings"]["series_name"] = series_name
    child["settings"]["book_number"] = str(next_number)
    child["settings"]["series_hook"] = creation["series_hook"]
    child["world_bible"] = str(parent.get("world_bible") or "")
    child["market_concept"] = {
        **result,
        "title": child["title"],
        "book_description": child["prompt"],
        "characters": characters,
        "recommended": {
            "age_range": child["settings"].get("audience"),
            "picture_book_type": child["settings"].get("picture_book_type"),
            "genre": child["settings"].get("genre"),
            "tone": child["settings"].get("tone"),
            "page_count": child["settings"].get("page_count"),
            "format_reason": "Inherited for consistency with the series.",
        },
        "series_source_title": parent.get("title"),
    }
    child["series_parent_id"] = parent["id"]
    child["series_root_id"] = root_id
    child["series_topic"] = topic
    child["series_adventure"] = adventure
    if bible.get("enabled"):
        child["series_bible"] = json.loads(json.dumps(bible))
        child["series_bible"]["series_name"] = series_name
        apply_series_bible(child, child["series_bible"])
    child["cover"].update(
        title=child["title"], subtitle=child["subtitle"],
        image_prompt=str(result.get("cover_direction") or ""))
    _copy_series_references(parent, child)
    child["history"].append({
        "at": now(),
        "action": f"Created as book {next_number} from the series continuity of {parent.get('title')}",
    })
    saved = save_project(child)
    studio_event(
        "series.book_created", project_id=saved["id"], parent_project_id=parent["id"],
        series_root_id=root_id, book_number=next_number, topic=topic,
        reference_count=len(saved.get("reference_images") or []),
    )
    return saved


KDP_METADATA_SHAPE = {
    "description": "Natural human-edited Amazon description, normally 150–300 words and maximum 4,000 characters",
    "keywords": [
        "Exactly 7 distinct multi-word customer search phrases"
    ],
    "categories": [
        "Exactly 3 relevant Amazon category paths for the selected marketplace"
    ],
    "bisac_subjects": [
        "Up to 3 accurate BISAC subject codes and labels for IngramSpark and Draft2Digital"
    ],
    "blurb_tags": [
        "Up to 10 concise accurate discovery tags for the Blurb listing"
    ],
    "age_range": "Reader-facing target age range",
    "reading_age_min": "Minimum age as a number",
    "reading_age_max": "Maximum age as a number",
    "grade_range": "Suggested grade range, or Not applicable",
    "primary_marketplace": "Recommended Amazon marketplace, e.g. Amazon.co.uk",
    "sexually_explicit": "Yes or No",
    "publishing_rights": "I own the copyright and hold the necessary publishing rights",
    "territories": "All territories (worldwide rights), if original content",
    "ai_generated_content": "Yes — AI-generated text and images; no AI translation",
    "drm_recommendation": "Apply DRM or DRM-free, with a short reason",
    "kdp_select_recommendation": "Enroll or do not enroll, with a short reason",
    "royalty_recommendation": "35% or 70%, with a short reason",
    "list_price": "Suggested numeric list price",
    "currency": "GBP, USD, EUR, etc.",
    "price_rationale": "Short editable pricing rationale",
    "contributors": "Other contributors or None",
    "publisher": "Publisher/imprint name or Independently published",
    "edition_number": "Edition number or 1",
    "release_timing": "Publish now or suggested release approach",
    "copyright_text": "Copyright page copy",
    "author_bio": "Editable author biography",
}


STORY_SHAPE = {
    "title": "Compelling book title",
    "subtitle": "Required compelling book subtitle; never blank",
    "story_summary": "Full synopsis",
    "character_bible": [{
        "name": "Name", "role": "Role", "appearance": "Exact reusable appearance",
        "personality": "Traits", "continuity_rules": "Never-changing details",
        "reference_image_prompt": "One-character full-body reference-sheet prompt",
    }],
    "world_bible": "Locations, palette, era, props and continuity",
    "narration": {
        "accent": "british or american",
        "gender": "female or male",
        "voice": "One exact supported Qwen3-TTS CustomVoice speaker",
        "delivery": "Specific expressive read-aloud direction",
        "speed": 0.92,
        "reason": "Why this narrator best suits this book and age group",
    },
    "cover": {
        "title": "Required cover title",
        "subtitle": "Required compelling cover subtitle; never blank",
        "image_prompt": "Precise scene-first cover prompt obeying the selected model budget",
    },
    "pages": [{
        "number": 1,
        "section": "introduction, story, or afterword",
        "heading": "Required short unique heading; never blank",
        "quote": "Concise page-specific famous quotation or original wisdom",
        "quote_author": "Accurate author/source when known, otherwise blank",
        "text": "Final page text",
        "dialogue": [], "image_prompt": "Precise scene-first illustration prompt obeying the selected model budget",
        "character_directions": [{
            "name": "Exact recurring character name",
            "depiction_mode": "living character, remains, portrait, statue, reflection, or vision",
            "identity": "Identity traits that come from this character's master reference",
            "wardrobe": "Exact clothing and accessories worn in this scene",
            "wardrobe_override": "yes only for an explicit disguise or costume change",
            "removed_items": "Default clothing removed for this scene, or blank",
            "action": "This character's exact visible action and interaction",
            "position": "Distinct position in the composition",
        }],
        "supporting_characters": [
            "Every additional visible person not represented by a master reference"
        ],
        "cast_complete": "yes when the directions and supporting list name the full visible cast",
        "continuity_objects": [{
            "id": "Stable identifier reused on every page containing this object",
            "literal_identity": "Exactly what the object physically is; never a euphemism",
            "appearance": "Stable visible construction, material and identifying details",
            "state": "Its exact state on this page",
            "holder_and_position": "Who holds it and where it appears",
        }],
        "negative_prompt": "Unwanted elements", "layout_note": "Text and image placement",
    }],
    "metadata": KDP_METADATA_SHAPE,
    "quality_report": {
        "story_coherent": "yes or no",
        "page_images_match": "yes or no",
        "human_voice": "yes only after removing canned or machine-like prose",
        "notes": ["Any remaining concrete concern; empty when ready"],
    },
}


def story_quality_instruction(project: dict, draft: dict) -> str:
    settings = project.get("settings") or {}
    count = expected_book_page_count(project)
    ancient_rules = ancient_wisdom_rules(project)
    ancient_context_rules = ancient_context_page_rules(project)
    editor_role = (
        "classical-literature source editor, historical adaptation editor and "
        "illustrated-book art director"
        if ancient_rules else
        "children's-book developmental editor and picture-book art director"
    )
    resolution_test = (
        "The climax, reversals, recognition, consequences and ending remain "
        "faithful to the named ancient source, even when the result is tragic or "
        "does not provide modern emotional closure."
        if ancient_rules else
        "The protagonist has a clear desire/problem, escalating attempts and "
        "setbacks, a meaningful climax/choice, and an earned emotionally "
        "satisfying resolution."
    )
    character_test = (
        "Create a precise bible for every story-important recurring character, up "
        f"to the supported maximum of {MAX_CHARACTER_REFERENCES}."
    )
    repair_guard = (
        "without changing or embellishing source facts"
        if ancient_rules else
        "while preserving the book's genuinely good premise and intended ending"
    )
    image_identity_test = (
        "name every visible character and put each character's physical form, "
        "wardrobe, action and position once in character_directions without "
        "repeating long appearance biographies in image_prompt"
        if settings.get("image_engine") in KLEIN_IMAGE_ENGINES else
        "name every visible character, repeat their identifying appearance, and "
        "specify each pose and action"
    )
    return f"""Act as a demanding senior {editor_role}. Rewrite and return this
entire draft as a publication-quality book. Preserve its genuinely good wording
and presentation, but repair every incoherent, abrupt, illogical, repetitive,
vague or underwritten part {repair_guard}.

{ancient_rules}
{ancient_context_rules}
{HUMAN_EDITORIAL_STANDARD}
{HUMAN_LISTING_STANDARD}

NON-NEGOTIABLE STORY TEST:
- Exactly {count} numbered content pages form one connected causal story.
- {resolution_test}
- Every page follows naturally from the previous page. No unexplained people,
  objects, location jumps, actions, solutions, pronouns or changes of species.
- Use natural read-aloud prose suitable for {settings.get('audience')} rather than
  isolated captions. Keep names, tense, facts and physical logic consistent.
- Start each page with meaningful content. Delete presenter stage directions and
  delivery filler that merely says how a guide looks, gazes, speaks, nods, smiles,
  pauses, sits or continues. Keep physical action only when it changes the plot,
  demonstrates an instruction or is otherwise necessary to understand the page.
- {character_test}

NON-NEGOTIABLE PICTURE/TEXT TEST:
- Treat each page's final words as the source of truth for its illustration.
- Each image_prompt must restate the exact visible action from those words,
  {image_identity_test}; specify emotion, interaction, important props, setting
  and consequence; and exclude anything not in the text. A generic location-only
  prompt fails.
- Make consecutive images visually varied while maintaining identity and world.
- The cover must promise this exact story without merely duplicating page 1.
- Follow these illustration-prompt output rules exactly:
{grok_local_image_prompt_rules(settings)}
- Finish quality_report with story_coherent=yes, page_images_match=yes and
  human_voice=yes only after you have actually fixed all failures and removed
  canned or machine-like prose. Do not merely comment on them.

Creator brief: {project.get('prompt')}
Required character/reference information: {json.dumps(project.get('character_bible') or [], ensure_ascii=False)}
Draft to edit:
{json.dumps(draft, ensure_ascii=False)}"""


def polish_story_result(project: dict, draft: dict) -> dict:
    polished = grok_structured(
        story_quality_instruction(project, draft), STORY_SHAPE, timeout=1200,
        web_search=(project.get("settings") or {}).get("book_type") == "ancient_wisdom")
    for key in ("title", "subtitle", "story_summary", "metadata", "narration"):
        if not polished.get(key) and draft.get(key):
            polished[key] = draft[key]
    if not polished.get("character_bible"):
        polished["character_bible"] = (
            draft.get("character_bible") or project.get("character_bible") or [])
    if not polished.get("world_bible"):
        polished["world_bible"] = draft.get("world_bible") or ""
    return polished


def normalize_metadata(metadata: dict, project: dict) -> dict:
    out = dict(metadata or {})
    out["description"] = str(out.get("description") or "")[:4000]
    out["keywords"] = [str(x).strip() for x in out.get("keywords", [])
                       if str(x).strip()][:7]
    out["categories"] = [str(x).strip() for x in out.get("categories", [])
                         if str(x).strip()][:3]
    out["bisac_subjects"] = [
        str(x).strip() for x in out.get("bisac_subjects", [])
        if str(x).strip()
    ][:3]
    out["blurb_tags"] = [
        str(x).strip() for x in out.get("blurb_tags", [])
        if str(x).strip()
    ][:10]
    out["primary_marketplace"] = (
        project.get("settings", {}).get("primary_marketplace") or "Amazon.co.uk")
    out.setdefault("sexually_explicit", "No")
    out.setdefault(
        "publishing_rights",
        "I own the copyright and hold the necessary publishing rights")
    out.setdefault("territories", "All territories (worldwide rights)")
    out.setdefault(
        "ai_generated_content",
        "Yes — AI-generated text and images; no AI translation")
    out.setdefault("publisher", "Independently published")
    out.setdefault("edition_number", "1")
    out.setdefault("contributors", "None")
    out.setdefault("release_timing", "Publish now")
    route = ((project.get("publishing") or {}).get("route") or {})
    if route.get("d2d_ebook", True) and not route.get("kdp_select", False):
        out["kdp_select_recommendation"] = (
            "Do not enrol in KDP Select while Draft2Digital distributes this "
            "ebook outside Amazon.")
    return out


def preserve_human_edited_description(existing: dict, generated: dict) -> dict:
    """Keep an author's edited sales copy when Grok refreshes other metadata."""
    existing = dict(existing or {})
    result = dict(generated or {})
    description = str(existing.get("description") or "").strip()
    if description and existing.get("description_source") == "manual":
        result["description"] = description[:4000]
        result["description_source"] = "manual"
        result["description_edited_at"] = str(
            existing.get("description_edited_at") or now())
        result.pop("description_generated_at", None)
    else:
        result["description_source"] = "grok"
        result["description_generated_at"] = now()
        result.pop("description_edited_at", None)
    return result


def normalize_publishing(value: dict | None, project: dict) -> dict:
    """Normalize per-platform editions without changing the creative book."""
    value = dict(value or {})
    settings = project.get("settings") or {}
    metadata = project.get("metadata") or {}
    route = dict(value.get("route") or {})
    isbn = dict(value.get("isbn") or {})
    editions = dict(value.get("editions") or {})
    pricing = dict(value.get("pricing") or {})
    tracking = dict(value.get("tracking") or {})

    normalized_route = {
        "kdp_ebook": bool(route.get("kdp_ebook", True)),
        "kdp_paperback": bool(route.get("kdp_paperback", True)),
        "kdp_expanded": bool(route.get("kdp_expanded", False)),
        "kdp_select": bool(route.get("kdp_select", False)),
        "ingram_print": bool(route.get("ingram_print", True)),
        "d2d_ebook": bool(route.get("d2d_ebook", True)),
        "d2d_amazon": bool(route.get("d2d_amazon", False)),
        "blurb_direct": bool(route.get("blurb_direct", False)),
        "blurb_global": bool(route.get("blurb_global", False)),
        "lulu_direct": bool(route.get("lulu_direct", False)),
    }
    ownership = str(isbn.get("ownership") or "publisher_owned")
    if ownership not in {"publisher_owned", "kdp_free", "ingram_free", "blurb_free"}:
        ownership = "publisher_owned"
    normalized_isbn = {
        "paperback": str(isbn.get("paperback") or "").strip()[:40],
        "hardcover": str(isbn.get("hardcover") or "").strip()[:40],
        "ebook": str(isbn.get("ebook") or "").strip()[:40],
        "blurb": str(isbn.get("blurb") or "").strip()[:40],
        "ownership": ownership,
        "imprint": str(isbn.get("imprint") or metadata.get("publisher")
                       or "").strip()[:120],
    }

    kdp = dict(editions.get("kdp") or {})
    ingram = dict(editions.get("ingram") or {})
    d2d = dict(editions.get("d2d") or {})
    blurb = dict(editions.get("blurb") or {})
    lulu = dict(editions.get("lulu") or {})
    blurb_trim = str(blurb.get("trim") or "7x7")
    if blurb_trim not in BLURB_EDITIONS:
        blurb_trim = "7x7"
    normalized_editions = {
        "kdp": {
            "binding": "paperback",
            "trim": str(settings.get("trim") or "8.5x8.5"),
            "interior": str(settings.get("print_interior") or "premium_colour"),
            "paper": "white",
            "bleed": True,
        },
        "ingram": {
            "binding": (
                str(ingram.get("binding") or "perfect_bound")
                if str(ingram.get("binding") or "perfect_bound")
                in {"perfect_bound", "case_laminate"} else "perfect_bound"
            ),
            "trim": str(ingram.get("trim") or settings.get("trim") or "8.5x8.5"),
            "interior": str(ingram.get("interior") or "premium_colour"),
            "cover_template_ready": bool(ingram.get("cover_template_ready", False)),
        },
        "d2d": {
            "format": "ebook",
            "exclude_amazon": not normalized_route["d2d_amazon"],
        },
        "blurb": {
            "trim": blurb_trim,
            "cover": (
                str(blurb.get("cover") or "hardcover_imagewrap")
                if str(blurb.get("cover") or "hardcover_imagewrap")
                in {"softcover", "hardcover_imagewrap", "hardcover_dust_jacket"}
                else "hardcover_imagewrap"
            ),
            "specifications_confirmed": bool(
                blurb.get("specifications_confirmed", False)),
        },
        "lulu": {
            "binding": (
                str(lulu.get("binding") or "perfect_bound")
                if str(lulu.get("binding") or "perfect_bound")
                in {"perfect_bound", "casewrap"} else "perfect_bound"
            ),
            "trim": str(settings.get("trim") or "8.5x8.5"),
            "interior": str(settings.get("print_interior") or "premium_colour"),
            "store_platform": (
                str(lulu.get("store_platform") or "shopify")
                if str(lulu.get("store_platform") or "shopify")
                in {"shopify", "woocommerce", "wix", "api", "manual"}
                else "shopify"
            ),
            "template_ready": bool(lulu.get("template_ready", False)),
            "product_connected": bool(lulu.get("product_connected", False)),
        },
    }

    normalized_pricing = {}
    for platform in PUBLISHING_PLATFORMS:
        row = dict(pricing.get(platform) or {})
        normalized_pricing[platform] = {
            "list_price": str(row.get("list_price") or
                              metadata.get("list_price") or "")[:30],
            "print_cost": str(row.get("print_cost") or "")[:30],
            "wholesale_discount": str(row.get("wholesale_discount") or
                                      ("55" if platform in {"ingram", "blurb"}
                                       else ""))[:12],
            "currency": str(row.get("currency") or metadata.get("currency")
                            or "GBP")[:8],
        }

    normalized_tracking = {}
    for platform in PUBLISHING_PLATFORMS:
        row = dict(tracking.get(platform) or {})
        status = str(row.get("status") or "not_started")
        if status not in PUBLISHING_STATUSES:
            status = "not_started"
        normalized_tracking[platform] = {
            "status": status,
            "account_ready": bool(row.get("account_ready", False)),
            "tax_ready": bool(row.get("tax_ready", False)),
            "proof_approved": bool(row.get("proof_approved", False)),
            "store_url": str(row.get("store_url") or "").strip()[:1000],
            "notes": str(row.get("notes") or "")[:4000],
        }
    return {
        "route": normalized_route,
        "isbn": normalized_isbn,
        "editions": normalized_editions,
        "pricing": normalized_pricing,
        "tracking": normalized_tracking,
        "print_preparation": {
            "upscale_1k_art": bool(
                (value.get("print_preparation") or {}).get(
                    "upscale_1k_art", True)),
            "target_ppi": 300,
        },
    }


def normalize_quote_author(value) -> str:
    author = str(value or "").strip().lstrip("-—– ")[:120]
    if author.casefold() in {
        "anonymous", "unknown", "original", "original wisdom",
        "original quote", "n/a", "none",
    }:
        return ""
    return author


def normalize_video_audio(value: dict | None, project: dict) -> dict:
    """Keep a safe, editable narration and book-video plan on every project."""
    value = dict(value or {})
    narrator = dict(value.get("narrator") or {})
    render = dict(value.get("render") or {})
    outputs = dict(value.get("outputs") or {})
    raw_direction = dict(value.get("direction") or {})
    raw_sound = dict(value.get("sound") or {})
    raw_sound_plan = dict(raw_sound.get("plan") or {})
    raw_script = dict(value.get("script") or {})
    marketplace = str((project.get("settings") or {}).get(
        "primary_marketplace") or "Amazon.co.uk").lower()
    default_accent = "american" if ".com" in marketplace else "british"
    accent = str(narrator.get("accent") or default_accent).lower()
    if accent not in {"american", "british"}:
        accent = default_accent
    gender = str(narrator.get("gender") or "female").lower()
    if gender not in {"female", "male"}:
        gender = "female"
    voice = LEGACY_KOKORO_TO_QWEN3.get(
        str(narrator.get("voice") or ""), str(narrator.get("voice") or ""))
    voice_info = QWEN3_NARRATORS.get(voice)
    if (not voice_info or voice_info["gender"] != gender
            or voice_info["accent"] != accent):
        voice = next((key for key, info in QWEN3_NARRATORS.items()
                      if info["gender"] == gender and info["accent"] == accent),
                     "Eleanor")
        voice_info = QWEN3_NARRATORS[voice]
    try:
        speed = max(0.82, min(1.12, float(narrator.get("speed") or 0.92)))
    except (TypeError, ValueError):
        speed = 0.92
    aspect = str(render.get("aspect") or "16:9")
    if aspect not in {"16:9", "9:16", "1:1"}:
        aspect = "16:9"
    resolution = str(render.get("resolution") or "1080p")
    if resolution not in {"1080p", "720p"}:
        resolution = "1080p"
    energy = str(render.get("energy") or "recommended")
    if energy not in {"recommended", "gentle", "balanced", "lively"}:
        energy = "recommended"
    transitions = str(render.get("transitions") or "dynamic")
    if transitions not in {"dynamic", "gentle", "cinematic", "playful"}:
        transitions = "dynamic"
    text_mode = str(render.get("text_mode") or "full")
    if text_mode not in {"full", "heading", "none"}:
        text_mode = "full"
    renderer = str(render.get("renderer") or "blender")
    if renderer not in {"blender", "fast"}:
        renderer = "blender"
    try:
        output_duration = max(0.0, float(outputs.get("duration_seconds") or 0))
    except (TypeError, ValueError):
        output_duration = 0.0
    try:
        output_scenes = max(0, int(outputs.get("scene_count") or 0))
    except (TypeError, ValueError):
        output_scenes = 0
    direction_scenes = [dict(row) for row in (
        raw_direction.get("scenes") or [])
        if isinstance(row, dict) and str(row.get("id") or "").strip()
    ][:140]
    script_scenes = []
    for raw_scene in raw_script.get("scenes") or []:
        if not isinstance(raw_scene, dict):
            continue
        scene_id = str(raw_scene.get("id") or "").strip()[:80]
        if not scene_id:
            continue
        segments = []
        for raw_segment in raw_scene.get("segments") or []:
            if not isinstance(raw_segment, dict):
                continue
            segment_text = re.sub(
                r"[ \t]+", " ", str(raw_segment.get("text") or ""))
            segment_text = re.sub(r"\n{3,}", "\n\n", segment_text).strip()[:1800]
            if not segment_text:
                continue
            try:
                pause_after = max(.12, min(
                    1.6, float(raw_segment.get("pause_after") or .34)))
            except (TypeError, ValueError):
                pause_after = .34
            segments.append({"text": segment_text,
                             "pause_after": round(pause_after, 2)})
        if segments:
            script_scenes.append({"id": scene_id, "segments": segments[:16]})
    sound_intensity = str(raw_sound.get("intensity") or "balanced")
    if sound_intensity not in {"none", "gentle", "balanced", "lively", "cinematic"}:
        sound_intensity = "balanced"
    sound_scenes = []
    for scene_index, raw_scene in enumerate(raw_sound_plan.get("scenes") or []):
        if not isinstance(raw_scene, dict):
            continue
        scene_id = str(raw_scene.get("id") or "").strip()[:80]
        if not scene_id:
            continue
        cues = []
        for cue_index, raw_cue in enumerate(raw_scene.get("cues") or []):
            if not isinstance(raw_cue, dict):
                continue
            cue_id = re.sub(
                r"[^a-zA-Z0-9_-]", "",
                str(raw_cue.get("id") or f"cue-{scene_index + 1}-{cue_index + 1}"),
            )[:80] or f"cue-{scene_index + 1}-{cue_index + 1}"
            kind = str(raw_cue.get("kind") or "foley").lower()
            if kind not in {"ambience", "foley", "impact", "nature", "comedy", "transition"}:
                kind = "foley"
            try:
                at = max(0.0, min(.96, float(raw_cue.get("at") or 0)))
            except (TypeError, ValueError):
                at = 0.0
            try:
                duration = max(.4, min(30.0, float(raw_cue.get("duration") or 2.0)))
            except (TypeError, ValueError):
                duration = 2.0
            try:
                volume = max(.02, min(1.0, float(raw_cue.get("volume") or .24)))
            except (TypeError, ValueError):
                volume = .24
            try:
                pan = max(-1.0, min(1.0, float(raw_cue.get("pan") or 0)))
            except (TypeError, ValueError):
                pan = 0.0
            audio = str(raw_cue.get("audio") or "")[:500]
            if audio and (not audio.startswith("media/") or "/" in audio[6:]):
                audio = ""
            cues.append({
                "id": cue_id,
                "kind": kind,
                "prompt": re.sub(r"\s+", " ", str(raw_cue.get("prompt") or "")).strip()[:450],
                "trigger": re.sub(r"\s+", " ", str(raw_cue.get("trigger") or "")).strip()[:180],
                "at": round(at, 3), "duration": round(duration, 2),
                "volume": round(volume, 3), "pan": round(pan, 2),
                "loop": bool(raw_cue.get("loop")),
                "enabled": raw_cue.get("enabled") is not False,
                "reason": str(raw_cue.get("reason") or "")[:500],
                "audio": audio,
                "generated_at": str(raw_cue.get("generated_at") or "")[:80],
                "seed": int(raw_cue.get("seed") or 0),
                "mastering_version": int(raw_cue.get("mastering_version") or 0),
            })
        try:
            scene_seconds = max(.1, min(180.0, float(
                raw_scene.get("seconds") or 10)))
        except (TypeError, ValueError):
            scene_seconds = 10.0
        sound_scenes.append({
            "id": scene_id,
            "heading": str(raw_scene.get("heading") or "")[:300],
            "seconds": round(scene_seconds, 2),
            "cues": cues[:8],
        })
    try:
        sound_master_volume = max(.05, min(
            1.0, float(raw_sound.get("master_volume") or .85)))
    except (TypeError, ValueError):
        sound_master_volume = .85
    return {
        "narrator": {
            "engine": voice_info.get("engine", NARRATION_TTS_ENGINE),
            "accent": accent, "gender": gender, "voice": voice,
            "delivery": str(narrator.get("delivery") or
                            "Warm, expressive story narration")[:300],
            "speed": speed,
            "reason": str(narrator.get("reason") or
                          "A clear, warm voice suitable for the selected audience.")[:800],
            "selected_by": str(narrator.get("selected_by") or "default")[:40],
        },
        "render": {
            "aspect": aspect, "resolution": resolution, "energy": energy,
            "transitions": transitions, "text_mode": text_mode,
            "renderer": renderer,
        },
        "direction": {
            "signature": str(raw_direction.get("signature") or "")[:80],
            "model": str(raw_direction.get("model") or "")[:300],
            "created_at": str(raw_direction.get("created_at") or "")[:80],
            "scenes": direction_scenes,
        },
        "script": {
            "signature": str(raw_script.get("signature") or "")[:80],
            "model": str(raw_script.get("model") or "")[:300],
            "created_at": str(raw_script.get("created_at") or "")[:80],
            "expert_voice": str(raw_script.get("expert_voice") or "")[:800],
            "approach": str(raw_script.get("approach") or "")[:1200],
            "scenes": script_scenes[:140],
        },
        "sound": {
            "intensity": sound_intensity,
            "ducking": raw_sound.get("ducking") is not False,
            "master_volume": round(sound_master_volume, 2),
            "plan": {
                "signature": str(raw_sound_plan.get("signature") or "")[:80],
                "model": str(raw_sound_plan.get("model") or "")[:300],
                "created_at": str(raw_sound_plan.get("created_at") or "")[:80],
                "scenes": sound_scenes[:140],
            },
        },
        "outputs": {
            "audio": str(outputs.get("audio") or "")[:500],
            "narration": str(outputs.get("narration") or outputs.get("audio") or "")[:500],
            "soundtrack": str(outputs.get("soundtrack") or "")[:500],
            "effects": str(outputs.get("effects") or "")[:500],
            "video": str(outputs.get("video") or "")[:500],
            "preview": str(outputs.get("preview") or "")[:500],
            "preview_created_at": str(
                outputs.get("preview_created_at") or "")[:80],
            "duration_seconds": round(output_duration, 2),
            "scene_count": output_scenes,
            "created_at": str(outputs.get("created_at") or "")[:80],
        },
    }


def _presenter_character_names(project: dict) -> list[str]:
    """Return recurring guides whose job is to present, rather than drive a plot."""
    presenter_role = re.compile(
        r"\b(?:teacher|master|guide|narrator|presenter|expert|mentor|coach|host|"
        r"instructor|practitioner|therapist|doctor|professor|sage|guru)\b", re.I)
    names = []
    for character in project.get("character_bible") or []:
        if not isinstance(character, dict):
            continue
        name = str(character.get("name") or "").strip()
        role = " ".join(str(character.get(key) or "") for key in (
            "role", "personality", "continuity_rules"))
        if name and presenter_role.search(role):
            names.append(name)
    return names


def _strip_empty_presenter_opening(text: str, presenter_names: list[str]) -> str:
    """Drop a short stage-direction paragraph that merely introduces a guide."""
    value = str(text or "").strip()
    if not value or not presenter_names:
        return value
    blocks = [part.strip() for part in re.split(r"\n\s*\n+", value)
              if part.strip()]
    if len(blocks) < 2:
        return value
    first = blocks[0]
    if not _is_empty_presenter_direction(first, presenter_names):
        return value
    return "\n\n".join(blocks[1:]).strip()


def _is_empty_presenter_direction(text: str, presenter_names: list[str]) -> bool:
    """Recognise delivery/posture filler without treating real fiction as filler."""
    first = str(text or "").strip()
    if not first or not presenter_names or len(first.split()) > 42:
        return False
    names = "|".join(re.escape(name) for name in presenter_names)
    if not re.match(rf"^(?:{names})(?:[’']s|\b)", first, re.I):
        return False
    # A genuine quotation or concrete teaching belongs to the content. The
    # disposable openings this targets only describe the guide's delivery,
    # expression, posture or transition into the next explanation.
    if re.search(r"[“\"]", first):
        return False
    framing = re.compile(
        r"\b(?:speaks?|speaking|continues?|begins?|explains?|teaches?|teaching|uses?|gaze|voice|"
        r"tone|presence|looks?|eyes?|nods?|smiles?|pauses?|gestures?|demonstrat(?:es|ing)|"
        r"sits?|stands?|rises?|settles?|rests?|places?|returns?|acknowledges?|"
        r"emphasises?|emphasizes?|brings?|gathers?|lists?|offers?|inviting|"
        r"stillness|breath|hands?|posture)\b", re.I)
    return bool(framing.search(first))


def _strip_leading_presenter_direction(text: str, presenter_names: list[str]) -> str:
    """Remove only a disposable opening sentence, preserving useful words after it."""
    value = str(text or "").strip()
    pieces = [part.strip() for part in re.findall(
        r".+?(?:[.!?…]+(?=\s|$)|$)", re.sub(r"\s+", " ", value))
              if part.strip()]
    if pieces and _is_empty_presenter_direction(pieces[0], presenter_names):
        return " ".join(pieces[1:]).strip()
    return value


def normalize_story(project: dict, data: dict) -> dict:
    existing_pages = {
        int(page.get("number", 0)): page
        for page in project.get("pages", [])
        if int(page.get("number", 0)) > 0
    }
    ancient_wisdom = (
        project.get("settings", {}).get("book_type") == "ancient_wisdom"
    )
    existing_story_pages = [
        page for page in project.get("pages", [])
        if str(page.get("section") or "story").lower() not in {
            "introduction", "afterword"
        }
    ]
    existing_introduction = next((
        page for page in project.get("pages", [])
        if str(page.get("section") or "").lower() == "introduction"
    ), {})
    existing_afterword = next((
        page for page in project.get("pages", [])
        if str(page.get("section") or "").lower() == "afterword"
    ), {})
    if not (project.get("series_template") or {}).get("lock_title"):
        project["title"] = str(data.get("title") or project["title"])
    project["subtitle"] = str(data.get("subtitle") or "")
    project["story_summary"] = str(data.get("story_summary") or "")
    generated_characters = data.get("character_bible") or []
    existing_by_name = {
        str(c.get("name", "")).strip().casefold(): c
        for c in project.get("character_bible", [])
    }
    existing_characters = project.get("character_bible", [])
    for index, character in enumerate(generated_characters):
        existing = existing_by_name.get(
            str(character.get("name", "")).strip().casefold(),
            existing_characters[index] if index < len(existing_characters) else {})
        for key in ("reference_image", "reference_file_id", "reference_slot"):
            if existing.get(key):
                character[key] = existing[key]
    project["character_bible"] = generated_characters
    project["world_bible"] = str(data.get("world_bible") or "")
    if isinstance(data.get("narration"), dict):
        current_video = dict(project.get("video_audio") or {})
        chosen = dict(data["narration"])
        chosen["selected_by"] = "grok"
        current_video["narrator"] = chosen
        project["video_audio"] = normalize_video_audio(current_video, project)
    project["quality_report"] = data.get("quality_report") or {}
    project["metadata"] = preserve_human_edited_description(
        project.get("metadata") or {},
        normalize_metadata(data.get("metadata") or {}, project))
    cover = data.get("cover") or {}
    project["cover"].update({
        "title": cover.get("title") or project["title"],
        "subtitle": cover.get("subtitle") or project["subtitle"],
        "image_prompt": cover.get("image_prompt") or "",
    })
    pages = []
    expected = expected_book_page_count(project)
    source = data.get("pages") or []
    if ancient_wisdom and source:
        introduction = next((
            page for page in source
            if str(page.get("section") or "").strip().lower() == "introduction"
        ), None)
        afterword = next((
            page for page in source
            if str(page.get("section") or "").strip().lower() == "afterword"
        ), None)
        if introduction and afterword:
            narrative = [
                page for page in source
                if page is not introduction and page is not afterword
            ]
            source = [introduction] + narrative[:expected - 2] + [afterword]
    local_art = project["settings"].get("image_engine") in ("hidream", "flux")
    default_negative = (
        "text, letters, signage, captions, labels, logo, watermark, distorted anatomy"
        if local_art else "unrequested writing, logo, watermark, distorted anatomy"
    )
    story_index = 0
    for i in range(expected):
        raw = source[i] if i < len(source) else {}
        raw_section = str(raw.get("section") or "").strip().lower()
        if ancient_wisdom:
            section = (
                "introduction" if i == 0
                else "afterword" if i == expected - 1
                else "story"
            )
            if section == "introduction":
                existing = existing_introduction
            elif section == "afterword":
                existing = existing_afterword
            else:
                existing = (
                    existing_story_pages[story_index]
                    if story_index < len(existing_story_pages) else {}
                )
                story_index += 1
        else:
            section = raw_section if raw_section in {
                "introduction", "story", "afterword"
            } else "story"
            existing = existing_pages.get(i + 1, {})
        existing_image = str(existing.get("image") or "")
        page_text = str(raw.get("text") or "")
        page_text = _strip_empty_presenter_opening(
            page_text, _presenter_character_names(project))
        pages.append({
            "number": i + 1, "section": section,
            "heading": str(raw.get("heading") or ""),
            "quote": str(raw.get("quote") or existing.get("quote") or ""),
            "quote_author": normalize_quote_author(
                raw.get("quote_author") or existing.get("quote_author") or ""),
            "text": page_text,
            "dialogue": raw.get("dialogue") or [],
            "image_prompt": str(raw.get("image_prompt") or ""),
            "character_directions": (
                raw.get("character_directions")
                if isinstance(raw.get("character_directions"), list)
                else existing.get("character_directions") or []
            ),
            "supporting_characters": (
                raw.get("supporting_characters")
                if isinstance(raw.get("supporting_characters"), list)
                else existing.get("supporting_characters") or []
            ),
            "cast_complete": str(
                raw.get("cast_complete") or existing.get("cast_complete") or ""
            ),
            "continuity_objects": (
                raw.get("continuity_objects")
                if isinstance(raw.get("continuity_objects"), list)
                else existing.get("continuity_objects") or []
            ),
            "negative_prompt": str(raw.get("negative_prompt") or
                                   default_negative),
            "layout_note": str(raw.get("layout_note") or ""),
            "image": existing_image,
            "approved": bool(existing.get("approved")) if existing_image else False,
            "text_approved": False,
        })
    project["pages"] = pages
    project["stage"] = "story"
    project["history"].append({"at": now(), "action": "Generated full book with Grok"})
    return project


def ensure_grok_titles(data: dict, expected_pages: int) -> dict:
    pages = data.get("pages") or []
    cover = data.get("cover") or {}
    missing = (
        not str(data.get("subtitle") or "").strip()
        or not str(cover.get("title") or "").strip()
        or not str(cover.get("subtitle") or "").strip()
        or any(not str((pages[i] if i < len(pages) else {}).get("heading") or "").strip()
               for i in range(expected_pages))
    )
    if not missing:
        return data
    context = {
        "title": data.get("title"), "subtitle": data.get("subtitle"),
        "summary": data.get("story_summary"),
        "cover": cover,
        "pages": [{"number": p.get("number"), "text": p.get("text"),
                   "heading": p.get("heading")} for p in pages],
    }
    completed = grok_structured(
        f"""Act as a senior publishing copywriter. Complete all required titling
for this book. Write a compelling book subtitle, a cover title and subtitle,
and one short, meaningful, unique heading for every page. No field may be
blank. Do not use generic placeholders such as Page 1, Untitled or Chapter.
{HUMAN_EDITORIAL_STANDARD}
Book context:\n""" + json.dumps(context, ensure_ascii=False),
        {
            "subtitle": "Required book subtitle",
            "cover_title": "Required cover title",
            "cover_subtitle": "Required cover subtitle",
            "page_headings": [
                {"number": i, "heading": "Required unique heading"}
                for i in range(1, expected_pages + 1)
            ],
        })
    data["subtitle"] = completed.get("subtitle") or data.get("subtitle")
    cover["title"] = completed.get("cover_title") or cover.get("title") or data.get("title")
    cover["subtitle"] = completed.get("cover_subtitle") or cover.get("subtitle")
    data["cover"] = cover
    headings = completed.get("page_headings") or []
    by_number = {int(x.get("number", 0)): x.get("heading") for x in headings}
    for i, page in enumerate(pages, 1):
        page["heading"] = page.get("heading") or by_number.get(i) or f"Part {i}"
    return data


def generate_story(project_id: str) -> dict:
    p = load_project(project_id)
    revision_snapshot(p, "before-full-generation")
    update_current_job("Grok is writing the complete first draft", 12)
    ancient_wisdom = (
        p.get("settings", {}).get("book_type") == "ancient_wisdom"
    )
    result = grok_structured(
        story_instruction(p), STORY_SHAPE, timeout=900,
        web_search=ancient_wisdom)
    if p.get("settings", {}).get("book_type") in {
            "picture_book", "early_reader", "chapter_book", "ancient_wisdom", "comic"}:
        update_current_job(
            "Grok is checking source fidelity and picture matching"
            if ancient_wisdom else
            "Grok is checking story logic and picture matching", 70)
        result = polish_story_result(p, result)
    result = ensure_grok_titles(result, expected_book_page_count(p))
    return save_project(normalize_story(p, result))


def quality_review_story(project_id: str) -> dict:
    p = load_project(project_id)
    if not p.get("pages"):
        raise ValueError("Generate or load a story first")
    update_current_job("Preparing the complete story for editorial review", 10)
    revision_snapshot(p, "before-quality-review")
    draft = {
        "title": p.get("title"), "subtitle": p.get("subtitle"),
        "story_summary": p.get("story_summary"),
        "character_bible": p.get("character_bible") or [],
        "world_bible": p.get("world_bible"), "cover": p.get("cover"),
        "narration": (p.get("video_audio") or {}).get("narrator") or {},
        "pages": p.get("pages"), "metadata": p.get("metadata") or {},
    }
    result = polish_story_result(p, draft)
    result = ensure_grok_titles(result, expected_book_page_count(p))
    reviewed = normalize_story(p, result)
    reviewed["history"].append({
        "at": now(),
        "action": "Quality-reviewed story logic and page-to-picture alignment",
    })
    return save_project(reviewed)


def populate_titles(project_id: str) -> dict:
    p = load_project(project_id)
    revision_snapshot(p, "before-title-population")
    ensure_grok_titles(p, len(p.get("pages") or []))
    p["history"].append({"at": now(), "action": "Populated missing titles with Grok"})
    return save_project(p)


def populate_quotes(project_id: str, page_number: int = 0) -> dict:
    p = load_project(project_id)
    pages = p.get("pages") or []
    if page_number:
        if page_number < 1 or page_number > len(pages):
            raise ValueError("page out of range")
        targets = [pages[page_number - 1]]
    else:
        targets = [page for page in pages if not str(page.get("quote") or "").strip()]
    if not targets:
        return p
    revision_snapshot(
        p, f"before-page-{page_number}-quote" if page_number else "before-missing-quotes")
    context = [{
        "number": page.get("number"),
        "heading": page.get("heading"),
        "page_subject": str(page.get("text") or "")[:1800],
        "existing_quote": page.get("quote") or "",
    } for page in targets]
    instruction = f"""Create one short quotation or piece of wisdom for the top of
each requested workbook page. It must be specific to that page's subject and the
book topic, not generic, and no two quotations may repeat. Keep each under 28 words.
Use an accurately worded and attributed public-domain famous quotation only when
you are confident both wording and author are correct. Otherwise write concise
original wisdom and leave quote_author blank. Never fabricate or guess an author,
source, or quotation. Do not include quotation marks or a leading dash in the JSON.
For original wisdom, avoid slogan-like symmetry, generic reassurance and polished
social-media aphorisms. Write one precise observation a thoughtful human editor
would keep.

{HUMAN_EDITORIAL_STANDARD}

Book: {p.get('title')}
Series focus: {(p.get('series_template') or {}).get('focus', p.get('story_summary', ''))}
Pages: {json.dumps(context, ensure_ascii=False)}"""
    shape = {"quotes": [{
        "number": int(page.get("number") or 0),
        "quote": "Concise quotation or original wisdom",
        "quote_author": "Accurate author/source when known, otherwise blank",
    } for page in targets]}
    result = grok_structured(instruction, shape, timeout=600)
    by_number = {
        int(item.get("number", 0)): item
        for item in (result.get("quotes") or [])
        if int(item.get("number", 0)) > 0
    }
    for page in targets:
        item = by_number.get(int(page.get("number", 0)), {})
        quote = str(item.get("quote") or "").strip().strip('“”\"')
        if quote:
            page["quote"] = quote[:350]
            page["quote_author"] = normalize_quote_author(item.get("quote_author"))
    p["history"].append({
        "at": now(),
        "action": (
            f"Regenerated quote for page {page_number} with Grok"
            if page_number else f"Added {len(targets)} missing page quotes with Grok"
        ),
    })
    return save_project(p)


def generate_metadata(project_id: str) -> dict:
    p = load_project(project_id)
    revision_snapshot(p, "before-kdp-listing")
    context = {
        "title": p.get("title"), "subtitle": p.get("subtitle"),
        "author": p.get("settings", {}).get("author"),
        "language": p.get("settings", {}).get("language"),
        "book_type": p.get("settings", {}).get("book_type"),
        "genre": p.get("settings", {}).get("genre"),
        "audience": p.get("settings", {}).get("audience"),
        "series_name": p.get("settings", {}).get("series_name"),
        "book_number": p.get("settings", {}).get("book_number"),
        "primary_marketplace": p.get("settings", {}).get(
            "primary_marketplace", "Amazon.co.uk"),
        "summary": p.get("story_summary"),
        "characters": p.get("character_bible"),
        "page_text": [page.get("text") for page in p.get("pages", [])],
        "publishing_route": normalize_publishing(
            p.get("publishing"), p).get("route"),
    }
    instruction = f"""Act as an ethical Amazon KDP metadata specialist and veteran
book-jacket copy editor. Create a
complete, accurate, conversion-focused listing worksheet for this exact book.
The description must be plain text, appealing to the intended buyer, contain
no reviews, unverifiable claims, prices, URLs, keyword stuffing or misleading
language, and be no more than 4,000 characters. Supply exactly seven distinct
multi-word search phrases and exactly three highly relevant Amazon category
path suggestions. Also supply up to three accurate BISAC subject codes with
labels for IngramSpark and Draft2Digital, plus up to ten concise Blurb discovery
tags. Do not invent content that is not in the book. Recommend the
remaining KDP choices responsibly. When Draft2Digital wide ebook distribution
is enabled, explicitly recommend against KDP Select because Select requires
ebook exclusivity. Mark both text and images as AI-generated,
with no AI translation.

{HUMAN_EDITORIAL_STANDARD}
{HUMAN_LISTING_STANDARD}

Book context:\n""" + json.dumps(context, ensure_ascii=False)
    generated_metadata = normalize_metadata(
        grok_structured(instruction, KDP_METADATA_SHAPE), p)
    p["metadata"] = preserve_human_edited_description(
        p.get("metadata") or {}, generated_metadata)
    p["history"].append({
        "at": now(), "action": "Generated multi-platform publishing metadata with Grok"})
    return save_project(p)


def regenerate_page(project_id: str, page_number: int, request: str) -> dict:
    p = load_project(project_id)
    idx = page_number - 1
    if idx < 0 or idx >= len(p["pages"]):
        raise ValueError("page out of range")
    revision_snapshot(p, f"before-page-{page_number}-revision")
    page = p["pages"][idx]
    context = {
        "title": p["title"], "summary": p["story_summary"],
        "characters": p["character_bible"], "world": p["world_bible"],
        "previous_page": p["pages"][idx - 1] if idx else None,
        "current_page": page,
        "next_page": p["pages"][idx + 1] if idx + 1 < len(p["pages"]) else None,
    }
    fallback_request = ("Improve clarity, pacing, age fit, continuity and "
                        "illustration specificity.")
    instruction = f"""Act as a meticulous book editor and art director. Revise page
{page_number} according to this request: {request or fallback_request}
Preserve continuity and do not rewrite adjacent pages. If revising the quote,
use an accurately attributed public-domain quotation only when confident;
otherwise write original wisdom and leave quote_author blank. Never invent an
attribution.
Keep this page's section role unchanged: {page.get('section') or 'story'}.
{HUMAN_EDITORIAL_STANDARD}
{ancient_wisdom_rules(p)}
{ancient_context_page_rules(p)}
Illustration-prompt output rules:
{grok_local_image_prompt_rules(p.get('settings') or {})}
Book context:
{json.dumps(context, ensure_ascii=False)}"""
    shape = {
        "heading": "Required short meaningful page heading; never blank",
        "quote": "Concise page-specific quotation or original wisdom",
        "quote_author": "Accurate author/source when known, otherwise blank",
        "text": "Revised final page text", "dialogue": [],
        "image_prompt": "Revised precise scene-first illustration prompt obeying the selected model budget",
        "character_directions": [{
            "name": "Exact recurring character name",
            "depiction_mode": "living character, remains, portrait, statue, reflection, or vision",
            "identity": "Master-reference identity traits",
            "wardrobe": "Exact current wardrobe",
            "wardrobe_override": "yes or no",
            "removed_items": "Removed default wardrobe, or blank",
            "action": "Exact visible action",
            "position": "Distinct composition position",
        }],
        "supporting_characters": ["Every other visible person, or an empty list"],
        "cast_complete": "yes",
        "continuity_objects": [{
            "id": "Stable object identifier",
            "literal_identity": "Exact physical identity, never a euphemism",
            "appearance": "Stable visible details",
            "state": "Exact current state",
            "holder_and_position": "Holder and composition position",
        }],
        "negative_prompt": "Unwanted elements", "layout_note": "Placement",
    }
    result = grok_structured(
        instruction, shape,
        web_search=p.get("settings", {}).get("book_type") == "ancient_wisdom")
    for key in (
            "heading", "quote", "quote_author", "text", "dialogue", "image_prompt",
            "character_directions", "supporting_characters", "cast_complete",
            "continuity_objects",
            "negative_prompt", "layout_note"):
        if key in result:
            page[key] = result[key]
    page["approved"] = False
    page["text_approved"] = False
    had_image = bool(page.get("image"))
    page["image"] = ""
    if had_image:
        bump_asset_epoch(p)
    p["history"].append({"at": now(), "action": f"Regenerated page {page_number}"})
    return save_project(p)


def regenerate_cover(project_id: str, request: str) -> dict:
    p = load_project(project_id)
    instruction = f"""Create improved Kindle cover copy and a precise, scene-first
image-generation prompt for this book, within the selected model's prompt budget.
User request: {request or 'Make it commercially compelling.'}
Title: {p['title']}; subtitle: {p.get('subtitle')}; summary: {p.get('story_summary')}
Characters: {json.dumps(p.get('character_bible'), ensure_ascii=False)}
Visual style: {style_label(p['settings'])}.
{image_text_policy(
    p['settings'].get('image_engine', 'grok'),
    bool(p.get('series_template'))
    or p['settings'].get('book_type') == 'educational')}
{grok_local_image_prompt_rules(p.get('settings') or {})}"""
    result = grok_structured(instruction + """
The title and subtitle are both required and must not be blank. The subtitle
should be concise, commercially appealing, and complement rather than repeat
the title.""" + HUMAN_EDITORIAL_STANDARD, {
        "title": p["title"], "subtitle": "Required compelling cover subtitle",
        "image_prompt": "Precise scene-first portrait cover prompt obeying the selected model budget",
    })
    had_image = bool(p["cover"].get("image"))
    p["cover"].update(result)
    p["cover"]["image"] = ""
    p["cover"]["approved"] = False
    if had_image:
        bump_asset_epoch(p)
    return save_project(p)


def download_image(item: dict) -> bytes:
    if item.get("b64_json"):
        return base64.b64decode(item["b64_json"])
    url = item.get("url")
    if not url:
        raise RuntimeError("xAI image response contained no URL or image data")
    req = urllib.request.Request(url, headers={"User-Agent": "oMLX-Kindle-Studio/1.0"})
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req, timeout=180) as response:
                return response.read()
        except urllib.error.HTTPError as e:
            if e.code not in (408, 425, 429, 500, 502, 503, 504) or attempt == 2:
                raise RuntimeError(f"Image download failed with HTTP {e.code}") from e
        except (urllib.error.URLError, TimeoutError, ConnectionError, OSError) as e:
            if attempt == 2:
                raise RuntimeError(
                    "The generated image download failed after three attempts: " + str(e)) from e
        time.sleep(2 ** attempt)
    raise RuntimeError("The generated image could not be downloaded")


def local_image_json(path: str, payload: dict | None = None, timeout=90) -> dict:
    url = LOCAL_IMAGE_BASE + path
    data = json.dumps(payload).encode() if payload is not None else None
    req = urllib.request.Request(
        url, data=data, method="POST" if payload is not None else "GET",
        headers={"Content-Type": "application/json",
                 "User-Agent": "oMLX-Kindle-Studio/1.0"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as response:
            return json.loads(response.read() or b"{}")
    except urllib.error.HTTPError as e:
        detail = e.read().decode("utf-8", "replace")[:1000]
        raise RuntimeError(f"Local image service {e.code}: {detail}") from e
    except (urllib.error.URLError, TimeoutError, ConnectionError, OSError) as e:
        raise RuntimeError(
            "The local image service is not available. Open oMLX and try again. " + str(e)) from e


def image_data_url(path: Path) -> str:
    mime = mimetypes.guess_type(path.name)[0] or "image/png"
    return f"data:{mime};base64," + base64.b64encode(path.read_bytes()).decode()


def possible_generated_text(path: Path) -> list[dict]:
    """Return likely embedded lettering without rejecting or regenerating art."""
    if not IMAGE_TEXT_DETECTOR.is_file():
        return []
    try:
        result = subprocess.run(
            [str(IMAGE_TEXT_DETECTOR), str(path)], capture_output=True,
            text=True, timeout=30, check=False,
        )
        if result.returncode != 0:
            studio_event(
                "illustration.text_check_failed", level="warning",
                detector_error=result.stderr.strip()[:500],
            )
            return []
        rows = json.loads(result.stdout or "[]")
    except Exception as error:
        studio_event(
            "illustration.text_check_failed", level="warning",
            detector_error=f"{type(error).__name__}: {error}",
        )
        return []
    detected = []
    for row in rows if isinstance(rows, list) else []:
        text = str(row.get("text") or "").strip()
        alphanumeric = re.sub(r"[^A-Za-z0-9]", "", text)
        confidence = float(row.get("confidence") or 0)
        if len(alphanumeric) < 2 or confidence < 0.30:
            continue
        detected.append({
            "text": text[:80], "confidence": round(confidence, 3),
            "x": round(float(row.get("x") or 0), 4),
            "y": round(float(row.get("y") or 0), 4),
            "width": round(float(row.get("width") or 0), 4),
            "height": round(float(row.get("height") or 0), 4),
        })
    return detected[:8]


def _art_record(project: dict, target: str) -> dict:
    if str(target) == "cover":
        return project.setdefault("cover", {})
    number = int(target)
    pages = project.get("pages") or []
    if number < 1 or number > len(pages):
        raise ValueError("That book page does not exist")
    return pages[number - 1]


def archive_image_version(record: dict, reason: str) -> None:
    image = str(record.get("image") or "")
    if not image:
        return
    versions = record.setdefault("image_versions", [])
    if versions and versions[-1].get("image") == image:
        return
    versions.append({"image": image, "at": now(), "reason": reason})
    if len(versions) > 30:
        del versions[:-30]


def repair_image_text(project_id: str, target: str) -> dict:
    """Repair only OCR-detected lettering locally; never call an image provider."""
    project = load_project(project_id)
    record = _art_record(project, target)
    source_rel = str(record.get("image") or "")
    source = image_abs(project, source_rel)
    if not source:
        raise ValueError("Generate or attach this illustration first")
    detections = possible_generated_text(source)
    if not detections:
        raise ValueError(
            "No generated text was detected, so the image was left unchanged")
    if not IMAGE_TEXT_REPAIR_PYTHON.is_file() or not IMAGE_TEXT_REPAIR_SCRIPT.is_file():
        raise RuntimeError("The local text-repair tool is not installed")
    revision_snapshot(project, f"before-{target}-local-text-repair")
    destination_name = (
        source.stem + "-text-repaired-" + uuid.uuid4().hex[:8] + ".png")
    destination = project_dir(project_id) / "images" / destination_name
    update_current_job("Repairing detected lettering on this Mac", 35)
    result = subprocess.run(
        [str(IMAGE_TEXT_REPAIR_PYTHON), str(IMAGE_TEXT_REPAIR_SCRIPT),
         str(source), str(destination), str(IMAGE_TEXT_DETECTOR)],
        capture_output=True, text=True, timeout=180, check=False,
    )
    try:
        report = json.loads(result.stdout or "{}")
    except ValueError:
        report = {}
    if result.returncode or not destination.is_file():
        destination.unlink(missing_ok=True)
        raise RuntimeError(
            str(report.get("error") or result.stderr.strip()
                or "The local text repair did not produce an image"))
    update_current_job("Checking the repaired image", 82)
    remaining = possible_generated_text(destination)
    archive_image_version(record, "Before local generated-text repair")
    record["image"] = "images/" + destination_name
    record["approved"] = False
    record["image_repair"] = {
        "at": now(), "method": "local OCR-guided inpainting",
        "source": source_rel, "detections_repaired": len(detections),
        "masked_percent": report.get("masked_percent", 0),
        "remaining_detections": len(remaining),
        "broad_repair": float(report.get("masked_percent") or 0) > 3.0,
    }
    if remaining:
        record["image_warning"] = {
            "message": (
                "Possible generated text remains after local repair. The repaired "
                "image was kept and no cloud generation was started."),
            "examples": ", ".join(repr(row["text"]) for row in remaining[:3]),
            "detections": remaining, "checked_at": now(),
        }
    elif record["image_repair"]["broad_repair"]:
        record["image_warning"] = {
            "message": (
                "The lettering was removed locally, but it covered a broad area. "
                "Inspect the repaired background closely or restore the previous "
                "image; no cloud generation was started."),
            "examples": (
                f"Local repair affected {record['image_repair']['masked_percent']}% "
                "of the illustration"),
            "detections": [], "checked_at": now(),
        }
    else:
        record.pop("image_warning", None)
    bump_asset_epoch(project)
    project.setdefault("history", []).append({
        "at": now(),
        "action": (
            f"Repaired detected generated text locally on {'cover' if target == 'cover' else f'page {target}'}; "
            "no cloud image call was made"),
    })
    save_project(project)
    record_local_image_use(
        "OpenCV local text repair", detail=f"repaired {len(detections)} OCR region(s)")
    studio_event(
        "illustration.text_repaired", project_id=project_id, target=str(target),
        source=source_rel, output=record["image"],
        repaired_detections=len(detections), remaining_detections=len(remaining),
        masked_percent=report.get("masked_percent", 0), cloud_cost_usd=0,
    )
    return {"project": project, "repair": record["image_repair"],
            "text_warning": record.get("image_warning")}


def restore_previous_image(project_id: str, target: str) -> dict:
    project = load_project(project_id)
    record = _art_record(project, target)
    versions = record.get("image_versions") or []
    while versions:
        previous = versions.pop()
        rel = str(previous.get("image") or "")
        if image_abs(project, rel):
            current = str(record.get("image") or "")
            if current and current != rel:
                versions.append({
                    "image": current, "at": now(),
                    "reason": "Before restoring an earlier image",
                })
            record["image"] = rel
            record["approved"] = False
            detected = possible_generated_text(image_abs(project, rel))
            if detected:
                record["image_warning"] = {
                    "message": (
                        "Possible generated text or signage detected in this restored image."),
                    "examples": ", ".join(
                        repr(row["text"]) for row in detected[:3]),
                    "detections": detected, "checked_at": now(),
                }
            else:
                record.pop("image_warning", None)
            bump_asset_epoch(project)
            save_project(project)
            return project
    raise ValueError("There is no earlier image version to restore")


def local_reference_entries(project: dict, target: str, engine: str) -> list[tuple[Path, str]]:
    """Choose ordered references for characters actually present in this scene."""
    folder = project_dir(project["id"])
    attached = []
    max_references = (
        MAX_CHARACTER_REFERENCES
        if engine in KLEIN_IMAGE_ENGINES or engine == "nano_banana_2"
        else MAX_GROK_REFERENCES
    )
    indexed_refs = list(enumerate(project.get("reference_images") or []))

    def reference_slot(item):
        index, reference = item
        try:
            return int(reference.get("slot", index)), index
        except (TypeError, ValueError):
            return index, index

    # Always send images in their stable character-slot order, even if a later
    # save or edit happens to reorder the reference_images JSON array.
    for original_index, ref in sorted(indexed_refs, key=reference_slot)[:max_references]:
        path = folder / str(ref.get("path") or "")
        if path.exists() and path.is_file():
            try:
                slot = int(ref.get("slot", original_index))
            except (TypeError, ValueError):
                slot = original_index
            character = (project.get("character_bible") or [])
            # The name captured when the portrait was attached is authoritative.
            # Slot lookup is only a legacy fallback for older saved books.
            character_name = str(ref.get("character_name") or "").strip()
            if not character_name and 0 <= slot < len(character):
                character_name = str(character[slot].get("name") or "").strip()
            character_name = character_name or str(
                f"reference {len(attached) + 1}").strip()
            attached.append((path, character_name))

    if target == "cover":
        scene_words = str((project.get("cover") or {}).get("image_prompt") or "")
        direction_names = set()
    else:
        page = project["pages"][int(target) - 1]
        # The image prompt and structured visual directions—not narrative prose—
        # decide who is visibly present. Prose can mention absent, dead or
        # remembered characters and previously caused them to be resurrected.
        scene_words = str(page.get("image_prompt") or "")
        direction_names = {
            str(item.get("name") or "").strip().casefold()
            for item in (page.get("character_directions") or [])
            if isinstance(item, dict) and str(item.get("name") or "").strip()
        }
    if direction_names:
        present = [
            entry for entry in attached if entry[1].casefold() in direction_names
        ]
    else:
        present = [entry for entry in attached
                   if entry[1] and re.search(
                       r"\b" + re.escape(entry[1]) + r"\b",
                       scene_words, flags=re.IGNORECASE)]
    # Pages must fail closed when no recurring name is visually requested;
    # otherwise all masters leak into unrelated scenes. Covers retain the old
    # all-character fallback when their art direction is intentionally generic.
    if target != "cover" or present:
        attached = present

    current = None
    if target == "cover":
        rel = (project.get("cover") or {}).get("image")
    else:
        rel = project["pages"][int(target) - 1].get("image")
    if rel:
        candidate = folder / rel
        if candidate.exists() and candidate.is_file():
            current = candidate

    existing = []
    cover_rel = (project.get("cover") or {}).get("image")
    if cover_rel:
        existing.append(folder / cover_rel)
    for page in project.get("pages") or []:
        if page.get("image"):
            existing.append(folder / page["image"])

    if engine == "flux":
        # Kontext is an edit/reference model, so prefer the image being revised,
        # then a character reference, then another existing book image.
        ordered = ([(current, "existing illustration")] if current else []) + attached + [
            (path, "existing illustration") for path in existing]
    else:
        # HiDream and Nano Banana can generate directly from text. Feeding a
        # previous page back as an implicit reference makes subsequent pages
        # repeat its composition; only explicit character references carry over.
        ordered = attached
    unique = []
    seen_paths = set()
    for path, name in ordered:
        if path and path.exists() and path.is_file() and path not in seen_paths:
            unique.append((path, name))
            seen_paths.add(path)
    # Keep two HiDream entries for description-based character distinction;
    # the unstable MLX edit path is deliberately not invoked below.
    limit = (
        1 if engine == "flux"
        else 2 if engine == "hidream"
        else MAX_CHARACTER_REFERENCES
        if engine in KLEIN_IMAGE_ENGINES or engine == "nano_banana_2"
        else MAX_GROK_REFERENCES
    )
    return unique[:limit]


def page_character_scene_contract(
        page: dict, reference_entries: list[tuple[Path, str]],
        compact: bool = False) -> str:
    """Create a numbered per-character contract that prevents attribute swaps."""
    directions = [
        item for item in (page.get("character_directions") or [])
        if isinstance(item, dict) and str(item.get("name") or "").strip()
    ]
    by_name = {
        str(item.get("name") or "").strip().casefold(): item
        for item in directions
    }
    lines = []
    mapped_names = []
    mapped_cast_labels = []
    scene_text = str(page.get("image_prompt") or "")

    def named_scene_clause(character_name: str) -> str:
        match = re.search(
            r"\b" + re.escape(character_name) + r"\b", scene_text,
            flags=re.IGNORECASE)
        if not match:
            return ""
        end = len(scene_text)
        for _, other_name in reference_entries:
            if other_name.casefold() == character_name.casefold():
                continue
            other = re.search(
                r"\b" + re.escape(other_name) + r"\b", scene_text[match.end():],
                flags=re.IGNORECASE)
            if other:
                end = min(end, match.end() + other.start())
        return re.sub(r"\s+", " ", scene_text[match.start():end]).strip(" ,.;")[:700]

    for index, (_, name) in enumerate(reference_entries, 1):
        mapped_names.append(name)
        item = by_name.get(name.casefold()) or {}
        depiction_mode = str(
            item.get("depiction_mode") or "living character").strip()
        mapped_cast_labels.append(
            f"{name} as {depiction_mode}" if item else name
        )
        identity = str(item.get("identity") or "").strip()
        wardrobe = str(item.get("wardrobe") or "").strip()
        removed = str(item.get("removed_items") or "").strip()
        action = str(item.get("action") or "").strip()
        position = str(item.get("position") or "").strip()
        override = str(item.get("wardrobe_override") or "").strip().casefold() in {
            "yes", "true", "1", "required", "override"
        }
        if compact:
            parts = [f"CAST {index} {name}: identity=image {index}"]
            if item:
                parts.append(f"form={depiction_mode}")
            if wardrobe:
                parts.append(("temporary costume=" if override else "wearing=")
                             + wardrobe)
            if override and removed:
                parts.append("remove=" + removed)
            if action:
                parts.append("action=" + action)
            if position:
                parts.append("position=" + position)
            if not item:
                inferred_clause = named_scene_clause(name)
                if inferred_clause:
                    parts.append("assignment=" + inferred_clause)
            lines.append("; ".join(parts) + ".")
            continue
        parts = [
            f"CAST {index} — {name} uses identity image {index} exclusively"
        ]
        if identity:
            parts.append(f"identity anchor: {identity}")
        if item:
            parts.append(f"depiction mode: {depiction_mode}")
            if depiction_mode.casefold() not in {
                    "living", "living character", "full-body living character"}:
                parts.append(
                    f"image {index} supplies {name}'s recognisable likeness only; "
                    f"depict {name} solely in the stated {depiction_mode} form"
                )
        if wardrobe:
            if override:
                parts.append(
                    f"temporary wardrobe override assigned only to {name}: {wardrobe}"
                )
            else:
                parts.append(f"current wardrobe: {wardrobe}")
        if override and removed:
            parts.append(f"replaced default items: {removed}")
        if action:
            parts.append(f"visible action: {action}")
        if position:
            parts.append(f"composition position: {position}")
        if not item:
            inferred_clause = named_scene_clause(name)
            parts.append(
                f"all appearance, wardrobe, action and position phrases attached "
                f"to the name {name} in the scene belong to {name} alone"
            )
            if inferred_clause:
                parts.append(
                    f"scene assignment bound only to {name}: {inferred_clause}"
                )
                if re.search(
                        r"\b(?:disguise|costume|gown|dress|wig|women['’]s clothing|"
                        r"woman['’]s clothes|maenad attire)\b",
                        inferred_clause, flags=re.IGNORECASE):
                    parts.append(
                        f"this scene assignment is a temporary wardrobe override "
                        f"for {name} alone"
                    )
        lines.append("; ".join(parts) + ".")

    supporting = [
        str(item).strip() for item in (page.get("supporting_characters") or [])
        if str(item).strip()
    ]
    complete = str(page.get("cast_complete") or "").strip().casefold() in {
        "yes", "true", "1", "complete"
    }
    if not complete and re.search(
            r"\b(?:other|extra|additional)\s+(?:people|persons|characters|figures)\b",
            str(page.get("negative_prompt") or ""), flags=re.IGNORECASE):
        complete = True
    cast = mapped_cast_labels + supporting
    if complete and cast:
        lines.append(
            "The complete visible cast consists of exactly: " + "; ".join(cast) + "."
        )
    elif supporting:
        lines.append(
            "Additional distinct supporting figures requested by this scene: "
            + "; ".join(supporting) + "."
        )
    if not lines:
        return ""
    heading = (
        "COMPACT CAST CONTRACT — do not swap any assignment:\n"
        if compact else
        "PER-CHARACTER SCENE CONTRACT — keep every identity, wardrobe, prop, "
        "action and position bound to its own named cast record:\n"
    )
    return heading + "\n".join(lines)


def _supporting_subject_quantity(value: str) -> int:
    """Estimate the number of visible figures in one supporting-cast record."""
    text = str(value or "").strip().casefold()
    if not text:
        return 0
    numeric = re.match(r"^\s*(\d+)\b", text)
    if numeric:
        return max(1, min(20, int(numeric.group(1))))
    quantities = {
        "a": 1, "an": 1, "one": 1, "two": 2, "three": 3, "four": 4,
        "five": 5, "six": 6, "seven": 7, "eight": 8, "nine": 9,
        "ten": 10, "eleven": 11, "twelve": 12,
    }
    first = re.match(r"^\s*([a-z]+)\b", text)
    return quantities.get(first.group(1), 1) if first else 1


def page_visible_subject_count(
        page: dict, reference_entries: list[tuple[Path, str]]) -> int:
    """Count the literal visible cast, including unnamed supporting figures."""
    names = {
        str(item.get("name") or "").strip().casefold()
        for item in (page.get("character_directions") or [])
        if isinstance(item, dict) and str(item.get("name") or "").strip()
    }
    if not names:
        names = {str(name or "").strip().casefold()
                 for _, name in reference_entries if str(name or "").strip()}
    supporting = sum(
        _supporting_subject_quantity(item)
        for item in (page.get("supporting_characters") or [])
        if str(item or "").strip()
    )
    return len(names) + supporting


def page_exact_cast_directive(
        page: dict, reference_entries: list[tuple[Path, str]]) -> str:
    """Put the page-only cast near the start so adjacent scenes cannot leak in."""
    names = [
        str(item.get("name") or "").strip()
        for item in (page.get("character_directions") or [])
        if isinstance(item, dict) and str(item.get("name") or "").strip()
    ]
    if not names:
        names = [str(name or "").strip() for _, name in reference_entries
                 if str(name or "").strip()]
    supporting = [
        str(item).strip() for item in (page.get("supporting_characters") or [])
        if str(item).strip()
    ]
    complete = str(page.get("cast_complete") or "").strip().casefold() in {
        "yes", "true", "1", "complete"
    }
    cast = names + supporting
    if not cast:
        return ""
    if complete:
        return (
            "EXACT VISIBLE CAST FOR THIS PAGE: " + "; ".join(cast) + ". "
            "Include every listed subject in the stated quantity and include no "
            "additional people, animals, creatures, bystanders or helpers."
        )
    return "VISIBLE SUBJECTS REQUESTED FOR THIS PAGE: " + "; ".join(cast) + "."


def page_continuity_object_contract(page: dict, compact: bool = False) -> str:
    """Keep a plot-critical object literal and stable across adjacent pages."""
    objects = [
        item for item in (page.get("continuity_objects") or [])
        if isinstance(item, dict) and str(item.get("id") or "").strip()
    ]
    if not objects:
        return ""
    lines = []
    for item in objects:
        if compact:
            identity = str(item.get("literal_identity") or item.get("id") or "").strip()
            details = []
            for key, label in (
                    ("appearance", "appearance"),
                    ("state", "current state"),
                    ("holder_and_position", "position")):
                value = str(item.get(key) or "").strip()
                if value:
                    details.append(f"{label}: {value}")
            line = f"Keep the plot-critical {identity} visually literal and consistent"
            if details:
                line += "; " + "; ".join(details)
            lines.append(line + ".")
            continue
        parts = [f"OBJECT {str(item.get('id') or '').strip()}"]
        for key, label in (
                ("literal_identity", "literal physical identity"),
                ("appearance", "stable appearance"),
                ("state", "state on this page"),
                ("holder_and_position", "holder and position")):
            value = str(item.get(key) or "").strip()
            if value:
                parts.append(f"{label}: {value}")
        parts.append("render literally; never substitute" if compact else
                     "render this same literal object rather than a symbolic or generic substitute")
        lines.append("; ".join(parts) + ".")
    if compact:
        return "\n".join(lines)
    return "PLOT-CRITICAL CONTINUITY OBJECT CONTRACT:\n" + "\n".join(lines)


def scene_explicitly_uninhabited(
        negative_prompt: str, reference_entries: list[tuple[Path, str]]) -> bool:
    """Only declare an empty scene when the prompt truly requests no people."""
    if reference_entries:
        return False
    return bool(re.search(
        r"\b(?:no people|no person|uninhabited|empty of people|"
        r"without people|deserted landscape)\b",
        str(negative_prompt or ""), flags=re.IGNORECASE))


def local_image_geometry(settings: dict, engine: str = "") -> tuple[int, int, str, str]:
    w_in, h_in = trim_size(settings)
    ratio = w_in / h_in
    candidates = {
        "1:1": 1.0, "2:3": 2 / 3, "3:4": 3 / 4, "4:5": 4 / 5,
        "3:2": 3 / 2, "4:3": 4 / 3,
    }
    aspect = min(candidates, key=lambda key: abs(candidates[key] - ratio))
    if engine in {"kontext", "hidream", "flux_2_klein_4b_local"}:
        # Both local engines are substantially faster near 1K. HiDream's native
        # presets are 4 MP and roughly double generation time; the PDF exporter
        # scales the finished art to trim size. Keep exact book proportions and
        # disable HiDream's trained-resolution snap in the request payload.
        local_sizes = {
            "1:1": (1024, 1024), "4:5": (1024, 1280),
            "2:3": (768, 1152), "3:4": (960, 1280),
            "3:2": (1152, 768), "4:3": (1280, 960),
        }
        width, height = local_sizes[aspect]
        return width, height, "", aspect
    if aspect == "1:1":
        return 2048, 2048, "1:1", "1:1"
    if aspect == "4:5":
        return 1792, 2304, "4:5", "4:5"
    if aspect == "2:3":
        return 1664, 2496, "", "2:3"
    if aspect == "3:4":
        return 1792, 2304, "", "3:4"
    if aspect == "3:2":
        return 2496, 1664, "3:2", "3:2"
    return 2304, 1792, "", "4:3"


def near_duplicate_project_image(project: dict, raw: bytes) -> str:
    """Return the matching image path when HiDream repeats a composition."""
    from PIL import Image

    def fingerprints(source) -> tuple[list[bool], list[bool]]:
        with Image.open(source) as image:
            grey = image.convert("L")
            average_image = grey.resize((16, 16), Image.Resampling.LANCZOS)
            average_pixels = list(average_image.getdata())
            average = sum(average_pixels) / len(average_pixels)
            average_hash = [value >= average for value in average_pixels]
            difference_image = grey.resize((17, 16), Image.Resampling.LANCZOS)
            difference_pixels = list(difference_image.getdata())
            difference_hash = [
                difference_pixels[row * 17 + col]
                > difference_pixels[row * 17 + col + 1]
                for row in range(16) for col in range(16)
            ]
            return average_hash, difference_hash

    try:
        new_average, new_difference = fingerprints(io.BytesIO(raw))
    except Exception:
        return ""
    rels = []
    cover_image = (project.get("cover") or {}).get("image")
    if cover_image:
        rels.append(cover_image)
    rels.extend(page.get("image") for page in project.get("pages") or []
                if page.get("image"))
    for rel in dict.fromkeys(rels):
        path = project_dir(project["id"]) / rel
        if not path.exists() or not path.is_file():
            continue
        try:
            old_average, old_difference = fingerprints(path)
        except Exception:
            continue
        average_distance = sum(a != b for a, b in zip(new_average, old_average))
        difference_distance = sum(
            a != b for a, b in zip(new_difference, old_difference))
        if average_distance <= 64 and difference_distance <= 72:
            return rel
    return ""


def run_local_image(engine: str, prompt: str, settings: dict,
                    references: list[str], reference_labels: list[str] | None = None) -> bytes:
    if current_job_canceled():
        raise ImageJobCanceled("Image generation stopped")
    service_engine = "kontext" if engine == "flux" else engine
    width, height, preset, aspect = local_image_geometry(settings, service_engine)
    payload = {
        "engine": service_engine,
        "prompt": (prompt[:6000] if service_engine in (
            "flux_2_klein_4b", "flux_2_klein_4b_local") else prompt[:2000]),
        "width": width, "height": height, "steps": 28,
        "seed": int(uuid.uuid4().hex[:8], 16),
        "snap": service_engine != "hidream",
        "reference_images": references,
        "reference_labels": list(reference_labels or [])[:len(references)],
        "nano_aspect_ratio": aspect, "nano_output_format": "png",
        "flux2_resolution": "1", "flux2_aspect_ratio": aspect,
        "flux2_output_format": "png", "flux2_output_quality": 100,
        "flux2_go_fast": False, "flux2_disable_safety_checker": False,
        "flux2_max_tokens": 1024,
    }
    if preset:
        payload["preset"] = preset
    if service_engine == "kontext":
        payload["reference_image"] = references[0] if references else ""
    cost_reservation = ""
    if service_engine == "nano_banana_2":
        cost_reservation = reserve_cloud_cost(
            "Replicate", "google/nano-banana-2", "replicate_image", 0.10,
            detail=(f"one 1K output image · {len(references)} reference image(s) · "
                    "conservative provider estimate"))
    elif service_engine == "flux_2_klein_4b":
        estimate = max(0.005, 0.001 * (len(references) + 1))
        cost_reservation = reserve_cloud_cost(
            "Replicate", "black-forest-labs/flux-2-klein-4b",
            "replicate_image", estimate,
            detail=(f"approximately {len(references) + 1} billed megapixel(s) · "
                    "provider estimate"))
    started = time.monotonic()
    studio_event(
        "image.local_requested", engine=service_engine, width=width, height=height,
        steps=payload["steps"], seed=payload["seed"], reference_count=len(references),
        reference_labels=payload["reference_labels"],
        reference_mapping_version=(
            "character-map-v2" if service_engine in KLEIN_IMAGE_ENGINES else ""
        ),
        prompt_chars=len(prompt), prompt_sha256=hashlib.sha256(
            prompt.encode()).hexdigest()[:16],
    )
    update_current_job(f"Starting {service_engine} image", 1)
    response = local_image_json("/generate", payload, timeout=120)
    job_id = response.get("job_id")
    if not job_id:
        raise RuntimeError("The local image service did not start the image job")
    studio_event("image.local_started", engine=service_engine,
                 local_image_job_id=job_id)
    parent_job_id = getattr(_job_context, "job_id", "")
    if parent_job_id:
        with _jobs_lock:
            parent = _jobs.get(parent_job_id)
            if parent:
                parent["local_image_job_id"] = job_id
    deadline = time.time() + 1200
    while time.time() < deadline:
        if current_job_canceled():
            try:
                local_image_json("/cancel", {}, timeout=15)
            except Exception:
                pass
            raise ImageJobCanceled("Image generation stopped")
        status = local_image_json(
            "/status?" + urllib.parse.urlencode({"id": job_id}), timeout=60)
        update_current_job(
            str(status.get("stage") or f"Creating with {service_engine}"),
            int(status.get("progress") or 0))
        if status.get("status") == "done":
            result = status.get("result") or {}
            image_url = result.get("url")
            if not image_url:
                raise RuntimeError("The local image job finished without an image")
            raw = download_image({"url": urllib.parse.urljoin(
                LOCAL_IMAGE_BASE + "/", image_url.lstrip("/"))})
            studio_event(
                "image.local_completed", engine=service_engine,
                local_image_job_id=job_id,
                duration_ms=round((time.monotonic() - started) * 1000),
                output_bytes=len(raw), width=result.get("width"),
                height=result.get("height"),
                reference_fallback=result.get("reference_fallback") or "",
            )
            if cost_reservation:
                settle_cloud_cost(
                    cost_reservation, estimated=True, status="estimated")
                cost_reservation = ""
            else:
                record_local_image_use(
                    service_engine,
                    detail=(
                        f"one local image · {result.get('width') or width}×"
                        f"{result.get('height') or height}"))
            return raw
        if status.get("status") == "canceled":
            raise ImageJobCanceled("Image generation stopped")
        if status.get("status") == "error" or status.get("error"):
            studio_event(
                "image.local_failed", level="error", engine=service_engine,
                local_image_job_id=job_id,
                duration_ms=round((time.monotonic() - started) * 1000),
                error=status.get("error") or "The local image job failed",
            )
            raise RuntimeError(status.get("error") or "The local image job failed")
        time.sleep(2)
    raise RuntimeError("The local image job timed out after 20 minutes")


def image_text_policy(engine: str, workbook_art: bool = False) -> str:
    if engine in ("flux_2_klein_4b", "flux_2_klein_4b_local"):
        return (
            "Fill the entire frame with one continuous natural illustrated scene. "
            "Use organic scenery, expressive character action, natural material "
            "textures and uninterrupted areas of foliage, sky, earth, water or "
            "solid-colour fabric. Keep the composition pictorial and immersive."
        )
    if engine in ("hidream", "flux"):
        return (
            "Render the requested people, scenery and ordinary objects as a finished "
            "full-page composition. Keep the sky, horizon, buildings, clothing and "
            "object surfaces natural, clean and visually uninterrupted. Use coherent "
            "material texture, organic shapes and uninterrupted colour throughout."
        )
    if engine == "nano_banana_2" and workbook_art:
        return (
            "Return one uninterrupted full-frame illustration, never a document, "
            "worksheet, poster, infographic or designed page. Keep every card, "
            "sheet of paper, tab, book, screen, wall and object surface visually "
            "blank and unmarked, using only natural material texture and colour. "
            "Translate written exercises and abstract concepts into the requested "
            "physical objects and scene; never typeset source copy, prompt "
            "instructions, headings, labels, captions, numbering or metadata."
        )
    return (
        "CLOUD-MODEL TEXT RULE: Readable signage, labels, or words within the scene "
        "are allowed when the page description requests them. Render requested "
        "wording accurately and legibly, but do not invent unnecessary writing, "
        "logos, or watermarks. Do not render the book title, subtitle, page heading, "
        "or body copy because the studio overlays those separately."
    )


def grok_local_image_prompt_rules(settings: dict) -> str:
    """Tell Grok how to author prompts that are safe for local image models."""
    engine = (settings or {}).get("image_engine")
    structural_rules = """For every page, return character_directions with one
separate record for each visible recurring character. Assign that named
character's exact wardrobe, action and composition position independently. If
the story explicitly puts a character into a disguise or different costume, set
wardrobe_override to yes, list the complete replacement wardrobe and list the
default clothing removed. Never transfer that costume, wig, armour, prop, action
or position to another named character. List every other visible person in
supporting_characters and set cast_complete to yes only after the full visible
cast is accounted for. Set depiction_mode precisely. A referenced identity shown
only as remains, a portrait, statue, reflection or vision is not a living
full-body cast member; state the physical form literally and bind only the needed
likeness to its master image. For every plot-critical object that continues
across pages, return a continuity_objects record with the same stable id and
literal_identity on every appearance. Never shorten a specific object into an
ambiguous word such as trophy, prize, thing or symbol. A restrained non-graphic
rendering may conceal wound detail, but it must retain the object's true identity
and state."""
    if engine in KLEIN_IMAGE_ENGINES:
        return """FLUX 2 KLEIN PROMPT BUDGET AND PRIORITY:
- The complete runtime prompt is limited to 1,024 tokens. Your image_prompt is
  only one component: normally use 140–220 tokens and never exceed 280 tokens.
- Begin immediately with the exact visible action and its visible result. Name
  only characters actually in frame, then state the setting and plot-critical
  objects, composition/camera, expression, lighting, colour and material style.
- Use concrete visual facts. Do not include narration, dialogue, backstory,
  abstract morals, alternatives, explanations, repeated synonyms or boilerplate.
- When master character images are supplied, do not repeat long appearance
  biographies inside image_prompt. Use exact names; put each character's current
  form, costume, action and position once in character_directions instead.
- Make the prompt self-contained and page-specific, but say every requirement
  once. Put the most important action and object state in the first two sentences.
""" + structural_rules + """
- image_prompt must contain only positive visual description: subjects, actions,
  setting, composition, light, colour, materials and natural open areas.
- Never place typography-control instructions inside image_prompt. Do not output
  the words or phrases: wordless, text, writing, letters, title, subtitle,
  caption, signage, logo, watermark, inscription, negative prompt, "no text",
  "no writing", "pure wordless illustration", or "space for text".
- Describe reserved areas only as "uncluttered natural open space" or "clean
  uninterrupted background". End with visual details, not a warning."""
    if engine not in (
            "hidream", "flux", "flux_2_klein_4b", "flux_2_klein_4b_local"):
        return (
            "Write each image_prompt as a precise visual scene description for the "
            "selected cloud image model.\n" + structural_rules
        )
    return """The selected illustrator is a local image model. Every returned
image_prompt must contain only positive visual scene description: subjects,
actions, setting, composition, light, colour, materials and natural open areas.
""" + structural_rules + """
Never place typography-control instructions inside image_prompt. In particular,
do not output the words or phrases: wordless, text, writing, letters, title,
subtitle, caption, signage, logo, watermark, inscription, negative prompt,
"no text", "no writing", "pure wordless illustration", or "space for text".
Do not quote any prohibited phrase. Describe reserved areas only as
"uncluttered natural open space" or "clean uninterrupted background". End each
image_prompt with visual details, not an instruction or warning."""


def local_wordless_description(value: str) -> str:
    """Remove prompt cues that make local Dev models invent typography.

    HiDream-O1 Dev runs at guidance_scale=0, so a conventional negative prompt
    has no effect. Positive, wordless scene descriptions are more dependable.
    The stored creative brief is left untouched; only the generation copy is
    made safer for local engines.
    """
    text = str(value or "")
    text = re.sub(
        r"\b(?:pure\s+)?wordless\s+(?:illustration|scene|scenery|artwork|art)\b"
        r"(?:\s+and\s+people\s+only)?",
        "finished visual composition", text, flags=re.IGNORECASE)
    text = re.sub(
        r"\b(?:wide\s+)?(?:safe\s+)?negative\s+space\s*"
        r"([^,.;]{0,45}?)\s+for\s+(?:the\s+)?(?:title\s+)?"
        r"(?:text|copy|caption|wording)\b",
        r"uncluttered natural open space \1", text, flags=re.IGNORECASE)
    text = re.sub(
        r"\bfor\s+(?:later\s+)?(?:title|text|copy|caption|wording)\s+"
        r"(?:overlay|placement)\b",
        "as uncluttered natural open space", text, flags=re.IGNORECASE)
    text = re.sub(
        r"\b(?:large\s+)?(?:safe\s+)?empty\s+([^,.;]{0,60}?)\s+for\s+"
        r"(?:title\s+lettering|subtitle|text|copy|caption|wording)"
        r"(?:\s+and\s+[^,.;]{0,35})?",
        r"uncluttered natural open space \1", text, flags=re.IGNORECASE)
    text = re.sub(
        r"\b(?:no|without)\s+(?:readable\s+)?"
        r"(?:writing|text|letters?|words?|titles?|captions?)\b"
        r"(?:\s+(?:or|and)\s+(?:writing|text|letters?|words?|titles?|captions?))*"
        r"(?:\s+(?:anywhere|on\s+any\s+surface|in\s+the\s+art(?:work)?))?",
        "clean uninterrupted surfaces", text, flags=re.IGNORECASE)
    text = re.sub(
        r"\b(?:no|without)\s+(?:readable\s+|real\s+)?"
        r"(?:logos?|watermarks?|inscriptions?|signage|price\s+labels?)\b",
        "clean uninterrupted surfaces", text, flags=re.IGNORECASE)
    text = re.sub(r"\bno\s+carvings\s+that\s+look\s+like\s+writing\b",
                  "simple shallow carvings", text, flags=re.IGNORECASE)
    text = re.sub(r"\bno\s+patterns\s+resembling\s+text\b",
                  "simple organic patterns", text, flags=re.IGNORECASE)
    text = re.sub(r"\bpages?\s+without\s+readable\s+text\b",
                  "pages showing simple colourful drawings", text,
                  flags=re.IGNORECASE)
    text = re.sub(r"\bno\s+real\s+text\b",
                  "a simple colourful drawing", text, flags=re.IGNORECASE)
    text = re.sub(r"\bsimple\s+non-letter\s+drawing\b",
                  "simple colourful drawing", text, flags=re.IGNORECASE)
    text = re.sub(
        r"\bno\s+[^,.;]{0,70}\b(?:writing|text|letters?|titles?|captions?|"
        r"signage|logos?|watermarks?|inscriptions?)\b",
        "clean uninterrupted surfaces", text, flags=re.IGNORECASE)
    text = re.sub(r"\btextured\s+paper\s+feel\b",
                  "traditional dry-media texture", text,
                  flags=re.IGNORECASE)
    text = re.sub(r"\bparchment(?:-toned)?\b", "warm ivory", text,
                  flags=re.IGNORECASE)
    text = re.sub(
        r"\bfor\s+(?:a\s+|the\s+)?(?:title|text|copy|caption|wording)\s+overlay\b",
        "as uninterrupted uncluttered space", text, flags=re.IGNORECASE)
    text = re.sub(
        r"\b(?:signage|signs?|labels?|posters?|menus?|screens?|displays?)\b",
        "plain unmarked surfaces", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()


def missing_required_reference_names(project: dict) -> list[str]:
    """Require masters for concept/series leads, not later supporting cast."""
    concept = project.get("market_concept") or {}
    recurring = concept.get("characters")
    if not isinstance(recurring, list) or not recurring:
        bible = normalize_series_bible(project.get("series_bible"), project)
        recurring = bible.get("characters") if bible.get("enabled") else []
    required_names = [
        str(character.get("name") or "").strip()
        for character in (recurring or [])[:MAX_CHARACTER_REFERENCES]
        if isinstance(character, dict) and str(character.get("name") or "").strip()
    ]
    attached_names = {
        str(reference.get("character_name") or "").strip().casefold()
        for reference in project.get("reference_images") or []
        if reference.get("path") and (project_dir(project["id"])
                                      / str(reference.get("path"))).is_file()
    }
    return [name for name in required_names if name.casefold() not in attached_names]


def generate_image(project_id: str, target: str, engine_override: str = "") -> dict:
    p = load_project(project_id)
    s = p["settings"]
    requested_override = str(engine_override or "").strip()
    engine = (requested_override if requested_override in IMAGE_ENGINE_IDS
              else s.get("image_engine", "grok"))
    if engine not in IMAGE_ENGINE_IDS:
        engine = "grok"
    if (engine != "hidream" and s.get("book_type") == "picture_book"
            and p.get("market_concept")
            and p.get("character_bible")):
        missing = missing_required_reference_names(p)
        if missing:
            raise ValueError(
                "Create or attach the character reference image first for: "
                + ", ".join(missing))
    studio_event(
        "illustration.started", project_id=project_id, target=str(target),
        engine=engine, title=p.get("title", ""),
        character_count=len(p.get("character_bible") or []),
        reference_count=len(p.get("reference_images") or []),
    )
    # Klein and Nano Banana need scene-only prompts. Sending the page's raw body
    # copy makes them typeset the workbook/story instead of illustrating it.
    scene_only_art = engine in (
        "hidream", "flux", "flux_2_klein_4b", "flux_2_klein_4b_local",
        "nano_banana_2")
    workbook_art = bool(p.get("series_template")) or s.get("book_type") == "educational"
    # Resolve the numbered reference-to-character map before composing the
    # prompt. Klein is easily confused when descriptions for absent recurring
    # characters are included alongside a smaller subset of reference images.
    reference_entries = local_reference_entries(p, target, engine)
    if engine == "flux_2_klein_4b_local" and len(reference_entries) > 3:
        names = ", ".join(name for _, name in reference_entries)
        studio_event(
            "illustration.local_flux_reference_limit", level="warning",
            project_id=project_id, target=str(target), engine=engine,
            reference_count=len(reference_entries), references=names,
            cloud_fallback_started=False,
        )
        raise ValueError(
            "Local FLUX is limited to three character references for reliable "
            f"identity matching. This scene needs {len(reference_entries)}: "
            f"{names}. Choose ‘Use Nano Banana 2 for this image’ instead. "
            "No cloud request or charge was started."
        )
    variation_id = uuid.uuid4().hex[:10]
    world_bible = p.get("world_bible") or ""
    if scene_only_art:
        world_bible = local_wordless_description(world_bible)
    if engine in KLEIN_IMAGE_ENGINES or engine == "nano_banana_2":
        # The image service prepends an explicit Image N -> character mapping.
        # Never add the full world bible here: it can contain writing-space cues,
        # props, supporting cast or events from other pages that leak into art.
        common = (
            f"STYLE: {style_label(s)}. Audience: {s.get('audience')}."
            " Preserve the established medium, broad palette and era, but use "
            "only the people, animals, props and action explicitly assigned to "
            "this page. Full-page professional composition."
        )
    else:
        common = (
            f"Original {'full-page' if scene_only_art else 'book'} illustration. "
            f"Visual style: {style_label(s)}. Audience: {s.get('audience')}. "
            f"Visual continuity: {world_bible}. "
            f"Character bible: {json.dumps(p.get('character_bible'), ensure_ascii=False)}. "
            "Maintain exact recurring character identity, clothing and palette. "
            "Professional publishable composition."
        )
        ref_notes = [
            r.get("label") for r in p.get("reference_images", []) if r.get("label")
        ]
        if ref_notes:
            common += " Reference image roles: " + "; ".join(ref_notes) + "."
    if target == "cover":
        scene = p["cover"]["image_prompt"]
        if scene_only_art:
            scene = local_wordless_description(scene)
            prompt = (
                "Create a full-frame cover-background illustration depicting this exact visible scene: "
                + scene
                + "\nDistinctive cover-background viewpoint with uncluttered "
                  "natural open space for later design work.\n"
                + common + "\n" + image_text_policy(engine, workbook_art)
            )
        else:
            prompt = (
                "COVER ART — follow this scene description first and exactly: "
                + scene
                + "\nCreate a distinctive front-cover composition, not an interior-page scene.\n"
                + common + "\n" + image_text_policy(engine, workbook_art)
            )
        had_existing_image = bool(p["cover"].get("image"))
        filename = f"cover-{variation_id}.png" if had_existing_image else "cover.png"
    else:
        number = int(target)
        page = p["pages"][number - 1]
        scene = page["image_prompt"]
        page_words = str(page.get("text") or "").strip()
        subject_count = page_visible_subject_count(page, reference_entries)
        if engine == "flux_2_klein_4b_local" and subject_count > 3:
            studio_event(
                "illustration.local_flux_subject_limit", level="warning",
                project_id=project_id, target=str(target), engine=engine,
                visible_subject_count=subject_count,
                recurring_characters=[
                    str(item.get("name") or "").strip()
                    for item in (page.get("character_directions") or [])
                    if isinstance(item, dict)
                    and str(item.get("name") or "").strip()
                ],
                supporting_characters=list(page.get("supporting_characters") or []),
                cloud_fallback_started=False,
            )
            raise ValueError(
                "This scene asks local FLUX to keep "
                f"{subject_count} visible characters or animals in sync. Its "
                "reliable limit is three. Choose ‘Use Nano Banana 2 for this "
                "image’ instead. No cloud request or charge was started."
            )
        exact_cast = page_exact_cast_directive(page, reference_entries)
        scene_contract = (
            page_character_scene_contract(
                page, reference_entries,
                compact=engine in KLEIN_IMAGE_ENGINES or engine == "nano_banana_2")
            if engine in KLEIN_IMAGE_ENGINES
            or engine in {"grok", "nano_banana_2"} else ""
        )
        object_contract = page_continuity_object_contract(
            page, compact=engine in KLEIN_IMAGE_ENGINES or engine == "nano_banana_2")
        if scene_contract or object_contract:
            studio_event(
                "illustration.scene_contract", project_id=project_id,
                target=str(target), engine=engine,
                characters=[name for _, name in reference_entries],
                continuity_objects=[
                    str(item.get("id") or "")
                    for item in (page.get("continuity_objects") or [])
                    if isinstance(item, dict) and str(item.get("id") or "").strip()
                ],
                wardrobe_overrides=[
                    str(item.get("name") or "")
                    for item in (page.get("character_directions") or [])
                    if isinstance(item, dict)
                    and str(item.get("wardrobe_override") or "").casefold()
                    in {"yes", "true", "1", "required", "override"}
                ],
                cast_complete=str(page.get("cast_complete") or ""),
            )
        if scene_only_art:
            scene = local_wordless_description(scene)
            prompt = (
                "Create one full-frame illustration depicting this exact visible scene: "
                + scene
                + (("\n" + exact_cast) if exact_cast else "")
                + (("\n" + scene_contract) if scene_contract else "")
                + (("\n" + object_contract) if object_contract else "")
                + "\nUse a unique viewpoint and clear visual storytelling.\n"
                + common + "\n" + image_text_policy(engine, workbook_art)
            )
            negative = str(page.get("negative_prompt") or "")
            if scene_explicitly_uninhabited(negative, reference_entries):
                prompt += (
                    "\nThis is a tranquil uninhabited scene containing scenery and "
                    "ordinary objects only."
                )
        else:
            prompt = (
                f"INTERIOR PAGE {number} — follow this scene description first and exactly: "
                + scene
                + "\nPAGE WORDS — VISUAL SOURCE OF TRUTH: " + page_words
                + "\nDepict the exact visible action and result in those words. Show only "
                  "the named characters who are present, with the correct props, poses, "
                  "expressions and spatial relationships. Do not substitute a generic scene."
                + (("\n" + scene_contract) if scene_contract else "")
                + (("\n" + object_contract) if object_contract else "")
                + "\nCreate a composition unique to this page; do not repeat the cover "
                  "or another page's scene.\n"
                + common + "\n" + image_text_policy(engine, workbook_art)
            )
        if not scene_only_art and page.get("negative_prompt"):
            prompt += "\nAvoid: " + page["negative_prompt"]
        had_existing_image = bool(page.get("image"))
        filename = (
            f"page-{number:03d}-{variation_id}.png"
            if had_existing_image else f"page-{number:03d}.png"
        )
    if had_existing_image:
        revision_snapshot(
            p, "before-cover-image-regeneration" if target == "cover"
            else f"before-page-{number}-image-regeneration")
        prompt += (
            f"\nREGENERATION VARIATION {variation_id}: Create a clearly new "
            "alternative illustration. Keep the established characters and art "
            "style, but substantially change the composition, poses, camera angle, "
            "background details and visual storytelling. Do not reproduce the "
            "previous image."
        )
    if (scene_only_art and engine not in KLEIN_IMAGE_ENGINES
            and engine != "nano_banana_2"):
        prompt += "\nFINAL VISUAL CHECK: " + image_text_policy(engine, workbook_art)
    studio_event(
        "illustration.prompt_built", project_id=project_id, target=str(target),
        engine=engine, prompt_chars=len(prompt),
        raw_page_words_included=bool(target != "cover" and not scene_only_art),
        workbook_art=workbook_art,
        reference_count=len(reference_entries),
    )
    duplicate_retry = False
    if engine == "grok":
        update_current_job("Preparing the Grok illustration", 8)
        w, h = trim_size(s)
        aspect = "1:1" if abs(w - h) < .25 else ("3:4" if h > w else "4:3")
        payload = {
            "model": IMAGE_MODEL, "prompt": prompt[:6000],
            "n": 1, "resolution": "1k", "aspect_ratio": aspect,
        }
        for ref in p.get("reference_images", []):
            if ref.get("key_tag") == api_key_tag() and ref.get("file_id"):
                continue
            local_ref = project_dir(project_id) / str(ref.get("path") or "")
            if local_ref.exists():
                ref["file_id"] = xai_upload_file(local_ref)
                ref["key_tag"] = api_key_tag()
        save_project(p)
        refs_by_path = {
            (project_dir(project_id) / str(r.get("path") or "")).resolve(): r
            for r in p.get("reference_images", []) if r.get("file_id")
        }
        refs = [
            refs_by_path[path.resolve()]
            for path, _ in reference_entries if path.resolve() in refs_by_path
        ][:MAX_GROK_REFERENCES]
        endpoint = "/images/generations"
        if refs:
            endpoint = "/images/edits"
            payload["images"] = [{"file_id": r["file_id"]} for r in refs]
        update_current_job("Waiting for Grok Imagine cloud illustration", 25)
        data = xai_json(XAI_BASE + endpoint, payload, timeout=600)
        items = data.get("data") or []
        if not items:
            raise RuntimeError("xAI returned no generated image")
        update_current_job("Downloading the finished illustration", 88)
        raw = download_image(items[0])
    else:
        references = [image_data_url(path) for path, _ in reference_entries]
        if engine == "hidream" and reference_entries:
            # The installed port's own STATE.md marks edit/multi-reference as
            # unfinished and degenerate for K=1 as well as K=2. Avoid wasting a
            # full failed pass; normal text-to-image uses the detailed character
            # bible and scene prompt reliably.
            references = []
            names = " and ".join(name for _, name in reference_entries)
            prompt = (
                f"CHARACTER DISTINCTION: Depict {names} as separate complete "
                "characters with visibly distinct species, faces, colours, clothing "
                "and permanent accessories. Keep every short tail short and every "
                "bushy tail bushy according to the character description.\n" + prompt
            )
            studio_event(
                "illustration.reference_conditioning_unavailable", level="warning",
                project_id=project_id,
                target=str(target), engine=engine,
                references=[name for _, name in reference_entries],
                delivery="description_only", reason="mlx_edit_path_degenerate",
            )
        if engine == "flux" and not references:
            # Kontext is reference-led. Make a private HiDream starter, then let
            # FLUX produce the selected final image from that reference.
            starter = run_local_image("hidream", prompt, s, [])
            references = ["data:image/png;base64," + base64.b64encode(starter).decode()]
        reference_labels = [name for _, name in reference_entries]
        if reference_labels and engine not in KLEIN_IMAGE_ENGINES:
            cast_names = ", ".join(reference_labels)
            prompt += (
                "\nThe numbered references identify these recurring characters: "
                f"{cast_names}. Render each mapped recurring character once when "
                "requested by the scene. Supporting people explicitly requested "
                "by the scene may also appear and must remain visually distinct "
                "from every mapped recurring character."
            )
        raw = run_local_image(engine, prompt, s, references, reference_labels)
        duplicate = near_duplicate_project_image(p, raw) if engine == "hidream" else ""
        if duplicate:
            duplicate_retry = True
            retry_prompt = (
                "IMPORTANT: Produce a radically different composition from every "
                "existing book image. Change the scene layout, focal objects, camera "
                "position and visual storytelling while following the requested page.\n"
                + prompt
            )
            raw = run_local_image(engine, retry_prompt, s, references, reference_labels)
    update_current_job("Saving the illustration", 97)
    path = project_dir(project_id) / "images" / filename
    path.write_bytes(raw)
    rel = "images/" + filename
    text_warning = None
    if engine in ("flux_2_klein_4b", "flux_2_klein_4b_local"):
        detections = possible_generated_text(path)
        if detections:
            examples = ", ".join(
                repr(row["text"]) for row in detections[:3])
            text_warning = {
                "message": (
                    "Possible generated text or signage detected. The image was kept "
                    "and no automatic regeneration was charged. Review it before approval."
                ),
                "examples": examples,
                "detections": detections,
                "checked_at": now(),
            }
            studio_event(
                "illustration.possible_text_detected", level="warning",
                project_id=project_id, target=str(target), engine=engine,
                detections=detections,
            )
    if target == "cover":
        archive_image_version(p["cover"], "Before cover image regeneration")
        p["cover"]["image"] = rel
        p["cover"]["approved"] = False
        if text_warning:
            p["cover"]["image_warning"] = text_warning
        else:
            p["cover"].pop("image_warning", None)
    else:
        target_page = p["pages"][int(target) - 1]
        archive_image_version(
            target_page, f"Before page {target} image regeneration")
        target_page["image"] = rel
        target_page["approved"] = False
        if text_warning:
            target_page["image_warning"] = text_warning
        else:
            target_page.pop("image_warning", None)
    bump_asset_epoch(p)
    p["stage"] = "illustrations"
    p["history"].append({
        "at": now(),
        "action": (
            f"Regenerated {'cover' if target == 'cover' else f'page {target}'} image"
            if had_existing_image
            else f"Generated {'cover' if target == 'cover' else f'page {target}'} image"
        ) + f" with {next(x['label'] for x in IMAGE_ENGINES if x['id'] == engine)}"
        + ("; automatically retried a near-duplicate" if duplicate_retry else ""),
        "warning": text_warning["message"] if text_warning else "",
    })
    save_project(p)
    studio_event(
        "illustration.saved", project_id=project_id, target=str(target),
        engine=engine, relative_path=rel, output_bytes=len(raw),
        regenerated=had_existing_image, duplicate_retry=duplicate_retry,
    )
    return {
        "project": p, "image": rel, "image_engine": engine,
        "text_warning": text_warning,
    }


def upload_reference(project_id: str, data: dict) -> dict:
    p = load_project(project_id)
    refs = p.setdefault("reference_images", [])
    character_name = str(data.get("character_name") or "").strip()
    character_index = int(data.get("character_index", -1))
    if character_index not in range(MAX_CHARACTER_REFERENCES):
        raise ValueError("Choose one of the book's first four character slots")
    # Slots, not editable names, identify references. This prevents a temporary
    # duplicate character name from replacing another character's image.
    matching = [
        i for i, ref in enumerate(refs)
        if int(ref.get("slot", -1)) == character_index
    ]
    if not matching and len(refs) >= MAX_CHARACTER_REFERENCES:
        raise ValueError("This book already has four master reference images")
    match = re.fullmatch(r"data:(image/(?:png|jpeg|webp));base64,(.+)",
                         data.get("data", ""), re.S)
    if not match:
        raise ValueError("Please choose a PNG, JPEG or WebP image")
    raw = base64.b64decode(match.group(2), validate=True)
    if len(raw) > 15 * 1024 * 1024:
        raise ValueError("Reference image must be smaller than 15 MB")
    ext = {"image/png": ".png", "image/jpeg": ".jpg", "image/webp": ".webp"}[match.group(1)]
    folder = project_dir(project_id) / "references"
    folder.mkdir(exist_ok=True)
    filename = (
        f"reference-{character_index + 1}-{uuid.uuid4().hex[:10]}-"
        f"{slug(data.get('name', 'image'))}{ext}"
    )
    path = folder / filename
    path.write_bytes(raw)
    # Keep the local reference even if xAI is temporarily unavailable. Grok
    # uploads it lazily when that engine is actually selected.
    file_id = ""
    key_tag = ""
    if api_key():
        try:
            file_id = xai_upload_file(path)
            key_tag = api_key_tag()
        except Exception:
            pass
    entry = {"name": data.get("name") or filename,
             "label": data.get("label") or f"Identity reference for {character_name}",
             "path": "references/" + filename, "file_id": file_id,
             "character_name": character_name, "key_tag": key_tag,
             "slot": character_index}
    if matching:
        refs[matching[0]] = entry
        for duplicate in reversed(matching[1:]):
            refs.pop(duplicate)
    else:
        refs.append(entry)
    characters = p.setdefault("character_bible", [])
    if 0 <= character_index < len(characters):
        characters[character_index]["reference_image"] = entry["path"]
        characters[character_index]["reference_file_id"] = file_id
        characters[character_index]["reference_slot"] = character_index
    bump_asset_epoch(p)
    save_project(p)
    studio_event(
        "reference.attached", project_id=project_id, slot=character_index,
        character_name=character_name, relative_path=entry["path"],
        image_bytes=len(raw), cloud_file_ready=bool(file_id), replaced=bool(matching),
    )
    return p


def generate_character_reference(project_id: str, character_index: int) -> dict:
    p = load_project(project_id)
    characters = p.get("character_bible") or []
    if (character_index < 0
            or character_index >= min(MAX_CHARACTER_REFERENCES, len(characters))):
        raise ValueError("Choose one of the book's first four characters")
    character = characters[character_index]
    name = str(character.get("name") or "").strip()
    if not name:
        raise ValueError("Give this character a name first")
    prompt = f"""Create the master identity reference image for {name}, used to keep
this character visually identical across a children's picture book.

Character role: {character.get('role')}
Exact appearance: {character.get('appearance')}
Personality and mannerisms: {character.get('personality')}
Never-changing continuity rules: {character.get('continuity_rules')}
Creator/reference direction: {character.get('reference_image_prompt')}
Book illustration style: {style_label(p.get('settings') or {})}

Show exactly one character and no other people, animals or creatures. Full body,
head to feet fully visible, front three-quarter view, relaxed neutral pose, clear
face and distinctive details, soft even studio lighting, simple neutral seamless
background. No scenery, props unless they are a permanent signature item, panels,
turnaround sheet, collage, border, title, caption, letters, logo or watermark."""
    update_current_job(
        f"Waiting for Grok Imagine to create {name}'s reference", 20)
    payload = {
        "model": IMAGE_MODEL, "prompt": prompt[:6000], "n": 1,
        "resolution": "1k", "aspect_ratio": "3:4",
    }
    result = xai_json(XAI_BASE + "/images/generations", payload, timeout=600)
    items = result.get("data") or []
    if not items:
        raise RuntimeError("xAI returned no character reference image")
    update_current_job("Saving and attaching the character reference", 88)
    raw = download_image(items[0])
    try:
        from PIL import Image
        source = Image.open(io.BytesIO(raw)).convert("RGB")
        converted = io.BytesIO()
        source.save(converted, format="PNG", optimize=True)
        raw = converted.getvalue()
    except Exception:
        pass
    encoded = "data:image/png;base64," + base64.b64encode(raw).decode()
    p = upload_reference(project_id, {
        "name": f"{slug(name)}-master-reference.png",
        "label": f"Master identity and appearance reference for {name}",
        "character_name": name,
        "character_index": character_index,
        "data": encoded,
    })
    p.setdefault("history", []).append({
        "at": now(), "action": f"Generated and attached {name}'s master reference"})
    return save_project(p)


def trim_size(settings: dict) -> tuple[float, float]:
    row = next((x for x in TRIMS if x[0] == settings.get("trim")), TRIMS[0])
    return row[2], row[3]


def is_28_day_workbook(project: dict) -> bool:
    """Recognise the workbook layout even if an older save lost its template."""
    if project.get("series_template"):
        return True
    settings = project.get("settings") or {}
    if settings.get("book_type") != "educational":
        return False
    title = str(project.get("title") or "").strip()
    series_name = str(settings.get("series_name") or "").strip().casefold()
    return bool(
        re.match(r"^28\s+days?\s+(?:of|to)\b", title, re.IGNORECASE)
        or series_name == SERIES_NAME.casefold()
    )


def image_abs(p: dict, rel: str) -> Path | None:
    if not rel:
        return None
    path = project_dir(p["id"]) / rel
    return path if path.exists() else None


def export_pdf(p: dict, out: Path, include_cover=True, bleed=False,
               pad_even=False,
               trim_override: tuple[float, float] | None = None) -> None:
    from PIL import Image, ImageOps
    from reportlab.lib.colors import Color, white
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    w_in, h_in = trim_override or trim_size(p["settings"])
    render_w = w_in + (.125 if bleed else 0)
    render_h = h_in + (.25 if bleed else 0)
    page_w, page_h = render_w * 72, render_h * 72
    regular = "BookArial"
    bold = "BookArialBold"
    try:
        if regular not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(
                regular, "/System/Library/Fonts/Supplemental/Arial.ttf"))
        if bold not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(
                bold, "/System/Library/Fonts/Supplemental/Arial Bold.ttf"))
    except Exception:
        regular, bold = "Helvetica", "Helvetica-Bold"

    doc = canvas.Canvas(str(out), pagesize=(page_w, page_h),
                        pageCompression=1, pdfVersion=(1, 7))
    doc.setTitle(p.get("title") or "Kindle Book")
    doc.setAuthor(p.get("settings", {}).get("author") or "Author")

    def wrap(text: str, font_name: str, font_size: float,
             max_width: float, preserve_blank=False) -> list[str]:
        lines = []
        for paragraph in (text or "").split("\n") or [""]:
            if not paragraph.strip():
                if preserve_blank:
                    lines.append("")
                continue
            words = paragraph.split()
            line = ""
            for word in words:
                candidate = (line + " " + word).strip()
                if line and pdfmetrics.stringWidth(
                        candidate, font_name, font_size) > max_width:
                    lines.append(line)
                    line = word
                else:
                    line = candidate
            if line:
                lines.append(line)
        return lines

    def draw_background(path: Path | None, cover=False, context=False) -> None:
        if not path:
            doc.setFillColor(
                Color(.20, .14, .20) if cover else Color(.97, .96, .94))
            doc.rect(0, 0, page_w, page_h, fill=1, stroke=0)
            return
        target = (max(1, int(render_w * 300)), max(1, int(render_h * 300)))
        with Image.open(path) as source:
            source = ImageOps.exif_transpose(source).convert("RGB")
            prepared = ImageOps.fit(
                source, target, method=Image.Resampling.LANCZOS,
                centering=(.5, .5))
            # Narrative story copy stays in the lower third. Ancient Wisdom's
            # introduction and afterword are deliberately substantial context
            # pages, so give their copy a readable lower two-thirds panel.
            workbook_page = is_28_day_workbook(p) and not cover
            panel_h = int(target[1] * (
                .36 if cover else .66 if context else .79
                if workbook_page else .34))
            mask = Image.linear_gradient("L").resize(
                (target[0], panel_h), Image.Resampling.BICUBIC)
            if cover:
                mask = mask.point(lambda value: int(18 + value * .80))
                overlay = Image.new("RGB", (target[0], panel_h), (18, 8, 15))
            else:
                # A gentle light fade keeps the words readable without washing
                # out the artwork beneath the story panel.
                mask = mask.point(lambda value: int(
                    (20 if context or workbook_page else 8)
                    + value * (.78 if workbook_page else .70 if context else .62)))
                overlay = Image.new("RGB", (target[0], panel_h), "white")
            prepared.paste(
                overlay, (0, target[1] - panel_h), mask)
            buffer = io.BytesIO()
            # The 300 PPI page background is photographic artwork. Embedding it
            # as a maximum-quality JPEG keeps upload packages manageable while
            # headings and story copy remain crisp vector text in the PDF.
            prepared.save(
                buffer, "JPEG", quality=96, subsampling=0, optimize=True,
                dpi=(300, 300))
        buffer.seek(0)
        doc.drawImage(
            ImageReader(buffer), 0, 0, width=page_w, height=page_h,
            preserveAspectRatio=False, mask="auto")

    def alpha(value: float) -> None:
        try:
            doc.setFillAlpha(value)
        except Exception:
            pass

    def reset_alpha() -> None:
        try:
            doc.setFillAlpha(1)
        except Exception:
            pass

    def draw_heading(title: str, cover=False, quote="", quote_author="") -> None:
        size = 30 if cover else 24
        max_width = page_w * .82
        lines = wrap(title, bold, size, max_width)
        while len(lines) > 2 and size > 16:
            size -= 1
            lines = wrap(title, bold, size, max_width)
        line_h = size * 1.18
        quote_size = 13
        quote_lines = wrap(f'"{quote}"' if quote else "", regular,
                           quote_size, page_w * .76)
        while len(quote_lines) > 3 and quote_size > 9:
            quote_size -= 1
            quote_lines = wrap(f'"{quote}"', regular,
                               quote_size, page_w * .76)
        author_lines = wrap(
            "- " + quote_author if quote and quote_author else "",
            regular, max(9, quote_size - 1), page_w * .72)
        quote_h = (
            len(quote_lines) * quote_size * 1.22
            + len(author_lines) * max(9, quote_size - 1) * 1.18 + 12
            if quote_lines else 0
        )
        box_h = max(line_h * len(lines) + quote_h + 22, 50)
        box_x, box_w = page_w * .06, page_w * .88
        box_y = page_h - box_h - page_h * .035
        doc.setFillColor(Color(.08, .04, .07))
        alpha(.66)
        doc.roundRect(box_x, box_y, box_w, box_h, 12,
                      fill=1, stroke=0)
        reset_alpha()
        doc.setFillColor(white)
        doc.setFont(bold, size)
        y = box_y + box_h - 16 - size
        for line in lines:
            doc.drawCentredString(page_w / 2, y, line)
            y -= line_h
        if quote_lines:
            y -= 4
            doc.setFont(regular, quote_size)
            for line in quote_lines:
                doc.drawCentredString(page_w / 2, y, line)
                y -= quote_size * 1.22
            if author_lines:
                author_size = max(9, quote_size - 1)
                doc.setFont(regular, author_size)
                for line in author_lines:
                    doc.drawCentredString(page_w / 2, y, line)
                    y -= author_size * 1.18

    def draw_body(body: str, cover=False, context=False) -> None:
        workbook_page = is_28_day_workbook(p) and not cover
        panel_h = page_h * (
            .36 if cover else .66 if context else .79
            if workbook_page else .34)
        font_name = bold
        size = 24 if context else 26 if workbook_page else 20
        # The workbook needs enough width for practical prompts, but the text
        # still stays inside KDP's 0.375-inch safe area after outer-edge bleed.
        max_width = page_w * (.875 if workbook_page else .84)
        max_height = panel_h * (
            .965 if workbook_page else .88 if context else .78)
        lines = wrap(body, font_name, size, max_width, preserve_blank=True)
        minimum_size = 20 if workbook_page else 7
        leading = 1.08 if workbook_page else 1.24
        while (lines and len(lines) * size * leading > max_height
               and size > minimum_size):
            size -= .5
            lines = wrap(body, font_name, size, max_width, preserve_blank=True)
        line_h = size * leading
        total_h = len(lines) * line_h
        y = (panel_h + total_h) / 2 - size
        if lines:
            # KDP measures the text safe zone from the trim line, not from the
            # outer edge of a bleed-sized PDF. With bleed, the bottom 0.125"
            # will be cut away and KDP then requires another 0.375" margin.
            # Clamp the lowest baseline far enough inward to include glyph
            # descenders, while leaving the full-bleed artwork untouched.
            trim_bleed = .125 if bleed else 0.0
            outside_margin = .375 if bleed else .25
            safe_bottom = ((trim_bleed + outside_margin) * 72
                           + max(2.5, size * .24))
            last_baseline = y - (len(lines) - 1) * line_h
            if last_baseline < safe_bottom:
                y += safe_bottom - last_baseline
        doc.setFillColor(white if cover else Color(.08, .06, .08))
        doc.setFont(font_name, size)
        for line in lines:
            if line:
                doc.drawCentredString(page_w / 2, y, line)
            y -= line_h

    def compose(path: Path | None, title: str, body: str,
                cover=False, quote="", quote_author="", context=False) -> None:
        draw_background(path, cover, context)
        draw_heading(title, cover, quote, quote_author)
        draw_body(body, cover, context)
        doc.showPage()

    if include_cover:
        compose(image_abs(p, p["cover"].get("image")),
                p["cover"].get("title") or p["title"],
                p["cover"].get("subtitle", ""), True)
    for page in p["pages"]:
        body = page.get("text", "")
        if page.get("dialogue"):
            body += "\n" + " ".join(map(str, page["dialogue"]))
        context_page = (
            p.get("settings", {}).get("book_type") == "ancient_wisdom"
            and str(page.get("section") or "").strip().lower()
            in {"introduction", "afterword"}
        )
        compose(image_abs(p, page.get("image")),
                page.get("heading", ""), body, False,
                page.get("quote", ""), page.get("quote_author", ""),
                context_page)
    if pad_even and len(p.get("pages", [])) % 2:
        doc.showPage()
    doc.save()


def export_print_cover(p: dict, out: Path, page_count: int) -> None:
    from PIL import Image, ImageEnhance, ImageFilter, ImageOps
    from reportlab.lib.colors import Color, white
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    w_in, h_in = trim_size(p["settings"])
    interior = p.get("settings", {}).get("print_interior", "premium_colour")
    rate = .002347 if interior == "premium_colour" else .002252
    spine_in = page_count * rate
    bleed_in = .125
    full_w_in = bleed_in + w_in + spine_in + w_in + bleed_in
    full_h_in = bleed_in + h_in + bleed_in
    full_w, full_h = full_w_in * 72, full_h_in * 72
    bleed, trim_w, trim_h = bleed_in * 72, w_in * 72, h_in * 72
    back_left = bleed
    back_right = back_left + trim_w
    spine_left = back_right
    front_left = spine_left + spine_in * 72

    regular, bold = "CoverArial", "CoverArialBold"
    try:
        if regular not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(
                regular, "/System/Library/Fonts/Supplemental/Arial.ttf"))
        if bold not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(
                bold, "/System/Library/Fonts/Supplemental/Arial Bold.ttf"))
    except Exception:
        regular, bold = "Helvetica", "Helvetica-Bold"

    doc = canvas.Canvas(str(out), pagesize=(full_w, full_h),
                        pageCompression=1, pdfVersion=(1, 7))
    doc.setTitle((p.get("title") or "Book") + " - Print Cover")
    cover_path = image_abs(p, p.get("cover", {}).get("image"))
    target = (max(1, int(full_w_in * 300)), max(1, int(full_h_in * 300)))
    background = Image.new("RGB", target, (70, 46, 63))
    if cover_path:
        with Image.open(cover_path) as source:
            source = ImageOps.exif_transpose(source).convert("RGB")
            panel_size = (max(1, int((w_in + bleed_in) * 300)),
                          max(1, int(full_h_in * 300)))
            front = ImageOps.fit(
                source, panel_size, Image.Resampling.LANCZOS)
            front_x = int(front_left / 72 * 300)
            background.paste(front, (front_x, 0))
            back = ImageOps.fit(
                source, (int((w_in + bleed_in) * 300), target[1]),
                Image.Resampling.LANCZOS)
            back = ImageEnhance.Brightness(
                back.filter(ImageFilter.GaussianBlur(16))).enhance(.36)
            background.paste(back, (0, 0))
    buffer = io.BytesIO()
    background.save(
        buffer, "JPEG", quality=96, subsampling=0, optimize=True,
        dpi=(300, 300))
    buffer.seek(0)
    doc.drawImage(ImageReader(buffer), 0, 0, full_w, full_h,
                  preserveAspectRatio=False, mask="auto")

    def wrap(text: str, font: str, size: float,
             max_width: float) -> list[str]:
        lines, line = [], ""
        for word in (text or "").split():
            candidate = (line + " " + word).strip()
            if line and pdfmetrics.stringWidth(
                    candidate, font, size) > max_width:
                lines.append(line)
                line = word
            else:
                line = candidate
        if line:
            lines.append(line)
        return lines

    def alpha(value: float) -> None:
        try:
            doc.setFillAlpha(value)
        except Exception:
            pass

    def reset_alpha() -> None:
        try:
            doc.setFillAlpha(1)
        except Exception:
            pass

    # Front-cover typography.
    title = p.get("cover", {}).get("title") or p.get("title") or ""
    subtitle = p.get("cover", {}).get("subtitle") or p.get("subtitle") or ""
    box_x = front_left + trim_w * .07
    box_w = trim_w * .86
    title_size = 28
    title_lines = wrap(title, bold, title_size, box_w - 24)
    while len(title_lines) > 3 and title_size > 18:
        title_size -= 1
        title_lines = wrap(title, bold, title_size, box_w - 24)
    title_h = len(title_lines) * title_size * 1.16 + 24
    title_y = full_h - bleed - title_h - 24
    doc.setFillColor(Color(.08, .04, .07))
    alpha(.68)
    doc.roundRect(box_x, title_y, box_w, title_h, 12, fill=1, stroke=0)
    reset_alpha()
    doc.setFillColor(white)
    doc.setFont(bold, title_size)
    y = title_y + title_h - title_size - 12
    for line in title_lines:
        doc.drawCentredString(front_left + trim_w / 2, y, line)
        y -= title_size * 1.16
    if subtitle:
        sub_size = 16
        sub_lines = wrap(subtitle, bold, sub_size, box_w - 20)
        sub_h = len(sub_lines) * sub_size * 1.18 + 20
        sub_y = bleed + 28
        doc.setFillColor(Color(.08, .04, .07))
        alpha(.72)
        doc.roundRect(box_x, sub_y, box_w, sub_h, 10, fill=1, stroke=0)
        reset_alpha()
        doc.setFillColor(white)
        doc.setFont(bold, sub_size)
        y = sub_y + sub_h - sub_size - 9
        for line in sub_lines:
            doc.drawCentredString(front_left + trim_w / 2, y, line)
            y -= sub_size * 1.18

    # Back-cover copy and reserved barcode area.
    safe_x = back_left + 28
    safe_w = trim_w - 56
    safe_top = full_h - bleed - 34
    doc.setFillColor(white)
    doc.setFont(bold, 21)
    doc.drawCentredString(back_left + trim_w / 2, safe_top,
                          p.get("title") or "")
    description = str(p.get("metadata", {}).get("description") or
                      p.get("story_summary") or "")
    body_size = 13
    body_lines = wrap(description, regular, body_size, safe_w)
    max_lines = max(1, int((trim_h - 190) / (body_size * 1.35)))
    while len(body_lines) > max_lines and body_size > 9:
        body_size -= .5
        body_lines = wrap(description, regular, body_size, safe_w)
        max_lines = max(1, int((trim_h - 190) / (body_size * 1.35)))
    doc.setFont(regular, body_size)
    y = safe_top - 42
    for line in body_lines[:max_lines]:
        doc.drawString(safe_x, y, line)
        y -= body_size * 1.35
    author = p.get("settings", {}).get("author") or ""
    if author:
        doc.setFont(bold, 13)
        doc.drawString(safe_x, bleed + 118, "By " + author)
    barcode_w, barcode_h = 2 * 72, 1.2 * 72
    barcode_x = back_right - barcode_w - 18
    barcode_y = bleed + 18
    doc.setFillColor(white)
    doc.roundRect(barcode_x, barcode_y, barcode_w, barcode_h,
                  4, fill=1, stroke=0)
    doc.setFillColor(Color(.38, .38, .38))
    doc.setFont(regular, 8)
    doc.drawCentredString(
        barcode_x + barcode_w / 2, barcode_y + barcode_h / 2,
        "Reserved for Amazon barcode")

    if page_count >= 79 and spine_in * 72 >= 18:
        doc.saveState()
        doc.translate(spine_left + spine_in * 36, full_h / 2)
        doc.rotate(90)
        spine_title = (p.get("title") or "")[:80]
        size = min(12, max(7, spine_in * 72 * .45))
        doc.setFillColor(white)
        doc.setFont(bold, size)
        doc.drawCentredString(0, -size / 3, spine_title)
        doc.restoreState()
    doc.save()


def export_docx(p: dict, out: Path) -> None:
    from docx import Document
    doc = Document()
    doc.add_heading(p["title"], 0)
    if p.get("subtitle"):
        doc.add_paragraph(p["subtitle"])
    doc.add_paragraph("By " + (p["settings"].get("author") or "Author"))
    doc.add_page_break()
    for page in p["pages"]:
        doc.add_heading(page.get("heading") or f"Page {page['number']}", level=1)
        if page.get("quote"):
            quote = doc.add_paragraph()
            quote_run = quote.add_run(f'"{page["quote"]}"')
            quote_run.italic = True
            if page.get("quote_author"):
                quote.add_run(f'\n- {page["quote_author"]}')
        if page.get("image"):
            img = image_abs(p, page["image"])
            if img:
                doc.add_picture(str(img), width=None)
        doc.add_paragraph(page.get("text", ""))
        for dialogue in page.get("dialogue") or []:
            doc.add_paragraph(str(dialogue))
    doc.save(out)


def xhtml_page(title: str, text: str, image_name: str | None,
               css_class="book-page", viewport: tuple[int, int] | None = None,
               body_size: int | None = None, quote: str = "",
               quote_author: str = "") -> str:
    image = (f'<img src="../images/{html.escape(image_name)}" alt="Illustration"/>'
             if image_name else "")
    paras = "".join(
        f"<p>{html.escape(line)}</p>" if line.strip()
        else '<p class="blank">&#160;</p>'
        for line in (text or "").split("\n")
    )
    viewport_meta = (
        f'<meta name="viewport" content="width={viewport[0]},height={viewport[1]}"/>'
        if viewport else "")
    body_style = f' style="--body-size:{body_size}px"' if body_size else ""
    quote_html = ""
    if quote:
        author = (f'<cite>- {html.escape(quote_author)}</cite>'
                  if quote_author else "")
        quote_html = (
            f'<blockquote class="page-quote"><p>"{html.escape(quote)}"</p>'
            f'{author}</blockquote>')
    return f"""<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml"><head>
<title>{html.escape(title)}</title>{viewport_meta}
<link rel="stylesheet" href="../styles/book.css" type="text/css"/>
</head><body class="{css_class}"><main>{image}<section class="copy">
<h1>{html.escape(title)}</h1>{quote_html}<div class="body-copy"{body_style}>{paras}</div>
</section></main></body></html>"""


def fixed_epub_body_size(text: str, viewport: tuple[int, int],
                         cover=False, context=False, workbook=False) -> int:
    max_size = 34 if cover else 42 if context else 48 if workbook else 38
    min_size = 40 if workbook else 14
    available_width = viewport[0] * (.875 if workbook else .86)
    available_height = viewport[1] * (
        .30 if cover else .66 if context else .79 if workbook else .34) - 48
    leading = 1.08 if workbook else 1.18
    logical_lines = (text or "").split("\n")
    for size in range(max_size, min_size - 1, -1):
        chars_per_line = max(8, int(available_width / (size * .54)))
        visual_lines = sum(
            1 if not line.strip() else
            max(1, math.ceil(len(line.expandtabs(4)) / chars_per_line))
            for line in logical_lines
        )
        if visual_lines * size * leading <= available_height:
            return size
    return min_size


def export_epub(p: dict, out: Path) -> None:
    fixed = p["settings"].get("layout") == "fixed"
    w_in, h_in = trim_size(p["settings"])
    viewport = (int(w_in * 150), int(h_in * 150)) if fixed else None
    entries: dict[str, bytes] = {}
    entries["META-INF/container.xml"] = b"""<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>"""
    if fixed:
        entries["OEBPS/styles/book.css"] = b"""html,body{margin:0;width:100%;height:100%;overflow:hidden}
body{font-family:serif;color:#111;background:#332334}main{position:relative;width:100vw;height:100vh;
box-sizing:border-box;overflow:hidden}main>img{position:absolute;inset:0;display:block;width:100%;
height:100%;object-fit:cover}.copy{position:absolute;inset:0;display:grid;grid-template-rows:auto auto 1fr;
box-sizing:border-box}.copy h1{align-self:start;text-align:center;font-size:38px;line-height:1.12;
margin:28px 5% 0;padding:14px 24px;border-radius:16px;color:#fff;background:rgba(20,12,18,.62);
text-shadow:0 2px 8px #000}.page-quote{margin:8px 9% 0;padding:10px 18px;border-radius:14px;
text-align:center;color:#fff;background:rgba(20,12,18,.62);text-shadow:0 2px 7px #000}
.page-quote p{font-family:Georgia,serif;font-size:22px;font-style:italic;font-weight:500;line-height:1.18;
margin:0;color:#fff}.page-quote cite{display:block;margin-top:6px;font-size:16px;font-style:normal}
.body-copy{align-self:end;height:34%;display:flex;flex-direction:column;
align-items:center;justify-content:center;padding:22px 7%;box-sizing:border-box;
overflow:hidden;background:linear-gradient(transparent 0%,rgba(255,255,255,.42) 24%,rgba(255,255,255,.68) 100%)}
p{text-align:center;font-size:var(--body-size,38px);font-weight:600;line-height:1.18;margin:0;color:#171219;
text-shadow:0 1px 1px rgba(255,255,255,.8)}p.blank{min-height:1.18em}
.context-page .body-copy{height:66%;padding:26px 7%;
background:linear-gradient(transparent 0%,rgba(255,255,255,.60) 16%,rgba(255,255,255,.90) 100%)}
.workbook-page .body-copy{height:79%;padding:24px 6.25%;
background:linear-gradient(transparent 0%,rgba(255,255,255,.64) 15%,rgba(255,255,255,.88) 100%)}
.workbook-page .body-copy p{line-height:1.08}.workbook-page .body-copy p.blank{min-height:1.08em}
.cover .copy{grid-template-rows:auto 1fr}.cover .copy h1{font-size:52px;margin-top:42px;
background:rgba(20,10,18,.68)}.cover .body-copy{height:36%;padding-bottom:44px;
background:linear-gradient(transparent,rgba(20,10,18,.86))}.cover p{font-size:var(--body-size,34px);color:#fff;
text-shadow:0 2px 8px #000}"""
    else:
        entries["OEBPS/styles/book.css"] = b"""body{margin:0;font-family:serif;color:#111;background:#fff}
main{padding:4%;box-sizing:border-box}img{display:block;width:100%;height:auto;margin:0 auto 1rem}
.copy{max-width:48em;margin:auto}h1{text-align:center;font-size:1.5em}
p{font-size:1.1em;line-height:1.45;margin:.5em 0}.page-quote{text-align:center;font-style:italic;
margin:1em auto;padding:.7em 1em;border-left:3px solid #bbb}.page-quote p{margin:0}.page-quote cite{display:block;
margin-top:.45em;font-style:normal;font-size:.9em}.cover main{padding:0}
.blank{min-height:1.45em}.cover .copy{text-align:center;padding:0 5% 5%}"""
    manifest = [
        '<item id="css" href="styles/book.css" media-type="text/css"/>',
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>',
    ]
    spine = []
    nav = []
    images_seen = {}

    def add_image(rel):
        if not rel:
            return None
        src = image_abs(p, rel)
        if not src:
            return None
        name = src.name
        if name not in images_seen:
            data = src.read_bytes()
            try:
                from PIL import Image
                with Image.open(io.BytesIO(data)) as im:
                    buf = io.BytesIO()
                    im.convert("RGB").save(buf, "JPEG", quality=92, optimize=True)
                    data = buf.getvalue()
                name = src.stem + ".jpg"
            except Exception:
                pass
            entries["OEBPS/images/" + name] = data
            iid = "img" + str(len(images_seen) + 1)
            images_seen[src.name] = (iid, name)
            manifest.append(f'<item id="{iid}" href="images/{name}" media-type="image/jpeg"/>')
        return images_seen[src.name][1]

    cover_img = add_image(p["cover"].get("image"))
    entries["OEBPS/text/cover.xhtml"] = xhtml_page(
        p["cover"].get("title") or p["title"],
        p["cover"].get("subtitle", ""), cover_img, "cover", viewport,
        fixed_epub_body_size(
            p["cover"].get("subtitle", ""), viewport, cover=True
        ) if fixed else None).encode()
    cover_image_id = ""
    if cover_img:
        for i, item in enumerate(manifest):
            if f'href="images/{cover_img}"' in item:
                cover_image_id = re.search(r'id="([^"]+)"', item).group(1)
                manifest[i] = item.replace("/>", ' properties="cover-image"/>')
                break
    manifest.append('<item id="cover" href="text/cover.xhtml" media-type="application/xhtml+xml"/>')
    spine.append('<itemref idref="cover"/>')
    nav.append('<li><a href="text/cover.xhtml">Cover</a></li>')
    workbook_page = is_28_day_workbook(p)
    for page in p["pages"]:
        pid = f"p{page['number']}"
        fname = f"page-{page['number']:03d}.xhtml"
        img = add_image(page.get("image"))
        body = page.get("text", "")
        if page.get("dialogue"):
            body += "\n" + "\n".join(map(str, page["dialogue"]))
        context_page = (
            p.get("settings", {}).get("book_type") == "ancient_wisdom"
            and str(page.get("section") or "").strip().lower()
            in {"introduction", "afterword"}
        )
        entries["OEBPS/text/" + fname] = xhtml_page(
            page.get("heading") or f"Page {page['number']}", body, img,
            ("book-page context-page" if context_page else
             "book-page workbook-page" if workbook_page else "book-page"), viewport,
            fixed_epub_body_size(
                body, viewport, context=context_page, workbook=workbook_page)
            if fixed else None,
            page.get("quote", ""), page.get("quote_author", "")).encode()
        manifest.append(f'<item id="{pid}" href="text/{fname}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="{pid}"/>')
        nav.append(f'<li><a href="text/{fname}">Page {page["number"]}</a></li>')
    entries["OEBPS/nav.xhtml"] = ("""<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml"><head><title>Contents</title></head>
<body><nav epub:type="toc" xmlns:epub="http://www.idpf.org/2007/ops"><h1>Contents</h1><ol>"""
        + "".join(nav) + "</ol></nav></body></html>").encode()
    ident = "urn:uuid:" + uuid.uuid5(uuid.NAMESPACE_URL, p["id"]).hex
    if fixed:
        orientation = "landscape" if w_in > h_in else "portrait"
        rendition = (
            '<meta property="rendition:layout">pre-paginated</meta>'
            '<meta property="rendition:spread">none</meta>'
            f'<meta property="rendition:orientation">{orientation}</meta>'
        )
    else:
        rendition = '<meta property="rendition:layout">reflowable</meta>'
    legacy_cover = (f'<meta name="cover" content="{cover_image_id}"/>'
                    if cover_image_id else "")
    entries["OEBPS/content.opf"] = f"""<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
<dc:identifier id="bookid">{ident}</dc:identifier><dc:title>{html.escape(p["title"])}</dc:title>
<dc:language>en</dc:language><dc:creator>{html.escape(p["settings"].get("author") or "Author")}</dc:creator>
<dc:date>{datetime.now().date().isoformat()}</dc:date>{legacy_cover}{rendition}</metadata>
<manifest>{''.join(manifest)}</manifest><spine page-progression-direction="ltr">{''.join(spine)}</spine>
</package>""".encode()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        for name, data in entries.items():
            z.writestr(name, data, compress_type=zipfile.ZIP_DEFLATED)


def choose_book_narrator(project_id: str) -> dict:
    project = load_project(project_id)
    if not project.get("pages"):
        raise ValueError("Generate the story before asking Grok to cast its narrator")
    update_current_job("Grok is casting the narrator for this exact book", 20)
    context = {
        "title": project.get("title"), "subtitle": project.get("subtitle"),
        "summary": project.get("story_summary"),
        "audience": (project.get("settings") or {}).get("audience"),
        "genre": (project.get("settings") or {}).get("genre"),
        "tone": (project.get("settings") or {}).get("tone"),
        "language": (project.get("settings") or {}).get("language"),
        "marketplace": (project.get("settings") or {}).get("primary_marketplace"),
        "sample_pages": [
            {"heading": page.get("heading"), "text": page.get("text")}
            for page in (project.get("pages") or [])[:6]
        ],
    }
    narrator = grok_structured(
        """Act as a children's audiobook casting director. Choose the single best
Qwen3-TTS native-English stable narrator for this finished book. Decide British or American and female or
male from the setting, language, target age, emotional tone, humour and read-aloud
rhythm. Do not default mechanically. Choose exactly one valid speaker whose accent
and gender match the request: Eleanor or Poppy (British female), Maya or Harper
(American female), Arthur or Oliver (British male), or Noah or Jack (American
male). Choose warm or lively to fit the book. Set a natural expressive speed
from 0.86 to 1.05 and give a concise useful performance direction.\n\nBook:\n"""
        + json.dumps(context, ensure_ascii=False),
        {"accent": "british or american", "gender": "female or male",
         "voice": "One exact valid voice id", "delivery": "Performance direction",
         "speed": 0.92, "reason": "Why this casting fits this specific book"},
        timeout=300)
    narrator["selected_by"] = "grok"
    video_audio = dict(project.get("video_audio") or {})
    video_audio["narrator"] = narrator
    # A new voice invalidates prior rendered media but never deletes it.
    video_audio["outputs"] = {}
    project["video_audio"] = normalize_video_audio(video_audio, project)
    project.setdefault("history", []).append({
        "at": now(), "action": "Grok selected the video and audiobook narrator"})
    update_current_job("Narrator casting is ready", 95)
    return save_project(project)


def _tts_json(path: str, payload: dict | None = None, timeout: int = 900) -> dict:
    url = TTS_BASE + path
    body = None if payload is None else json.dumps(payload).encode("utf-8")
    request = urllib.request.Request(
        url, data=body,
        headers={"Content-Type": "application/json"} if body is not None else {},
        method="POST" if body is not None else "GET")
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            data = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", "replace")[:500]
        raise RuntimeError(f"Local TTS service rejected the request: {detail}") from exc
    except Exception as exc:
        raise RuntimeError(
            "Local TTS Studio is unavailable at " + TTS_BASE
            + ". Start TTS Studio and try again: " + str(exc)) from exc
    if isinstance(data, dict) and data.get("error"):
        raise RuntimeError("Local TTS failed: " + str(data["error"]))
    return data


def _orchestrator_json(path: str, payload: dict | None = None,
                       timeout: int = 1200) -> dict:
    """Call the local oMLX Orchestrator with useful studio-facing errors."""
    url = ORCHESTRATOR_BASE + path
    body = None if payload is None else json.dumps(payload).encode("utf-8")
    request = urllib.request.Request(
        url, data=body,
        headers={"Content-Type": "application/json"} if body is not None else {},
        method="POST" if body is not None else "GET")
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            data = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", "replace")[:900]
        raise RuntimeError(
            "oMLX Orchestrator could not create the Qwen direction: " + detail
        ) from exc
    except Exception as exc:
        raise RuntimeError(
            "oMLX Orchestrator is unavailable at " + ORCHESTRATOR_BASE
            + ". Start it and try again: " + str(exc)) from exc
    if not isinstance(data, dict):
        raise RuntimeError("oMLX Orchestrator returned an invalid response")
    if data.get("error"):
        raise RuntimeError("Qwen direction failed: " + str(data["error"]))
    return data


def _run_media_command(command: list[str], label: str) -> None:
    if command and command[0] in {"ffmpeg", "ffprobe", "blender"}:
        command = [_media_tool(command[0]), *command[1:]]
    process = subprocess.run(command, capture_output=True, text=True)
    if process.returncode:
        detail = (process.stderr or process.stdout or "unknown error")[-1400:]
        raise RuntimeError(f"{label} failed: {detail}")


def _media_tool(name: str) -> str:
    """Resolve Homebrew media tools when launchd supplies a minimal PATH."""
    env_name = "OMLX_" + name.upper()
    candidates = [
        os.environ.get(env_name, ""), shutil.which(name) or "",
        f"/opt/homebrew/bin/{name}", f"/usr/local/bin/{name}",
        f"/usr/bin/{name}",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file() and os.access(candidate, os.X_OK):
            return candidate
    raise FileNotFoundError(
        f"{name} is required for Video & Audio but could not be found. "
        f"Expected it at /opt/homebrew/bin/{name}.")


def _media_seconds(path: Path) -> float:
    process = subprocess.run(
        [_media_tool("ffprobe"), "-v", "error", "-show_entries", "format=duration",
         "-of", "default=noprint_wrappers=1:nokey=1", str(path)],
        capture_output=True, text=True)
    if process.returncode:
        raise RuntimeError("Could not measure narration: " + process.stderr[-500:])
    return max(.1, float(process.stdout.strip()))


def _raw_book_video_scenes(project: dict) -> list[dict]:
    author = str((project.get("settings") or {}).get("author") or "").strip()
    presenter_names = _presenter_character_names(project)
    cover = project.get("cover") or {}
    cover_words = ". ".join(filter(None, [
        str(cover.get("title") or project.get("title") or ""),
        str(cover.get("subtitle") or project.get("subtitle") or ""),
        "Written by " + author if author else "",
    ]))
    scenes = [{
        "id": "cover", "heading": str(cover.get("title") or project.get("title") or ""),
        "text": str(cover.get("subtitle") or project.get("subtitle") or ""),
        "narration": cover_words, "image": cover.get("image"),
        "image_prompt": str(cover.get("image_prompt") or ""),
    }]
    for page in project.get("pages") or []:
        dialogue = page.get("dialogue") or []
        words = _strip_empty_presenter_opening(
            str(page.get("text") or ""), presenter_names)
        if dialogue:
            words += ("\n\n" if words else "") + "\n".join(map(str, dialogue))
        heading = str(page.get("heading") or f"Page {page.get('number')}").strip()
        narration = (heading + ".\n" + words).strip() if heading else words
        scenes.append({
            "id": f"page-{page.get('number')}", "heading": heading,
            "text": words, "narration": narration, "image": page.get("image"),
            "image_prompt": str(page.get("image_prompt") or ""),
        })
    return scenes


def _strip_narration_production_language(text: str) -> str:
    """Remove book-making commentary from words intended to be spoken aloud."""
    production = re.compile(
        r"\b(?:claymation|illustrat(?:ion|ions|ed|ive)|image prompts?|"
        r"visual style|render(?:ed|ing)?|generated (?:image|picture|art)|"
        r"page layout|video production|book design|publishing workflow|"
        r"ai[- ]generated)\b", re.I)
    pieces = [part.strip() for part in re.findall(
        r".+?(?:[.!?…]+(?=\s|$)|$)", re.sub(r"\s+", " ", str(text or "")))
              if part.strip()]
    kept = [piece for piece in pieces if not production.search(piece)]
    return " ".join(kept).strip()


def _fallback_spoken_segments(text: str) -> list[dict]:
    """Make short performance beats if the saved Grok script is incomplete."""
    clean = _strip_narration_production_language(text)
    sentences = [part.strip() for part in re.findall(
        r".+?(?:[.!?…]+(?=\s|$)|$)", clean) if part.strip()]
    segments = []
    for sentence in sentences:
        words = sentence.split()
        while len(words) > 24:
            split_at = min(20, len(words))
            # Prefer a natural clause boundary near the target length.
            for index in range(min(22, len(words) - 1), 8, -1):
                if words[index - 1].endswith((",", ";", ":", "—")):
                    split_at = index
                    break
            piece, words = words[:split_at], words[split_at:]
            segments.append({"text": " ".join(piece), "pause_after": .28})
        if words:
            ending = " ".join(words)
            pause = .58 if ending.endswith(("?", "!")) else .42
            segments.append({"text": ending, "pause_after": pause})
    return segments or ([{"text": clean, "pause_after": .42}] if clean else [])


def _narration_script_signature(project: dict) -> str:
    scenes = _raw_book_video_scenes(project)
    payload = {
        "title": project.get("title"),
        "summary": project.get("story_summary"),
        "settings": {key: (project.get("settings") or {}).get(key) for key in (
            "book_type", "genre", "audience", "language", "tone")},
        "scenes": [{"id": row["id"], "heading": row["heading"],
                    "text": row["text"]} for row in scenes],
        "version": NARRATION_SCRIPT_VERSION,
    }
    return hashlib.sha256(json.dumps(
        payload, sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:24]


def _ensure_narration_script(project: dict) -> dict:
    """Create a subject-expert spoken adaptation with deliberate pause beats."""
    raw_scenes = _raw_book_video_scenes(project)
    presenter_names = _presenter_character_names(project) + [
        "The narrator", "The guide", "The teacher", "The presenter"]
    expected_ids = [row["id"] for row in raw_scenes if row["id"] != "cover"]
    signature = _narration_script_signature(project)
    video_audio = normalize_video_audio(project.get("video_audio"), project)
    script = video_audio.get("script") or {}
    script_ids = [row.get("id") for row in script.get("scenes") or []]
    if script.get("signature") == signature and script_ids == expected_ids:
        return script

    settings = project.get("settings") or {}
    update_current_job(
        "Grok is adapting the book into a natural expert narration", 4)
    result = grok_structured(
        f"""Act as a senior audiobook script editor and an accomplished human
practitioner-teacher in the exact subject below. Adapt each supplied page into
natural spoken English for the audio/video edition.

Choose the most credible expert voice for this subject. For meditation, write
with the lived understanding, quiet authority and practical precision of a Zen
and meditation master with decades of personal practice and teaching experience,
while remaining faithful to any specifically named source tradition. For yoga,
use an equally seasoned yoga teacher and scholar. For other fields, use the
equivalent deeply experienced expert. For fiction and picture books, use the
instincts of a master oral storyteller rather than turning the story into a
lesson or lecture; preserve character voice, dialogue, humour and emotional
truth. Never claim qualifications or say that the
narrator personally holds a title; demonstrate expertise through the writing.

MANDATORY CONTENT RULES:
- Stay entirely with the subject, story, teaching or action on that exact page.
- Preserve every essential fact, causal link, safety qualification and source
  distinction. Do not invent research, quotations, experiences or credentials.
- Do not talk about the book as a product or describe its pictures. Never mention
  claymation, illustrations, visual style, rendering, image generation, prompts,
  pages, layout, video production, AI, metadata, publishing or how it was made.
- Do not say “in this book”, “on this page”, “the reader will”, “we will explore”,
  “let us delve into”, or use formulaic recaps and generic motivational filler.
- Begin with the actual story, fact, instruction or insight. Never speak a stage
  direction, narrator tag or presenter cue. Do not describe how a recurring guide
  gazes, speaks, continues, nods, smiles, pauses, sits, stands, breathes, rests
  their hands or delivers an explanation. Do not use a person's name merely to
  wrap the next piece of exposition. Retain physical action only when it changes
  the plot, demonstrates an instruction, interacts with an important object or is
  otherwise necessary to understand the content.
- Make it sound genuinely spoken by a thoughtful expert: varied sentence length,
  contractions where natural, precise vocabulary, warm restraint and concrete
  explanation. Avoid repetitive openings and conclusions.
- Do not repeat the supplied heading inside the segment text; the heading is
  narrated separately.
- Keep each scene close to the original information and approximate spoken
  length. This is an expert spoken adaptation, not a summary or advertisement.

PERFORMANCE AND PAUSE RULES:
- Divide each scene into short thought-complete performance segments, normally
  8–22 spoken words and never more than 28. Do not split a name or phrase merely
  to hit a count.
- Give every segment a pause_after in seconds. Use 0.16–0.30 for a continuing
  clause, 0.32–0.55 after a complete sentence, 0.60–0.90 for reflection or an
  emotional turn, and 0.95–1.35 only for a major revelation or section close.
- Vary pauses purposefully. Punctuation and pause duration must agree. Avoid a
  monotonous identical pause after every sentence and avoid theatrical overuse
  of ellipses or em dashes.

Return every input scene exactly once and in the identical id order.

{HUMAN_EDITORIAL_STANDARD}

Book context:\n""" + json.dumps({
            "title": project.get("title"),
            "subtitle": project.get("subtitle"),
            "summary": project.get("story_summary"),
            "book_type": settings.get("book_type"),
            "genre": settings.get("genre"),
            "audience": settings.get("audience"),
            "tone": settings.get("tone"),
            "language": settings.get("language"),
            "source_fidelity_notes": (project.get("quality_report") or {}).get("notes"),
            "scenes": [{"id": row["id"], "heading": row["heading"],
                        "text": row["text"]}
                       for row in raw_scenes if row["id"] != "cover"],
        }, ensure_ascii=False),
        {
            "expert_voice": "The exact experienced human expert voice chosen",
            "approach": "Concise explanation of the natural spoken approach",
            "scenes": [{
                "id": "Exact supplied scene id",
                "segments": [{
                    "text": "One natural thought-complete spoken segment",
                    "pause_after": 0.42,
                }],
            }],
        }, timeout=900)

    returned = result.get("scenes") or []
    by_id = {str(row.get("id") or ""): row for row in returned
             if isinstance(row, dict)}
    scenes = []
    for raw_scene in raw_scenes:
        if raw_scene["id"] == "cover":
            continue
        source_row = by_id.get(raw_scene["id"]) or {}
        cleaned = []
        for segment in source_row.get("segments") or []:
            if not isinstance(segment, dict):
                continue
            text = _strip_narration_production_language(segment.get("text") or "")
            text = _strip_leading_presenter_direction(text, presenter_names)
            if not text:
                continue
            try:
                pause = max(.12, min(1.6, float(
                    segment.get("pause_after") or .4)))
            except (TypeError, ValueError):
                pause = .4
            # An unexpectedly long segment loses the regular sync anchors, so
            # split it safely even if Grok overlooked the word-count rule.
            if len(text.split()) > 28:
                cleaned.extend(_fallback_spoken_segments(text))
            else:
                cleaned.append({"text": text, "pause_after": round(pause, 2)})
        if not cleaned:
            cleaned = _fallback_spoken_segments(raw_scene["text"])
        scenes.append({"id": raw_scene["id"], "segments": cleaned[:16]})

    script = {
        "signature": signature, "model": TEXT_MODEL, "created_at": now(),
        "expert_voice": str(result.get("expert_voice") or
                            "Experienced subject-matter teacher")[:800],
        "approach": str(result.get("approach") or
                        "Natural, precise spoken explanation with deliberate pauses")[:1200],
        "scenes": scenes,
    }
    video_audio["script"] = script
    video_audio["outputs"] = {}
    # Spoken wording changes scene direction and sound-effect placement.
    video_audio["direction"] = {}
    video_audio["sound"]["plan"] = {}
    project["video_audio"] = normalize_video_audio(video_audio, project)
    project.setdefault("history", []).append({
        "at": now(), "action": "Grok created a natural expert narration with timed pauses"})
    save_project(project)
    return project["video_audio"]["script"]


def _book_video_scenes(project: dict) -> list[dict]:
    scenes = _raw_book_video_scenes(project)
    video_audio = normalize_video_audio(project.get("video_audio"), project)
    script_by_id = {row["id"]: row for row in (
        video_audio.get("script") or {}).get("scenes") or []}
    for scene in scenes:
        if scene["id"] == "cover":
            cover_parts = [part.strip() for part in re.split(
                r"\.\s+", scene["narration"]) if part.strip()]
            scene["segments"] = [
                {"text": part + ("." if not part.endswith((".", "!", "?")) else ""),
                 "pause_after": .48 if index < len(cover_parts) - 1 else .7,
                 "kind": "heading"}
                for index, part in enumerate(cover_parts)]
            continue
        saved = script_by_id.get(scene["id"])
        if not saved:
            continue
        body = [dict(segment, kind="body") for segment in saved.get("segments") or []]
        # Page/section headings remain visible in the book and video, but are
        # not spoken aloud. Reading every heading sounded artificial and broke
        # the flow between illustrated scenes.
        scene["segments"] = body
        scene["text"] = " ".join(segment["text"] for segment in body).strip()
        scene["narration"] = scene["text"]
    return scenes


def _qwen3_narration_instruction(project: dict, narrator: dict,
                                 scene: dict, segment: dict) -> str:
    """Create an audible-performance direction, never reader-facing prose."""
    accent = "American" if narrator.get("accent") == "american" else "British"
    audience = str((project.get("settings") or {}).get("audience") or
                   "young listeners")
    delivery = re.sub(r"\s+", " ", str(
        narrator.get("delivery") or
        "Warm, expressive, natural children's audiobook narration")).strip()
    kind = str(segment.get("kind") or "body")
    moment = (
        "Give this brief heading a clear, inviting page-turn lift without "
        "announcing that it is a heading."
        if kind == "heading" else
        "Let the emotion, energy and emphasis follow the meaning of this exact "
        "story moment; vary the melody and rhythm naturally."
    )
    return (
        f"{delivery}. Perform as a professional {accent} English "
        f"{narrator.get('gender', 'female')} children's audiobook storyteller "
        f"for {audience}. {moment} Use meaningful phrase-level pauses, crisp "
        "diction and warm human timing. Observe every punctuation mark. Never "
        "sound monotone, robotic, breathless or overacted. Speak only the exact "
        "supplied words: do not add introductions, stage directions, commentary "
        "or production language."
    )[:1800]


def _ensure_book_audio(project: dict, preview_pages: int = 0,
                       require_images: bool = True) -> tuple[list[dict], Path]:
    _ensure_narration_script(project)
    scenes = _book_video_scenes(project)
    for original_index, scene in enumerate(scenes):
        scene["audio_index"] = original_index
    if preview_pages:
        scenes = [scene for scene in scenes
                  if str(scene.get("id") or "").startswith("page-")][
                      :max(1, preview_pages)]
    missing = ([scene["id"] for scene in scenes if not image_abs(
        project, str(scene.get("image") or ""))] if require_images else [])
    if require_images and missing:
        raise ValueError(
            "Create the missing book pictures first: " + ", ".join(missing[:10]))
    narrator = normalize_video_audio(project.get("video_audio"), project)["narrator"]
    narration_engine = str(narrator.get("engine") or NARRATION_TTS_ENGINE)
    narration_seed = int(hashlib.sha256(
        f"{project.get('id')}:{narrator.get('voice')}".encode()
    ).hexdigest()[:8], 16)
    media_dir = project_dir(project["id"]) / "media"
    audio_dir = media_dir / "audio"
    audio_dir.mkdir(parents=True, exist_ok=True)
    _tts_json("/health", timeout=15)

    def silence_for(seconds: float) -> Path:
        milliseconds = max(120, min(1600, round(seconds * 1000)))
        path = audio_dir / f"pause-{milliseconds:04d}ms.wav"
        if not path.exists():
            _run_media_command([
                "ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i",
                "anullsrc=r=48000:cl=stereo", "-t", f"{milliseconds / 1000:.3f}",
                "-c:a", "pcm_s16le", str(path)], "Creating a narration pause")
        return path

    def collect_tts_file(source_url: str, destination: Path) -> None:
        if not source_url.startswith("/"):
            raise RuntimeError("Qwen3-TTS did not return an audio file")
        partial = destination.with_suffix(destination.suffix + ".partial")
        try:
            with urllib.request.urlopen(
                    TTS_BASE + source_url, timeout=120) as response:
                partial.write_bytes(response.read())
            if not partial.exists() or partial.stat().st_size < 1024:
                raise RuntimeError("Qwen3-TTS returned an empty audio file")
            partial.replace(destination)
        finally:
            partial.unlink(missing_ok=True)

    def generate_one(item: dict) -> None:
        result = _tts_json("/generate", {
            "engine": narration_engine, "text": item["text"],
            "voice": narrator["voice"], "speed": narrator["speed"],
            "instruct": item["instruct"], "seed": narration_seed,
            "temperature": 0.85, "top_p": 0.95, "top_k": 50,
            "format": "wav", "sample_rate": 48000, "stereo": True,
            "loudness": "youtube", "trim": True, "takes": 1,
        })
        collect_tts_file(str(result.get("url") or ""), item["part"])

    for preview_index, scene in enumerate(scenes):
        index = int(scene.get("audio_index", preview_index))
        segments = scene.get("segments") or [{
            "text": scene["narration"], "pause_after": .42,
            "kind": "body"}]
        prepared = []
        for segment_index, segment in enumerate(segments):
            segment_text = str(segment.get("text") or "").strip()
            instruct = _qwen3_narration_instruction(
                project, narrator, scene, segment)
            segment_signature = hashlib.sha256(json.dumps({
                "engine": narration_engine,
                "text": segment_text, "voice": narrator["voice"],
                "speed": narrator["speed"],
                "instruct": instruct,
                "seed": narration_seed,
                "version": NARRATION_AUDIO_VERSION,
            }, sort_keys=True).encode()).hexdigest()[:12]
            part = audio_dir / (
                f"{index:03d}-part-{segment_index:02d}-{segment_signature}.wav")
            prepared.append({
                "segment": segment, "text": segment_text,
                "instruct": instruct, "part": part,
                "segment_index": segment_index,
            })

        missing_parts = [item for item in prepared if not item["part"].exists()]
        for batch_start in range(0, len(missing_parts), NARRATION_TTS_BATCH_SIZE):
            batch = missing_parts[batch_start:batch_start + NARRATION_TTS_BATCH_SIZE]
            progress = 7 + int(preview_index / max(1, len(scenes)) * 55)
            first_number = int(batch[0]["segment_index"]) + 1
            last_number = int(batch[-1]["segment_index"]) + 1
            label = (f"thought {first_number}" if first_number == last_number
                     else f"thoughts {first_number}–{last_number}")
            update_current_job(
                f"Qwen3-TTS is performing {scene['id']} · {label} of "
                f"{len(segments)}", progress)
            try:
                result = _tts_json("/generate-batch", {
                    "engine": narration_engine,
                    "texts": [item["text"] for item in batch],
                    "voice": narrator["voice"], "speed": narrator["speed"],
                    "seed": narration_seed,
                    "temperature": 0.85, "top_p": 0.95, "top_k": 50,
                    "format": "wav", "sample_rate": 48000, "stereo": True,
                    "loudness": "youtube", "trim": True,
                })
                rendered = sorted(result.get("items") or [],
                                  key=lambda row: int(row.get("index", -1)))
                if len(rendered) != len(batch):
                    raise RuntimeError("batch returned the wrong number of clips")
                for item, output in zip(batch, rendered):
                    collect_tts_file(str(output.get("url") or ""), item["part"])
            except Exception as batch_error:
                # Older TTS services and an exceptional low-memory batch both
                # retain the proven single-clip path rather than failing a book.
                studio_event("narration.batch_fallback", level="warning",
                             scene=scene["id"], items=len(batch),
                             error=str(batch_error)[:500])
                for item in batch:
                    if not item["part"].exists():
                        generate_one(item)

        parts = [(item["segment"], item["part"], _media_seconds(item["part"]))
                 for item in prepared]

        signature = hashlib.sha256(json.dumps({
            "parts": [part.name for _, part, _ in parts],
            "pauses": [segment.get("pause_after") for segment, _, _ in parts],
            "version": NARRATION_AUDIO_VERSION,
        }, sort_keys=True).encode()).hexdigest()[:12]
        audio = audio_dir / f"{index:03d}-performed-{signature}.wav"
        cursor = 0.0
        scene["sync_blocks"] = []
        concat_rows = []
        for segment_index, (segment, part, speech_seconds) in enumerate(parts):
            speech_end = cursor + speech_seconds
            pause = (max(.12, min(1.6, float(
                segment.get("pause_after") or .4)))
                if segment_index < len(parts) - 1 else 0.0)
            scene["sync_blocks"].append({
                "text": str(segment.get("text") or ""),
                "kind": str(segment.get("kind") or "body"),
                "start": round(cursor, 4),
                "speech_end": round(speech_end, 4),
                "end": round(speech_end + pause, 4),
            })
            concat_rows.append(f"file '{part}'")
            if pause:
                concat_rows.append(f"file '{silence_for(pause)}'")
            cursor = speech_end + pause
        if not audio.exists():
            manifest = audio_dir / f"scene-{index:03d}-{signature}-concat.txt"
            manifest.write_text("\n".join(concat_rows) + "\n", encoding="utf-8")
            _run_media_command([
                "ffmpeg", "-y", "-loglevel", "error", "-f", "concat",
                "-safe", "0", "-i", str(manifest), "-ar", "48000", "-ac", "2",
                "-c:a", "pcm_s16le", str(audio)],
                f"Assembling the performed pauses for {scene['id']}")
        scene["audio_path"] = audio
        scene["seconds"] = _media_seconds(audio)

    signature = hashlib.sha256("|".join(
        str(scene["audio_path"].name) for scene in scenes).encode()).hexdigest()[:12]
    prefix = "narration-preview" if preview_pages else "narration"
    combined = media_dir / f"{prefix}-{signature}.m4a"
    if not combined.exists():
        silence = audio_dir / "silence-650ms.wav"
        if not silence.exists():
            _run_media_command([
                "ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i",
                "anullsrc=r=48000:cl=stereo", "-t", "0.65", "-c:a", "pcm_s16le",
                str(silence)], "Creating narration pauses")
        concat = audio_dir / f"concat-{signature}.txt"
        rows = []
        for index, scene in enumerate(scenes):
            rows.append(f"file '{scene['audio_path']}'")
            if index < len(scenes) - 1:
                rows.append(f"file '{silence}'")
        concat.write_text("\n".join(rows) + "\n", encoding="utf-8")
        _run_media_command([
            "ffmpeg", "-y", "-loglevel", "error", "-f", "concat", "-safe", "0",
            "-i", str(concat), "-c:a", "aac", "-b:a", "256k", str(combined)],
            "Combining the audiobook")
    return scenes, combined


def generate_book_audio_preview(project_id: str) -> dict:
    """Render only pages 1–2 so a narrator can be auditioned cheaply."""
    project = load_project(project_id)
    if not project.get("pages"):
        raise ValueError("Generate the story before testing its narrator")
    scenes, combined = _ensure_book_audio(
        project, preview_pages=2, require_images=False)
    video_audio = normalize_video_audio(project.get("video_audio"), project)
    video_audio["outputs"]["preview"] = "media/" + combined.name
    video_audio["outputs"]["preview_created_at"] = now()
    project["video_audio"] = video_audio
    project.setdefault("history", []).append({
        "at": now(), "action": (
            f"Tested {video_audio['narrator']['voice']} on the first "
            f"{len(scenes)} pages")})
    update_current_job("Two-page narration test ready", 96)
    return save_project(project)


def generate_book_audio(project_id: str) -> dict:
    project = load_project(project_id)
    if not project.get("pages"):
        raise ValueError("Generate the story before creating its narration")
    revision_snapshot(project, "before-video-audio-render")
    scenes, combined = _ensure_book_audio(project)
    duration = _media_seconds(combined)
    video_audio = normalize_video_audio(project.get("video_audio"), project)
    video_audio["outputs"].update({
        "audio": "media/" + combined.name,
        "narration": "media/" + combined.name,
        "duration_seconds": duration, "scene_count": len(scenes),
        "created_at": now(),
    })
    project["video_audio"] = video_audio
    project.setdefault("history", []).append({
        "at": now(), "action": f"Created Qwen3-TTS BF16 narration for {len(scenes)} scenes"})
    update_current_job("Narration mastered and ready", 96)
    return save_project(project)


def _video_dimensions(render: dict) -> tuple[int, int]:
    sizes = {
        ("16:9", "1080p"): (1920, 1080), ("16:9", "720p"): (1280, 720),
        ("9:16", "1080p"): (1080, 1920), ("9:16", "720p"): (720, 1280),
        ("1:1", "1080p"): (1080, 1080), ("1:1", "720p"): (720, 720),
    }
    return sizes[(render["aspect"], render["resolution"])]


def _video_font(pixels: int, bold: bool = False):
    from PIL import ImageFont
    candidates = [
        ("/System/Library/Fonts/Avenir Next.ttc", 2 if bold else 7),
        ("/System/Library/Fonts/SFNSRounded.ttf", 0),
        ("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 0),
    ]
    for path, index in candidates:
        try:
            return ImageFont.truetype(path, pixels, index=index)
        except (OSError, TypeError):
            continue
    return ImageFont.load_default()


def _video_wrap(draw, text: str, selected_font, max_width: int) -> list[str]:
    lines = []
    for paragraph in str(text or "").splitlines() or [""]:
        words = paragraph.split()
        if not words:
            continue
        line = ""
        for word in words:
            candidate = (line + " " + word).strip()
            box = draw.textbbox((0, 0), candidate, font=selected_font)
            if line and box[2] - box[0] > max_width:
                lines.append(line)
                line = word
            else:
                line = candidate
        if line:
            lines.append(line)
    return lines


def _video_caption_chunks(text: str, max_words: int) -> list[str]:
    """Split prose into short, sentence-aware captions for spoken-word video."""
    clean = re.sub(r"\s+", " ", str(text or "")).strip()
    if not clean:
        return []
    sentences = [part.strip() for part in re.findall(
        r".+?(?:[.!?…]+(?=\s|$)|$)", clean) if part.strip()]
    pieces = []
    for sentence in sentences or [clean]:
        words = sentence.split()
        if len(words) <= max_words:
            pieces.append(sentence)
            continue
        group_count = math.ceil(len(words) / max_words)
        base_size, larger_groups = divmod(len(words), group_count)
        start = 0
        for group_index in range(group_count):
            group_size = base_size + (1 if group_index < larger_groups else 0)
            pieces.append(" ".join(words[start:start + group_size]))
            start += group_size
    chunks = []
    current = ""
    for piece in pieces:
        candidate = (current + " " + piece).strip()
        if current and len(candidate.split()) > max_words:
            chunks.append(current)
            current = piece
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks


def _video_caption_weight(text: str) -> float:
    """Estimate spoken time while allowing a little extra for punctuation."""
    value = str(text or "")
    return max(1.0, len(value.split()) + value.count(",") * .35
               + len(re.findall(r"[.!?…]", value)) * .8)


def _blender_single_line_captions(
        text: str, kind: str, width: int, height: int,
        max_words: int) -> list[dict]:
    """Create phrases that Blender can render as exactly one visual line.

    Word count alone is not a safe width estimate for proportional fonts. A
    phrase containing wide letters or long words can make Blender wrap it,
    while the rolling-caption animation still advances by one line height. The
    wrapped half then collides with the next strip. Measure with the same Avenir
    font and keep extra width in reserve for Blender's slightly different text
    metrics and shadow.
    """
    from PIL import Image, ImageDraw

    clean = re.sub(r"\s+", " ", str(text or "")).strip()
    if not clean:
        return []
    font_size = max(30, round(height * (
        .045 if kind == "heading" else .052)))
    selected_font = _video_font(font_size)
    measuring_canvas = Image.new("L", (max(1, width), max(1, height)))
    draw = ImageDraw.Draw(measuring_canvas)
    # Blender receives a 90%-wide wrap box. Plan within 76% so font-engine
    # differences, punctuation and the drop shadow cannot trigger a late wrap.
    maximum_width = max(160, round(width * .76))
    initial = ([clean] if kind == "heading" else
               _video_caption_chunks(clean, max_words))
    lines = []
    for phrase in initial:
        wrapped = _video_wrap(draw, phrase, selected_font, maximum_width)
        lines.extend(wrapped or [phrase])
    return [{"text": line, "kind": kind, "font_size": font_size}
            for line in lines if line]


def _detect_video_focus(image_path: Path) -> dict:
    """Find faces and salient subjects locally for safe video camera framing."""
    if not VIDEO_FOCUS_DETECTOR.is_file():
        return {}
    try:
        process = subprocess.run(
            [str(VIDEO_FOCUS_DETECTOR), str(image_path)],
            capture_output=True, text=True, timeout=20)
        if process.returncode:
            studio_event("video_focus_detection_failed", level="warning",
                         image=str(image_path), error=(process.stderr or "")[-500:])
            return {}
        result = json.loads(process.stdout)
        return result if isinstance(result, dict) else {}
    except Exception as exc:
        studio_event("video_focus_detection_failed", level="warning",
                     image=str(image_path), error=str(exc)[:500])
        return {}


def _normalised_focus_boxes(rows) -> list[dict]:
    boxes = []
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        try:
            width = max(0.0, min(1.0, float(row.get("width") or 0)))
            height = max(0.0, min(1.0, float(row.get("height") or 0)))
            x = max(0.0, min(1.0, float(row.get("x") or .5)))
            y = max(0.0, min(1.0, float(row.get("y") or .5)))
            confidence = max(0.0, min(1.0, float(
                row.get("confidence") if row.get("confidence") is not None else 1)))
        except (TypeError, ValueError):
            continue
        if width > .005 and height > .005 and confidence >= .35:
            boxes.append({"x": x, "y": y, "width": width,
                          "height": height, "confidence": confidence})
    return boxes


def _video_focus_crop(source, size: tuple[int, int], detected: dict,
                      scene: dict):
    """Cover-crop an illustration while protecting detected heads and people."""
    from PIL import Image

    target_width, target_height = size
    source_width, source_height = source.size
    target_ratio = target_width / target_height
    source_ratio = source_width / source_height
    faces = _normalised_focus_boxes(detected.get("faces"))
    people = _normalised_focus_boxes(detected.get("people"))
    salient = _normalised_focus_boxes(detected.get("salient"))

    if source_ratio > target_ratio:
        crop_height = source_height
        crop_width = max(1, round(crop_height * target_ratio))
    else:
        crop_width = source_width
        crop_height = max(1, round(crop_width / target_ratio))

    focal_x, focal_y = .5, .40
    focal_kind = "default"
    if faces:
        # Keep all detected heads together. Their centre belongs in the upper
        # third, leaving the lower third available for rolling narration.
        weights = [max(.05, row["confidence"] * row["width"] * row["height"])
                   for row in faces]
        focal_x = sum(row["x"] * weight for row, weight in zip(faces, weights)) / sum(weights)
        focal_y = sum(row["y"] * weight for row, weight in zip(faces, weights)) / sum(weights)
        focal_kind = "face"
    elif people:
        person = max(people, key=lambda row: row["confidence"] * row["width"] * row["height"])
        prompt = str(scene.get("image_prompt") or "").lower()
        unusual_pose = any(term in prompt for term in (
            "inversion", "up the wall", "up a wall", "reclin", "lying",
            "bridge pose", "upside down", "headstand", "shoulder stand"))
        if not unusual_pose and person["height"] > person["width"] * 1.12:
            focal_x = person["x"]
            focal_y = max(0.0, person["y"] - person["height"] * .36)
            focal_kind = "estimated_head"
        elif salient:
            subject = max(salient, key=lambda row: row["confidence"])
            focal_x, focal_y = subject["x"], subject["y"]
            focal_kind = "salient_pose"
        else:
            focal_x, focal_y = person["x"], person["y"]
            focal_kind = "person"
    elif salient:
        subject = max(salient, key=lambda row: row["confidence"])
        focal_x, focal_y = subject["x"], subject["y"]
        focal_kind = "salient"

    # Place an actual or estimated head at about 31% of the screen height.
    # Other subjects sit nearer the centre. Clamp the crop to the source edges.
    desired_screen_y = .31 if focal_kind in {"face", "estimated_head"} else .45
    left = round(focal_x * source_width - .5 * crop_width)
    top = round(focal_y * source_height - desired_screen_y * crop_height)
    left = max(0, min(left, source_width - crop_width))
    top = max(0, min(top, source_height - crop_height))

    # Hard-protect every detected head with generous top/side breathing room.
    if faces:
        margin_y = crop_height * .075
        margin_x = crop_width * .06
        face_left = min((row["x"] - row["width"] / 2) * source_width for row in faces)
        face_right = max((row["x"] + row["width"] / 2) * source_width for row in faces)
        face_top = min((row["y"] - row["height"] / 2) * source_height for row in faces)
        face_bottom = max((row["y"] + row["height"] / 2) * source_height for row in faces)
        if face_right - face_left + 2 * margin_x <= crop_width:
            left = min(left, round(face_left - margin_x))
            left = max(left, round(face_right + margin_x - crop_width))
        if face_bottom - face_top + 2 * margin_y <= crop_height:
            top = min(top, round(face_top - margin_y))
            top = max(top, round(face_bottom + margin_y - crop_height))
        left = max(0, min(left, source_width - crop_width))
        top = max(0, min(top, source_height - crop_height))

    cropped = source.crop((left, top, left + crop_width, top + crop_height))
    canvas = cropped.resize(size, Image.Resampling.LANCZOS)

    def transformed(rows: list[dict]) -> list[dict]:
        values = []
        for row in rows:
            x = (row["x"] * source_width - left) / crop_width
            y = (row["y"] * source_height - top) / crop_height
            width = row["width"] * source_width / crop_width
            height = row["height"] * source_height / crop_height
            if -.15 <= x <= 1.15 and -.15 <= y <= 1.15:
                values.append(dict(row, x=round(x, 5), y=round(y, 5),
                                   width=round(width, 5), height=round(height, 5)))
        return values

    points = transformed(faces)
    if not points:
        synthetic_width = .12 if focal_kind == "estimated_head" else .20
        synthetic_height = .14 if focal_kind == "estimated_head" else .22
        points = [{
            "x": round((focal_x * source_width - left) / crop_width, 5),
            "y": round((focal_y * source_height - top) / crop_height, 5),
            "width": synthetic_width, "height": synthetic_height,
            "confidence": .5, "synthetic": True,
        }]
    return canvas, {
        "kind": focal_kind, "faces": transformed(faces),
        "points": points, "people": transformed(people),
    }


def _render_book_video_slide(project: dict, scene: dict, out: Path,
                             size: tuple[int, int], text_mode: str) -> dict:
    """Create a clean full-bleed, subject-aware plate for camera animation."""
    from PIL import Image, ImageEnhance, ImageOps
    source_path = image_abs(project, str(scene.get("image") or ""))
    if not source_path:
        raise ValueError("Missing picture for " + scene["id"])
    with Image.open(source_path) as opened:
        source = ImageOps.exif_transpose(opened).convert("RGB")
        detected = _detect_video_focus(source_path)
        canvas, focus = _video_focus_crop(source, size, detected, scene)
        canvas = ImageEnhance.Color(canvas).enhance(1.025)
        canvas = ImageEnhance.Contrast(canvas).enhance(1.015)
        canvas.save(out, "PNG", optimize=True)
        return focus


def _render_book_video_caption(scene: dict, out: Path,
                               size: tuple[int, int], body: str,
                               include_heading: bool, cover: bool = False) -> None:
    """Render one transparent, broadcast-safe lower-third caption layer."""
    from PIL import Image, ImageDraw
    width, height = size
    overlay = Image.new("RGBA", size, (0, 0, 0, 0))
    pixels = overlay.load()
    gradient_top = int(height * (.43 if cover else .55))
    for y in range(gradient_top, height):
        progress = (y - gradient_top) / max(1, height - gradient_top)
        alpha = int(210 * (progress ** 1.35))
        for x in range(width):
            pixels[x, y] = (8, 10, 18, alpha)
    draw = ImageDraw.Draw(overlay, "RGBA")
    side = int(width * (.075 if width > height else .065))
    bottom = int(height * .075)
    max_width = width - side * 2
    heading = str(scene.get("heading") or "").strip()

    if cover:
        title_size = max(42, int(min(width, height) * .075))
        title_font = _video_font(title_size, bold=True)
        title_lines = _video_wrap(draw, heading, title_font, max_width)[:3]
        subtitle_size = max(26, int(min(width, height) * .036))
        subtitle_font = _video_font(subtitle_size)
        subtitle_lines = _video_wrap(draw, body, subtitle_font, max_width)[:3]
        total = (len(title_lines) * int(title_size * 1.08)
                 + (int(subtitle_size * .55) if subtitle_lines else 0)
                 + len(subtitle_lines) * int(subtitle_size * 1.2))
        y = height - bottom - total
        for line in title_lines:
            draw.text((side + 3, y + 4), line, font=title_font,
                      fill=(0, 0, 0, 150))
            draw.text((side, y), line, font=title_font,
                      fill=(255, 252, 245, 255))
            y += int(title_size * 1.08)
        if subtitle_lines:
            y += int(subtitle_size * .55)
            for line in subtitle_lines:
                draw.text((side + 2, y + 3), line, font=subtitle_font,
                          fill=(0, 0, 0, 150))
                draw.text((side, y), line, font=subtitle_font,
                          fill=(255, 255, 255, 242))
                y += int(subtitle_size * 1.2)
        overlay.save(out, "PNG", optimize=True)
        return

    heading_size = max(22, int(min(width, height) * .027))
    body_size = max(34, int(min(width, height) * (.050 if width > height else .041)))
    body_font = _video_font(body_size, bold=True)
    body_lines = _video_wrap(draw, body, body_font, max_width)
    max_lines = 4 if width > height else 6
    while len(body_lines) > max_lines and body_size > 24:
        body_size -= 2
        body_font = _video_font(body_size, bold=True)
        body_lines = _video_wrap(draw, body, body_font, max_width)
    body_lines = body_lines[:max_lines]
    line_height = int(body_size * 1.18)
    heading_height = int(heading_size * 1.5) if include_heading and heading else 0
    y = height - bottom - len(body_lines) * line_height - heading_height
    if include_heading and heading:
        heading_font = _video_font(heading_size, bold=True)
        box = draw.textbbox((0, 0), heading, font=heading_font)
        pill_w = min(max_width, box[2] - box[0] + int(heading_size * 1.25))
        pill_h = int(heading_size * 1.35)
        draw.rounded_rectangle((side, y, side + pill_w, y + pill_h),
                               radius=pill_h // 2, fill=(255, 255, 255, 218))
        draw.text((side + int(heading_size * .6), y + int(heading_size * .13)),
                  heading, font=heading_font, fill=(22, 24, 32, 255))
        y += heading_height
    for line in body_lines:
        draw.text((side + 3, y + 4), line, font=body_font,
                  fill=(0, 0, 0, 170), stroke_width=2, stroke_fill=(0, 0, 0, 130))
        draw.text((side, y), line, font=body_font,
                  fill=(255, 255, 255, 255))
        y += line_height
    overlay.save(out, "PNG", optimize=True)


def _render_book_video_backdrop(out: Path, size: tuple[int, int]) -> None:
    """Make the subtle fixed reading gradient used behind Blender captions."""
    from PIL import Image, ImageDraw
    width, height = size
    overlay = Image.new("RGBA", size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay, "RGBA")
    top = int(height * .68)
    for y in range(top, height):
        progress = (y - top) / max(1, height - top)
        alpha = int(178 * (progress ** 1.45))
        draw.line((0, y, width, y), fill=(7, 10, 18, alpha))
    overlay.save(out, "PNG", optimize=True)


def _blender_caption_timeline(scene: dict, seconds: float, lead: float,
                              size: tuple[int, int], text_mode: str) -> list[dict]:
    """Create a three-line rolling caption timeline from the spoken script."""
    heading = str(scene.get("heading") or "").strip()
    if text_mode == "none":
        return []
    if text_mode == "heading":
        return ([{"text": heading, "start": lead,
                  "end": lead + seconds, "kind": "heading"}]
                if heading else [])
    width, height = size
    # Each phrase is one readable line. New lines enter at the bottom, move the
    # previous two upward, then gently fade the oldest line on its fourth step.
    max_words = 8 if width >= height else 6
    phrases = []
    starts = []
    sync_blocks = [row for row in scene.get("sync_blocks") or []
                   if isinstance(row, dict) and str(row.get("text") or "").strip()]
    if sync_blocks:
        # Each short performed thought has an exact measured audio start/end.
        # Caption estimates restart here every few seconds, so cadence differences
        # can never accumulate across a page or the rest of the book.
        for block in sync_blocks:
            kind = "heading" if block.get("kind") == "heading" else "body"
            block_phrases = _blender_single_line_captions(
                str(block.get("text") or ""), kind,
                width, height, max_words)
            if not block_phrases:
                continue
            block_weights = [_video_caption_weight(row["text"])
                             for row in block_phrases]
            block_total = sum(block_weights) or 1.0
            block_start = lead + max(0.0, float(block.get("start") or 0))
            block_end = lead + max(
                float(block.get("start") or 0) + .1,
                float(block.get("speech_end") or block.get("end") or 0))
            block_cursor = block_start
            for phrase_index, (row, weight) in enumerate(zip(
                    block_phrases, block_weights)):
                phrases.append(row)
                starts.append(block_cursor)
                if phrase_index < len(block_phrases) - 1:
                    block_cursor += (block_end - block_start) * weight / block_total
    else:
        if heading:
            phrases.extend(_blender_single_line_captions(
                heading, "heading", width, height, max_words))
        phrases.extend(_blender_single_line_captions(
            str(scene.get("text") or ""), "body", width, height, max_words))
    if not phrases:
        return []
    end_of_voice = lead + seconds
    if not starts:
        weights = [_video_caption_weight(row["text"]) for row in phrases]
        total = sum(weights) or 1.0
        cursor = lead
        for index, (row, weight) in enumerate(zip(phrases, weights)):
            starts.append(cursor)
            if index < len(phrases) - 1:
                cursor += seconds * weight / total
    timeline = []
    for index, row in enumerate(phrases):
        scroll_steps = starts[index + 1:index + 4]
        # A line is fully visible in the three-line window. When the next line
        # would make it the fourth, retain it only through the gentle fade.
        end = (min(end_of_voice, scroll_steps[2] + .48)
               if len(scroll_steps) == 3 else end_of_voice)
        timeline.append({
            "text": row["text"],
            "start": round(starts[index], 3),
            "end": round(max(starts[index] + .55, end), 3),
            "kind": row["kind"],
            "font_size": row["font_size"],
            "motion": "scroll_up",
            "scroll_steps": [round(value, 3) for value in scroll_steps],
        })
    return timeline


def _video_direction_signature(project: dict, scenes: list[dict],
                               render: dict) -> str:
    """Fingerprint the creative inputs so stale Qwen direction is never reused."""
    payload = {
        "title": project.get("title"),
        "audience": (project.get("settings") or {}).get("audience"),
        "tone": (project.get("settings") or {}).get("tone"),
        "render": {key: render.get(key) for key in (
            "aspect", "energy", "transitions", "text_mode", "renderer")},
        "scenes": [{
            "id": scene.get("id"), "heading": scene.get("heading"),
            "text": scene.get("text"), "image_prompt": scene.get("image_prompt"),
        } for scene in scenes],
        # Version 3 adds image-aware face protection and head-anchored moves.
        # Invalidate older text-only camera plans when a video is recreated.
        "version": 3,
    }
    return hashlib.sha256(json.dumps(
        payload, sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:24]


def _effective_video_energy(project: dict, render: dict) -> str:
    energy = str(render.get("energy") or "recommended")
    if energy != "recommended":
        return energy
    audience = str((project.get("settings") or {}).get("audience") or "")
    return ("lively" if any(age in audience for age in (
        "2–4", "2-4", "4–6", "4-6")) else "balanced")


def _request_video_direction(project: dict, scenes: list[dict],
                             render: dict) -> dict:
    """Ask Qwen3.6, through oMLX Orchestrator, for safe Blender shot beats."""
    settings = project.get("settings") or {}
    signature = _video_direction_signature(project, scenes, render)
    update_current_job(
        "oMLX Orchestrator is asking Qwen3.6 to direct every scene", 58)
    payload_scenes = []
    for scene in scenes:
        seconds = scene.get("seconds")
        if not seconds:
            seconds = max(3.5, min(45.0, len(str(
                scene.get("narration") or "").split()) / 2.35 + .8))
        payload_scenes.append({
            "id": scene.get("id"), "heading": scene.get("heading"),
            "text": scene.get("text"),
            "image_prompt": scene.get("image_prompt"),
            "seconds": round(float(seconds), 2),
        })
    response = _orchestrator_json("/animation-plan", {
        "title": project.get("title"),
        "audience": settings.get("audience"),
        "tone": settings.get("tone"),
        "energy": _effective_video_energy(project, render),
        "aspect": render.get("aspect"),
        "scenes": payload_scenes,
        "release_model": True,
    })
    directed = response.get("scenes") or []
    expected_ids = [str(scene.get("id") or "") for scene in scenes]
    directed_ids = [str(row.get("id") or "") for row in directed
                    if isinstance(row, dict)]
    if directed_ids != expected_ids:
        raise RuntimeError(
            "Qwen direction did not match the book's scenes in order")
    update_current_job(
        f"Qwen3.6 directed {len(directed)} scenes and released its memory", 61)
    return {
        "signature": signature,
        "model": str(response.get("model") or "Qwen3.6-35B"),
        "created_at": now(),
        "scenes": directed,
    }


def direct_book_video(project_id: str) -> dict:
    """Create and save Qwen's shot plan without rendering the full video."""
    project = load_project(project_id)
    if not project.get("pages"):
        raise ValueError("Generate the story before asking Qwen to direct it")
    revision_snapshot(project, "before-video-direction")
    _ensure_narration_script(project)
    scenes = _book_video_scenes(project)
    video_audio = normalize_video_audio(project.get("video_audio"), project)
    video_audio["direction"] = _request_video_direction(
        project, scenes, video_audio["render"])
    # Direction changes invalidate the prior video, but not its narration.
    video_audio["outputs"]["video"] = ""
    project["video_audio"] = video_audio
    project.setdefault("history", []).append({
        "at": now(),
        "action": f"Qwen3.6 directed {len(scenes)} Blender scenes",
    })
    update_current_job("Qwen's Blender direction is ready", 96)
    return save_project(project)


def _sound_plan_signature(project: dict, scenes: list[dict], sound: dict) -> str:
    payload = {
        "title": project.get("title"),
        "audience": (project.get("settings") or {}).get("audience"),
        "tone": (project.get("settings") or {}).get("tone"),
        "intensity": sound.get("intensity"),
        "scenes": [{
            "id": scene.get("id"), "heading": scene.get("heading"),
            "text": scene.get("text"), "image_prompt": scene.get("image_prompt"),
            "seconds": round(float(scene.get("seconds") or 0), 2),
        } for scene in scenes],
        "version": 1,
    }
    return hashlib.sha256(json.dumps(
        payload, sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:24]


def _request_sound_plan(project: dict, scenes: list[dict], sound: dict) -> dict:
    """Ask local Qwen for restrained, editable cues using final narration timing."""
    settings = project.get("settings") or {}
    update_current_job(
        "Qwen3.6 is placing sound effects around the narration", 60)
    payload_scenes = [{
        "id": scene.get("id"), "heading": scene.get("heading"),
        # Spoken content is authoritative. Production-style image prompts can
        # otherwise make the model ask for nonsensical "clay figure" sounds.
        "text": scene.get("text"),
        "seconds": round(float(scene.get("seconds") or 10), 2),
    } for scene in scenes]
    response = _orchestrator_json("/sound-plan", {
        "title": project.get("title"),
        "audience": settings.get("audience"), "tone": settings.get("tone"),
        "intensity": sound.get("intensity") or "balanced",
        "scenes": payload_scenes, "release_model": True,
    })
    directed = response.get("scenes") or []
    expected_ids = [str(scene.get("id") or "") for scene in scenes]
    directed_ids = [str(row.get("id") or "") for row in directed
                    if isinstance(row, dict)]
    if directed_ids != expected_ids:
        raise RuntimeError("Qwen sound direction did not match the book's scenes")
    update_current_job(
        f"Qwen3.6 planned sound for {len(directed)} scenes and released its memory", 92)
    return {
        "signature": _sound_plan_signature(project, scenes, sound),
        "model": str(response.get("model") or "Qwen3.6-35B"),
        "created_at": now(), "scenes": directed,
    }


def plan_book_sound_effects(project_id: str) -> dict:
    project = load_project(project_id)
    if not project.get("pages"):
        raise ValueError("Generate the story before planning its sound effects")
    revision_snapshot(project, "before-sound-effects-plan")
    scenes, narration = _ensure_book_audio(project)
    video_audio = normalize_video_audio(project.get("video_audio"), project)
    sound = video_audio["sound"]
    sound["plan"] = _request_sound_plan(project, scenes, sound)
    video_audio["sound"] = sound
    video_audio["outputs"].update({
        "audio": "media/" + narration.name,
        "narration": "media/" + narration.name,
        "soundtrack": "", "effects": "", "video": "",
        "duration_seconds": _media_seconds(narration),
        "scene_count": len(scenes), "created_at": now(),
    })
    project["video_audio"] = video_audio
    project.setdefault("history", []).append({
        "at": now(), "action": f"Qwen3.6 planned local sound effects for {len(scenes)} scenes"})
    update_current_job("The editable sound-effects plan is ready", 97)
    return save_project(project)


def _find_sound_cue(project: dict, scene_id: str, cue_id: str):
    video_audio = normalize_video_audio(project.get("video_audio"), project)
    for scene in video_audio["sound"]["plan"]["scenes"]:
        if scene["id"] != scene_id:
            continue
        for cue in scene["cues"]:
            if cue["id"] == cue_id:
                return video_audio, scene, cue
    raise ValueError("That sound effect is no longer in this book's sound plan")


def _generate_local_sound_cue(project: dict, scene: dict, cue: dict,
                              progress: int = 45) -> None:
    direct_runtime = (SFX_PYTHON.exists() and os.access(SFX_PYTHON, os.X_OK)
                      and SFX_SCRIPT.exists())
    wrapper_runtime = SFX_CLI.exists() and os.access(SFX_CLI, os.X_OK)
    if not direct_runtime and not wrapper_runtime:
        raise FileNotFoundError(
            "Stable Audio 3 Small SFX is not installed at " + str(SFX_CLI))
    prompt = re.sub(r"\s+", " ", str(cue.get("prompt") or "")).strip()
    if not prompt:
        raise ValueError("Enter a description for this sound effect first")
    duration = max(.4, min(30.0, float(cue.get("duration") or 2.0)))
    seed = int(uuid.uuid4().hex[:8], 16)
    generation_prompt = (
        "Gentle child-friendly non-startling sound effect, natural real-world "
        "scale, soft onset, controlled dynamics, no sudden loud peak. " + prompt
    )
    signature = hashlib.sha256(json.dumps({
        "prompt": prompt, "duration": duration, "seed": seed,
        "model": "stable-audio-3-sm-sfx",
        "version": SFX_MASTERING_VERSION,
    }, sort_keys=True).encode()).hexdigest()[:12]
    media_dir = project_dir(project["id"]) / "media"
    media_dir.mkdir(parents=True, exist_ok=True)
    safe_scene = re.sub(r"[^a-zA-Z0-9_-]", "-", scene["id"])[:50]
    safe_cue = re.sub(r"[^a-zA-Z0-9_-]", "-", cue["id"])[:50]
    output = media_dir / f"sfx-{safe_scene}-{safe_cue}-{signature}.wav"
    raw_output = output.with_name(output.stem + "-raw.wav")
    lease = None
    job_id = getattr(_job_context, "job_id", "") or "sfx-" + uuid.uuid4().hex[:10]
    try:
        lease = acquire_lease(
            "sfx:stable-audio-3", job_id,
            waiting=lambda: update_current_job(
                "Waiting for another local model before creating sound", 2),
        )
        unload_omlx_models(stage=lambda message, pct: update_current_job(message, pct))
        update_current_job(
            f"Stable Audio is creating {scene['id']} · {cue.get('kind', 'effect')}",
            progress)
        # Launchd intentionally supplies a minimal PATH, so the repository's
        # convenience wrapper cannot reliably discover Homebrew `uv`. The
        # installed virtual environment is self-contained; invoking it
        # directly is also faster and avoids an unnecessary setup check.
        command = ([str(SFX_PYTHON), str(SFX_SCRIPT)]
                   if direct_runtime else [str(SFX_CLI)]) + [
            "--prompt", generation_prompt,
            "--negative-prompt",
            "intelligible speech, narration, dialogue, singing, melody, music, "
            "sudden loud peak, harsh transient, jump scare, alarm, scream, "
            "screech, blast, explosion, frightening sound, distortion, clipping",
            "--dit", "sm-sfx", "--decoder", "same-s",
            "--seconds", f"{duration:.2f}", "--steps", "8", "--cfg", "3.0",
            "--seed", str(seed), "--out", str(raw_output),
        ]
        _run_media_command(command, f"Generating {scene['id']} sound effect")
        fade = min(.10, duration / 5)
        fade_out = max(fade, duration - fade)
        _run_media_command([
            "ffmpeg", "-y", "-loglevel", "error", "-i", str(raw_output),
            "-af", (
                "highpass=f=35,lowpass=f=14000,"
                "loudnorm=I=-24:LRA=5:TP=-4,"
                f"afade=t=in:st=0:d={fade:.3f},"
                f"afade=t=out:st={fade_out:.3f}:d={fade:.3f}"
            ),
            "-ar", "44100", "-ac", "2", "-c:a", "pcm_s16le", str(output),
        ], f"Safety-mastering {scene['id']} sound effect")
        raw_output.unlink(missing_ok=True)
    finally:
        release_lease(lease)
    if not output.exists():
        raise RuntimeError("Stable Audio finished without creating the sound file")
    cue["audio"] = "media/" + output.name
    cue["seed"] = seed
    cue["generated_at"] = now()
    cue["mastering_version"] = SFX_MASTERING_VERSION
    studio_event(
        "sound_effect.generated", scene_id=scene["id"], cue_id=cue["id"],
        kind=cue.get("kind"), duration=duration, prompt_chars=len(prompt),
        model="stable-audio-3-sm-sfx",
        mastering_version=SFX_MASTERING_VERSION)


def generate_book_sound_effect(project_id: str, scene_id: str,
                               cue_id: str) -> dict:
    project = load_project(project_id)
    revision_snapshot(project, f"before-sound-effect-{scene_id}-{cue_id}")
    video_audio, scene, cue = _find_sound_cue(project, scene_id, cue_id)
    _generate_local_sound_cue(project, scene, cue)
    video_audio["outputs"].update({"soundtrack": "", "effects": "", "video": ""})
    project["video_audio"] = video_audio
    project.setdefault("history", []).append({
        "at": now(), "action": f"Created local sound effect for {scene_id}"})
    update_current_job("The local sound effect is ready to preview", 97)
    return save_project(project)


def generate_all_book_sound_effects(project_id: str) -> dict:
    project = load_project(project_id)
    video_audio = normalize_video_audio(project.get("video_audio"), project)
    scenes = video_audio["sound"]["plan"]["scenes"]
    pending = [(scene, cue) for scene in scenes for cue in scene["cues"]
               if cue.get("enabled") and cue.get("prompt") and (
                   not cue.get("audio")
                   or int(cue.get("mastering_version") or 0)
                   < SFX_MASTERING_VERSION)]
    if not scenes:
        raise ValueError("Create a Qwen sound plan before generating effects")
    if not pending:
        update_current_job("Every enabled sound effect has already been created", 96)
        return project
    revision_snapshot(project, "before-generating-all-sound-effects")
    for index, (scene, cue) in enumerate(pending):
        progress = 8 + int(index / max(1, len(pending)) * 82)
        _generate_local_sound_cue(project, scene, cue, progress)
        project["video_audio"] = video_audio
        save_project(project)
    video_audio["outputs"].update({"soundtrack": "", "effects": "", "video": ""})
    project["video_audio"] = video_audio
    project.setdefault("history", []).append({
        "at": now(), "action": f"Created {len(pending)} local sound effects"})
    update_current_job(f"All {len(pending)} local sound effects are ready", 97)
    return save_project(project)


def _render_scene_effects(project: dict, scene: dict, cues: list[dict],
                          out: Path, master_volume: float) -> None:
    seconds = max(.1, float(scene.get("seconds") or 1))
    available = []
    for cue in cues:
        path = project_dir(project["id"]) / str(cue.get("audio") or "")
        if cue.get("enabled") and path.exists() and path.is_file():
            available.append((cue, path))
    if not available:
        _run_media_command([
            "ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i",
            "anullsrc=r=48000:cl=stereo", "-t", f"{seconds:.3f}",
            "-c:a", "pcm_s16le", str(out)], "Creating a silent effects scene")
        return
    command = ["ffmpeg", "-y", "-loglevel", "error"]
    for cue, path in available:
        if cue.get("loop"):
            # Never use an infinite input loop. Some WAV demuxers reset their
            # timestamps when looping, which can defeat downstream atrim and
            # cause ffmpeg to write until the disk is full. Calculate the small
            # finite repeat count actually needed for this cue instead.
            cue_duration = min(
                seconds, max(.4, float(cue.get("duration") or 2)))
            source_duration = max(.05, _media_seconds(path))
            repeats = max(0, min(
                60, math.ceil(cue_duration / source_duration) - 1))
            if repeats:
                command += ["-stream_loop", str(repeats)]
        command += ["-i", str(path)]
    filters = []
    labels = []
    for index, (cue, _path) in enumerate(available):
        cue_duration = min(seconds, max(.4, float(cue.get("duration") or 2)))
        start = min(seconds - .05, max(0.0, float(cue.get("at") or 0) * seconds))
        fade = min(.12, cue_duration / 5)
        fade_out = max(fade, cue_duration - fade)
        requested_volume = max(.01, min(
            1.0, float(cue.get("volume") or .2)))
        # Generated effects have already been safety-mastered to a controlled
        # loudness. Treat the cue value as creative prominence rather than
        # multiplying it as a raw percentage a second time; the old approach
        # made a 20% cue roughly 18 dB quieter before narration ducking.
        kind_floor = {
            "ambience": .15, "nature": .22, "transition": .25,
            "foley": .38, "impact": .42, "comedy": .35,
        }.get(str(cue.get("kind") or "foley"), .34)
        audible_gain = min(.90, kind_floor + requested_volume * 1.55)
        kind_boost = {
            "ambience": 1.30, "nature": 1.50, "transition": 1.65,
            "foley": 1.80, "impact": 1.85, "comedy": 1.75,
        }.get(str(cue.get("kind") or "foley"), 1.70)
        volume = min(1.35, audible_gain * master_volume * kind_boost)
        pan = max(-1.0, min(1.0, float(cue.get("pan") or 0)))
        filters.append(
            f"[{index}:a]aresample=48000,aformat=sample_fmts=fltp:channel_layouts=stereo,"
            f"atrim=start=0:end={cue_duration:.3f},asetpts=PTS-STARTPTS,"
            f"afade=t=in:st=0:d={fade:.3f},afade=t=out:st={fade_out:.3f}:d={fade:.3f},"
            f"volume={volume:.4f},stereotools=balance_out={pan:.3f},"
            f"adelay={int(start * 1000)}:all=1[s{index}]"
        )
        labels.append(f"[s{index}]")
    if len(labels) == 1:
        filters.append(
            f"{labels[0]}apad=whole_dur={seconds:.3f},"
            f"atrim=start=0:end={seconds:.3f}[effects]")
    else:
        filters.append(
            "".join(labels) + f"amix=inputs={len(labels)}:duration=longest:"
            f"dropout_transition=0:normalize=0,apad=whole_dur={seconds:.3f},"
            f"atrim=start=0:end={seconds:.3f}[effects]")
    # PCM stereo at 48 kHz is 192,000 bytes/second. This hard file cap is
    # deliberately independent of the filter graph so a future filter mistake
    # cannot consume the SSD.
    maximum_bytes = int((seconds + .5) * 192_000 + 1_048_576)
    partial = out.with_name(out.stem + ".partial.wav")
    partial.unlink(missing_ok=True)
    command += [
        "-filter_complex", ";".join(filters), "-map", "[effects]",
        "-t", f"{seconds:.3f}", "-fs", str(maximum_bytes),
        "-ar", "48000", "-ac", "2", "-c:a", "pcm_s16le", str(partial)]
    try:
        _run_media_command(command, f"Mixing effects for {scene['id']}")
        actual_seconds = _media_seconds(partial)
        actual_bytes = partial.stat().st_size
        if actual_seconds > seconds + .25 or actual_bytes > maximum_bytes:
            raise RuntimeError(
                f"Effects safety check rejected {scene['id']}: "
                f"{actual_seconds:.2f}s, {actual_bytes} bytes")
        partial.replace(out)
    finally:
        partial.unlink(missing_ok=True)


def _mix_scene_soundtrack(narration: Path, effects: Path, out: Path,
                          seconds: float, ducking: bool) -> None:
    filters = [
        "[0:a]aresample=48000,aformat=sample_fmts=fltp:channel_layouts=stereo[n]",
        "[1:a]aresample=48000,aformat=sample_fmts=fltp:channel_layouts=stereo[e]",
    ]
    effect_label = "e"
    if ducking:
        filters.append(
            "[e][n]sidechaincompress=threshold=0.15:ratio=1.35:attack=10:"
            "release=180:makeup=1[ducked]")
        effect_label = "ducked"
    filters.append(
        f"[n][{effect_label}]amix=inputs=2:duration=longest:normalize=0,"
        f"alimiter=limit=0.94,apad,atrim=duration={seconds:.3f}[master]")
    _run_media_command([
        "ffmpeg", "-y", "-loglevel", "error", "-i", str(narration),
        "-i", str(effects), "-filter_complex", ";".join(filters),
        "-map", "[master]", "-ar", "48000", "-ac", "2",
        "-c:a", "pcm_s16le", str(out)], "Mixing narration and sound effects")


def _concat_audio_scenes(paths: list[Path], out: Path, work: Path,
                         label: str) -> None:
    silence = work / "soundtrack-silence-650ms.wav"
    if not silence.exists():
        _run_media_command([
            "ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i",
            "anullsrc=r=48000:cl=stereo", "-t", "0.65", "-c:a", "pcm_s16le",
            str(silence)], "Creating soundtrack pauses")
    manifest = work / (out.stem + "-concat.txt")
    rows = []
    for index, path in enumerate(paths):
        rows.append(f"file '{path}'")
        if index < len(paths) - 1:
            rows.append(f"file '{silence}'")
    manifest.write_text("\n".join(rows) + "\n", encoding="utf-8")
    _run_media_command([
        "ffmpeg", "-y", "-loglevel", "error", "-f", "concat", "-safe", "0",
        "-i", str(manifest), "-c:a", "aac", "-b:a", "256k", str(out)], label)


def _ensure_video_sound_assets(project: dict) -> dict:
    """Automatically create every missing sound asset required by a video."""
    scenes, narration = _ensure_book_audio(project)
    video_audio = normalize_video_audio(project.get("video_audio"), project)
    sound = video_audio["sound"]
    plan = sound["plan"]
    expected_ids = [str(scene.get("id") or "") for scene in scenes]
    plan_ids = [str(scene.get("id") or "") for scene in plan.get("scenes") or []]
    expected_signature = _sound_plan_signature(project, scenes, sound)
    if (plan_ids != expected_ids
            or plan.get("signature") != expected_signature):
        update_current_job(
            "The sound plan is missing or out of date — Qwen is rebuilding it",
            18)
        sound["plan"] = _request_sound_plan(project, scenes, sound)
        plan = sound["plan"]
        video_audio["sound"] = sound
        video_audio["outputs"].update({
            "audio": "media/" + narration.name,
            "narration": "media/" + narration.name,
            "soundtrack": "", "effects": "", "video": "",
        })
        project["video_audio"] = video_audio
        save_project(project)

    pending = [
        (scene, cue)
        for scene in plan.get("scenes") or []
        for cue in scene.get("cues") or []
        if cue.get("enabled") and cue.get("prompt") and (
            not cue.get("audio")
            or int(cue.get("mastering_version") or 0)
            < SFX_MASTERING_VERSION)
    ]
    if pending:
        update_current_job(
            f"Creating {len(pending)} missing child-safe sound effects locally",
            63)
        for index, (scene, cue) in enumerate(pending):
            progress = 63 + int(index / max(1, len(pending)) * 17)
            _generate_local_sound_cue(project, scene, cue, progress)
            project["video_audio"] = video_audio
            save_project(project)
        video_audio["outputs"].update({
            "soundtrack": "", "effects": "", "video": "",
        })
        project.setdefault("history", []).append({
            "at": now(),
            "action": (
                f"Automatically created {len(pending)} missing sound effects "
                "for the video"),
        })
    project["video_audio"] = video_audio
    return project


def _ensure_book_soundtrack(project: dict) -> tuple[list[dict], Path, Path, Path]:
    scenes, narration = _ensure_book_audio(project)
    video_audio = normalize_video_audio(project.get("video_audio"), project)
    sound = video_audio["sound"]
    plan_by_id = {row["id"]: row for row in sound["plan"]["scenes"]}
    signature_data = {
        "narration": narration.name, "ducking": sound["ducking"],
        "master_volume": sound["master_volume"],
        "cues": [{
            "scene": row["id"], "id": cue["id"], "audio": cue.get("audio"),
            "enabled": cue.get("enabled"), "at": cue.get("at"),
            "volume": cue.get("volume"), "pan": cue.get("pan"),
            "duration": cue.get("duration"), "loop": cue.get("loop"),
        } for row in sound["plan"]["scenes"] for cue in row["cues"]],
        "version": SFX_MIX_VERSION,
    }
    signature = hashlib.sha256(json.dumps(
        signature_data, sort_keys=True).encode()).hexdigest()[:12]
    media_dir = project_dir(project["id"]) / "media"
    work = media_dir / "sound-work"
    work.mkdir(parents=True, exist_ok=True)
    mixed_paths = []
    effects_paths = []
    for index, scene in enumerate(scenes):
        update_current_job(
            f"Mixing sound for {scene['id']} · {index + 1} of {len(scenes)}",
            30 + int(index / max(1, len(scenes)) * 50))
        narration_path = Path(scene["audio_path"])
        effect_path = work / f"effects-{index:03d}-{signature}.wav"
        mixed_path = work / f"soundtrack-{index:03d}-{signature}.wav"
        plan_scene = plan_by_id.get(scene["id"], {})
        if not effect_path.exists():
            _render_scene_effects(
                project, scene, plan_scene.get("cues") or [], effect_path,
                float(sound["master_volume"]))
        if not mixed_path.exists():
            _mix_scene_soundtrack(
                narration_path, effect_path, mixed_path,
                float(scene["seconds"]), bool(sound["ducking"]))
        scene["narration_path"] = narration_path
        scene["effects_path"] = effect_path
        scene["audio_path"] = mixed_path
        effects_paths.append(effect_path)
        mixed_paths.append(mixed_path)
    soundtrack = media_dir / f"soundtrack-{signature}.m4a"
    effects = media_dir / f"effects-{signature}.m4a"
    if not soundtrack.exists():
        _concat_audio_scenes(mixed_paths, soundtrack, work, "Creating final soundtrack")
    if not effects.exists():
        _concat_audio_scenes(effects_paths, effects, work, "Creating effects-only track")
    return scenes, soundtrack, effects, narration


def mix_book_soundtrack(project_id: str) -> dict:
    project = load_project(project_id)
    if not project.get("pages"):
        raise ValueError("Generate the story before creating its soundtrack")
    revision_snapshot(project, "before-soundtrack-mix")
    scenes, soundtrack, effects, narration = _ensure_book_soundtrack(project)
    video_audio = normalize_video_audio(project.get("video_audio"), project)
    video_audio["outputs"].update({
        "audio": "media/" + soundtrack.name,
        "narration": "media/" + narration.name,
        "soundtrack": "media/" + soundtrack.name,
        "effects": "media/" + effects.name,
        "video": "", "duration_seconds": _media_seconds(soundtrack),
        "scene_count": len(scenes), "created_at": now(),
    })
    project["video_audio"] = video_audio
    project.setdefault("history", []).append({
        "at": now(), "action": "Mixed narration and local sound effects"})
    update_current_job("The professional soundtrack is ready", 97)
    return save_project(project)


def generate_book_video(project_id: str) -> dict:
    project = load_project(project_id)
    if not project.get("pages"):
        raise ValueError("Generate the story before creating its video")
    revision_snapshot(project, "before-video-render")
    project = _ensure_video_sound_assets(project)
    scenes, combined_audio, effects_audio, narration_audio = _ensure_book_soundtrack(project)
    video_audio = normalize_video_audio(project.get("video_audio"), project)
    render = video_audio["render"]
    width, height = _video_dimensions(render)
    media_dir = project_dir(project_id) / "media"
    work = media_dir / "video-work"
    work.mkdir(parents=True, exist_ok=True)
    transition_sets = {
        "gentle": ["dissolve", "fade"],
        "cinematic": ["dissolve", "fadeblack", "smoothleft"],
        "playful": ["dissolve", "smoothleft", "circleopen"],
        "dynamic": ["dissolve", "fade", "smoothleft"],
    }
    transitions = transition_sets[render["transitions"]]
    energy = _effective_video_energy(project, render)
    movement = {"gentle": .035, "balanced": .055, "lively": .075}[energy]
    transition_duration = {"gentle": .70, "balanced": .56, "lively": .46}[energy]
    # Do not start a new page's words while its incoming picture is still
    # crossfading over the previous page. The small settling beat also makes the
    # page change easier for young viewers to understand.
    narration_lead = transition_duration + .18
    narration_tail = .95
    segment_paths = []
    durations = []
    use_blender = render.get("renderer") == "blender"
    direction_map = {}
    if use_blender:
        signature = _video_direction_signature(project, scenes, render)
        direction = video_audio.get("direction") or {}
        if (direction.get("signature") != signature
                or len(direction.get("scenes") or []) != len(scenes)):
            direction = _request_video_direction(project, scenes, render)
            video_audio["direction"] = direction
            project["video_audio"] = video_audio
            save_project(project)
        direction_map = {
            str(row.get("id") or ""): row
            for row in direction.get("scenes") or [] if isinstance(row, dict)
        }
    blender_jobs = []
    blender_visuals = []
    blender_visual_signatures = []
    blender_backdrop = work / f"caption-backdrop-{width}x{height}.png"
    if use_blender and render["text_mode"] != "none" and not blender_backdrop.exists():
        _render_book_video_backdrop(blender_backdrop, (width, height))
    for index, scene in enumerate(scenes):
        update_current_job(
            f"Animating {scene['id']} · {index + 1} of {len(scenes)}",
            62 + int(index / max(1, len(scenes)) * 23))
        slide = work / f"slide-{index:03d}.png"
        segment = work / f"scene-{index:03d}.mp4"
        plate_focus = _render_book_video_slide(
            project, scene, slide, (width, height), render["text_mode"])
        duration = narration_lead + float(scene["seconds"]) + narration_tail
        durations.append(duration)
        frames = max(1, int(duration * 30))
        step = movement / frames
        motion_kind = index % 4
        if motion_kind == 0:
            zoom = f"min(1.0+on*{step:.8f},{1 + movement:.5f})"
            x = "iw/2-(iw/zoom/2)"
        elif motion_kind == 1:
            zoom = f"{1 + movement:.5f}"
            x = f"(iw-iw/zoom)*on/{frames}"
        elif motion_kind == 2:
            zoom = f"max({1 + movement:.5f}-on*{step:.8f},1.0)"
            x = "iw/2-(iw/zoom/2)"
        else:
            zoom = f"{1 + movement:.5f}"
            x = f"(iw-iw/zoom)*(1-on/{frames})"
        if use_blender:
            directed = direction_map.get(scene["id"], {})
            visual = work / f"blender-visual-{index:03d}.mp4"
            cover_overlay = ""
            captions = []
            if scene["id"] == "cover":
                cover_path = work / f"blender-cover-{width}x{height}.png"
                _render_book_video_caption(
                    scene, cover_path, (width, height),
                    str(scene.get("text") or ""), True, True)
                cover_overlay = str(cover_path)
            else:
                captions = _blender_caption_timeline(
                    scene, float(scene["seconds"]), narration_lead,
                    (width, height), render["text_mode"])
                if render["text_mode"] != "full":
                    caption_motion = str(
                        directed.get("caption_motion") or "rise")
                    captions = [dict(caption, motion=caption_motion)
                                for caption in captions]
            directed_pace = str(directed.get("pace") or energy)
            directed_movement = {
                "gentle": .04, "balanced": .06, "lively": .085,
            }.get(directed_pace, movement)
            blender_job = {
                "id": scene["id"], "plate": str(slide),
                "output": str(visual),
                "frames_dir": str(work / f"blender-frames-{index:03d}"),
                "ffmpeg": _media_tool("ffmpeg"),
                "width": width, "height": height, "fps": 30,
                "duration": round(duration, 4), "movement": directed_movement,
                "pattern": motion_kind,
                "beats": directed.get("beats") or [],
                "focus": plate_focus,
                "pace": directed_pace,
                "font": "/System/Library/Fonts/Avenir Next.ttc",
                "backdrop": (str(blender_backdrop)
                             if captions and render["text_mode"] != "none" else ""),
                "cover_overlay": cover_overlay, "captions": captions,
            }
            source_art = image_abs(project, str(scene.get("image") or ""))
            source_stat = source_art.stat() if source_art else None
            visual_signature = hashlib.sha256(json.dumps({
                "job": {key: value for key, value in blender_job.items()
                        if key not in {"plate", "output", "frames_dir", "ffmpeg"}},
                "art": {
                    "path": str(scene.get("image") or ""),
                    "size": source_stat.st_size if source_stat else 0,
                    "modified": source_stat.st_mtime_ns if source_stat else 0,
                },
                "render_version": 2,
            }, sort_keys=True).encode()).hexdigest()[:16]
            visual_cache = visual.with_suffix(".render-signature")
            try:
                cached_signature = visual_cache.read_text(encoding="utf-8").strip()
            except OSError:
                cached_signature = ""
            if not visual.exists() or cached_signature != visual_signature:
                blender_job["cache_signature"] = visual_signature
                blender_job["cache_file"] = str(visual_cache)
                blender_jobs.append(blender_job)
            blender_visuals.append(visual)
            blender_visual_signatures.append(visual_signature)
            segment_paths.append(segment)
            continue
        vf = (f"zoompan=z='{zoom}':x='{x}':y='ih/2-(ih/zoom/2)':"
              f"d=1:s={width}x{height}:fps=30,format=yuv420p")

        caption_specs = []
        synced_caption_windows = []
        text_mode = render["text_mode"]
        if scene["id"] == "cover":
            caption_specs = [(str(scene.get("text") or ""), True, True)]
        elif text_mode == "full":
            blocks = [row for row in scene.get("sync_blocks") or []
                      if isinstance(row, dict) and str(row.get("text") or "").strip()]
            if blocks:
                caption_specs = [(str(row["text"]), False, False) for row in blocks]
                for block_index, row in enumerate(blocks):
                    start = narration_lead + float(row.get("start") or 0)
                    end = (narration_lead + float(blocks[block_index + 1].get("start") or 0)
                           if block_index + 1 < len(blocks)
                           else narration_lead + float(scene["seconds"]))
                    synced_caption_windows.append((start, max(start + .2, end)))
            else:
                max_words = 18 if width > height else (14 if height > width else 16)
                chunks = _video_caption_chunks(str(scene.get("text") or ""), max_words)
                caption_specs = [(chunk, chunk_index == 0, False)
                                 for chunk_index, chunk in enumerate(chunks)]
        elif text_mode == "heading":
            caption_specs = [(str(scene.get("heading") or ""), False, False)]

        caption_paths = []
        for caption_index, (body, include_heading, is_cover) in enumerate(caption_specs):
            caption = work / f"caption-{index:03d}-{caption_index:02d}.png"
            _render_book_video_caption(
                scene, caption, (width, height), body, include_heading, is_cover)
            caption_paths.append(caption)

        command = [
            "ffmpeg", "-y", "-loglevel", "error", "-framerate", "30", "-loop", "1",
            "-i", str(slide), "-i", str(scene["audio_path"]),
        ]
        for caption in caption_paths:
            command += ["-framerate", "30", "-loop", "1", "-i", str(caption)]
        filters = [f"[0:v]{vf}[motion]"]
        previous = "motion"
        if caption_specs:
            if synced_caption_windows:
                windows = synced_caption_windows
            elif scene["id"] == "cover" or text_mode == "heading":
                windows = [(narration_lead, narration_lead + float(scene["seconds"]))]
            else:
                weights = [_video_caption_weight(spec[0]) for spec in caption_specs]
                weight_total = sum(weights) or 1
                cursor = narration_lead
                windows = []
                for caption_index, weight in enumerate(weights):
                    end = (narration_lead + float(scene["seconds"])
                           if caption_index == len(weights) - 1 else
                           cursor + float(scene["seconds"]) * weight / weight_total)
                    windows.append((cursor, end))
                    cursor = end
            for caption_index, (start, end) in enumerate(windows):
                fade = min(.16, max(.04, (end - start) / 5))
                fade_out = max(start + fade, end - fade)
                filters.append(
                    f"[{caption_index + 2}:v]format=rgba,"
                    f"fade=t=in:st={start:.3f}:d={fade:.3f}:alpha=1,"
                    f"fade=t=out:st={fade_out:.3f}:d={fade:.3f}:alpha=1"
                    f"[caption{caption_index}]")
                filters.append(
                    f"[{previous}][caption{caption_index}]overlay=0:0:"
                    f"enable='between(t,{start:.3f},{end:.3f})'[captioned{caption_index}]")
                previous = f"captioned{caption_index}"
        filters.append(
            f"[1:a]adelay={int(narration_lead * 1000)}:all=1,"
            f"apad=pad_dur={narration_tail:.3f}[a]")
        command += [
            "-filter_complex", ";".join(filters),
            "-map", f"[{previous}]", "-map", "[a]", "-t", f"{duration:.3f}",
            "-c:v", "libx264", "-preset", "veryfast", "-crf", "18",
            "-c:a", "aac", "-b:a", "256k", "-movflags", "+faststart", str(segment)]
        _run_media_command(command, f"Animating {scene['id']}")
        segment_paths.append(segment)

    if use_blender:
        # Two independent Blender workers keep the Mac's CPU/GPU fed without
        # changing a single render or encoding quality setting. More workers
        # contend for unified memory and become slower on long books.
        if blender_jobs:
            worker_count = 2 if len(blender_jobs) >= 4 else 1
            job_groups = [blender_jobs[index::worker_count]
                          for index in range(worker_count)]
            job_paths = []
            for worker_index, group in enumerate(job_groups):
                job_path = work / f"blender-storybook-job-{worker_index + 1}.json"
                job_path.write_text(json.dumps({"scenes": group}, indent=2),
                                    encoding="utf-8")
                job_paths.append(job_path)
            update_current_job(
                f"Blender is rendering {len(blender_jobs)} changed story scenes "
                f"with {worker_count} worker{'s' if worker_count > 1 else ''}", 72)

            def render_blender_job(job_path: Path) -> None:
                _run_media_command([
                    "blender", "--background", "--factory-startup", "--python",
                    str(Path(__file__).with_name("blender_storybook_renderer.py")),
                    "--", "--job", str(job_path)],
                    "Blender storybook rendering")

            if worker_count == 1:
                render_blender_job(job_paths[0])
            else:
                from concurrent.futures import ThreadPoolExecutor
                with ThreadPoolExecutor(max_workers=worker_count) as executor:
                    futures = [executor.submit(render_blender_job, path)
                               for path in job_paths]
                    for future in futures:
                        future.result()
            for job in blender_jobs:
                output = Path(job["output"])
                if output.exists():
                    Path(job["cache_file"]).write_text(
                        str(job["cache_signature"]), encoding="utf-8")
        else:
            update_current_job(
                f"Reusing {len(blender_visuals)} unchanged professional scenes", 80)
        for index, (scene, visual, segment, duration) in enumerate(zip(
                scenes, blender_visuals, segment_paths, durations)):
            update_current_job(
                f"Mastering Blender scene {index + 1} of {len(scenes)}", 82)
            if not visual.exists():
                raise RuntimeError("Blender did not create " + str(visual))
            segment_signature = hashlib.sha256(json.dumps({
                "visual": blender_visual_signatures[index],
                "audio": Path(scene["audio_path"]).name,
                "duration": round(duration, 4),
                "lead": narration_lead, "tail": narration_tail,
                "master_version": 2,
            }, sort_keys=True).encode()).hexdigest()[:16]
            segment_cache = segment.with_suffix(".master-signature")
            try:
                cached_segment = segment_cache.read_text(
                    encoding="utf-8").strip()
            except OSError:
                cached_segment = ""
            if not segment.exists() or cached_segment != segment_signature:
                _run_media_command([
                    "ffmpeg", "-y", "-loglevel", "error", "-i", str(visual),
                    "-i", str(scene["audio_path"]), "-filter_complex",
                    f"[1:a]adelay={int(narration_lead * 1000)}:all=1,"
                    f"apad=pad_dur={narration_tail:.3f}[a]",
                    "-map", "0:v:0", "-map", "[a]", "-t", f"{duration:.3f}",
                    "-c:v", "copy", "-c:a", "aac", "-b:a", "256k",
                    "-movflags", "+faststart", str(segment)],
                    f"Mastering Blender scene {scene['id']}")
                segment_cache.write_text(segment_signature, encoding="utf-8")

    final_signature = hashlib.sha256(json.dumps({
        "segments": [{
            "name": path.name, "size": path.stat().st_size,
            "modified": path.stat().st_mtime_ns,
        } for path in segment_paths],
        "durations": [round(value, 4) for value in durations],
        "transition_duration": transition_duration,
        "transitions": [str((direction_map.get(
            scenes[index]["id"], {}) if use_blender else {}).get(
                "transition") or transitions[(index - 1) % len(transitions)])
            for index in range(1, len(scenes))],
        "master_version": 2,
    }, sort_keys=True).encode()).hexdigest()[:16]
    final = media_dir / (
        f"{slug(project.get('title') or 'book')}-video-{final_signature}.mp4")
    if not final.exists() and len(segment_paths) == 1:
        shutil.copy2(segment_paths[0], final)
    elif not final.exists():
        command = ["ffmpeg", "-y", "-loglevel", "error"]
        for path in segment_paths:
            command += ["-i", str(path)]
        video_parts = [
            f"[{index}:v]fps=30,settb=AVTB,setpts=PTS-STARTPTS[vb{index}]"
            for index in range(len(segment_paths))
        ]
        scene_starts = [0.0]
        for index in range(1, len(segment_paths)):
            scene_starts.append(
                scene_starts[-1] + durations[index - 1] - transition_duration)
        audio_parts = [
            f"[{index}:a]aresample=async=1:first_pts=0,asetpts=PTS-STARTPTS,"
            f"adelay={int(scene_starts[index] * 1000)}:all=1[ab{index}]"
            for index in range(len(segment_paths))
        ]
        previous_video = "vb0"
        elapsed = durations[0]
        for index in range(1, len(segment_paths)):
            offset = elapsed - transition_duration * index
            transition = str((direction_map.get(
                scenes[index]["id"], {}) if use_blender else {}).get(
                    "transition") or transitions[(index - 1) % len(transitions)])
            video_parts.append(
                f"[{previous_video}][vb{index}]xfade=transition={transition}:"
                f"duration={transition_duration:.2f}:offset={offset:.3f}[vx{index}]")
            previous_video = f"vx{index}"
            elapsed += durations[index]
        final_duration = scene_starts[-1] + durations[-1]
        audio_inputs = "".join(f"[ab{index}]" for index in range(len(segment_paths)))
        audio_parts.append(
            f"{audio_inputs}amix=inputs={len(segment_paths)}:duration=longest:"
            f"dropout_transition=0:normalize=0,alimiter=limit=0.95,"
            f"atrim=duration={final_duration:.3f}[audio_master]")
        command += [
            "-filter_complex", ";".join(video_parts + audio_parts),
            "-map", f"[{previous_video}]", "-map", "[audio_master]",
            "-c:v", "libx264", "-preset", "medium", "-crf", "18",
            "-c:a", "aac", "-b:a", "256k", "-movflags", "+faststart", str(final)]
        update_current_job("Adding the final transitions and mastering audio", 88)
        _run_media_command(command, "Finishing the book video")
    duration = _media_seconds(final)
    video_audio["outputs"].update({
        "audio": "media/" + combined_audio.name,
        "narration": "media/" + narration_audio.name,
        "soundtrack": "media/" + combined_audio.name,
        "effects": "media/" + effects_audio.name,
        "video": "media/" + final.name,
        "duration_seconds": duration, "scene_count": len(scenes),
        "created_at": now(),
    })
    project["video_audio"] = video_audio
    project.setdefault("history", []).append({
        "at": now(), "action": f"Created animated narrated video with {len(scenes)} scenes"})
    update_current_job("Video and narration are ready", 97)
    return save_project(project)


def publishing_preflight(project: dict) -> dict:
    """Return plain-language, platform-specific readiness checks."""
    from PIL import Image

    p = project
    publishing = normalize_publishing(p.get("publishing"), p)
    route = publishing["route"]
    isbn = publishing["isbn"]
    metadata = p.get("metadata") or {}
    pages = p.get("pages") or []
    page_count = len(pages) + (len(pages) % 2)
    w_in, h_in = trim_size(p.get("settings") or {})
    trim_id = str((p.get("settings") or {}).get("trim") or "")
    interior = str((p.get("settings") or {}).get(
        "print_interior") or "premium_colour")
    missing_art = []
    active_art = [("cover", (p.get("cover") or {}).get("image"))]
    active_art.extend((f"page {page.get('number')}", page.get("image"))
                      for page in pages)
    effective_ppi = []
    for label, rel in active_art:
        path = image_abs(p, str(rel or ""))
        if not path:
            missing_art.append(label)
            continue
        try:
            with Image.open(path) as source:
                ppi = min(source.width / max(.01, w_in),
                          source.height / max(.01, h_in))
                effective_ppi.append({
                    "target": label, "ppi": round(ppi),
                    "pixels": f"{source.width} × {source.height}",
                })
        except Exception:
            missing_art.append(label)
    lowest_ppi = min((row["ppi"] for row in effective_ppi), default=0)
    metadata_ready = (
        bool(str(metadata.get("description") or "").strip())
        and len(metadata.get("keywords") or []) == 7
        and len(metadata.get("categories") or []) == 3
    )
    bisac_ready = bool(metadata.get("bisac_subjects"))

    def check(state: str, title: str, detail: str) -> dict:
        return {"state": state, "title": title, "detail": detail}

    global_issues = []
    if route["kdp_expanded"] and route["ingram_print"]:
        global_issues.append(check(
            "block", "Overlapping print distribution",
            "Turn off KDP Expanded Distribution when IngramSpark distributes "
            "this paperback. Keep KDP direct Amazon sales enabled."))
    if route["kdp_select"] and route["d2d_ebook"]:
        global_issues.append(check(
            "block", "Conflicting ebook exclusivity",
            "KDP Select requires ebook exclusivity. Turn off KDP Select or turn "
            "off Draft2Digital wide ebook distribution."))
    if route["d2d_ebook"] and route["d2d_amazon"] and route["kdp_ebook"]:
        global_issues.append(check(
            "block", "Duplicate Amazon ebook route",
            "Amazon is selected in both KDP and Draft2Digital. Keep the direct KDP "
            "listing and exclude Amazon in Draft2Digital."))
    if route["blurb_global"] and route["ingram_print"]:
        global_issues.append(check(
            "block", "Duplicate Ingram route",
            "Blurb Global Distribution also uses Ingram. Use Blurb for direct "
            "premium sales only when IngramSpark already handles wide print."))
    if (route["ingram_print"] and route["kdp_paperback"]
            and isbn["ownership"] != "publisher_owned"):
        global_issues.append(check(
            "block", "Publisher-owned paperback ISBN needed",
            "The same exact paperback can use your own ISBN across KDP and "
            "IngramSpark. A platform-provided free ISBN cannot be shared."))
    if (route["ingram_print"] and route["kdp_paperback"]
            and not isbn["paperback"]):
        global_issues.append(check(
            "warn", "Paperback ISBN not entered",
            "Enter the publisher-owned paperback ISBN before requesting the final "
            "Ingram cover template and submitting either print edition."))

    kdp_minimum = 72 if interior == "standard_colour" else 24
    kdp_checks = [
        check("pass" if trim_id == "8.5x8.5" else "warn",
              "Square paperback trim",
              ("8.5 × 8.5 inches selected." if trim_id == "8.5x8.5" else
               f"Current trim is {w_in:g} × {h_in:g} inches.")),
        check("pass" if page_count >= kdp_minimum else "block",
              "KDP colour page count",
              f"{page_count} interior pages; this colour setting requires at least "
              f"{kdp_minimum}."),
        check("pass" if not missing_art else "block", "Complete illustrations",
              ("Cover and every page have artwork." if not missing_art else
               "Missing: " + ", ".join(missing_art[:8]))),
        check("pass" if metadata_ready else "block", "Amazon listing metadata",
              ("Description, seven keywords and three categories are ready."
               if metadata_ready else
               "Generate or complete the description, seven keywords and three categories.")),
        check("pass" if lowest_ppi >= 300 else "warn", "Source artwork resolution",
              (f"Lowest effective source resolution is {lowest_ppi} PPI."
               if lowest_ppi else "No active artwork could be measured.")
              + ("" if lowest_ppi >= 300 else
                 " Export will enlarge it to the 300 PPI print canvas, but this does "
                 "not create the same detail as native 300 PPI artwork.")),
    ]
    ingram_checks = [
        check("pass" if trim_id == "8.5x8.5" else "warn", "Ingram trim edition",
              ("8.5 × 8.5 is selected for the shared print edition."
               if trim_id == "8.5x8.5" else
               "Confirm this trim is available for the chosen Ingram binding and paper.")),
        check("pass" if page_count >= 18 else "block", "Ingram page count",
              f"{page_count} interior pages; at least 18 are required for this profile."),
        check("pass" if isbn["ownership"] == "publisher_owned"
              and bool(isbn["paperback"]) else "block", "Print ISBN and imprint",
              ("Publisher-owned paperback ISBN is recorded."
               if isbn["ownership"] == "publisher_owned" and isbn["paperback"]
               else "Enter a publisher-owned ISBN and matching imprint.")),
        check("pass" if bisac_ready else "warn", "BISAC subjects",
              "BISAC subjects are ready." if bisac_ready else
              "Generate or enter at least one BISAC subject."),
        check("pass" if publishing["editions"]["ingram"]["cover_template_ready"]
              else "block", "Official Ingram cover template",
              ("You confirmed the ISBN-specific template is ready."
               if publishing["editions"]["ingram"]["cover_template_ready"] else
               "Request the official ISBN-, binding-, paper- and page-count-specific "
               "cover template before final cover submission.")),
    ]
    d2d_checks = [
        check("pass", "Ebook route", "Draft2Digital is configured for EPUB only; "
              "its grayscale non-square print service is not used for this edition."),
        check("pass" if not route["d2d_amazon"] else "block",
              "Amazon excluded", "Amazon is excluded from Draft2Digital."
              if not route["d2d_amazon"] else
              "Exclude Amazon because the ebook is published directly through KDP."),
        check("pass" if bisac_ready else "warn", "BISAC subject",
              "BISAC metadata is ready." if bisac_ready else
              "Generate or enter at least one BISAC subject."),
    ]
    blurb_edition = publishing["editions"]["blurb"]
    blurb_label = BLURB_EDITIONS[blurb_edition["trim"]]["label"]
    blurb_checks = [
        check("pass", "Separate premium edition", blurb_label +
              " selected; Blurb does not offer an exact 8.5 × 8.5 photo-book trim."),
        check("pass" if page_count >= 20 and page_count % 2 == 0 else "block",
              "Blurb page count", f"{page_count} even interior pages."),
        check("pass" if blurb_edition["specifications_confirmed"] else "block",
              "Blurb specification calculator",
              ("You confirmed the current Blurb PDF and cover specifications."
               if blurb_edition["specifications_confirmed"] else
               "Confirm the exact current exported-page and cover measurements in "
               "Blurb's specification calculator before final upload.")),
        check("pass" if not route["blurb_global"] else "block",
              "Direct-sale route", "Blurb direct sales avoid a second Ingram listing."
              if not route["blurb_global"] else
              "Turn off Blurb Global Distribution while IngramSpark handles wide print."),
    ]
    lulu_edition = publishing["editions"]["lulu"]
    title_length = len(str(p.get("title") or "").strip())
    lulu_checks = [
        check("pass" if trim_id == "8.5x8.5" else "warn",
              "Lulu square trim",
              ("8.5 × 8.5 Square is selected; its bleed PDF is 8.75 × 8.75 inches."
               if trim_id == "8.5x8.5" else
               f"Current trim is {w_in:g} × {h_in:g} inches. Confirm this exact "
               "format in Lulu's calculator before ordering.")),
        check("pass" if not missing_art else "block", "Complete illustrations",
              ("Cover and every page have artwork." if not missing_art else
               "Missing: " + ", ".join(missing_art[:8]))),
        check("pass" if lowest_ppi >= 300 else "warn", "300 PPI artwork",
              (f"Lowest effective source resolution is {lowest_ppi} PPI."
               if lowest_ppi else "No active artwork could be measured.")
              + ("" if lowest_ppi >= 300 else
                 " The print export is rendered at 300 PPI, but native high-resolution "
                 "artwork gives the best result.")),
        check("pass" if 0 < title_length <= 80 else "block",
              "Store-compatible title length",
              (f"The title is {title_length} characters."
               if 0 < title_length <= 80 else
               f"The title is {title_length} characters; Lulu Direct store products "
               "must use a title of 80 characters or fewer.")),
        check("pass" if lulu_edition["template_ready"] else "block",
              "Official Lulu cover template",
              ("You confirmed the final binding-, trim- and page-count-specific "
               "template is ready." if lulu_edition["template_ready"] else
               "Confirm the final format and download Lulu's custom cover template. "
               "The wrap cover must be one PDF containing back, spine and front.")),
        check("pass" if lulu_edition["product_connected"] else "block",
              "Store product connected",
              ("The Lulu Project is connected to the selected store channel."
               if lulu_edition["product_connected"] else
               "After approving the Lulu Project, connect it to Shopify, "
               "WooCommerce, Wix or your Print API integration.")),
    ]

    def platform(label: str, enabled: bool, checks: list[dict], package: str) -> dict:
        return {
            "label": label, "enabled": enabled, "checks": checks,
            "ready": enabled and not any(row["state"] == "block" for row in checks),
            "package": package,
        }

    return {
        "page_count": page_count,
        "trim_size": f"{w_in:g} × {h_in:g} inches",
        "lowest_effective_ppi": lowest_ppi,
        "image_resolution": effective_ppi,
        "missing_art": missing_art,
        "global_issues": global_issues,
        "platforms": {
            "kdp": platform("Amazon KDP", route["kdp_ebook"] or
                            route["kdp_paperback"], kdp_checks,
                            "EPUB, ebook cover, manuscript PDF and KDP cover PDF"),
            "ingram": platform("IngramSpark", route["ingram_print"], ingram_checks,
                               "Interior PDF, cover-design assets, metadata and template checklist"),
            "d2d": platform("Draft2Digital", route["d2d_ebook"], d2d_checks,
                            "EPUB, portrait ebook cover and wide metadata"),
            "blurb": platform("Blurb", route["blurb_direct"] or
                              route["blurb_global"], blurb_checks,
                              "Square-edition layout proof, cover assets and specification checklist"),
            "lulu": platform("Lulu Direct", route["lulu_direct"], lulu_checks,
                             "Full-bleed interior PDF, front-cover asset, metadata and store-connection checklist"),
        },
    }


def export_marketing_cover(project: dict, out: Path,
                           size: tuple[int, int]) -> None:
    """Create a high-resolution front cover asset from the approved cover art."""
    from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont, ImageOps

    width, height = size
    cover_path = image_abs(project, (project.get("cover") or {}).get("image", ""))
    if cover_path:
        with Image.open(cover_path) as source_image:
            source = ImageOps.exif_transpose(source_image).convert("RGB")
            background = ImageOps.fit(
                source, size, method=Image.Resampling.LANCZOS).filter(
                    ImageFilter.GaussianBlur(max(12, width // 55)))
            background = ImageEnhance.Brightness(background).enhance(.66)
            if abs(width - height) < width * .08:
                foreground = ImageOps.fit(
                    source, size, method=Image.Resampling.LANCZOS)
                background.paste(foreground, (0, 0))
            else:
                hero_h = min(width, int(height * .66))
                foreground = ImageOps.fit(
                    source, (width, hero_h), method=Image.Resampling.LANCZOS)
                background.paste(foreground, (0, (height - hero_h) // 2))
    else:
        background = Image.new("RGB", size, (72, 43, 64))

    draw = ImageDraw.Draw(background, "RGBA")
    font_regular = "/System/Library/Fonts/Supplemental/Arial.ttf"
    font_bold = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"

    def font(path: str, px: int):
        try:
            return ImageFont.truetype(path, px)
        except OSError:
            return ImageFont.load_default()

    def lines_for(text: str, selected_font, max_width: int,
                  max_lines: int) -> list[str]:
        words = str(text or "").split()
        lines = []
        line = ""
        for word in words:
            candidate = (line + " " + word).strip()
            bounds = draw.textbbox((0, 0), candidate, font=selected_font)
            if line and bounds[2] - bounds[0] > max_width:
                lines.append(line)
                line = word
            else:
                line = candidate
        if line:
            lines.append(line)
        if len(lines) > max_lines:
            lines = lines[:max_lines]
            lines[-1] = lines[-1].rstrip(" .") + "..."
        return lines

    title = str(project.get("title") or "Untitled Book")
    subtitle = str(project.get("subtitle") or "")
    author = str((project.get("settings") or {}).get("author") or "")
    title_size = max(42, int(width * .064))
    title_font = font(font_bold, title_size)
    title_lines = lines_for(title, title_font, int(width * .84), 3)
    subtitle_size = max(26, int(width * .030))
    subtitle_font = font(font_regular, subtitle_size)
    subtitle_lines = lines_for(subtitle, subtitle_font, int(width * .82), 2)
    title_panel_h = int(height * (.24 if subtitle_lines else .19))
    draw.rectangle((0, 0, width, title_panel_h), fill=(20, 10, 18, 185))
    y = int(height * .035)
    for line in title_lines:
        bounds = draw.textbbox((0, 0), line, font=title_font)
        draw.text(((width - (bounds[2] - bounds[0])) / 2, y), line,
                  font=title_font, fill=(255, 255, 255, 255),
                  stroke_width=max(1, width // 900), stroke_fill=(0, 0, 0, 150))
        y += int(title_size * 1.12)
    if subtitle_lines:
        y += int(subtitle_size * .20)
        for line in subtitle_lines:
            bounds = draw.textbbox((0, 0), line, font=subtitle_font)
            draw.text(((width - (bounds[2] - bounds[0])) / 2, y), line,
                      font=subtitle_font, fill=(255, 255, 255, 240))
            y += int(subtitle_size * 1.15)
    if author:
        author_font = font(font_bold, max(28, int(width * .032)))
        author_text = "by " + author
        bounds = draw.textbbox((0, 0), author_text, font=author_font)
        panel_h = int(height * .105)
        draw.rectangle((0, height - panel_h, width, height),
                       fill=(20, 10, 18, 188))
        draw.text(((width - (bounds[2] - bounds[0])) / 2,
                   height - panel_h / 2 - (bounds[3] - bounds[1]) / 2),
                  author_text, font=author_font, fill=(255, 255, 255, 255))
    background.save(out, "JPEG", quality=96, subsampling=0, dpi=(300, 300),
                    optimize=True)


def platform_metadata(project: dict, platform: str, preflight: dict) -> dict:
    publishing = normalize_publishing(project.get("publishing"), project)
    metadata = project.get("metadata") or {}
    settings = project.get("settings") or {}
    common = {
        "platform": PUBLISHING_PLATFORMS[platform]["label"],
        "title": project.get("title"), "subtitle": project.get("subtitle"),
        "author": settings.get("author"), "language": settings.get("language"),
        "series_name": settings.get("series_name"),
        "book_number": settings.get("book_number"),
        "edition_number": metadata.get("edition_number", "1"),
        "description": metadata.get("description", ""),
        "reading_age": metadata.get("age_range", settings.get("audience", "")),
        "grade_range": metadata.get("grade_range", ""),
        "publisher_imprint": publishing["isbn"]["imprint"],
        "contributors": metadata.get("contributors", "None"),
        "author_bio": metadata.get("author_bio", ""),
        "copyright_text": metadata.get("copyright_text", ""),
        "publishing_rights": metadata.get("publishing_rights", ""),
        "territories": metadata.get("territories", ""),
        "ai_generated_content": metadata.get("ai_generated_content", ""),
        "release_timing": metadata.get("release_timing", ""),
        "pricing": publishing["pricing"][platform],
        "preflight": preflight["platforms"][platform],
    }
    if platform == "kdp":
        common.update({
            "paperback_isbn": publishing["isbn"]["paperback"],
            "keywords": metadata.get("keywords", []),
            "categories": metadata.get("categories", []),
            "kdp_select": publishing["route"]["kdp_select"],
            "expanded_distribution": publishing["route"]["kdp_expanded"],
        })
    elif platform in {"ingram", "d2d"}:
        common.update({
            "paperback_isbn": (publishing["isbn"]["paperback"]
                               if platform == "ingram" else ""),
            "ebook_isbn": publishing["isbn"]["ebook"],
            "bisac_subjects": metadata.get("bisac_subjects", []),
            "keywords": metadata.get("keywords", []),
        })
    elif platform == "blurb":
        common.update({
            "blurb_isbn": publishing["isbn"]["blurb"],
            "tags": metadata.get("blurb_tags", []),
            "edition": publishing["editions"]["blurb"],
        })
    elif platform == "lulu":
        common.update({
            "paperback_isbn": publishing["isbn"]["paperback"],
            "bisac_subjects": metadata.get("bisac_subjects", []),
            "keywords": metadata.get("keywords", []),
            "edition": publishing["editions"]["lulu"],
            "store_url": publishing["tracking"]["lulu"]["store_url"],
        })
    return common


def platform_readme(project: dict, platform: str, preflight: dict) -> str:
    publishing = normalize_publishing(project.get("publishing"), project)
    page_count = preflight["page_count"]
    header = (
        f"{PUBLISHING_PLATFORMS[platform]['label']} publishing package\n"
        f"Book: {project.get('title')}\n"
        f"Prepared: {datetime.now().isoformat(timespec='seconds')}\n\n"
    )
    checks = "\n".join(
        f"[{row['state'].upper()}] {row['title']}: {row['detail']}"
        for row in preflight["platforms"][platform]["checks"])
    if platform == "kdp":
        steps = f"""
FILES
- EPUB: upload as the Kindle ebook manuscript.
- Ebook cover JPG: upload as the Kindle cover image.
- Manuscript PDF: interior-only, full-bleed paperback file.
- Book Cover PDF: back, spine and front paperback cover.

SETUP
- Paperback trim: {preflight['trim_size']}
- Interior pages: {page_count}
- Keep Expanded Distribution off when IngramSpark distributes this ISBN.
- Do not enable KDP Select while the ebook is distributed by Draft2Digital.
"""
    elif platform == "ingram":
        steps = f"""
FILES
- Interior PDF: full-bleed interior prepared at the selected book trim.
- Front cover JPG: 300 PPI design asset for the official cover template.
- Metadata JSON: copy-ready title, BISAC, pricing and ISBN worksheet.

FINAL COVER REQUIRED
IngramSpark cover dimensions depend on the ISBN, binding, paper, colour and
final page count. Request the official template after the interior is final:
{PUBLISHING_PLATFORMS['ingram']['template_url']}
Place the supplied front-cover asset into that template and retain its barcode.
Do not submit a KDP wraparound cover as an Ingram cover.
"""
    elif platform == "d2d":
        steps = """
FILES
- EPUB: upload as the finished ebook file.
- Ebook cover JPG: 1600 x 2400 portrait cover.
- Metadata JSON: BISAC, keywords and listing worksheet.

DISTRIBUTION
This profile is ebook-only. Exclude Amazon because this book is published
directly through KDP. Draft2Digital Print is not used for square colour books.
"""
    elif platform == "blurb":
        edition = publishing["editions"]["blurb"]
        label = BLURB_EDITIONS[edition["trim"]]["label"]
        steps = f"""
FILES
- Layout proof PDF: a visual adaptation to {label}.
- Square front cover JPG: 300 PPI design asset.
- Metadata JSON: listing tags, description, pricing and edition record.

SPECIFICATION CONFIRMATION REQUIRED
Blurb does not offer an exact 8.5 x 8.5 photo-book edition. Cover and exported
page dimensions also vary by size, cover and paper. Confirm the current exact
measurements in Blurb's specification tool before final upload:
{PUBLISHING_PLATFORMS['blurb']['spec_url']}
Use Blurb direct sales only when IngramSpark already handles wide distribution.
"""
    else:
        edition = publishing["editions"]["lulu"]
        channel = edition["store_platform"].replace("_", " ").title()
        steps = f"""
FILES
- Interior PDF: single-page, full-bleed interior at {preflight['trim_size']}.
- Front cover JPG: 300 PPI design asset for Lulu's custom cover template.
- Metadata JSON: copy-ready listing, pricing, ISBN and store worksheet.

FINAL COVER REQUIRED
Use Lulu's custom template for the final page count, binding and trim. Build one
PDF containing the back cover, spine and front cover. Do not upload a KDP or
Ingram wrap cover because their spine calculations may differ.
Calculator and template route:
{PUBLISHING_PLATFORMS['lulu']['template_url']}

STORE FULFILMENT
- Selected connection: {channel}
- Create and approve the Lulu Project, then connect it to the matching product.
- Order and approve a physical proof before accepting customer orders.
- Confirm live print cost, shipping, taxes, per-order and storefront fees.
- Lulu PDF settings: {PUBLISHING_PLATFORMS['lulu']['spec_url']}
"""
    return header + steps.strip() + "\n\nPREFLIGHT\n" + checks + "\n"


def build_platform_package(project: dict, platform: str, stamp: str) -> dict:
    if platform not in PUBLISHING_PLATFORMS:
        raise ValueError("Unknown publishing platform")
    p = project
    preflight = publishing_preflight(p)
    prefix = EXPORT_DIR / f"{slug(p['title'])}-{stamp}-{platform}"
    files = []
    outputs = {}

    metadata_path = prefix.parent / (prefix.name + "-metadata.json")
    metadata_path.write_text(json.dumps(
        platform_metadata(p, platform, preflight), ensure_ascii=False, indent=2),
        encoding="utf-8")
    readme = prefix.parent / (prefix.name + "-README.txt")
    readme.write_text(platform_readme(p, platform, preflight), encoding="utf-8")
    files.extend([metadata_path, readme])

    if platform == "kdp":
        epub = prefix.with_suffix(".epub")
        ebook_cover = prefix.parent / (prefix.name + "-ebook-cover.jpg")
        export_epub(p, epub)
        export_marketing_cover(p, ebook_cover, (1600, 2560))
        files.extend([epub, ebook_cover])
        outputs.update(epub=str(epub), ebook_cover=str(ebook_cover))
        if preflight["platforms"]["kdp"]["checks"][1]["state"] != "block":
            manuscript = prefix.parent / (prefix.name + "-manuscript.pdf")
            cover = prefix.parent / (prefix.name + "-book-cover.pdf")
            export_pdf(p, manuscript, include_cover=False, bleed=True, pad_even=True)
            export_print_cover(p, cover, preflight["page_count"])
            files.extend([manuscript, cover])
            outputs.update(manuscript_pdf=str(manuscript), cover_pdf=str(cover))
    elif platform == "ingram":
        interior = prefix.parent / (prefix.name + "-interior-bleed.pdf")
        front_cover = prefix.parent / (prefix.name + "-front-cover-300ppi.jpg")
        export_pdf(p, interior, include_cover=False, bleed=True, pad_even=True)
        export_marketing_cover(p, front_cover, (2550, 2550))
        files.extend([interior, front_cover])
        outputs.update(interior_pdf=str(interior), front_cover=str(front_cover),
                       cover_template_url=PUBLISHING_PLATFORMS["ingram"][
                           "template_url"])
    elif platform == "d2d":
        epub = prefix.with_suffix(".epub")
        ebook_cover = prefix.parent / (prefix.name + "-ebook-cover-1600x2400.jpg")
        export_epub(p, epub)
        export_marketing_cover(p, ebook_cover, (1600, 2400))
        files.extend([epub, ebook_cover])
        outputs.update(epub=str(epub), ebook_cover=str(ebook_cover))
    elif platform == "blurb":
        publishing = normalize_publishing(p.get("publishing"), p)
        selected = publishing["editions"]["blurb"]["trim"]
        layout_trim = BLURB_EDITIONS[selected]["layout_trim"]
        proof = prefix.parent / (prefix.name + f"-{selected}-layout-proof.pdf")
        cover = prefix.parent / (prefix.name + f"-{selected}-front-cover.jpg")
        export_pdf(p, proof, include_cover=False, pad_even=True,
                   trim_override=layout_trim)
        cover_px = 2100 if selected == "7x7" else 3600
        export_marketing_cover(p, cover, (cover_px, cover_px))
        files.extend([proof, cover])
        outputs.update(layout_proof=str(proof), front_cover=str(cover),
                       specification_url=PUBLISHING_PLATFORMS["blurb"]["spec_url"])
    else:
        w_in, h_in = trim_size(p.get("settings") or {})
        interior = prefix.parent / (prefix.name + "-interior-bleed.pdf")
        front_cover = prefix.parent / (prefix.name + "-front-cover-300ppi.jpg")
        export_pdf(p, interior, include_cover=False, bleed=True, pad_even=True)
        export_marketing_cover(
            p, front_cover,
            (int(round(w_in * 300)), int(round(h_in * 300))))
        files.extend([interior, front_cover])
        outputs.update(
            interior_pdf=str(interior), front_cover=str(front_cover),
            cover_template_url=PUBLISHING_PLATFORMS["lulu"]["template_url"],
            pdf_settings_url=PUBLISHING_PLATFORMS["lulu"]["spec_url"])

    package = prefix.with_suffix(".zip")
    with zipfile.ZipFile(package, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in files:
            archive.write(path, path.name)
    platform_check = dict(preflight["platforms"][platform])
    platform_check["global_issues"] = preflight["global_issues"]
    if any(row["state"] == "block" for row in preflight["global_issues"]):
        platform_check["ready"] = False
    outputs.update({
        "platform": platform, "platform_label": PUBLISHING_PLATFORMS[platform]["label"],
        "bundle": str(package), "metadata": str(metadata_path),
        "readme": str(readme), "preflight": platform_check,
    })
    return outputs


def _export_project_unlocked(project_id: str) -> dict:
    p = load_project(project_id)
    build_signature = publishing_build_signature(p)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    base = EXPORT_DIR / f"{slug(p['title'])}-{stamp}"
    proof_pdf = base.parent / (base.name + "-review.pdf")
    docx = base.with_suffix(".docx")
    bundle = base.with_suffix(".zip")
    preflight = publishing_preflight(p)
    publishing = normalize_publishing(p.get("publishing"), p)
    route = publishing["route"]
    enabled = []
    if route["kdp_ebook"] or route["kdp_paperback"]:
        enabled.append("kdp")
    if route["ingram_print"]:
        enabled.append("ingram")
    if route["d2d_ebook"]:
        enabled.append("d2d")
    if route["blurb_direct"] or route["blurb_global"]:
        enabled.append("blurb")
    if route["lulu_direct"]:
        enabled.append("lulu")
    update_current_job("Rendering the complete visual review", 8)
    export_pdf(p, proof_pdf)
    update_current_job("Building the editable Word manuscript", 18)
    export_docx(p, docx)
    platform_packages = {}
    for index, platform_name in enumerate(enabled):
        progress = 24 + int(index / max(1, len(enabled)) * 60)
        update_current_job(
            "Building " + PUBLISHING_PLATFORMS[platform_name]["label"]
            + " package", progress)
        platform_packages[platform_name] = build_platform_package(
            p, platform_name, stamp)
    manifest = base.parent / (base.name + "-publishing-manifest.json")
    manifest.write_text(json.dumps({
        "project_id": project_id, "title": p.get("title"), "created_at": now(),
        "build_version": PUBLISHING_BUILD_VERSION,
        "build_signature": build_signature,
        "publishing": publishing, "preflight": preflight,
        "platform_packages": platform_packages,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    update_current_job("Packaging the multi-platform publishing workspace", 90)
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as z:
        for path in [proof_pdf, docx, manifest]:
            z.write(path, path.name)
        for platform_name, row in platform_packages.items():
            platform_zip = Path(row["bundle"])
            z.write(platform_zip, f"platforms/{platform_zip.name}")
        z.write(project_file(project_id), "project.json")
        for image in (project_dir(project_id) / "images").glob("*"):
            if image.is_file():
                z.write(image, "images/" + image.name)
        for reference in (project_dir(project_id) / "references").glob("*"):
            if reference.is_file():
                z.write(reference, "references/" + reference.name)
    p["stage"] = "exported"
    p["build"] = {
        "version": PUBLISHING_BUILD_VERSION,
        "signature": build_signature,
        "built_at": now(),
        "bundle": str(bundle),
        "manifest": str(manifest),
        "platforms": enabled,
    }
    publishing_outputs = p.setdefault("publishing_outputs", {})
    for platform_name, row in platform_packages.items():
        publishing_outputs[platform_name] = {
            "version": PUBLISHING_BUILD_VERSION,
            "signature": build_signature,
            "built_at": p["build"]["built_at"],
            "bundle": row.get("bundle", ""),
            "ready": bool((row.get("preflight") or {}).get("ready")),
        }
    p["history"].append({
        "at": now(), "action": "Exported multi-platform publishing packages: "
        + ", ".join(PUBLISHING_PLATFORMS[name]["label"] for name in enabled)
    })
    save_project(p)
    studio_event(
        "export.package_created", project_id=project_id,
        bundle=str(bundle), proof_pdf=str(proof_pdf),
        platforms=enabled, preflight_blockers=sum(
            row["state"] == "block" for row in preflight["global_issues"]),
    )
    kdp = platform_packages.get("kdp") or {}
    return {
        "epub": kdp.get("epub", ""), "pdf": str(proof_pdf),
        "manuscript_pdf": kdp.get("manuscript_pdf", ""),
        "cover_pdf": kdp.get("cover_pdf", ""),
        "print_setup": {
            "print_eligible": (
                preflight["platforms"]["kdp"]["ready"]
                and not any(row["state"] == "block"
                            for row in preflight["global_issues"])),
            "page_count": preflight["page_count"],
            "trim_size": preflight["trim_size"],
            "warning": "Review the platform preflight before submission.",
        },
        "docx": str(docx), "metadata": str(manifest),
        "bundle": str(bundle), "platform_packages": platform_packages,
        "preflight": preflight,
    }


def export_project(project_id: str) -> dict:
    """Build one book while serialising duplicate-title output filenames."""
    project = load_project(project_id)
    output_key = slug(str(project.get("title") or "Untitled Book"))
    with _export_name_locks_guard:
        build_lock = _export_name_locks.setdefault(output_key, threading.Lock())
    with build_lock:
        return _export_project_unlocked(project_id)


def export_platform(project_id: str, platform: str) -> dict:
    p = load_project(project_id)
    build_signature = publishing_build_signature(p)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    update_current_job(
        "Building " + PUBLISHING_PLATFORMS.get(platform, {}).get(
            "label", "publishing") + " package", 20)
    result = build_platform_package(p, platform, stamp)
    p.setdefault("history", []).append({
        "at": now(), "action": "Exported " + result["platform_label"] + " package"
    })
    p.setdefault("publishing_outputs", {})[platform] = {
        "version": PUBLISHING_BUILD_VERSION,
        "signature": build_signature,
        "built_at": now(), "bundle": result.get("bundle", ""),
        "ready": bool((result.get("preflight") or {}).get("ready")),
    }
    save_project(p)
    studio_event(
        "export.platform_package_created", project_id=project_id,
        platform=platform, bundle=result.get("bundle"),
        ready=(result.get("preflight") or {}).get("ready", False),
    )
    return result


def _tail_text(path: Path, max_chars=50000) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="replace")[-max_chars:]
    except OSError:
        return ""


def diagnostic_bundle(project_id: str = "") -> tuple[str, bytes]:
    """Build a local, credential-free support bundle for fast troubleshooting."""
    project = None
    if project_id:
        project = load_project(project_id)
    kindle_events = read_events(EVENT_LOG.path, 1500)
    image_events = read_events(IMAGE_EVENT_LOG, 1500)
    if project_id:
        correlated_jobs = {
            row.get("job_id") for row in kindle_events
            if row.get("project_id") == project_id and row.get("job_id")
        }
        kindle_events = [
            row for row in kindle_events
            if row.get("project_id") in ("", project_id)
            or row.get("job_id") in correlated_jobs
        ]
        local_jobs = {
            row.get("local_image_job_id") for row in kindle_events
            if row.get("local_image_job_id")
        }
        if local_jobs:
            image_events = [
                row for row in image_events
                if row.get("job_id") in local_jobs
            ] or image_events[-300:]
        else:
            image_events = image_events[-300:]
    with _jobs_lock:
        recent_jobs = []
        for jid, job in list(_jobs.items())[-100:]:
            recent_jobs.append({
                "job_id": jid, **{k: job.get(k) for k in (
                    "status", "stage", "progress", "kind", "target", "label",
                    "provider", "started_at", "updated_at", "finished_at", "error",
                    "local_image_job_id",
                )}
            })
    project_summary = {}
    if project:
        project_summary = {
            "id": project.get("id"), "title": project.get("title"),
            "stage": project.get("stage"), "created_at": project.get("created_at"),
            "updated_at": project.get("updated_at"), "settings": project.get("settings"),
            "page_count": len(project.get("pages") or []),
            "pages_with_images": sum(bool(p.get("image")) for p in project.get("pages") or []),
            "cover_has_image": bool((project.get("cover") or {}).get("image")),
            "character_count": len(project.get("character_bible") or []),
            "reference_count": len(project.get("reference_images") or []),
            "history_tail": (project.get("history") or [])[-30:],
        }
    summary = {
        "created_at": now(), "service": "Kindle Book Studio",
        "python": sys.version, "text_model": TEXT_MODEL,
        "cloud_image_model": IMAGE_MODEL, "xai_key_configured": bool(api_key()),
        "local_image_base": LOCAL_IMAGE_BASE,
        "paths": {"library": str(LIBRARY_DIR), "exports": str(EXPORT_DIR)},
        "optional_tools": {
            "epubcheck": shutil.which("epubcheck") or "not installed",
            "ace": shutil.which("ace") or "not installed",
            "kindle_create": "/Applications/Kindle Create.app"
            if Path("/Applications/Kindle Create.app").exists() else "not installed",
            "kindle_previewer": "/Applications/Kindle Previewer 3.app"
            if Path("/Applications/Kindle Previewer 3.app").exists()
            else "not installed",
        },
        "project": project_summary, "recent_jobs": recent_jobs,
        "event_counts": {"kindle": len(kindle_events), "image": len(image_events)},
        "privacy": (
            "API keys, authorization headers, base64 images and full prompts are excluded."
        ),
    }
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "diagnostics-summary.json",
            json.dumps(summary, ensure_ascii=False, indent=2),
        )
        archive.writestr(
            "kindle-events.jsonl",
            "\n".join(json.dumps(row, ensure_ascii=False) for row in kindle_events) + "\n",
        )
        archive.writestr(
            "image-events.jsonl",
            "\n".join(json.dumps(row, ensure_ascii=False) for row in image_events) + "\n",
        )
        for name, path in (
            ("kindle-service-errors.txt", BASE_DIR / "service.err"),
            ("image-service-errors.txt", Path("/Users/joebains/omlx-image/service.err")),
        ):
            archive.writestr(name, _tail_text(path))
    filename = "kindle-studio-diagnostics"
    if project:
        filename += "-" + slug(project.get("title", "book"))
    filename += "-" + datetime.now().strftime("%Y%m%d-%H%M%S") + ".zip"
    studio_event(
        "diagnostics.created", project_id=project_id, filename=filename,
        kindle_events=len(kindle_events), image_events=len(image_events),
    )
    return filename, out.getvalue()


JOB_ACTIVITY = {
    "generate_story": ("Writing and planning the book", "Grok cloud", "task"),
    "generate_market_concept": ("Researching an original concept", "Grok web + cloud", "task"),
    "create_series_continuation": ("Creating the next book in this series", "Grok cloud", "task"),
    "quality_review_story": ("Reviewing story and picture logic", "Grok cloud", "task"),
    "populate_titles": ("Writing missing titles", "Grok cloud", "task"),
    "populate_quotes": ("Writing quotations and wisdom", "Grok cloud", "task"),
    "generate_metadata": ("Creating the publishing listing", "Grok cloud", "task"),
    "regenerate_cover": ("Revising the cover direction", "Grok cloud", "task"),
    "regenerate_page": ("Revising a story page", "Grok cloud", "task"),
    "generate_image": ("Creating a book illustration", "Image provider", "image"),
    "generate_character_reference": ("Creating a character reference", "Grok Imagine cloud", "image"),
    "repair_image_text": ("Removing detected image text", "This Mac", "image"),
    "export_project": ("Building the book package", "This Mac", "export"),
    "export_platform": ("Building a platform package", "This Mac", "export"),
    "rebuild_outdated_books": ("Updating the book library", "This Mac", "export"),
    "cleanup_old_builds": ("Cleaning old publishing builds", "This Mac", "export"),
    "complete_missing_series_books": (
        "Completing the 28-day series", "Grok + local FLUX + This Mac", "image"),
    "choose_book_narrator": ("Casting the book narrator", "Grok cloud", "task"),
    "plan_book_sound_effects": ("Planning the story sound effects", "Qwen3.6 local", "audio"),
    "generate_book_sound_effect": ("Creating a local sound effect", "Stable Audio 3 local", "audio"),
    "generate_all_book_sound_effects": ("Creating local sound effects", "Stable Audio 3 local", "audio"),
    "mix_book_soundtrack": ("Mixing narration and sound effects", "This Mac", "audio"),
    "generate_book_audio": ("Creating expressive Qwen3-TTS narration", "Qwen3-TTS BF16 local", "audio"),
    "generate_book_audio_preview": ("Testing the first two pages", "Qwen3-TTS BF16 local", "audio"),
    "generate_book_video": ("Animating the narrated book", "This Mac + Qwen3-TTS", "video"),
}


def gpu_user_label(owner: str, gpu_util) -> str:
    """Turn the shared local-model lease into one compact resource label."""
    owner = str(owner or "").strip()
    if owner.startswith("image:"):
        engine = owner.split(":", 1)[1]
        return {
            "hidream": "HiDream local",
            "kontext": "FLUX.1 local",
            "flux_2_klein_4b_local": "FLUX 2 Klein local",
        }.get(engine, f"Local image · {engine}")
    if owner.startswith("song:"):
        return "ACE-Step music"
    if owner.startswith("sfx:"):
        return "Stable Audio 3 sound effects"
    if owner.startswith("video:"):
        return "Local video model"
    if owner:
        return owner.replace(":", " · ")[:48]
    try:
        return "Idle" if float(gpu_util or 0) < 5 else "macOS / another app"
    except (TypeError, ValueError):
        return "Idle"


def activity_snapshot() -> dict:
    """Small, result-free summary used by the always-visible header monitor."""
    with _jobs_lock:
        rows = []
        for job_id, job in _jobs.items():
            rows.append({
                "id": job_id, "status": job.get("status"),
                "function": job.get("function") or "",
                "label": job.get("label") or "Studio task",
                "provider": job.get("provider") or "",
                "kind": job.get("kind") or "task",
                "target": job.get("target") or "",
                "stage": job.get("stage") or "working",
                "progress": int(job.get("progress") or 0),
                "started_at": job.get("started_at") or "",
                "updated_at": job.get("updated_at") or "",
                "finished_at": job.get("finished_at") or "",
                "error": str(job.get("error") or "")[:240],
                "queue_label": str(job.get("queue_label") or "")[:160],
                "details": dict(job.get("details") or {}),
            })
    active = sorted(
        (row for row in rows if row["status"] == "running"),
        key=lambda row: row["started_at"])
    recent = sorted(
        (row for row in rows if row["status"] in {"done", "error", "canceled"}),
        key=lambda row: row["finished_at"], reverse=True)[:1]
    return {"active_jobs": active, "active_count": len(active), "recent": recent}


def run_job(job_id: str, fn, *args):
    function_name = getattr(fn, "__name__", "")
    label, provider, kind = JOB_ACTIVITY.get(
        function_name, ("Working in Kindle Studio", "", "task"))
    is_image = kind == "image"
    target = str(args[-1]) if is_image and args else ""
    project_id = ("" if function_name in {
        "rebuild_outdated_books", "cleanup_old_builds"
    } else str(args[0]) if args else "")
    started_at = now()
    started_clock = time.monotonic()
    with _jobs_lock:
        _jobs[job_id] = {
            "status": "running",
            "stage": "preparing image" if is_image else "working",
            "progress": 1 if is_image else 5,
            "kind": kind, "target": target, "label": label,
            "provider": provider, "started_at": started_at,
            "function": function_name,
            "updated_at": started_at, "project_id": project_id,
        }
    _job_context.job_id = job_id
    _job_context.project_id = project_id
    _job_context.function_name = function_name
    _job_context.stage_prefix = ""
    _job_context.progress_range = None
    _job_context.progress_monotonic = False
    studio_event(
        "job.started", function=function_name, label=label, provider=provider,
        kind=kind, target=target,
    )
    try:
        result = fn(*args)
        canceled_after_result = False
        with _jobs_lock:
            current = _jobs.get(job_id) or {}
            if current.get("cancel_requested"):
                canceled_after_result = True
                current.update(status="canceled", stage="stopped", progress=0,
                               error=None, cancel_requested=True,
                               finished_at=now(), updated_at=now())
            else:
                current.update(status="done", stage="complete", progress=100,
                               result=result, finished_at=now(), updated_at=now())
        studio_event(
            "job.canceled" if canceled_after_result else "job.completed",
            level="warning" if canceled_after_result else "info",
            function=function_name, kind=kind, target=target,
            duration_ms=round((time.monotonic() - started_clock) * 1000),
            result_type=type(result).__name__,
        )
    except Exception as e:
        raw_error = f"{type(e).__name__}: {e}"
        if "used all available credits" in raw_error or "monthly spending limit" in raw_error:
            shown_error = (
                "xAI could not start this request because the account has no "
                "available API credits or has reached its monthly spending limit. "
                "Add credits or raise the spending limit in the xAI Console, then try again."
            )
        else:
            shown_error = raw_error
        with _jobs_lock:
            current = _jobs.get(job_id) or {}
            canceled = isinstance(e, ImageJobCanceled) or current.get("cancel_requested")
            if canceled:
                current.update(status="canceled", stage="stopped", progress=0,
                               error=None, cancel_requested=True,
                               finished_at=now(), updated_at=now())
                studio_event(
                    "job.canceled", level="warning", function=function_name,
                    kind=kind, target=target,
                    duration_ms=round((time.monotonic() - started_clock) * 1000),
                )
            else:
                print(f"[{now()}] job {job_id} failed: {raw_error}",
                      file=sys.stderr, flush=True)
                traceback.print_exc(file=sys.stderr)
                current.update(status="error", stage="needs attention",
                               progress=100, error=shown_error,
                               finished_at=now(), updated_at=now())
                EVENT_LOG.exception(
                    "job.failed", e, job_id=job_id, project_id=project_id,
                    function=function_name, kind=kind, target=target,
                    duration_ms=round((time.monotonic() - started_clock) * 1000),
                )
    finally:
        # A cloud provider may have accepted a request even when its response or
        # downstream download failed. Preserve that possible cost visibly.
        with _cost_lock:
            unsettled = [
                reservation_id
                for reservation_id, row in _cost_reservations.items()
                if row.get("job_id") == job_id
            ]
        for reservation_id in unsettled:
            settle_cloud_cost(
                reservation_id, estimated=True, status="uncertain")
        _job_context.job_id = ""
        _job_context.project_id = ""
        _job_context.function_name = ""
        _job_context.stage_prefix = ""
        _job_context.progress_range = None
        _job_context.progress_monotonic = False


def start_job(fn, *args) -> str:
    jid = "job_" + uuid.uuid4().hex[:12]
    EVENT_LOG.event(
        "job.queued", job_id=jid, project_id=str(args[0]) if args else "",
        function=getattr(fn, "__name__", ""),
    )
    threading.Thread(target=run_job, args=(jid, fn, *args), daemon=True).start()
    return jid


class Handler(BaseHTTPRequestHandler):
    server_version = "omlx-kindle/1.0"

    def log_message(self, *args):
        pass

    def cors(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, PUT, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")

    def json(self, code: int, payload):
        body = json.dumps(payload, ensure_ascii=False).encode()
        try:
            self.send_response(code)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.cors()
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        except (BrokenPipeError, ConnectionResetError):
            return

    def body(self) -> dict:
        n = int(self.headers.get("Content-Length", "0") or 0)
        try:
            return json.loads(self.rfile.read(n) or b"{}")
        except Exception:
            return {}

    def do_OPTIONS(self):
        self.send_response(204); self.cors(); self.end_headers()

    def do_GET(self):
        path = urllib.parse.urlparse(self.path).path
        try:
            if path in ("/", "/index.html"):
                return self.serve_file(STATIC_DIR / "index.html")
            if path == "/health":
                return self.json(200, {
                    "ok": True, "service": "Kindle Book Studio",
                    "grok_key_configured": bool(api_key()), "text_model": TEXT_MODEL,
                    "image_model": IMAGE_MODEL, "projects": len(list_projects()),
                })
            if path == "/api/diagnostics/download":
                query = urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)
                project_id = str((query.get("project") or [""])[0])
                filename, payload = diagnostic_bundle(project_id)
                return self.serve_download(filename, payload, "application/zip")
            if path == "/api/presets":
                return self.json(200, {
                    "book_types": [{"id": x, "label": y} for x, y in BOOK_TYPES],
                    "picture_book_types": [
                        {"id": x, "label": y} for x, y in PICTURE_BOOK_TYPES],
                    "popular_kids_blueprints": POPULAR_KIDS_BLUEPRINTS,
                    "genres": GENRES,
                    "image_styles": [{"id": x, "label": y} for x, y in IMAGE_STYLES],
                    "trims": [{"id": x[0], "label": x[1]} for x in TRIMS],
                    "reading_levels": READING_LEVELS, "page_counts": PAGE_COUNTS,
                    "series_name": SERIES_NAME,
                    "series_books": series_books_catalog(),
                    "text_model": TEXT_MODEL, "image_model": IMAGE_MODEL,
                    "image_engines": IMAGE_ENGINES,
                    "publishing_platforms": PUBLISHING_PLATFORMS,
                    "publishing_statuses": sorted(PUBLISHING_STATUSES),
                    "narration_engine": NARRATION_TTS_ENGINE,
                    "narrators": [
                        {"id": key, **value}
                        for key, value in QWEN3_NARRATORS.items()
                    ],
                    # Compatibility for an already-open Studio tab. It will
                    # receive the new Qwen speakers even before a hard refresh.
                    "kokoro_narrators": [
                        {"id": key, **value}
                        for key, value in QWEN3_NARRATORS.items()
                    ],
                    "blurb_editions": [
                        {"id": key, **value}
                        for key, value in BLURB_EDITIONS.items()
                    ],
                })
            if path == "/api/projects":
                return self.json(200, {"projects": list_projects()})
            if path == "/api/library/status":
                return self.json(200, library_status())
            if path == "/api/series/status":
                return self.json(200, series_catalog_status())
            if path == "/api/costs":
                query = urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)
                project_id = str((query.get("project") or [""])[0])
                return self.json(200, cloud_cost_snapshot(project_id))
            if path == "/api/image-monitor":
                activity = activity_snapshot()
                try:
                    stats = local_image_json("/stats", timeout=3)
                    health = local_image_json("/health", timeout=3)
                    lease = health.get("model_memory_lease") or {}
                    owner = str(lease.get("owner") or "")
                    image_job = {}
                    if owner.startswith("image:") and lease.get("job_id"):
                        try:
                            status = local_image_json(
                                "/status?" + urllib.parse.urlencode({
                                    "id": lease["job_id"]}), timeout=3)
                            image_job = {
                                "active": status.get("status") in (
                                    "queued", "running"),
                                "engine": owner.split(":", 1)[-1],
                                "stage": status.get("stage") or "creating image",
                                "progress": int(status.get("progress") or 0),
                            }
                        except Exception:
                            image_job = {"active": True, "engine": "local",
                                         "stage": "creating image", "progress": 0}
                    return self.json(200, {
                        "available": True, **stats,
                        "queue": int(health.get("queue") or 0),
                        "image_job": image_job,
                        "gpu_user": gpu_user_label(
                            owner, stats.get("gpu_util")),
                        **activity,
                    })
                except Exception:
                    return self.json(200, {
                        "available": False, "queue": 0, "image_job": {},
                        **activity,
                    })
            if path.startswith("/api/projects/"):
                parts = path.strip("/").split("/")
                pid = parts[2]
                if len(parts) == 3:
                    return self.json(200, load_project(pid))
                if len(parts) == 4 and parts[3] == "publishing-preflight":
                    return self.json(200, publishing_preflight(load_project(pid)))
                if len(parts) == 5 and parts[3] == "images":
                    rel = "images/" + os.path.basename(parts[4])
                    target = image_abs(load_project(pid), rel)
                    if not target:
                        return self.json(404, {"error": "image not found"})
                    return self.serve_file(target)
                if len(parts) == 5 and parts[3] == "references":
                    rel = "references/" + os.path.basename(parts[4])
                    target = project_dir(pid) / rel
                    if not target.exists() or not target.is_file():
                        return self.json(404, {"error": "reference image not found"})
                    return self.serve_file(target)
                if len(parts) == 5 and parts[3] == "media":
                    name = os.path.basename(parts[4])
                    target = project_dir(pid) / "media" / name
                    if not target.exists() or not target.is_file():
                        return self.json(404, {"error": "media file not found"})
                    return self.serve_file(target)
            if path.startswith("/api/jobs/"):
                jid = path.rsplit("/", 1)[-1]
                with _jobs_lock:
                    job = dict(_jobs.get(jid) or {})
                return self.json(200 if job else 404, job or {"error": "unknown job"})
            return self.json(404, {"error": "not found"})
        except Exception as e:
            return self.json(500, {"error": f"{type(e).__name__}: {e}"})

    def serve_file(self, path: Path):
        if not path.exists() or not path.is_file():
            return self.json(404, {"error": "not found"})
        data = path.read_bytes()
        try:
            self.send_response(200)
            self.send_header("Content-Type", mimetypes.guess_type(path.name)[0] or
                             "application/octet-stream")
            self.send_header("Cache-Control", "no-store")
            self.cors()
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
        except (BrokenPipeError, ConnectionResetError):
            return

    def serve_download(self, filename: str, data: bytes,
                       content_type="application/octet-stream"):
        try:
            self.send_response(200)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
            self.send_header("Cache-Control", "no-store")
            self.cors()
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
        except (BrokenPipeError, ConnectionResetError):
            return

    def do_POST(self):
        path = urllib.parse.urlparse(self.path).path
        data = self.body()
        try:
            if path == "/api/images/stop":
                return self.json(200, cancel_image_jobs())
            if path == "/api/cost-settings":
                return self.json(200, {
                    "settings": save_cloud_cost_settings(data),
                    "costs": cloud_cost_snapshot(str(data.get("project_id") or "")),
                })
            if path == "/api/library/rebuild-outdated":
                return self.json(202, {
                    "job_id": start_job(
                        rebuild_outdated_books,
                        max(2, min(4, int(data.get("parallel_books") or 2))))})
            if path == "/api/library/cleanup-old-builds":
                return self.json(202, {"job_id": start_job(cleanup_old_builds)})
            if path == "/api/series/titles":
                entry = add_series_book(data)
                return self.json(201, {
                    "entry": entry, "status": series_catalog_status()})
            if path == "/api/series/build-missing":
                return self.json(202, {"job_id": start_job(
                    complete_missing_series_books,
                    str(data.get("source_project_id") or ""),
                    str(data.get("series_name") or SERIES_NAME),
                    max(2, min(4, int(data.get("parallel_books") or 2))))})
            if path == "/api/projects":
                return self.json(201, new_project(data))
            if path.startswith("/api/projects/"):
                parts = path.strip("/").split("/")
                pid = parts[2]
                action = parts[3] if len(parts) > 3 else ""
                if action == "publication":
                    return self.json(200, set_project_publication(pid, data))
                if action == "review-approval":
                    return self.json(200, set_project_review_approval(pid, data))
                if action == "delete":
                    return self.json(200, delete_project_to_trash(
                        pid, str(data.get("confirm_title") or "")))
                if action == "generate":
                    return self.json(202, {"job_id": start_job(generate_story, pid)})
                if action == "quality-review":
                    return self.json(202, {"job_id": start_job(
                        quality_review_story, pid)})
                if action == "generate-market-concept":
                    return self.json(202, {"job_id": start_job(
                        generate_market_concept, pid)})
                if action == "create-series-book":
                    return self.json(202, {"job_id": start_job(
                        create_series_continuation, pid,
                        data.get("topic", ""), data.get("adventure", ""))})
                if action == "populate-titles":
                    return self.json(202, {"job_id": start_job(populate_titles, pid)})
                if action == "populate-quotes":
                    return self.json(202, {"job_id": start_job(
                        populate_quotes, pid, int(data.get("page", 0) or 0))})
                if action == "generate-metadata":
                    return self.json(202, {"job_id": start_job(generate_metadata, pid)})
                if action == "regenerate-cover":
                    return self.json(202, {"job_id": start_job(
                        regenerate_cover, pid, data.get("request", ""))})
                if action == "regenerate-page":
                    return self.json(202, {"job_id": start_job(
                        regenerate_page, pid, int(data["page"]), data.get("request", ""))})
                if action == "generate-image":
                    return self.json(202, {"job_id": start_job(
                        generate_image, pid, str(data["target"]),
                        str(data.get("engine_override") or ""))})
                if action == "repair-image-text":
                    return self.json(202, {"job_id": start_job(
                        repair_image_text, pid, str(data["target"]))})
                if action == "restore-previous-image":
                    return self.json(200, restore_previous_image(
                        pid, str(data["target"])))
                if action == "clear-artwork":
                    return self.json(200, clear_project_artwork(
                        pid, str(data.get("reason") or ""),
                        bool(data.get("disconnect_inherited_series"))))
                if action == "generate-character-reference":
                    return self.json(202, {"job_id": start_job(
                        generate_character_reference, pid,
                        int(data.get("character_index", -1)))})
                if action == "upload-reference":
                    return self.json(200, upload_reference(pid, data))
                if action == "export":
                    return self.json(202, {"job_id": start_job(export_project, pid)})
                if action == "export-platform":
                    return self.json(202, {"job_id": start_job(
                        export_platform, pid, str(data.get("platform") or ""))})
                if action == "choose-narrator":
                    return self.json(202, {"job_id": start_job(
                        choose_book_narrator, pid)})
                if action == "direct-video":
                    return self.json(202, {"job_id": start_job(
                        direct_book_video, pid)})
                if action == "plan-sound-effects":
                    return self.json(202, {"job_id": start_job(
                        plan_book_sound_effects, pid)})
                if action == "generate-sound-effect":
                    return self.json(202, {"job_id": start_job(
                        generate_book_sound_effect, pid,
                        str(data.get("scene_id") or ""),
                        str(data.get("cue_id") or ""))})
                if action == "generate-all-sound-effects":
                    return self.json(202, {"job_id": start_job(
                        generate_all_book_sound_effects, pid)})
                if action == "mix-soundtrack":
                    return self.json(202, {"job_id": start_job(
                        mix_book_soundtrack, pid)})
                if action == "generate-audio":
                    return self.json(202, {"job_id": start_job(
                        generate_book_audio, pid)})
                if action == "generate-audio-preview":
                    return self.json(202, {"job_id": start_job(
                        generate_book_audio_preview, pid)})
                if action == "generate-video":
                    return self.json(202, {"job_id": start_job(
                        generate_book_video, pid)})
            return self.json(404, {"error": "not found"})
        except Exception as e:
            return self.json(500, {"error": f"{type(e).__name__}: {e}"})

    def do_PUT(self):
        path = urllib.parse.urlparse(self.path).path
        data = self.body()
        try:
            if path.startswith("/api/projects/"):
                pid = path.strip("/").split("/")[2]
                current = load_project(pid)
                incoming_id = str(data.get("id") or "").strip()
                if incoming_id and incoming_id != pid:
                    return self.json(409, {"error": (
                        "This browser tab tried to save one book into another. "
                        "Nothing was saved; reload the intended book and try again.")})
                identity_error = protected_series_identity_error(current, data)
                if identity_error:
                    return self.json(409, {"error": identity_error})
                current_metadata = current.get("metadata") or {}
                incoming_metadata = data.get("metadata") or {}
                if isinstance(incoming_metadata, dict):
                    current_description = str(
                        current_metadata.get("description") or "")
                    incoming_description = str(
                        incoming_metadata.get("description") or "")
                    if incoming_description != current_description:
                        stale_save = bool(
                            data.get("updated_at")
                            and data.get("updated_at") != current.get("updated_at"))
                        if (stale_save and current_description
                                and current_metadata.get(
                                    "description_source") == "manual"):
                            incoming_metadata["description"] = current_description
                            incoming_metadata["description_source"] = "manual"
                            incoming_metadata["description_edited_at"] = str(
                                current_metadata.get("description_edited_at") or now())
                        else:
                            incoming_metadata["description_source"] = "manual"
                            incoming_metadata["description_edited_at"] = now()
                        incoming_metadata.pop("description_generated_at", None)
                    elif current_metadata.get("description_source") == "manual":
                        incoming_metadata["description_source"] = "manual"
                        incoming_metadata["description_edited_at"] = str(
                            current_metadata.get("description_edited_at") or now())
                        incoming_metadata.pop("description_generated_at", None)
                    data["metadata"] = incoming_metadata
                revision_snapshot(current, "manual-save")
                data["id"] = pid
                data.setdefault("created_at", current.get("created_at", now()))
                data.setdefault("history", current.get("history", []))
                return self.json(200, save_project(
                    preserve_server_assets(current, data)))
            return self.json(404, {"error": "not found"})
        except Exception as e:
            return self.json(500, {"error": f"{type(e).__name__}: {e}"})


if __name__ == "__main__":
    print(f"Kindle Book Studio ready at http://{HOST}:{PORT}", flush=True)
    EVENT_LOG.event(
        "service.started", host=HOST, port=PORT, text_model=TEXT_MODEL,
        image_model=IMAGE_MODEL, project_count=len(list_projects()),
        xai_key_configured=bool(api_key()),
    )
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
